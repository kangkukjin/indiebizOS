"""Create/import documents into owner-controlled storage; never overwrite a name."""
import io
from pathlib import Path

from office_sessions import MAX_BYTES, owner
from office_store import identifier

TEMPLATES = {
    "blank": ("새 문서", []),
    "report": ("보고서", [("개요", "목적과 핵심 내용을 작성하세요."), ("조사 결과", "확인한 사실과 근거를 작성하세요."), ("결론", "다음 행동과 담당자를 정리하세요.")]),
    "minutes": ("회의록", [("회의 정보", "일시 / 참석자 / 안건"), ("논의 내용", "주요 논의를 작성하세요."), ("결정 사항", "결정 / 담당 / 기한")]),
    "proposal": ("제안서", [("배경", "해결할 문제"), ("제안", "방법과 기대 효과"), ("계획", "일정과 비용")]),
    "letter": ("공문", [("수신", "수신자 / 참조"), ("제목", "공문 제목"), ("내용", "요청 사항을 작성하세요.")]),
    "application": ("신청서", [("신청인", "성명 / 연락처"), ("신청 내용", "신청 사항과 이유"), ("확인", "작성일 / 서명")]),
    "contract": ("계약서 초안", [("당사자", "계약 당사자"), ("계약 내용", "목적 / 기간 / 대금 / 의무"), ("서명", "검토 후 서명하세요.")]),
    "resume": ("이력서", [("인적 사항", "이름 / 연락처"), ("경력", "기간 / 조직 / 주요 업무"), ("학력과 역량", "학력 / 자격 / 기술")]),
}


def import_bytes(app, filename, data):
    owner()
    if not filename or Path(filename).name != filename or "/" in filename or "\\" in filename:
        raise ValueError("문서 파일명을 확인하세요")
    if len(data) > MAX_BYTES:
        raise ValueError("문서 크기 상한은 25MB입니다")
    folder = app.store.root / "files" / identifier()
    folder.mkdir(parents=True)
    path = folder / filename
    with path.open("xb") as output:
        output.write(data)
    return app.open(str(path))


def create(app, title="새 문서", format="docx", template="blank"):
    owner()
    if template not in TEMPLATES or format not in {"docx", "md", "txt", "html", "typ", "tex", "pdf"}:
        raise ValueError("지원하지 않는 새 문서 형식 또는 템플릿입니다")
    if not isinstance(title, str) or not title.strip() or len(title) > 120:
        raise ValueError("문서 제목은 1~120자로 입력하세요")
    _, sections = TEMPLATES[template]
    if format == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for heading, text in sections:
            doc.add_heading(heading, 1)
            doc.add_paragraph(text)
        if not sections:
            doc.add_paragraph("")
        buffer = io.BytesIO(); doc.save(buffer); data = buffer.getvalue()
    elif format == "pdf":
        import pymupdf
        doc = pymupdf.open(); doc.new_page(); data = doc.tobytes(); doc.close()
    else:
        text = "# " + title + "\n\n" + "\n\n".join("## " + h + "\n\n" + t for h, t in sections)
        if format == "txt":
            text = text.replace("## ", "").replace("# ", "")
        elif format == "html":
            from html import escape
            text = '<!doctype html><meta charset="utf-8"><h1>' + escape(title) + '</h1>' + ''.join('<h2>'+escape(h)+'</h2><p>'+escape(t)+'</p>' for h,t in sections)
        elif format == "typ":
            text = "= " + title + "\n\n"
        elif format == "tex":
            text = "\\documentclass{article}\n\\begin{document}\nWrite here.\n\\end{document}\n"
        data = text.encode()
    return import_bytes(app, title.strip() + "." + format, data)
