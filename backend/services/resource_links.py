"""Version-pinned sheet imports and document selections for lecture materials.

All references retain immutable bytes in the shared office store. Refresh is
explicit; neither checking a source nor delivering a lecture rewrites a report.
"""
import tempfile
import time
from pathlib import Path

from office_store import digest, identifier
from office_sessions import DocumentConflict, DocumentUnsupported, read_bytes


def package_module(package, module):
    from common.pkg_utils import load_sibling
    root = Path(__file__).resolve().parents[2]
    anchor = root / "data/packages/installed/tools" / package / "handler.py"
    return load_sibling(anchor, module)


class ResourceLinks:
    def __init__(self, workspace):
        self.app = workspace
        self.store = workspace.store
        self.resources = workspace.resources

    def refresh(self, resource_id, expected_revision):
        with self.store.lock():
            row = self.resources.get(resource_id)
            if row["revision_id"] != expected_revision:
                raise DocumentConflict("자료 버전이 바뀌었습니다")
            if row["session_id"]:
                raise DocumentConflict("활성 편집 세션은 먼저 저장·닫기를 완료하세요")
            data = read_bytes(Path(row["source_uri"]))
            key = self.store.blob(data)
            if key == row["source_sha256"]:
                return row
            revision = {"id": identifier(), "document_id": resource_id,
                        "parent_revision_id": row["revision_id"], "blob": key,
                        "created_at": time.time()}
            row.update(revision_id=revision["id"], source_sha256=key)
            with self.store.connect() as conn:
                self.store.put("document", row, conn)
                self.store.put("revision", revision, conn)
                self.store.event(resource_id, {"type": "external_revision", "revision_id": revision["id"]}, conn)
            return row

    def sheet(self, resource_id, revision_id, sheet, cell_range):
        row, revision, data = self.resources.revision(resource_id, revision_id)
        if row["source_format"] not in {"xlsx", "xlsm"}:
            raise DocumentUnsupported("현재 범위 가져오기는 XLSX/XLSM만 지원합니다")
        reader = package_module("system_essentials", "sheet_range_ops")
        # Read the pinned bytes, never a changing source path. The existing
        # package validates ZIP limits, XML safety and the 10000-cell bound.
        with tempfile.TemporaryDirectory(prefix="sheet-reference-") as directory:
            path = Path(directory) / ("source." + row["source_format"])
            path.write_bytes(data)
            result = reader.op_range({"path": str(path), "sheet": sheet, "range": cell_range})
        ref = {"resource_id": resource_id, "revision_id": revision_id,
               "selector": {"kind": "sheet_range", "sheet": result["sheet"], "range": cell_range.upper()},
               "provenance": {"source_uri": row["source_uri"], "source_sha256": revision["blob"],
                              "calculation_state": "stored_cache_unverified"},
               "items": result["items"]}
        return ref

    def _reference(self, reference_id):
        link = self.store.get("resource_link", reference_id)
        self.resources.get(link["target_id"])
        self.resources.get(link["resource_id"])
        return link

    def links(self, document_id):
        self.resources.get(document_id)
        return [r for r in self.store.list("resource_link") if r["target_id"] == document_id]

    def status(self, reference_id):
        ref = self._reference(reference_id)
        source = self.resources.get(ref["resource_id"])
        current = digest(read_bytes(Path(source["source_uri"])))
        return {"reference": ref, "source_changed": current != ref["provenance"]["source_sha256"],
                "observed_sha256": current, "automatic_update": False}

    @staticmethod
    def markdown_table(ref):
        from openpyxl.utils.cell import range_boundaries, get_column_letter
        c1, r1, c2, r2 = range_boundaries(ref["selector"]["range"])
        cells = {v["cell"]: v for v in ref["items"]}
        def value(address):
            cell = cells[address]
            v = cell["value"]
            if cell["formula"] and v is None:
                return "[계산값 없음]"
            if v is None:
                return ""
            return str(v).replace("\\", "\\\\").replace("|", "\\|").replace("\r", " ").replace("\n", " ")
        lines = ["| " + " | ".join(get_column_letter(c) for c in range(c1, c2 + 1)) + " |",
                 "| " + " | ".join("---" for _ in range(c1, c2 + 1)) + " |"]
        for r in range(r1, r2 + 1):
            lines.append("| " + " | ".join(value(f"{get_column_letter(c)}{r}") for c in range(c1, c2 + 1)) + " |")
        return "\n".join(lines) + "\n"

    def sheet_proposal(self, document_id, snapshot_id, start, end, selected_sha256,
                       resource_id, revision_id, sheet, cell_range, linked=False):
        target = self.resources.get(document_id)
        if target["source_format"] not in {"md", "markdown"}:
            raise DocumentUnsupported("현재 표 삽입은 Markdown 소스 문서에서 지원합니다")
        if type(linked) is not bool:
            raise ValueError("linked는 boolean입니다")
        ref = self.sheet(resource_id, revision_id, sheet, cell_range)
        ref.update(id=identifier(), target_id=document_id, linked=linked, created_at=time.time())
        proposal = self.app.propose(document_id, snapshot_id, start, end, selected_sha256,
                                    self.markdown_table(ref))
        proposal["resource_reference"] = ref
        self.store.put("proposal", proposal)
        return proposal

    def deliver(self, document_id, revision_id, start, end, selected_sha256, lecture_id, operation_id):
        row, revision, data = self.resources.revision(document_id, revision_id)
        if not row.get("encoding"):
            raise DocumentUnsupported("강의 선택 전달은 현재 소스 문서에서 지원합니다")
        text = data.decode(row["encoding"], errors="strict")
        if type(start) is not int or type(end) is not int or not 0 <= start < end <= len(text):
            raise ValueError("전달할 원본 버전의 문자 범위를 지정하세요")
        if digest(text[start:end].encode()) != selected_sha256:
            raise DocumentConflict("저장된 원본과 고정한 선택이 다릅니다. 다시 선택하세요")
        with self.store.lock():
            op, cached = self.app._operation(document_id, operation_id,
                ["lecture_material", revision_id, start, end, lecture_id])
            if cached is not None:
                return cached
            lecture = package_module("lecture_workspace", "lecture_store")
            deck = lecture.read_deck(lecture_id)
            reference = {"id": op["id"], "resource_id": document_id, "revision_id": revision_id,
                         "selector": {"kind": "text", "start": start, "end": end},
                         "provenance": {"source_uri": row["source_uri"], "source_sha256": revision["blob"],
                                        "selected_sha256": digest(text[start:end].encode())}}
            entries = [m for m in deck.get("materials", [])
                       if m.get("source_reference", {}).get("id") == op["id"]]
            if entries:
                entry = entries[0]
            else:
                if op.get("status") == "delivering":
                    raise DocumentConflict("강의 전달 결과가 불명확합니다. 강의 재료를 확인하세요. 중복 생성하지 않습니다")
                op["status"] = "delivering"
                self.store.put("operation", op)
                entry = lecture.add_material_from_text(lecture_id, text[start:end],
                    "document-" + op["id"] + ".md", source_reference=reference)
            result = {"lecture_id": lecture_id, "material": entry, "reference": reference}
            op.update(status="completed", result=result)
            self.store.put("operation", op)
            return result
