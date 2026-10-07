"""격자 엔진(브라우저 안 Univer)과 XLSX 사이의 투영 — docs/SPREADSHEET_APP_ON_IBL_PLAN_2026_10_07.md §4.

계산 권위는 격자다. 이 모듈은 셀·수식·서식·병합·열폭·행높이·틀 고정을 양방향으로 옮기고,
격자가 계산한 값을 수식 셀의 캐시(<v>)로 심는다 — 그래야 `[self:read]`·엑셀·LibreOffice 가
파일만 보고도 최신값을 읽는다(엑셀은 fullCalcOnLoad 로 열며 다시 계산한다).

grid.v1 통화:
  {"version": 1, "date_system": "1900", "sheets": [{
      "name": str, "hidden": bool, "tab_color": "#RRGGBB"|null, "gridlines": bool,
      "freeze": {"rows": n, "cols": n},
      "cols": {"0": {"w": px, "hidden": bool}}, "rows": {"0": {"h": px, "hidden": bool}},
      "merges": ["A1:B2"],
      "cells": {"A1": {"v": 값, "t": "n|s|b|e", "f": "=수식", "ref": "C1:C3"(배열 수식 범위), "s": {Univer IStyleData 부분집합}}}
  }]}

1차 기록기 = openpyxl 왕복: 원본을 열어 각 시트의 셀·병합·치수만 바꾼다(유효성·조건부 서식·정의 이름·인쇄 설정은
시트 객체에 남아 보존). openpyxl 이 잃는 부품(차트·그림·피벗·주석)이 있는 파일은 capabilities 가 처음부터
사무 엔진으로 보낸다(`grid_blockers`).
"""
import io
import json
import re
import zipfile
from datetime import date, datetime, time as dtime

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.formula import ArrayFormula

MAX_CELLS = 2_000_000
MAX_TEXT = 32767
GRID_BLOCKING_PARTS = ('xl/drawings/', 'xl/charts/', 'xl/pivottables/', 'xl/pivotcache/', 'xl/comments', 'xl/media/', 'xl/embeddings/', 'xl/slicers/', 'xl/ctrlprops/', 'xl/threadedcomments/')
# 엑셀이 파일 안에서 접두사를 요구하는 새 함수들 — 격자는 접두사 없이 계산하고, 파일에는 붙여 쓴다.
XLFN = ('LET', 'LAMBDA', 'XLOOKUP', 'XMATCH', 'FILTER', 'SORT', 'SORTBY', 'UNIQUE', 'SEQUENCE', 'RANDARRAY', 'IFS', 'SWITCH',
        'MAXIFS', 'MINIFS', 'TEXTJOIN', 'CONCAT', 'IFNA', 'STDEV.S', 'STDEV.P', 'VAR.S', 'VAR.P', 'RANK.EQ', 'RANK.AVG',
        'PERCENTILE.INC', 'PERCENTILE.EXC', 'QUARTILE.INC', 'QUARTILE.EXC', 'MODE.SNGL', 'AGGREGATE', 'DAYS', 'ISOWEEKNUM',
        'TEXTSPLIT', 'TEXTBEFORE', 'TEXTAFTER', 'VSTACK', 'HSTACK', 'TOCOL', 'TOROW', 'TAKE', 'DROP', 'CHOOSECOLS', 'CHOOSEROWS',
        'BYROW', 'BYCOL', 'MAP', 'REDUCE', 'SCAN', 'MAKEARRAY', 'ARRAYTOTEXT', 'VALUETOTEXT', 'NUMBERVALUE', 'FORMULATEXT')
_XLFN_RE = re.compile(r'(?<![A-Za-z0-9_.])(' + '|'.join(re.escape(n) for n in sorted(XLFN, key=len, reverse=True)) + r')\s*\(')
_PREFIX_RE = re.compile(r'_xl(?:fn|ws|pm|op)\.')
BORDER_TO_GRID = {'thin': 1, 'hair': 2, 'dotted': 3, 'dashed': 4, 'dashDot': 5, 'dashDotDot': 6, 'double': 7, 'medium': 8,
                  'mediumDashed': 9, 'mediumDashDot': 10, 'mediumDashDotDot': 11, 'slantDashDot': 12, 'thick': 13}
BORDER_TO_XLSX = {v: k for k, v in BORDER_TO_GRID.items()}
H_TO_GRID = {'left': 1, 'center': 2, 'right': 3, 'justify': 4, 'distributed': 6, 'centerContinuous': 2}
H_TO_XLSX = {1: 'left', 2: 'center', 3: 'right', 4: 'justify', 5: 'justify', 6: 'distributed'}
V_TO_GRID = {'top': 1, 'center': 2, 'bottom': 3}
V_TO_XLSX = {1: 'top', 2: 'center', 3: 'bottom'}
CHAR_PX = 7  # Calibri 11 기준 한 글자 폭(px). 엑셀 열 너비 단위 ↔ px 근사.


def grid_blockers(part_names):
    """격자가 그리지 못하는(저장하면 잃는) 부품 — 있으면 사무 엔진으로."""
    found = set()
    for name in part_names:
        low = name.lower()
        for part in GRID_BLOCKING_PARTS:
            if low.startswith(part):
                found.add(part.rstrip('/').split('/')[-1])
    return sorted(found)


def _color(c):
    """openpyxl Color → '#RRGGBB' (테마·인덱스 색은 모른다 → None)."""
    if c is None or c.type != 'rgb' or not isinstance(c.rgb, str):
        return None
    rgb = c.rgb
    if len(rgb) == 8:
        rgb = rgb[2:]
    return '#' + rgb.lower() if len(rgb) == 6 else None


def _argb(hexcolor):
    h = (hexcolor or '').lstrip('#')
    return 'FF' + h.upper() if len(h) == 6 else None


def _style_of(cell):
    s = {}
    f = cell.font
    if f is not None:
        if f.b: s['bl'] = 1
        if f.i: s['it'] = 1
        if f.u: s['ul'] = {'s': 1}
        if f.strike: s['st'] = {'s': 1}
        if f.sz and float(f.sz) != 11: s['fs'] = float(f.sz)
        if f.name and f.name != 'Calibri': s['ff'] = f.name
        col = _color(f.color)
        if col and col != '#000000': s['cl'] = {'rgb': col}
    if cell.fill is not None and cell.fill.fill_type == 'solid':
        col = _color(cell.fill.fgColor)
        if col: s['bg'] = {'rgb': col}
    bd = {}
    for key, side in (('t', cell.border.top), ('b', cell.border.bottom), ('l', cell.border.left), ('r', cell.border.right)):
        if side is not None and side.style in BORDER_TO_GRID:
            bd[key] = {'s': BORDER_TO_GRID[side.style], 'cl': {'rgb': _color(side.color) or '#000000'}}
    if bd: s['bd'] = bd
    a = cell.alignment
    if a is not None:
        if a.horizontal in H_TO_GRID: s['ht'] = H_TO_GRID[a.horizontal]
        if a.vertical in V_TO_GRID: s['vt'] = V_TO_GRID[a.vertical]
        if a.wrap_text: s['tb'] = 3
    if cell.number_format and cell.number_format != 'General':
        s['n'] = {'pattern': cell.number_format}
    return s or None


def _value_of(cell):
    """(값, 타입, 수식, 배열범위). 날짜는 시리얼로. 수식 셀의 값은 모른다(격자가 계산)."""
    v = cell.value
    if isinstance(v, ArrayFormula):
        return None, None, _PREFIX_RE.sub('', v.text or ''), v.ref
    if cell.data_type == 'f' or (isinstance(v, str) and v.startswith('=')):
        return None, None, _PREFIX_RE.sub('', v), None
    if v is None:
        return None, None, None, None
    if isinstance(v, bool):
        return v, 'b', None, None
    if isinstance(v, (int, float)):
        return v, 'n', None, None
    if isinstance(v, (datetime, date, dtime)):
        return to_excel(v), 'n', None, None
    if cell.data_type == 'e':
        return str(v), 'e', None, None
    return str(v), 's', None, None


def to_grid(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data))
    sheets = []
    for ws in wb.worksheets:
        cells = {}
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                v, t, f, ref = _value_of(cell)
                s = _style_of(cell)
                if v is None and f is None and s is None:
                    continue
                entry = {}
                if f is not None:
                    entry['f'] = f
                    if ref: entry['ref'] = ref
                elif v is not None:
                    entry['v'] = v; entry['t'] = t
                if s: entry['s'] = s
                cells[cell.coordinate] = entry
                count += 1
                if count > MAX_CELLS:
                    raise ValueError('격자 상한(셀 200만)을 넘는 통합문서입니다. 사무 편집기로 여세요')
        cols = {}
        for letter, dim in ws.column_dimensions.items():
            if dim.width is None and not dim.hidden:
                continue
            start, end = column_index_from_string(letter), dim.max or column_index_from_string(letter)
            for idx in range(start, max(start, min(end, 16384)) + 1):
                entry = {}
                if dim.width is not None: entry['w'] = round(float(dim.width) * CHAR_PX + 5)
                if dim.hidden: entry['hidden'] = True
                cols[str(idx - 1)] = entry
        rows = {}
        for num, dim in ws.row_dimensions.items():
            if dim.height is None and not dim.hidden:
                continue
            entry = {}
            if dim.height is not None: entry['h'] = round(float(dim.height) * 96 / 72)
            if dim.hidden: entry['hidden'] = True
            rows[str(int(num) - 1)] = entry
        freeze = {'rows': 0, 'cols': 0}
        if ws.freeze_panes:
            col, row = coordinate_from_string(ws.freeze_panes)
            freeze = {'rows': row - 1, 'cols': column_index_from_string(col) - 1}
        tab = ws.sheet_properties.tabColor
        sheets.append({
            'name': ws.title, 'hidden': ws.sheet_state != 'visible', 'tab_color': _color(tab) if tab is not None else None,
            'gridlines': ws.sheet_view.showGridLines is not False, 'freeze': freeze, 'cols': cols, 'rows': rows,
            'merges': [str(r) for r in ws.merged_cells.ranges], 'cells': cells,
        })
    return {'version': 1, 'date_system': '1904' if wb.epoch.year == 1904 else '1900', 'sheets': sheets}


def _check(grid):
    if not isinstance(grid, dict) or grid.get('version') != 1 or not isinstance(grid.get('sheets'), list) or not grid['sheets']:
        raise ValueError('grid.v1 통화가 아닙니다')
    names = set()
    total = 0
    for sheet in grid['sheets']:
        name = sheet.get('name')
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name):
            raise ValueError('시트 이름을 확인하세요(1~31자, \\ / * ? : [ ] 제외)')
        if name in names:
            raise ValueError('시트 이름이 겹칩니다: ' + name)
        names.add(name)
        cells = sheet.get('cells') or {}
        if not isinstance(cells, dict):
            raise ValueError('cells 는 주소→셀 매핑입니다')
        total += len(cells)
        if total > MAX_CELLS:
            raise ValueError('격자 상한(셀 200만)을 넘었습니다')
        for addr, c in cells.items():
            if not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', addr) or not isinstance(c, dict):
                raise ValueError('셀 주소 또는 셀 통화를 확인하세요: ' + str(addr))
            v = c.get('v')
            if v is not None and type(v) not in (str, int, float, bool):
                raise ValueError('셀 값은 문자·숫자·불리언입니다: ' + addr)
            if isinstance(v, str) and len(v) > MAX_TEXT:
                raise ValueError('셀 문자열 상한은 32767자입니다: ' + addr)
            if isinstance(v, float) and v != v:
                raise ValueError('NaN 은 셀 값이 될 수 없습니다: ' + addr)
            f = c.get('f')
            if f is not None and (not isinstance(f, str) or not f.startswith('=') or len(f) > 8192):
                raise ValueError('수식은 = 로 시작하는 8192자 이내 문자열입니다: ' + addr)
            if f and re.search(r"\[[^\]]+\][A-Za-z0-9 _.']*!|https?:|WEBSERVICE|IMAGE\s*\(|RTD\s*\(|\|", f, re.I):
                raise ValueError('외부 접근 수식은 저장하지 않습니다: ' + addr)


def _xlsx_formula(f):
    return _XLFN_RE.sub(lambda m: '_xlfn.' + m.group(1) + '(', f)


def _apply_style(cell, s):
    if not isinstance(s, dict) or not s:
        return
    cl = (s.get('cl') or {}).get('rgb') if isinstance(s.get('cl'), dict) else None
    font_kw = {}
    if s.get('bl'): font_kw['b'] = True
    if s.get('it'): font_kw['i'] = True
    if isinstance(s.get('ul'), dict) and s['ul'].get('s'): font_kw['u'] = 'single'
    if isinstance(s.get('st'), dict) and s['st'].get('s'): font_kw['strike'] = True
    if isinstance(s.get('fs'), (int, float)) and 1 <= s['fs'] <= 409: font_kw['sz'] = float(s['fs'])
    if isinstance(s.get('ff'), str) and s['ff']: font_kw['name'] = s['ff'][:64]
    if _argb(cl): font_kw['color'] = _argb(cl)
    if font_kw:
        cell.font = Font(**font_kw)
    bg = (s.get('bg') or {}).get('rgb') if isinstance(s.get('bg'), dict) else None
    if _argb(bg):
        cell.fill = PatternFill(fill_type='solid', fgColor=_argb(bg), bgColor=_argb(bg))
    bd = s.get('bd')
    if isinstance(bd, dict) and bd:
        sides = {}
        for key, attr in (('t', 'top'), ('b', 'bottom'), ('l', 'left'), ('r', 'right')):
            side = bd.get(key)
            if isinstance(side, dict) and side.get('s') in BORDER_TO_XLSX:
                sides[attr] = Side(style=BORDER_TO_XLSX[side['s']], color=_argb((side.get('cl') or {}).get('rgb')) or 'FF000000')
        if sides:
            cell.border = Border(**sides)
    align_kw = {}
    if s.get('ht') in H_TO_XLSX: align_kw['horizontal'] = H_TO_XLSX[s['ht']]
    if s.get('vt') in V_TO_XLSX: align_kw['vertical'] = V_TO_XLSX[s['vt']]
    if s.get('tb') == 3: align_kw['wrap_text'] = True
    if align_kw:
        cell.alignment = Alignment(**align_kw)
    n = s.get('n')
    if isinstance(n, dict) and isinstance(n.get('pattern'), str) and n['pattern'] and len(n['pattern']) <= 255:
        cell.number_format = n['pattern']


def from_grid(base: bytes | None, grid: dict) -> bytes:
    """격자 통화 → XLSX. base(원본/초안)가 있으면 그 통합문서 위에 시트 내용만 다시 쓴다."""
    _check(grid)
    if base:
        wb = load_workbook(io.BytesIO(base))
    else:
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
    wanted = [s['name'] for s in grid['sheets']]
    for ws in list(wb.worksheets):
        if ws.title not in wanted:
            wb.remove(ws)
    for index, sheet in enumerate(grid['sheets']):
        ws = wb[sheet['name']] if sheet['name'] in wb.sheetnames else wb.create_sheet(sheet['name'])
        wb.move_sheet(ws, offset=index - wb.index(ws))
        # 셀·병합·치수는 통째로 격자 것으로. 시트 수준 객체(유효성·조건부 서식·인쇄 설정·표)는 그대로 둔다.
        ws._cells.clear()
        ws.merged_cells.ranges.clear() if hasattr(ws.merged_cells, 'ranges') else None
        ws.column_dimensions.clear()
        ws.row_dimensions.clear()
        for addr, c in (sheet.get('cells') or {}).items():
            cell = ws[addr]
            f = c.get('f')
            if f:
                text = _xlsx_formula(f)
                ref = c.get('ref')
                cell.value = ArrayFormula(ref, text) if isinstance(ref, str) and re.fullmatch(r'[A-Z]{1,3}[0-9]+(:[A-Z]{1,3}[0-9]+)?', ref) else text
            else:
                v = c.get('v')
                if v is not None:
                    if c.get('t') == 's' and not isinstance(v, str):
                        v = str(v)
                    cell.value = v
                    if c.get('t') == 's' and isinstance(v, str):
                        cell.data_type = 's'  # 숫자처럼 보이는 문자(식별자 00123)를 문자로 고정
            _apply_style(cell, c.get('s'))
        for rng in sheet.get('merges') or []:
            if isinstance(rng, str) and re.fullmatch(r'[A-Z]{1,3}[0-9]+:[A-Z]{1,3}[0-9]+', rng):
                ws.merge_cells(rng)
        for key, entry in (sheet.get('cols') or {}).items():
            if not str(key).isdigit() or not isinstance(entry, dict): continue
            dim = ws.column_dimensions[get_column_letter(int(key) + 1)]
            if isinstance(entry.get('w'), (int, float)) and entry['w'] > 0: dim.width = max(0.5, (float(entry['w']) - 5) / CHAR_PX)
            if entry.get('hidden'): dim.hidden = True
        for key, entry in (sheet.get('rows') or {}).items():
            if not str(key).isdigit() or not isinstance(entry, dict): continue
            dim = ws.row_dimensions[int(key) + 1]
            if isinstance(entry.get('h'), (int, float)) and entry['h'] > 0: dim.height = float(entry['h']) * 72 / 96
            if entry.get('hidden'): dim.hidden = True
        fr = sheet.get('freeze') or {}
        rows, cols = int(fr.get('rows') or 0), int(fr.get('cols') or 0)
        ws.freeze_panes = f'{get_column_letter(cols + 1)}{rows + 1}' if (rows or cols) else None
        ws.sheet_state = 'hidden' if sheet.get('hidden') else 'visible'
        ws.sheet_view.showGridLines = bool(sheet.get('gridlines', True))
        tab = sheet.get('tab_color')
        ws.sheet_properties.tabColor = _argb(tab) if isinstance(tab, str) and _argb(tab) else None
    if not any(ws.sheet_state == 'visible' for ws in wb.worksheets):
        wb.worksheets[0].sheet_state = 'visible'
    wb.active = next((i for i, ws in enumerate(wb.worksheets) if ws.sheet_state == 'visible'), 0)
    wb.calculation.fullCalcOnLoad = True
    buffer = io.BytesIO()
    wb.save(buffer)
    return _inject_cached_values(buffer.getvalue(), grid)


def _inject_cached_values(data: bytes, grid: dict) -> bytes:
    """수식 셀에 격자가 계산한 값을 <v> 로 심고, 바이트를 결정적으로 만든다(같은 격자 = 같은 파일 해시).
    openpyxl 은 저장 시각·ZIP 항목 시각을 매번 새로 쓰므로 그대로 두면 바뀐 것이 없어도 초안 해시가 달라진다."""
    from defusedxml.ElementTree import fromstring
    from xml.etree import ElementTree as ET
    import spreadsheet_files as files
    computed = {s['name']: {a: c.get('v') for a, c in (s.get('cells') or {}).items() if c.get('f') and c.get('v') is not None}
                for s in grid['sheets']}
    NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', NS)
    ET.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')
    meta = files.inspect(data)
    patched = {}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for sheet in meta['sheets']:
            values = computed.get(sheet['name'])
            if not values:
                continue
            root = fromstring(z.read(sheet['path']))
            changed = False
            for c in root.iter('{%s}c' % NS):
                v = values.get(c.get('r'))
                if v is None or c.find('{%s}f' % NS) is None:
                    continue
                for old in c.findall('{%s}v' % NS):
                    c.remove(old)
                node = ET.SubElement(c, '{%s}v' % NS)
                if isinstance(v, bool):
                    c.set('t', 'b'); node.text = '1' if v else '0'
                elif isinstance(v, (int, float)):
                    c.attrib.pop('t', None); node.text = repr(float(v)) if isinstance(v, float) else str(v)
                elif isinstance(v, str) and re.fullmatch(r'#[A-Z0-9/!?]+', v):
                    c.set('t', 'e'); node.text = v
                else:
                    c.set('t', 'str'); node.text = str(v)
                changed = True
            if changed:
                patched[sheet['path']] = ET.tostring(root, xml_declaration=True, encoding='UTF-8')
        if 'docProps/core.xml' in z.namelist():
            core = z.read('docProps/core.xml').decode('utf-8', 'replace')
            patched['docProps/core.xml'] = re.sub(r'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', r'\g<1>2026-01-01T00:00:00Z\g<2>', core).encode('utf-8')
        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
            for item in z.infolist():
                info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = item.external_attr
                target.writestr(info, patched.get(item.filename, z.read(item.filename)))
    return out.getvalue()


def engine_state(grid: dict) -> str:
    """`[self:workspace]` 스냅샷 계약의 engine_state — 시트별 사용 범위의 값·수식·서식(플러그인 state() 와 같은 꼴)."""
    out = []
    for sheet in grid['sheets']:
        cells = sheet.get('cells') or {}
        if not cells:
            out.append({'name': sheet['name'], 'address': 'A1', 'values': [[None]], 'formulas': [['']], 'format': [['General']]})
            continue
        max_c = max_r = 0
        for addr in cells:
            col, row = coordinate_from_string(addr)
            max_c = max(max_c, column_index_from_string(col)); max_r = max(max_r, row)
        values = [[None] * max_c for _ in range(max_r)]
        formulas = [[''] * max_c for _ in range(max_r)]
        fmt = [['General'] * max_c for _ in range(max_r)]
        for addr, c in cells.items():
            col, row = coordinate_from_string(addr)
            i, j = row - 1, column_index_from_string(col) - 1
            values[i][j] = c.get('v')
            formulas[i][j] = c.get('f') or ''
            n = (c.get('s') or {}).get('n') if isinstance(c.get('s'), dict) else None
            if isinstance(n, dict) and n.get('pattern'): fmt[i][j] = n['pattern']
        out.append({'name': sheet['name'], 'address': f'A1:{get_column_letter(max_c)}{max_r}', 'values': values, 'formulas': formulas, 'format': fmt})
    return json.dumps(out, ensure_ascii=False, separators=(',', ':'))


def range_cells(sheet: dict, address: str):
    """격자 시트에서 A1 범위의 셀을 2차원으로(없는 셀은 None)."""
    c1, r1, c2, r2 = range_boundaries(address)
    cells = sheet.get('cells') or {}
    return [[cells.get(f'{get_column_letter(c)}{r}') for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
