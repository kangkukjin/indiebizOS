"""Printable forms generated only for new workbooks, never existing user files."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

PRINT_FORMS = {'print_quotation': '견적서', 'invoice': '청구서'}


def populate(book, template):
    """Build an editable A4 form with explicit input cells and protected totals."""
    title = PRINT_FORMS[template]
    sheet = book.active
    sheet.title = title
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4, header=.15, footer=.15)
    sheet.print_options.horizontalCentered = True
    sheet.print_area = 'A1:F31'
    sheet.print_title_rows = '1:9'
    # ONLYOFFICE 9.3.1 preserves &N in XLSX but emits the current page as
    # the PDF total (1/1, 2/2, 3/3). Use truthful page numbers until fixed.
    sheet.oddFooter.center.text = '&P 쪽'
    sheet.oddFooter.center.size = 9
    sheet.freeze_panes = 'C10'
    for col, width in {'A': 7, 'B': 32, 'C': 10, 'D': 10, 'E': 16, 'F': 19}.items():
        sheet.column_dimensions[col].width = width
    border = Border(bottom=Side(style='hair', color='CBD5E1'))
    for row in sheet.iter_rows(min_row=1, max_row=31, min_col=1, max_col=6):
        for cell in row:
            cell.font = Font(name='Noto Sans CJK KR', size=10, color='20312B')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = border
        sheet.row_dimensions[row[0].row].height = 23
    sheet.merge_cells('A1:F2')
    sheet['A1'] = title
    sheet['A1'].font = Font(name='Noto Sans CJK KR', size=24, bold=True, color='245B48')
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    for address, value in {'A4': '받는 분', 'D4': '공급자', 'A5': '발행일',
                           'D5': '문서 번호', 'A6': '연락처', 'D6': '지급 기한' if template == 'invoice' else '유효 기한'}.items():
        sheet[address] = value
    for merged in ('B4:C4', 'E4:F4', 'B5:C5', 'E5:F5', 'B6:C6', 'E6:F6', 'A8:F8'):
        sheet.merge_cells(merged)
    sheet['A8'] = '아래 내역과 같이 청구합니다.' if template == 'invoice' else '아래 내역과 같이 견적합니다.'
    inputs = ['B4', 'E4', 'B5', 'E5', 'B6', 'E6']
    for address in ('B5', 'E6'):
        sheet[address].number_format = 'yyyy-mm-dd'
    for col, label in enumerate(('번호', '품목 / 내용', '단위', '수량', '단가', '금액'), 1):
        cell = sheet.cell(9, col, label)
        cell.font = Font(name='Noto Sans CJK KR', size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245B48')
    for row in range(10, 25):
        sheet.cell(row, 1, '=ROW()-9')
        sheet.cell(row, 6, f'=IF(OR(D{row}="",E{row}=""),"",ROUND(D{row}*E{row},0))')
        for col in ('B', 'C', 'D', 'E'):
            inputs.append(f'{col}{row}')
        for col in ('E', 'F'):
            sheet[f'{col}{row}'].number_format = '#,##0;[Red]-#,##0'
    for row, label, formula in ((25, '공급가액', '=SUM(F10:F24)'),
                                (26, '세액', '=ROUND(F25*D26,0)'),
                                (27, '합계', '=SUM(F25:F26)')):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.cell(row, 1, label)
        sheet.cell(row, 6, formula).number_format = '#,##0;[Red]-#,##0'
        sheet.cell(row, 6).font = Font(name='Noto Sans CJK KR', size=12, bold=True)
    sheet['D26'] = 0
    sheet['D26'].number_format = '0%'
    sheet['E26'] = '세율'
    inputs.append('D26')
    sheet.merge_cells('A29:F31')
    sheet['A29'] = '비고: 세율은 거래 조건에 맞게 입력하세요. 기본값은 0%입니다.'
    inputs.append('A29')
    for address in inputs:
        sheet[address].protection = Protection(locked=False)
        sheet[address].fill = PatternFill('solid', fgColor='EFF7F2')
    units = DataValidation(type='list', formula1='"개,건,시간,일,월,식"', allow_blank=True)
    units.errorTitle = '단위 확인'
    units.error = '목록에서 단위를 선택하세요.'
    units.showErrorMessage = True
    units.showDropDown = False
    sheet.add_data_validation(units)
    units.add('C10:C24')
    # No password: users can explicitly unprotect to add rows, logos or alter layout.
    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False
    help_sheet = book.create_sheet('사용 안내')
    help_sheet['A1'] = title + ' 사용 안내'
    help_sheet['A2'] = '연녹색 셀에 거래 정보를 입력합니다. 합계 셀은 수식으로 계산됩니다.'
    help_sheet['A3'] = '세율은 기본 0%입니다. 거래 조건에 맞는 세율을 직접 입력하세요.'
    help_sheet['A4'] = '행 추가·로고 삽입: 보호 탭에서 암호 없이 시트 보호를 해제한 뒤 편집합니다.'
    help_sheet['A5'] = '품목 사이의 셀을 선택하고 Ctrl+Shift+더하기 → 전체 행으로 삽입합니다. 기존 품목 영역 안에 추가하세요.'
    help_sheet['A6'] = '바로 위 품목의 A~F 셀을 새 행에 복사한 뒤 품목·단위·수량·단가만 바꿉니다. 번호·금액 수식과 입력 서식을 이어받습니다.'
    help_sheet['A7'] = '합계·단위 목록·인쇄 영역을 확인하고 필요하면 시트 보호를 다시 켜세요. 맨 아래에 덧붙일 때는 합계 범위도 확인하세요.'
    help_sheet['A8'] = 'PDF 설정의 범위에서 활성 시트는 양식만, 모든 시트는 이 안내까지 출력합니다.'
    help_sheet.column_dimensions['A'].width = 90
    for row in range(1, 9):
        cell = help_sheet.cell(row, 1)
        cell.font = Font(name='Noto Sans CJK KR', size=11, bold=row == 1)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        help_sheet.row_dimensions[row].height = 48 if row > 1 else 32
    help_sheet.print_area = 'A1:A8'
    help_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    help_sheet.page_setup.fitToWidth = 1
    help_sheet.page_setup.fitToHeight = 1
    return sheet
