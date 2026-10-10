"""Pages conversion identity, crash recovery and input boundaries (no live engine)."""
import io
import threading
import zipfile
from pathlib import Path

import boot_paths  # noqa: F401
import pytest
from docx import Document
import document_formats
import document_imports
import document_office
from document_pages import ImportFailure, snapshot
from document_workspace import DocumentWorkspace
from office_store import digest
from workspace_sessions import Workspace, kind_of_path


def pages(path, text='original'):
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('Index/Document.iwa', text.encode())
    return path


def docx_bytes(text='Converted 한글'):
    d = Document(); d.add_paragraph(text)
    b = io.BytesIO(); d.save(b)
    return b.getvalue()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    app = DocumentWorkspace(tmp_path / 'state')
    source = pages(tmp_path / 'source.pages')
    calls = []
    def convert(*args):
        calls.append(args)
        return docx_bytes()
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    monkeypatch.setattr(document_office, 'available', lambda: True)
    return app, source, calls


def test_auto_open_reuses_saved_copy_and_unsaved_session(setup):
    app, source, calls = setup
    original = source.read_bytes()
    first = app.open(source)
    row = first['document']
    assert row['source_format'] == 'docx'
    assert first['conversion']['source_resource'] != row['id']
    assert not first['conversion']['reused']
    assert 'Converted 한글' in Document(row['source_uri']).paragraphs[0].text
    session = app.acquire(row['id'], 'window')['session']
    session.update(blob=app.store.blob(docx_bytes('unsaved')), state='draft', session_revision=1)
    app.store.put('session', session)
    again = DocumentWorkspace(app.store.root).open(source)
    assert again['document']['id'] == row['id']
    assert again['session']['blob'] == session['blob']
    assert again['conversion']['reused'] and len(calls) == 1
    app.save(row['id'], session['id'], 'window', session['engine_epoch'], 1, 'save', row['revision_id'])
    assert Document(row['source_uri']).paragraphs[0].text == 'unsaved'
    assert source.read_bytes() == original
    assert app.open(source)['document']['id'] == row['id']


def test_source_change_keeps_working_copy_until_refresh(setup):
    app, source, calls = setup
    first = app.open(source)
    pages(source, 'new')
    old = app.open(source)
    assert old['document']['id'] == first['document']['id']
    assert old['conversion']['source_changed']
    new = app.open(source, conversion='refresh', request_id='refresh-1')
    assert new['document']['id'] != old['document']['id']
    assert not new['conversion']['source_changed']
    assert app.open(source, conversion='refresh', request_id='refresh-1')['document']['id'] == new['document']['id']
    assert len(calls) == 2 and app.open(source)['document']['id'] == new['document']['id']
    assert Path(old['document']['source_uri']).exists()


def test_working_copy_external_change_is_not_overwritten(setup):
    app, source, calls = setup
    d = app.open(source)['document']
    Path(d['source_uri']).write_bytes(docx_bytes('external'))
    with pytest.raises(ImportFailure, match='working_copy_conflict'):
        app.open(source)
    assert len(calls) == 1
    assert Document(d['source_uri']).paragraphs[0].text == 'external'


def test_package_snapshot_stable_and_document_routing(setup, tmp_path):
    app, _, calls = setup
    folder = tmp_path / 'package.pages'; (folder / 'Index').mkdir(parents=True)
    (folder / 'Index/Document.iwa').write_bytes(b'package')
    assert snapshot(folder) == snapshot(folder)
    assert kind_of_path(folder) == 'document'
    ws = Workspace(app.store.root)
    d = ws.open(folder)
    assert d['kind'] == 'document' and d['conversion']['source_format'] == 'pages'
    assert ws.read(d['resource'])['projection']
    assert not (folder / '.git').exists()
    assert ws.open(folder)['resource'] == d['resource'] and len(calls) == 1


def test_source_mode_never_acquires_native_editor(setup):
    app, source, calls = setup
    result = app.open(source, conversion='source')
    assert result['document']['source_format'] == 'pages'
    assert not result['capabilities']['save']
    assert result['capabilities']['conversion_available']
    assert not calls
    with pytest.raises(ValueError):
        app.acquire(result['document']['id'], 'w')


def test_bad_output_is_not_registered_and_retry_succeeds(setup, monkeypatch):
    app, source, calls = setup
    convert = document_formats.convert_bytes
    monkeypatch.setattr(document_formats, 'convert_bytes', lambda *a: b'not docx')
    with pytest.raises(ImportFailure, match='invalid_conversion_output'):
        app.open(source)
    assert len(app.list()) == 1
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    assert app.open(source)['document']['source_format'] == 'docx'


def test_concurrent_open_is_fenced(setup, monkeypatch):
    app, source, _ = setup
    entered, finish = threading.Event(), threading.Event()
    results = []
    def convert(*args):
        entered.set(); assert finish.wait(5)
        return docx_bytes()
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    thread = threading.Thread(target=lambda: results.append(app.open(source)))
    thread.start()
    try:
        assert entered.wait(5)
        with pytest.raises(ImportFailure, match='conversion_in_progress'):
            app.open(source)
    finally:
        finish.set(); thread.join(5)
    assert len(results) == 1
    assert app.open(source)['document']['id'] == results[0]['document']['id']


def test_crash_after_output_preparation_recovers_without_conversion(setup, monkeypatch):
    app, source, calls = setup
    finish = document_imports._finish
    monkeypatch.setattr(document_imports, '_finish', lambda *a: (_ for _ in ()).throw(OSError('crash')))
    with pytest.raises(OSError, match='crash'):
        app.open(source)
    monkeypatch.setattr(document_imports, '_finish', finish)
    result = app.open(source)
    assert result['document']['source_format'] == 'docx' and len(calls) == 1
    assert len(list((app.store.root / 'files').rglob('*.docx'))) == 1


def test_changed_source_during_conversion_is_reported(setup, monkeypatch):
    app, source, _ = setup
    def convert(*a):
        pages(source, 'modified while converting')
        return docx_bytes('old snapshot')
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    result = app.open(source)
    assert result['conversion']['source_changed']
    p = result['document']['provenance']
    assert p['snapshot_sha256'] != digest(source.read_bytes())


@pytest.mark.parametrize('name', ['../escape', '/absolute', 'a\\b', 'C:/escape'])
def test_unsafe_zip_paths_rejected(tmp_path, name):
    source = tmp_path / 'bad.pages'
    with zipfile.ZipFile(source, 'w') as z:
        z.writestr('Index/Document.iwa', b'data'); z.writestr(name, b'x')
    with pytest.raises(ImportFailure, match='invalid_pages'):
        snapshot(source)


def test_fake_file_and_package_link_rejected(tmp_path):
    source = tmp_path / 'fake.pages'; source.write_bytes(b'fake')
    with pytest.raises(ImportFailure, match='invalid_pages'):
        snapshot(source)
    folder = tmp_path / 'link.pages'; folder.mkdir()
    (folder / 'data').symlink_to(source)
    with pytest.raises(ImportFailure, match='invalid_pages'):
        snapshot(folder)


def test_expired_attempt_does_not_publish_late_response(setup, monkeypatch):
    app, source, _ = setup
    def supersede(*args):
        row = next(d for d in app.list() if d['source_format'] == 'pages')
        row['conversion_attempt'] = 'new-attempt'
        app.store.put('document', row)
        return docx_bytes()
    monkeypatch.setattr(document_formats, 'convert_bytes', supersede)
    with pytest.raises(ImportFailure, match='conversion_superseded'):
        app.open(source)
    assert len(app.list()) == 1


def test_manual_candidate_selection(setup):
    app, source, _ = setup
    first = app.open(source)
    second = app.open(source, conversion='refresh', request_id='second')
    source_id = first['conversion']['source_resource']
    row = app._doc(source_id); row.pop('working_resource'); app.store.put('document', row)
    with pytest.raises(ImportFailure, match='working_copy_choice'):
        app.open(source)
    chosen = app.open(source, working_resource=first['document']['id'])
    assert chosen['document']['id'] == first['document']['id']
    assert app.open(source)['document']['id'] != second['document']['id']


def test_workspace_member_rejected_before_open(setup, monkeypatch):
    import importlib.util
    import principal
    root = Path(__file__).resolve().parents[1]
    path = root / 'data/packages/installed/tools/system_essentials/essentials_workspace.py'
    spec = importlib.util.spec_from_file_location('workspace_member_gate', path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    monkeypatch.setattr(principal, 'is_owner', lambda: False)
    with pytest.raises(PermissionError):
        module.op_open({'path': str(setup[1]), 'conversion': 'auto'})
    assert not setup[2]


def test_http_open_returns_actual_docx_and_source_query_stays_source(setup, monkeypatch):
    import api_documents
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    app, path, _ = setup
    monkeypatch.setattr(api_documents, 'service', lambda: app)
    server = FastAPI(); server.include_router(api_documents.router)
    with TestClient(server) as client:
        result = client.post('/documents/open', json={'path': str(path)}).json()
        assert result['document']['source_format'] == 'docx'
        source = client.get('/documents/' + result['conversion']['source_resource']).json()
        assert source['document']['source_format'] == 'pages'
        assert source['conversion_candidates'][0]['resource'] == result['document']['id']


def test_refresh_requires_id_and_foreign_candidate_is_rejected(setup, tmp_path):
    app, source, _ = setup
    with pytest.raises(ValueError, match='request_id'):
        app.open(source, conversion='refresh')
    other = tmp_path / 'unrelated.docx'; other.write_bytes(docx_bytes())
    candidate = app.open(other)['document']['id']
    with pytest.raises(ValueError, match='Pages에서 변환한'):
        app.open(source, working_resource=candidate)


def test_source_deleted_during_conversion_reports_unknown(setup, monkeypatch):
    app, source, _ = setup
    def convert(*args):
        source.unlink()
        return docx_bytes()
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    assert app.open(source)['conversion']['source_changed'] is None


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
