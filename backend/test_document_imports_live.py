"""Real Pages input → ONLYOFFICE → DOCX editor → save/reopen and PDF output."""
import os
import shutil
import socket
import threading
import time
from pathlib import Path

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'backend/fixtures/pages'


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='로컬 문서 엔진 실인수')
def test_pages_live_roundtrip(tmp_path, monkeypatch):
    import api_documents
    import document_office
    import document_formats
    import uvicorn
    import pymupdf
    import zipfile
    from docx import Document
    from document_workspace import DocumentWorkspace
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from playwright.sync_api import sync_playwright, expect
    from test_document_app_support import mount_document_app, open_document

    workspace = DocumentWorkspace(tmp_path / 'workspace')
    app = FastAPI(); app.include_router(api_documents.router)
    monkeypatch.setattr(api_documents, 'service', lambda: workspace)
    mount_document_app(app, workspace)
    @app.get('/launcher/auth/session')
    def auth():
        return {'authenticated': True, 'external': False}
    @app.get('/')
    def index():
        return HTMLResponse((ROOT / 'frontend/dist/index.html').read_text().replace(
            '<html', '<html data-indiebiz-surface="remote"', 1))
    app.mount('/assets', StaticFiles(directory=ROOT / 'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('0.0.0.0', 0)); port = sock.getsockname()[1]
    cfg = document_office.settings()
    assert cfg
    cfg = {**cfg, 'callback_origin': f'http://192.168.5.2:{port}'}
    monkeypatch.setattr(document_formats, 'settings', lambda: cfg)
    monkeypatch.setattr(document_office, 'settings', lambda: cfg)
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    until = time.monotonic() + 10
    while not server.started and time.monotonic() < until:
        time.sleep(.05)
    try:
        source = tmp_path / 'simple.pages'; shutil.copyfile(FIXTURES / source.name, source)
        before = source.read_bytes()
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1600, 'height': 1100})
            errors = []; page.on('pageerror', lambda e: errors.append(str(e)))
            try:
                open_document(page, port, source)
                status = page.locator('.office-document').get_by_role('status')
                expect(status).to_have_text('편집 준비 완료', timeout=180000)
                expect(page.get_by_label('Pages 변환 정보')).to_contain_text('Pages에서 변환한 작업 사본')
                result = workspace.open(source); d = result['document']
                assert result['conversion']['reused']
                text = '\n'.join(p.text for p in Document(d['source_uri']).paragraphs)
                assert '원본에서 보존할 한글 문장입니다.' in text and 'English paragraph' in text
                page.locator('.office-document').get_by_role('button', name='⚙ 도구', exact=True).click()
                frame = page.frame_locator('iframe[name="frameEditor"]')
                area = frame.locator('#area_id')
                tip = frame.get_by_text('확인', exact=True)
                if tip.count() and tip.first.is_visible():
                    tip.first.click()
                area.focus(); page.keyboard.press('Control+Home')
                page.keyboard.type('PAGES EDIT ', delay=30); page.keyboard.press('Control+s')
                expect(status).to_contain_text('편집기 반영됨', timeout=30000)
                page.get_by_role('button', name='DOCX 저장', exact=True).click()
                expect(status).to_contain_text('DOCX 저장됨', timeout=90000)
                assert 'PAGES EDIT' in '\n'.join(p.text for p in Document(d['source_uri']).paragraphs)
                # Selection/AI replacement uses the DOCX engine's existing proposal contract.
                expect(page.get_by_role('button', name='문구 선택 고정', exact=True)).to_be_enabled(timeout=20000)
                page.locator('iframe[name="frameEditor"]').click(position={'x': 600, 'y': 310})
                page.keyboard.type('ANCHOR', delay=30)
                for _ in range(6):
                    page.keyboard.press('Shift+ArrowLeft')
                page.get_by_role('button', name='문구 선택 고정', exact=True).click()
                expect(status).to_contain_text('선택 위치를 고정했습니다', timeout=20000)
                page.get_by_label('사무 문서 교체 문구').fill('REVIEWED ')
                page.get_by_role('button', name='사무 문서 제안 만들기', exact=True).click()
                expect(page.get_by_role('button', name='변경 추적으로 적용', exact=True)).to_be_enabled(timeout=30000)
                page.get_by_role('button', name='변경 추적으로 적용', exact=True).click()
                expect(status).to_contain_text('변경 추적으로 적용했습니다', timeout=30000)
                area.focus(); page.keyboard.press('Control+s')
                page.get_by_role('button', name='DOCX 저장', exact=True).click()
                expect(status).to_contain_text('DOCX 저장됨', timeout=90000)
                with zipfile.ZipFile(d['source_uri']) as z:
                    xml = z.read('word/document.xml').decode()
                assert 'REVIEWED' in xml and '<w:ins' in xml
                open_document(page, port, source, fresh=True)
                expect(status).to_have_text('편집 준비 완료', timeout=120000)
                assert workspace.open(source)['document']['id'] == d['id']
                assert source.read_bytes() == before
                page.screenshot(path=str(tmp_path / 'pages-editor.png'))
                assert not errors
            except Exception:
                print('Browser failure:', page.locator('body').inner_text()[-5000:])
                page.screenshot(path=str(tmp_path / 'failure.png'))
                raise
            finally:
                browser.close()
        # A genuine package-directory input has identical archive members.
        package = tmp_path / 'package.pages'; package.mkdir()
        with zipfile.ZipFile(source) as z:
            z.extractall(package)  # trusted synthetic fixture
        assert workspace.open(package)['document']['source_format'] == 'docx'
        assert not (package / '.git').exists()
        complex_source = tmp_path / 'complex.pages'
        shutil.copyfile(FIXTURES / 'complex.pages', complex_source)
        complex_before = complex_source.read_bytes()
        converted = workspace.open(complex_source)['document']
        doc = Document(converted['source_uri'])
        assert any('병합 표 제목' in t.cell(0, 0).text for t in doc.tables)
        assert any('행1 열0' in t.cell(1, 0).text for t in doc.tables)
        with zipfile.ZipFile(converted['source_uri']) as z:
            all_xml = '\n'.join(z.read(n).decode() for n in z.namelist() if n.endswith('.xml'))
            assert all(s in all_xml for s in ['머리말 보존', '꼬리말 보존', '각주 내용 보존', '텍스트 상자 보존'])
            assert any(n.startswith('word/media/') for n in z.namelist())
            # ONLYOFFICE 9.3.1.10 loses horizontal merges in this Pages fixture.
            # Report the compatibility gap; never present this as full fidelity.
            print(f'Pages table merge retained: {"gridSpan" in all_xml or "hMerge" in all_xml}')
        out = document_formats.convert(workspace, converted['id'], 'pdf', converted['revision_id'])
        with pymupdf.open(out['document']['source_uri']) as pdf, pymupdf.open(FIXTURES / 'complex.pdf') as baseline:
            actual = ''.join(p.get_text() for p in pdf).replace(' ', '').replace('\n', '')
            for expected in ['한글본문과', 'Englishtext', '병합표제목', '머리말보존', '꼬리말보존', '각주내용보존', '텍스트상자보존']:
                assert expected in actual, expected
            print(f'Pages complex PDF: baseline={len(baseline)} pages, converted={len(pdf)} pages')
            pdf[0].get_pixmap().save(str(tmp_path / 'complex-converted.png'))
            baseline[0].get_pixmap().save(str(tmp_path / 'complex-baseline.png'))
        assert complex_source.read_bytes() == complex_before
        print(f'Pages real acceptance artifacts: {tmp_path}')
    finally:
        server.should_exit = True; thread.join(10)


@pytest.mark.system
def test_cancelled_conversion_keeps_previous_editor_and_draft(tmp_path, monkeypatch):
    """A late conversion response must not navigate away from the user's draft."""
    import api_documents
    import document_formats
    import document_office
    import uvicorn
    from document_workspace import DocumentWorkspace
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from playwright.sync_api import sync_playwright, expect
    from test_document_app_support import mount_document_app, open_document
    from test_document_imports import pages, docx_bytes

    workspace = DocumentWorkspace(tmp_path / 'state')
    entered, release, completed = threading.Event(), threading.Event(), threading.Event()
    def convert(*args):
        entered.set()
        assert release.wait(20)
        return docx_bytes()
    original_open = workspace.open
    def observe(path, *args, **kwargs):
        result = original_open(path, *args, **kwargs)
        if str(path).endswith('.pages'):
            completed.set()
        return result
    monkeypatch.setattr(workspace, 'open', observe)
    monkeypatch.setattr(document_formats, 'convert_bytes', convert)
    monkeypatch.setattr(document_office, 'available', lambda: False)
    monkeypatch.setattr(api_documents, 'service', lambda: workspace)
    app = FastAPI(); app.include_router(api_documents.router)
    mount_document_app(app, workspace)
    @app.get('/launcher/auth/session')
    def auth():
        return {'authenticated': True, 'external': False}
    @app.get('/')
    def index():
        return HTMLResponse((ROOT / 'frontend/dist/index.html').read_text().replace(
            '<html', '<html data-indiebiz-surface="remote"', 1))
    app.mount('/assets', StaticFiles(directory=ROOT / 'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('127.0.0.1', 0)); port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    try:
        until = time.monotonic() + 10
        while not server.started and time.monotonic() < until:
            time.sleep(.02)
        text_path = tmp_path / 'previous.txt'; text_path.write_text('original')
        pages_path = pages(tmp_path / 'delayed.pages')
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                open_document(page, port, text_path)
                editor = page.get_by_role('textbox', name='문서 원문', exact=True)
                expect(editor).to_have_value('original')
                editor.fill('keep this unsaved draft')
                open_document(page, port, pages_path)
                assert entered.wait(5)
                expect(page.get_by_role('dialog')).to_be_visible()
                page.get_by_role('dialog').get_by_role('button', name='취소', exact=True).click()
                release.set(); assert completed.wait(10)
                page.wait_for_timeout(200)  # deliver the completed HTTP response to React
                expect(editor).to_have_value('keep this unsaved draft')
                expect(page.get_by_role('dialog')).to_have_count(0)
                # A successful retry flushes the previous editor before changing surfaces.
                editor.fill('last edit before switch')
                open_document(page, port, pages_path)
                expect(page.get_by_text('delayed_converted.docx —', exact=False)).to_be_visible()
                previous = workspace.open(text_path)
                assert previous['text'] == 'last edit before switch'
                assert text_path.read_text() == 'original'
            finally:
                browser.close()
    finally:
        release.set(); server.should_exit = True; thread.join(10)


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
