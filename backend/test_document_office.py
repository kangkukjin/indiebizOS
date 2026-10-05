"""Native draft fencing, signed callbacks, immutable references and OCR copies."""
import io
from pathlib import Path

import boot_paths  # noqa: F401
import jwt
import pytest
from docx import Document

import document_office as engine_module
from document_office import OfficeEngine
from document_workspace import DocumentWorkspace, DocumentConflict


def arguments(d, s):
    return dict(document_id=d['id'], session_id=s['id'], client_id=s['client_id'],
                epoch=s['engine_epoch'], expected=s['session_revision'])


def docx(text):
    doc = Document(); doc.add_paragraph(text)
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


@pytest.fixture
def office(tmp_path, monkeypatch):
    cfg = dict(url='http://127.0.0.1:8093', callback_origin='http://127.0.0.1:8765', secret='test-' * 10)
    monkeypatch.setattr(engine_module, 'settings', lambda: cfg)
    path = tmp_path / 'original.docx'; path.write_bytes(docx('Original'))
    app = DocumentWorkspace(tmp_path / 'state')
    d = app.open(path)['document']; s = app.acquire(d['id'], 'window')['session']
    engine = OfficeEngine(app)
    setup = engine.config(**arguments(d, s), browser_origin='http://127.0.0.1:8765',
                          browser_api_origin='http://127.0.0.1:8765')
    return app, path, d, setup['session'], engine, cfg


def signed(cfg, s, **changes):
    body = dict(key=s['engine_key'], status=6, filetype='docx', url=cfg['url']+'/file', userdata='capture')
    body.update(changes)
    return {**body, 'token': jwt.encode(body, cfg['secret'], algorithm='HS256')}


def test_callback_is_draft_duplicate_and_explicit_save(office, monkeypatch):
    app, path, d, s, engine, cfg = office
    original = path.read_bytes(); edited = docx('Edited')
    monkeypatch.setattr(engine_module, 'download', lambda *a: edited)
    body = signed(cfg, s)
    engine.callback(s['id'], s['ticket'], body)
    next_s = app.detail(d['id'])['session']
    assert next_s['session_revision'] == 1 and next_s['state'] == 'draft'
    assert path.read_bytes() == original
    engine.callback(s['id'], s['ticket'], body)
    assert app.detail(d['id'])['session']['session_revision'] == 1
    app.save(**arguments(d, next_s), expected_revision=d['revision_id'], operation_id='save')
    assert path.read_bytes() == edited
    assert len(app.versions(d['id'])) == 2


def test_invalid_signature_and_type_cannot_publish(office, monkeypatch):
    app, path, d, s, engine, cfg = office
    def unexpected(*args):
        pytest.fail('invalid callback must never download')
    monkeypatch.setattr(engine_module, 'download', unexpected)
    body = signed(cfg, s); body['status'] = 2
    with pytest.raises(PermissionError):
        engine.callback(s['id'], s['ticket'], body)
    with pytest.raises(ValueError):
        engine.callback(s['id'], s['ticket'], signed(cfg, s, filetype='pdf'))
    assert app.detail(d['id'])['session']['session_revision'] == 0


def test_late_download_cannot_overwrite_reclaimed_session(office, monkeypatch):
    app, path, d, s, engine, cfg = office
    def late(*args):
        app.reclaim(d['id'], 'other-window', s['engine_epoch'])
        return docx('Late old writer')
    monkeypatch.setattr(engine_module, 'download', late)
    with pytest.raises(DocumentConflict):
        engine.callback(s['id'], s['ticket'], signed(cfg, s))
    assert app.detail(d['id'])['session']['blob'] == d['source_sha256']


def test_download_origin_rejects_external_and_credentials():
    for url in ('https://example.org/file', 'http://user@127.0.0.1:8093/file', 'http://127.0.0.1:8765/file'):
        with pytest.raises(PermissionError):
            engine_module.download(url, {'url': 'http://127.0.0.1:8093'})


def test_out_of_order_engine_callbacks_keep_newest_draft(office, monkeypatch):
    app, path, d, s, engine, cfg = office
    s['capture_requests'] = {'old': 1, 'new': 2}
    app.store.put('session', s)
    newer = docx('Newer draft')
    monkeypatch.setattr(engine_module, 'download', lambda *a: newer)
    engine.callback(s['id'], s['ticket'], signed(cfg, s, userdata='new'))
    monkeypatch.setattr(engine_module, 'download', lambda *a: docx('Late older draft'))
    engine.callback(s['id'], s['ticket'], signed(cfg, s, userdata='old'))
    current = app.detail(d['id'])['session']
    assert app.store.bytes(current['blob']) == newer
    assert current['session_revision'] == 1
    assert set(current['capture_ids']) == {'old', 'new'}


def test_office_proposal_rejects_changed_document(office):
    from document_office_ai import approve
    app, path, d, s, engine, cfg = office
    proposal = {'id': 'proposal', 'document_id': d['id'], 'session_id': s['id'],
                'epoch': s['engine_epoch'], 'blob': 'different-snapshot'}
    app.store.put('office_proposal', proposal)
    with pytest.raises(DocumentConflict):
        approve(app, **arguments(d, s), proposal_id=proposal['id'])


def test_office_restore_fences_old_engine(office, monkeypatch):
    app, path, d, s, engine, cfg = office
    monkeypatch.setattr(engine_module, 'download', lambda *a: docx('Edited'))
    engine.callback(s['id'], s['ticket'], signed(cfg, s))
    latest = app.detail(d['id'])['session']
    restored = app.restore(**arguments(d, latest), operation_id='restore', revision_id=d['revision_id'])
    assert restored['session']['engine_epoch'] != s['engine_epoch']
    with pytest.raises(DocumentConflict):
        engine.content(s['id'], s['ticket'])
    assert path.read_bytes() == app.store.bytes(d['source_sha256'])


def test_sheet_reference_uses_pinned_bytes(tmp_path):
    from openpyxl import Workbook
    from resource_links import ResourceLinks
    path = tmp_path / 'source.xlsx'
    book = Workbook(); book.active['A1'] = 'Original'; book.save(path)
    app = DocumentWorkspace(tmp_path / 'state'); d = app.open(path)['document']
    book.active['A1'] = 'External change'; book.save(path)
    ref = ResourceLinks(app).sheet(d['id'], d['revision_id'], book.active.title, 'A1:B2')
    assert ref['items'][0]['value'] == 'Original'
    assert 'Original' in ResourceLinks.markdown_table(ref)


def test_ibl_uses_same_session_and_blocks_code_publication(tmp_path, monkeypatch):
    """[self:workspace](2026-10-05) — 옛 document 세션 op 를 흡수한 한 낱말이 같은 서비스·같은 관문을 지난다."""
    from resource_links import package_module
    from workspace_sessions import Workspace
    bridge = package_module('system_essentials', 'essentials_workspace')
    path = tmp_path / 'source.md'; path.write_text('Original')
    app = Workspace(tmp_path / 'state')
    monkeypatch.setattr(bridge, 'service', lambda: app)
    opened = bridge.op_open({'path': str(path)})
    resource = opened['resource']
    assert opened['kind'] == 'document' and opened['items'][0]['resource'] == resource
    proposal = bridge.op_propose({'resource': resource, 'selector': {'start': 0, 'end': 8}, 'replacement': 'Changed'})['proposal']
    applied = bridge.op_apply({'resource': resource, 'proposal': proposal})
    assert applied['applied'] is True and applied['text'] == 'Changed'
    request = {'resource': resource, '_path_guard': lambda *a: None, '_code_path': lambda *a: True}
    with pytest.raises(PermissionError):
        bridge.op_save(request)
    assert path.read_text() == 'Original'
    request['_code_path'] = lambda *a: False
    result = bridge.op_save(request)
    assert result['state'] == 'saved'
    assert path.read_text() == 'Changed'


def test_local_ocr_correction_searchable_copy(tmp_path):
    import shutil
    import pymupdf
    from document_pdf import ocr, correct, export
    if not shutil.which('tesseract'):
        pytest.skip('local Tesseract required')
    path = tmp_path / 'scan.pdf'
    with pymupdf.open() as text:
        p = text.new_page(); p.insert_text((60, 100), 'ORIGINAL SCANNED DOCUMENT', fontsize=24)
        image = p.get_pixmap(dpi=160).tobytes('png')
    with pymupdf.open() as scan:
        p = scan.new_page(); p.insert_image(p.rect, stream=image); scan.save(path)
    original = path.read_bytes()
    app = DocumentWorkspace(tmp_path / 'state'); d = app.open(path)['document']
    result = ocr(app, d['id'], d['revision_id'], [1], 'eng')
    assert any(w['text'] == 'ORIGINAL' for w in result['pages'][0]['words'])
    word = result['pages'][0]['words'][0]
    correct(app, d['id'], result['id'], [{'page': 1, 'word': word['id'], 'text': 'CORRECTED'}])
    output = export(app, d['id'], result['id'])
    with pymupdf.open(output['document']['source_uri']) as pdf:
        assert 'CORRECTED' in pdf[0].get_text()
    assert path.read_bytes() == original


def test_pdf_page_draft_fences_old_editor(tmp_path, monkeypatch):
    import pymupdf
    from document_pdf import change_pages
    monkeypatch.setattr(engine_module, 'available', lambda: True)
    path = tmp_path / 'pages.pdf'
    with pymupdf.open() as pdf:
        for _ in range(3):
            pdf.new_page()
        pdf.save(path)
    original = path.read_bytes()
    app = DocumentWorkspace(tmp_path / 'state'); d = app.open(path)['document']
    s = app.acquire(d['id'], 'window')['session']
    next_s = change_pages(app, **arguments(d, s), pages=[2], action='rotate')['session']
    assert next_s['engine_epoch'] != s['engine_epoch']
    with pytest.raises(DocumentConflict):
        change_pages(app, **arguments(d, s), pages=[1], action='delete')
    with pymupdf.open(stream=app.store.bytes(next_s['blob']), filetype='pdf') as pdf:
        assert pdf[1].rotation == 90 and len(pdf) == 3
    assert path.read_bytes() == original
    with pytest.raises(ValueError):
        change_pages(app, **arguments(d, next_s), pages=[1, 2, 3], action='delete')


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__] + sys.argv[1:]))
