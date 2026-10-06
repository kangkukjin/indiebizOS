"""Real local ONLYOFFICE + Chromium native-file round trip (opt-in engine fixture)."""
import os
import socket
import threading
import time
from pathlib import Path

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.system
@pytest.mark.skipif(os.environ.get("INDIEBIZ_OFFICE_LIVE_TEST") != "1", reason="로컬 ONLYOFFICE 실인수 명시 실행")
def test_office_docx_roundtrip(tmp_path, monkeypatch):
    import api_documents
    import document_office
    import document_formats
    import uvicorn
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from playwright.sync_api import sync_playwright, expect
    from docx import Document
    from document_workspace import DocumentWorkspace

    workspace = DocumentWorkspace(tmp_path / "workspace")
    monkeypatch.setattr(api_documents, "service", lambda: workspace)
    app = FastAPI(); app.include_router(api_documents.router)
    from test_document_app_support import mount_document_app, open_document
    mount_document_app(app, workspace)
    @app.get("/launcher/auth/session")
    def auth():
        return {"authenticated": True, "external": False}
    @app.get("/health")
    def health():
        return {"status": "ok"}
    html = (ROOT / "frontend/dist/index.html").read_text().replace("<html", '<html data-indiebiz-surface="remote"', 1)
    @app.get("/")
    def index():
        return HTMLResponse(html)
    app.mount("/assets", StaticFiles(directory=ROOT / "frontend/dist/assets"))
    sock = socket.socket(); sock.bind(("0.0.0.0", 0)); port = sock.getsockname()[1]
    cfg = document_office.settings()
    assert cfg, "먼저 로컬 문서 엔진을 설치하세요"
    cfg = {**cfg, "callback_origin": f"http://192.168.5.2:{port}"}
    monkeypatch.setattr(document_office, "settings", lambda: cfg)
    monkeypatch.setattr(document_formats, "settings", lambda: cfg)
    server = uvicorn.Server(uvicorn.Config(app, log_level="error"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True); thread.start()
    source = tmp_path / "인수보고서.docx"
    doc = Document(); doc.add_heading("문서앱 인수", 0); doc.add_paragraph("보존할 원래 문장"); doc.add_table(rows=2, cols=2).cell(0, 0).text="표 보존"; doc.save(source)
    try:
        until = time.monotonic()+10
        while not server.started and time.monotonic()<until:
            time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": 1600, "height": 1100})
            try:
                errors=[]; page.on("pageerror", lambda e: errors.append(str(e)))
                messages=[]; page.on("console", lambda e: messages.append(e.type+": "+e.text))
                open_document(page, port, source)
                page.locator('.office-document').get_by_role('button', name='⚙ 도구', exact=True).click()  # 저장·내보내기 등은 ⚙ 안에 접혀 있다
                expect(page.locator(".office-document").get_by_role("status")).to_have_text("편집 준비 완료", timeout=120000)
                frame = page.frame_locator('iframe[name="frameEditor"]')
                area = frame.locator('#area_id')
                tip = frame.get_by_text('확인', exact=True)
                if tip.count() and tip.first.is_visible():
                    tip.first.click()
                area.focus()
                page.keyboard.press("Control+Home")
                page.keyboard.type("ACCEPTANCE EDIT ", delay=40)
                page.keyboard.press("Control+s")
                expect(page.locator('.office-document').get_by_role('status')).to_contain_text('편집기 반영됨', timeout=20000)
                page.locator('.office-document').get_by_role('button', name='원본 저장', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_have_text('저장됨 · 원본 파일 기록 확인', timeout=90000)
                reopened = Document(source)
                assert 'ACCEPTANCE EDIT' in '\n'.join(p.text for p in reopened.paragraphs)
                assert reopened.tables[0].cell(0,0).text == '표 보존'
                expect(page.get_by_role('button', name='문구 선택 고정', exact=True)).to_be_enabled(timeout=20000)
                tip = frame.get_by_text('확인', exact=True)
                if tip.count() and tip.first.is_visible():
                    tip.first.click()
                page.locator('iframe[name="frameEditor"]').click(position={'x':600,'y':310})
                page.keyboard.type('ANCHOR', delay=40)
                for _ in range(6):
                    page.keyboard.press('Shift+ArrowLeft')
                page.get_by_role('button', name='문구 선택 고정', exact=True).click()
                page.get_by_label('사무 문서 교체 문구').fill('REVIEWED ')
                page.get_by_role('button', name='사무 문서 제안 만들기', exact=True).click()
                expect(page.get_by_role('button', name='변경 추적으로 적용', exact=True)).to_be_enabled(timeout=30000)
                page.get_by_role('button', name='변경 추적으로 적용', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_contain_text('변경 추적으로 적용했습니다', timeout=20000)
                area.focus(); page.keyboard.press('Control+s')
                page.locator('.office-document').get_by_role('button', name='원본 저장', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_have_text('저장됨 · 원본 파일 기록 확인', timeout=90000)
                import zipfile
                with zipfile.ZipFile(source) as saved:
                    xml = saved.read('word/document.xml').decode()
                assert 'REVIEWED' in xml and '<w:ins' in xml and '<w:del' in xml
                page.locator('.office-document').get_by_role('button', name='내보내기', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_contain_text('변환 사본 저장됨:', timeout=90000)
                import pymupdf
                pdfs = list((tmp_path / 'workspace/files').rglob('*.pdf'))
                assert pdfs
                with pymupdf.open(pdfs[0]) as pdf:
                    assert pdf.page_count >= 1
                # Other advertised native formats must survive the same UI
                # publication boundary; producing a conversion alone is not an
                # edit/save acceptance test.
                for format in ('odt', 'rtf'):
                    original_path = tmp_path / ('base-' + format + '.docx')
                    original_path.write_bytes(docx_bytes('Native format acceptance'))
                    base = workspace.open(original_path)['document']
                    converted = document_formats.convert(workspace, base['id'], format, base['revision_id'])
                    path = Path(converted['document']['source_uri'])
                    previous = path.read_bytes()
                    open_document(page, port, path)
                    page.locator('.office-document').get_by_role('button', name='⚙ 도구', exact=True).click()  # 저장·내보내기 등은 ⚙ 안에 접혀 있다
                    expect(page.locator('.office-document').get_by_role('status')).to_have_text('편집 준비 완료', timeout=60000)
                    area = page.frame_locator('iframe[name="frameEditor"]').locator('#area_id')
                    area.focus(); page.keyboard.type('NATIVE EDIT ', delay=40); page.keyboard.press('Control+s')
                    expect(page.locator('.office-document').get_by_role('status')).to_contain_text('편집기 반영됨', timeout=20000)
                    page.locator('.office-document').get_by_role('button', name='원본 저장', exact=True).click()
                    expect(page.locator('.office-document').get_by_role('status')).to_have_text('저장됨 · 원본 파일 기록 확인', timeout=60000)
                    document_office.validate(path.read_bytes(), format)
                    assert path.read_bytes() != previous
                    if format == 'odt':
                        with zipfile.ZipFile(path) as odt:
                            assert b'NATIVE EDIT' in odt.read('content.xml')
                    else:
                        current = workspace.detail(converted['document']['id']); session = current['session']
                        roundtrip = document_formats.convert(workspace, current['document']['id'], 'docx',
                            current['document']['revision_id'], session['id'], session['client_id'],
                            session['engine_epoch'], session['session_revision'])
                        readback = Document(roundtrip['document']['source_uri'])
                        assert 'NATIVE EDIT' in '\n'.join(p.text for p in readback.paragraphs)
                page.screenshot(path='/tmp/document-office-acceptance.png')
                open_document(page, port, pdfs[0])
                page.locator('.office-document').get_by_role('button', name='⚙ 도구', exact=True).click()  # 저장·내보내기 등은 ⚙ 안에 접혀 있다
                expect(page.locator('.office-document').get_by_role('status')).to_have_text('편집 준비 완료', timeout=60000)
                with pymupdf.open(pdfs[0]) as pdf:
                    rotation = pdf[0].rotation
                page.get_by_role('button', name='PDF 쪽 변경', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_have_text('편집 준비 완료', timeout=60000)
                page.locator('.office-document').get_by_role('button', name='원본 저장', exact=True).click()
                expect(page.locator('.office-document').get_by_role('status')).to_have_text('저장됨 · 원본 파일 기록 확인', timeout=60000)
                with pymupdf.open(pdfs[0]) as pdf:
                    assert pdf[0].rotation == (rotation + 90) % 360
                assert not errors, errors
                page.screenshot(path='/tmp/document-pdf-acceptance.png')
            except Exception:
                page.screenshot(path='/tmp/document-office-failure.png')
                Path('/tmp/document-office-ui.txt').write_text('\n'.join(errors+messages)+'\n'+page.locator('body').inner_text()+'\nFRAMES\n'+'\n'.join(f.url.split('/documents/engine-io/')[0].split('?')[0] for f in page.frames))
                raise
            finally:
                browser.close()
    finally:
        server.should_exit=True; thread.join(timeout=10); sock.close()


def docx_bytes(text):
    import io
    from docx import Document
    document = Document(); document.add_paragraph(text)
    stream = io.BytesIO(); document.save(stream)
    return stream.getvalue()


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
