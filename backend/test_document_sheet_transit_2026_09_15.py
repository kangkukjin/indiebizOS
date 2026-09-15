"""문서·장부의 보존/실계산 및 대중교통 계약 회귀. 외부 교통 호출은 대역."""
import importlib.util
import json
import os
from pathlib import Path
import zipfile

import boot_paths  # noqa: F401
from docx import Document
from lxml import etree
import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / 'data/packages/installed/tools'


def load(name, package='system_essentials'):
    spec = importlib.util.spec_from_file_location('feature_' + name, PKG/package/(name+'.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def params(path, **kw):
    return {'path': str(path), '_project_path': str(path.parent), '_path_guard': lambda *a: None, **kw}


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture
def proposal(tmp_path):
    path = tmp_path/'proposal.docx'
    d = Document()
    p = d.add_paragraph()
    p.add_run('납기: ')
    p.add_run('9월 ').bold = True
    p.add_run('30일').italic = True
    p.add_run('까지 납품합니다.')
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '품목'
    table.cell(0, 1).text = '금액'
    table.cell(1, 0).text = '제작'
    table.cell(1, 1).text = '100만원'
    d.sections[0].header.paragraphs[0].text = '제안서 1판'
    d.sections[0].footer.paragraphs[0].text = '보존할 꼬리말'
    d.save(path)
    return path


def test_docx_tracked_replacement_preserves_original_and_parts(proposal):
    mod = load('docx_edit_ops')
    before = proposal.read_bytes()
    snapshot = mod.op_inspect(params(proposal))
    row = next(x for x in snapshot['items'] if '납기:' in x['text'])
    out = mod.op_edit(params(proposal, expected_sha256=snapshot['sha256'], edits=[
        {'block_id': row['block_id'], 'old_string': '9월 30일', 'new_string': '10월 15일', 'comment': '납기 변경'}]))
    assert proposal.read_bytes() == before
    oldparts, newparts = parts(proposal), parts(out['path'])
    for n, b in oldparts.items():
        if n not in ('word/document.xml', 'word/settings.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'):
            assert newparts[n] == b, n
    root = etree.fromstring(newparts['word/document.xml'])
    ns = mod.NS
    assert ''.join(root.xpath('//w:del/w:r/w:delText/text()', namespaces=ns)) == '9월 30일'
    assert root.xpath('//w:ins/w:r/w:t/text()', namespaces=ns) == ['10월 15일']
    assert root.xpath('//w:ins/w:r/w:rPr/w:b', namespaces=ns)
    assert root.xpath('//w:commentRangeStart/@w:id', namespaces=ns) == root.xpath('//w:commentReference/@w:id', namespaces=ns)
    assert '납기 변경' in newparts['word/comments.xml'].decode()
    assert out['changed_count'] == 1


def test_docx_clean_table_header_batch(proposal):
    mod = load('docx_edit_ops')
    snapshot = mod.op_inspect(params(proposal))
    edits = [{'block_id': r['block_id'], 'old_string': r['text'], 'new_string': new}
             for old, new in [('100만원', '120만원'), ('제안서 1판', '제안서 2판')]
             for r in snapshot['items'] if r['text'] == old]
    out = mod.op_edit(params(proposal, expected_sha256=snapshot['sha256'], edits=edits, track_changes=False))
    d = Document(out['path'])
    assert d.tables[0].cell(1, 1).text == '120만원'
    assert d.sections[0].header.paragraphs[0].text == '제안서 2판'
    assert d.sections[0].footer.paragraphs[0].text == '보존할 꼬리말'


@pytest.mark.parametrize('failure', ['stale', 'missing', 'duplicate_block', 'guard', 'existing_output', 'non_boolean', 'ambiguous'])
def test_docx_rejection_never_changes_source(proposal, failure):
    mod = load('docx_edit_ops')
    if failure == 'ambiguous':
        d = Document(proposal)
        d.paragraphs[0].text = '같은 같은'
        d.save(proposal)
    before = proposal.read_bytes()
    snap = mod.op_inspect(params(proposal))
    row = snap['items'][0]
    edit = {'block_id': row['block_id'], 'old_string': '같은' if failure == 'ambiguous' else '9월 30일', 'new_string': '새 납기'}
    kw = params(proposal, expected_sha256=snap['sha256'], edits=[edit])
    if failure == 'stale': kw['expected_sha256'] = 'old'
    if failure == 'missing': edit['old_string'] = '없는 문구'
    if failure == 'duplicate_block': kw['edits'].append(edit)
    if failure == 'guard': kw['_path_guard'] = lambda *a: 'denied'
    if failure == 'existing_output': kw['output'] = str(proposal)
    if failure == 'non_boolean': kw['track_changes'] = 'false'
    with pytest.raises(ValueError): mod.op_edit(kw)
    assert proposal.read_bytes() == before
    assert not proposal.with_name('proposal_edited.docx').exists()


def test_docx_complex_field_is_readable_but_not_editable(proposal):
    d = Document(proposal)
    p = d.paragraphs[0]._p
    mod = load('docx_edit_ops')
    etree.SubElement(p[0], mod.tag('fldChar'), {mod.tag('fldCharType'): 'begin'})
    d.save(proposal)
    assert mod.op_inspect(params(proposal))['items'][0]['editable'] is False


@pytest.fixture
def workbook(tmp_path):
    p = tmp_path/'book.xlsx'
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '매출'
    ws.append(['단가', '수량', '합계'])
    ws.append([100, 3, '=A2*B2'])
    ws.append([200, 4, '=A3*B3'])
    ws['C4'] = '=SUM(C2:C3)'
    ws['A2'].font = openpyxl.styles.Font(name='Arial', italic=True)
    wb.create_sheet('원본').append(['유지', 123])
    chart = openpyxl.chart.BarChart()
    chart.add_data(openpyxl.chart.Reference(ws, min_col=1, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, 'F1')
    wb.save(p)
    return p


def test_sheet_range_write_preserves_other_sheet_chart_and_styles(workbook):
    mod = load('sheet_range_ops')
    snapshot = mod.op_range(params(workbook, range='A2:C3'))
    assert snapshot['items'][2]['formula'] == '=A2*B2'
    out = mod.op_range_write(params(workbook, expected_sha256=snapshot['sha256'], range='A2:B3',
                                     values=[[110, 5], [220, 6]], format={'bold': True, 'number_format': '#,##0'}))
    before, after = parts(workbook), parts(out['path'])
    for name in before:
        if name not in ('xl/worksheets/sheet1.xml', 'xl/styles.xml', 'xl/workbook.xml'):
            assert before[name] == after[name], name
    wb = openpyxl.load_workbook(out['path'])
    assert wb['매출']['A2'].value == 110
    assert wb['매출']['A2'].font.bold and wb['매출']['A2'].font.italic
    assert wb['매출']['C2'].value == '=A2*B2'
    assert wb['매출']['A2'].number_format == '#,##0'
    assert len(wb['매출']._charts) == 1
    assert wb['원본']['B1'].value == 123


@pytest.mark.parametrize('patch', [
    {'range': 'B2:A1'}, {'range': 'A:A'}, {'range': 'XFE1'}, {'range': 'A1:A20000'},
    {'values': [[1]]}, {'values': [[float('nan'), 1]]}, {'format': {'bold': 'false'}},
    {'format': {'unknown': 3}}, {'expected_sha256': 'old'},
])
def test_sheet_invalid_write_is_atomic(workbook, patch):
    mod = load('sheet_range_ops')
    before = workbook.read_bytes()
    snap = mod.op_range(params(workbook, range='A2:B2'))
    kw = params(workbook, expected_sha256=snap['sha256'], range='A2:B2', values=[[1, 2]])
    kw.update(patch)
    with pytest.raises(ValueError): mod.op_range_write(kw)
    assert workbook.read_bytes() == before
    assert not workbook.with_name('book_range.xlsx').exists()


def test_sheet_new_column_and_formula(workbook):
    mod = load('sheet_range_ops')
    snap = mod.op_range(params(workbook, range='D1:D2'))
    out = mod.op_range_write(params(workbook, expected_sha256=snap['sha256'], range='D1:D2', values=[['부가세'], ['=C2*0.1']]))
    wb = openpyxl.load_workbook(out['path'])
    assert wb['매출']['D2'].value == '=C2*0.1'
    assert wb['매출'].max_column >= 4


def test_calculate_missing_engine(workbook, monkeypatch):
    mod = load('sheet_range_ops')
    monkeypatch.setattr(mod, '_soffice', lambda: None)
    assert mod.op_calculate(params(workbook))['error_type'] == 'dependency_missing'


def test_calculate_real_engine_caches_only(workbook):
    mod = load('sheet_range_ops')
    if not mod._soffice(): pytest.skip('LibreOffice not installed')
    before = parts(workbook)
    out = mod.op_calculate(params(workbook))
    assert out['success'], out
    wb = openpyxl.load_workbook(out['path'], data_only=True)
    assert wb['매출']['C4'].value == 1100
    after = parts(out['path'])
    for name, data in before.items():
        if name != 'xl/worksheets/sheet1.xml': assert after[name] == data, name
    assert openpyxl.load_workbook(out['path'])['매출']['C4'].value == '=SUM(C2:C3)'


def test_calculate_formula_error_not_success(workbook):
    mod = load('sheet_range_ops')
    if not mod._soffice(): pytest.skip('LibreOffice not installed')
    wb = openpyxl.load_workbook(workbook)
    wb['매출']['D1'] = '=1/0'
    wb.save(workbook)
    out = mod.op_calculate(params(workbook))
    assert not out['success'] and out['recalculated']
    assert any(e['error'] == '#DIV/0!' for e in out['errors'])
    assert Path(out['path']).exists()


def urban():
    return {'result': {'searchType': 0, 'path': [{'pathType': 3, 'info': {'totalTime': 40, 'totalWalk': 600, 'payment': 1550},
          'subPath': [{'trafficType': 3, 'sectionTime': 5, 'distance': 300},
                      {'trafficType': 1, 'sectionTime': 15, 'lane': [{'name': '1호선'}]},
                      {'trafficType': 2, 'sectionTime': 15, 'lane': [{'busNo': '100'}]},
                      {'trafficType': 3, 'sectionTime': 5, 'distance': 300}]}]}}


def test_transit_normalization_retains_segments():
    mod = load('transit_routes', 'location-services')
    out = mod.normalize(urban())
    assert out['success']
    row = out['items'][0]
    assert row['transfer_count'] == 1 and row['duration_min'] == 40
    assert row['walking_distance_m'] == 600 and row['fare_krw'] == 1550
    assert len(row['segments']) == 4 and not out['live_arrival']


@pytest.mark.parametrize('payload,kind', [
    ({'error': {'code': -99}}, 'no_route'), ({'error': {'code': 500}}, 'provider_error'),
    ({}, 'invalid_response'), ([], 'invalid_response'),
    ({'result': {'searchType': 0, 'path': []}}, 'no_route'),
    ({'result': {'searchType': 1, 'path': [{'info': {'totalTime': 60}}]}}, 'incomplete_route'),
])
def test_transit_failure_not_car_route(payload, kind):
    out = load('transit_routes', 'location-services').normalize(payload)
    assert not out['success'] and out['error_type'] == kind and out['mode'] == 'transit'


def test_transit_partial_malformed_route():
    payload = urban()
    payload['result']['path'].append({'info': {}})
    out = load('transit_routes', 'location-services').normalize(payload)
    assert out['success'] and out['partial'] and len(out['items']) == 1 and out['rejected']


def test_transit_provider_request(monkeypatch):
    mod = load('transit_routes', 'location-services')
    monkeypatch.setattr(mod, 'check_api_key', lambda s: (True, ''))
    calls = []
    monkeypatch.setattr(mod, 'api_call', lambda *a, **kw: calls.append((a, kw)) or urban())
    out = mod.route({'origin': '126.97,37.55', 'destination': '127.02,37.50'}, lambda _: None)
    assert out['success']
    assert calls[0][1]['params']['SX'] == 126.97 and calls[0][1]['params']['SearchType'] == 0


def test_transit_missing_key_and_invalid_coords(monkeypatch):
    mod = load('transit_routes', 'location-services')
    monkeypatch.setattr(mod, 'check_api_key', lambda s: (False, 'missing'))
    assert mod.route({'origin': '126,37', 'destination': '127,38'}, lambda _: None)['error_type'] == 'credentials_missing'
    assert mod.route({'origin': 'nan,37', 'destination': '127,38'}, lambda _: None)['error_type'] == 'invalid_params'


def test_dispatch_and_default_car_preserved(monkeypatch):
    mod = load('handler', 'location-services')
    from types import SimpleNamespace
    monkeypatch.setattr(mod, 'kakao_navigation', lambda **kw: {'error': 'car sentinel'})
    ctx = SimpleNamespace(tool_name='kakao_navigation')
    assert json.loads(mod.execute({'origin': 'A', 'destination': 'B'}, ctx))['error'] == 'car sentinel'
    assert not json.loads(mod.execute({'origin': 'A', 'destination': 'B', 'mode': 'flying'}, ctx))['success']


def test_docx_accept_and_reject_restore_full_text(proposal):
    mod = load('docx_edit_ops')
    snap = mod.op_inspect(params(proposal))
    out = mod.op_edit(params(proposal, expected_sha256=snap['sha256'], edits=[
        {'block_id': snap['items'][0]['block_id'], 'old_string': '9월 30일', 'new_string': '10월 15일'}]))
    root = etree.fromstring(parts(out['path'])['word/document.xml'])
    p = root.find('w:body/w:p', mod.NS)
    assert ''.join(p.xpath('.//w:t[not(ancestor::w:del)]/text()', namespaces=mod.NS)) == '납기: 10월 15일까지 납품합니다.'
    assert ''.join(p.xpath('.//w:t[not(ancestor::w:ins)]/text() | .//w:delText/text()', namespaces=mod.NS)) == '납기: 9월 30일까지 납품합니다.'
    settings = etree.fromstring(parts(out['path'])['word/settings.xml'])
    assert settings.index(settings.find(mod.tag('trackRevisions'))) < settings.index(settings.find(mod.tag('defaultTabStop')))


def test_sheet_active_sheet_and_format_only(workbook):
    mod = load('sheet_range_ops')
    wb = openpyxl.load_workbook(workbook)
    wb.active = 1
    wb.save(workbook)
    snap = mod.op_range(params(workbook, range='B1'))
    assert snap['sheet'] == '원본' and snap['items'][0]['value'] == 123
    out = mod.op_range_write(params(workbook, expected_sha256=snap['sha256'], range='B1', format={'bold': True}))
    changed = openpyxl.load_workbook(out['path'])
    assert changed['원본']['B1'].value == 123 and changed['원본']['B1'].font.bold
    assert parts(out['path'])['xl/worksheets/sheet1.xml'] == parts(workbook)['xl/worksheets/sheet1.xml']


@pytest.mark.parametrize('kind', ['merged', 'protected', 'array'])
def test_sheet_structural_rejection(workbook, kind):
    mod = load('sheet_range_ops')
    wb = openpyxl.load_workbook(workbook)
    if kind == 'merged': wb.active.merge_cells('A2:B2')
    elif kind == 'protected': wb.active.protection.sheet = True
    else:
        from openpyxl.worksheet.formula import ArrayFormula
        wb.active['C2'] = ArrayFormula('C2:C3', '=A2:A3*B2:B3')
    wb.save(workbook)
    snapshot = mod.op_range(params(workbook, range='A2:B2'))
    with pytest.raises(ValueError):
        mod.op_range_write(params(workbook, expected_sha256=snapshot['sha256'], range='A2:B2', values=[[1, 2]]))
    assert not workbook.with_name('book_range.xlsx').exists()


def test_public_ibl_document_and_sheet_dispatch(proposal, workbook):
    from ibl_engine import execute_ibl as execute

    def execute_ibl(call, directory):
        from ibl_parser import parse
        code = '[' + call['_node'] + ':' + call['action'] + ']' + json.dumps(
            {k: v for k, v in call.items() if k not in ('_node', 'action')}, ensure_ascii=False)
        result = execute(parse(code)[0], directory)
        return json.loads(result) if isinstance(result, str) else result

    doc = execute_ibl({'_node': 'self', 'action': 'document', 'op': 'inspect', 'path': str(proposal)}, str(proposal.parent))
    assert doc['success'] and doc['items'][0]['block_id'] == 'word/document.xml#p0', doc
    out = execute_ibl({'_node': 'self', 'action': 'document', 'op': 'edit', 'path': str(proposal),
                      'expected_sha256': doc['sha256'], 'edits': [{'block_id': doc['items'][0]['block_id'],
                      'old_string': '9월 30일', 'new_string': '10월 15일'}]}, str(proposal.parent))
    assert out['success'] and Path(out['path']).exists(), out
    sheet = execute_ibl({'_node': 'self', 'action': 'sheet', 'op': 'range', 'path': str(workbook), 'range': 'A2:C2'}, str(workbook.parent))
    assert sheet['success'] and sheet['items'][2]['formula'] == '=A2*B2', sheet
    written = execute_ibl({'_node': 'self', 'action': 'sheet', 'op': 'range_write', 'path': str(workbook),
                          'range': 'A2', 'values': [[555]], 'expected_sha256': sheet['sha256']}, str(workbook.parent))
    assert written['success'] and openpyxl.load_workbook(written['path']).active['A2'].value == 555, written


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
