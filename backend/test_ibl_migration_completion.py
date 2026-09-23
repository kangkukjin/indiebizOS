"""Migration through actual member/document boundaries and durable effect receipts."""
import boot_paths  # noqa: F401
import json
from pathlib import Path
import threading
from types import SimpleNamespace
import pytest
from ibl_v2_adapters import Adapter
from ibl_v2_compile import compile_program
from ibl_v2_runtime import Runtime
from ibl_v2_entry import handle_request
from ibl_run_journal import Journal
from ibl_v2_ir import Fault


def effect(fn):
    return Adapter({'version': 1, 'params': {'n': 'Number'}, 'result': 'Number',
                    'effects': ['write_external']}, fn)


def test_resume_skips_completed_calls_and_distinguishes_identical_occurrences(tmp_path):
    calls = []
    stop = threading.Event()
    def write(rt, args):
        calls.append(args['n'])
        if len(calls) == 1:
            stop.set()
        return args['n']
    plan = compile_program('[test:write]{n:1}; [test:write]{n:1}; [test:write]{n:2}', {'test:write': effect(write)})
    with Journal(tmp_path, plan.fingerprint) as j:
        first = Runtime(plan, journal=j, cancel_check=stop.is_set).run()
    assert first['diagnostic']['code'] == 'CANCELLED'
    with Journal(tmp_path, plan.fingerprint, first['resume']) as j:
        second = Runtime(plan, journal=j).run()
    assert second['success'] and calls == [1, 1, 2]
    assert sum(e['kind'] == 'receipt_reused' for e in second['evidence']) == 1
    with Journal(tmp_path, plan.fingerprint, first['resume']) as j:
        third = Runtime(plan, journal=j).run()
    assert third['success'] and calls == [1, 1, 2]
    with pytest.raises(Fault, match='다릅니다'):
        with Journal(tmp_path, 'changed', first['resume']):
            pass


def test_resume_parallel_nested_occurrences_have_stable_identity(tmp_path):
    calls = []
    source = '[repeat:2]{[table:each]{items:[1,1],parallel:2}{[test:write]{n:$it}}}'
    plan = compile_program(source, {'test:write': effect(lambda rt, a: calls.append(a['n']) or a['n'])})
    assert not plan.issues
    with Journal(tmp_path, plan.fingerprint) as j:
        first = Runtime(plan, journal=j).run()
    with Journal(tmp_path, plan.fingerprint, first['resume']) as j:
        second = Runtime(plan, journal=j).run()
    assert first['success'] and second['success'] and len(calls) == 4
    assert sum(e['kind'] == 'receipt_reused' for e in second['evidence']) == 4


def test_crash_after_effect_started_is_not_retried_and_lock_excludes_second_runner(tmp_path):
    calls = []
    def crash(rt, args):
        calls.append(1)
        raise KeyboardInterrupt()
    plan = compile_program('[test:write]{n:1}', {'test:write': effect(crash)})
    with Journal(tmp_path, 'p') as j:
        resume = {'run_id': j.run_id}
        with pytest.raises(Fault, match='진행 중'):
            with Journal(tmp_path, 'p', resume):
                pass
        with pytest.raises(KeyboardInterrupt):
            Runtime(plan, journal=j).run()
    with Journal(tmp_path, 'p', resume) as j:
        result = Runtime(plan, journal=j).run()
    assert result['diagnostic']['code'] == 'EFFECT_UNCERTAIN' and calls == [1]
    with pytest.raises(Fault):
        with Journal(tmp_path, 'p', {'run_id': '../private'}):
            pass


def test_resume_preserves_caught_failure_and_finally_lifecycle(tmp_path):
    calls, stop = [], threading.Event()
    def write(rt, args):
        calls.append(args['n'])
        if args['n'] == 1:
            stop.set()
        return args['n']
    plan = compile_program('[try]{[test:write]{n:1}; [test:write]{n:2}}[finally]{[test:write]{n:3}}', {'test:write': effect(write)})
    with Journal(tmp_path, 'p') as j:
        result = Runtime(plan, journal=j, cancel_check=stop.is_set).run()
    assert calls == [1, 3]
    with pytest.raises(Fault) as exc:
        with Journal(tmp_path, 'p', result['resume']):
            pass
    assert exc.value.code == 'RESUME_CLEANED_UP'


def test_public_resume_pins_inputs_and_does_not_repeat_write(tmp_path, monkeypatch):
    import ibl_run_journal
    monkeypatch.setattr(ibl_run_journal, 'journal_root', lambda _: tmp_path / 'runs')
    code = '[self:write]{path:"outputs/a.txt",content:$text}'
    request = {'edition': 2, 'code': code, 'inputs': {'text': 'first'}}
    first = handle_request(request, str(tmp_path))
    assert first['success'], first
    path = tmp_path / 'outputs/a.txt'
    path.write_text('external change')
    again = handle_request({**request, 'resume': first['resume']}, str(tmp_path))
    assert again['success'] and path.read_text() == 'external change'
    changed = handle_request({**request, 'inputs': {'text': 'different'}, 'resume': first['resume']}, str(tmp_path))
    assert changed['diagnostic']['code'] == 'RESUME_CHANGED'


def test_pdf_docx_xlsx_native_read_and_table_data(tmp_path):
    import fitz
    from docx import Document
    from openpyxl import Workbook
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), 'page one')
    pdf.new_page().insert_text((72, 72), 'page two')
    pdf.save(tmp_path / 'doc.pdf')
    pdf.close()
    doc = Document()
    doc.add_paragraph('word text')
    doc.save(tmp_path / 'doc.docx')
    wb = Workbook()
    wb.active.title = 'first'
    second = wb.create_sheet('chosen')
    second.append(['name', 'qty'])
    second.append(['A', 7])
    wb.save(tmp_path / 'doc.xlsx')
    for code, phrase in [('[self:read]{path:"doc.pdf",pages:"2"}', 'page two'),
                         ('[self:read]{path:"doc.docx",extract_images:false}', 'word text'),
                         ('[self:read]{path:"doc.xlsx",sheet:"chosen"}', 'A')]:
        result = handle_request({'edition': 2, 'code': code}, str(tmp_path))
        assert result['success'], result
        assert phrase in result['value']['text']
        assert isinstance(result['value']['blocks'], list)
    assert result['value']['data']['table']['rows'] == [['A', 7]]


def test_member_current_authoring_libraries_and_private_journal(tmp_path, monkeypatch):
    import principal as p
    import member_runtime as mr
    import member_profile
    import member_bridge
    from member_runner import MemberRunner
    monkeypatch.setattr(member_profile, '_package_open', lambda *a, **kw: True)
    monkeypatch.setattr(member_bridge, 'connected', lambda _: True)
    monkeypatch.setattr(member_bridge, 'request', lambda command, **kw: {'success': True, 'path': command.get('path')})
    runner = MemberRunner.__new__(MemberRunner)
    runner.project_path = tmp_path
    runner.config = {'_member_libraries': ['#!ibl edition=2\n[def:double]($x){return $x*2}',
                                         '[def:old]{[self:write]{path:"old.txt",content:"$text"}}']}
    with p.narrow(p.member('A', 4, 'dev')), mr.turn_scope(tmp_path, 'dev', 'task', threading.Event(), {}):
        out = json.loads(runner._member_tool('execute_ibl', {'code': '[fn:double]{x:$x}', 'inputs': {'x': 3}}))
        assert out['edition'] == 2 and out['value'] == 6, out
        legacy = json.loads(runner._member_tool('execute_ibl', {'code': '[fn:old]{text:"member"}'}))
        assert legacy['success'], legacy
        assert list((tmp_path / 'ibl_runs').glob('*.sqlite'))
        prompt = runner._build_agent_prompt_split('member role')[0]
        assert '($목록,$단가)' in prompt and 'data/common_prompts' not in prompt
        denied = json.loads(runner._member_tool('execute_ibl', {'code': '[self:write]{path:"a",content:"$file:0"}', 'files':['secret']}))
        assert not denied['success']


def test_remote_negotiates_and_sends_one_lossless_call(monkeypatch):
    import requests
    from ibl_remote_call import forward, PROTOCOL
    posts = []
    monkeypatch.setattr(requests, 'get', lambda *a, **kw: SimpleNamespace(status_code=200, json=lambda:{'call_protocols':[PROTOCOL]}))
    monkeypatch.setattr(requests, 'post', lambda url, **kw: posts.append(kw['json']) or SimpleNamespace(status_code=200, json=lambda:{'success':True, 'value':{'error':'business', 'items':[]}}))
    result = forward('http://device', 'self', 'script', {'id':'x', 'args':{'literal':'$x'}, '_ibl_edition':2}, 'a', 'phone')
    assert result['value']['error'] == 'business'
    assert len(posts) == 1 and '_ibl_edition' not in posts[0]['params']
    monkeypatch.setattr(requests, 'post', lambda *a, **kw: (_ for _ in ()).throw(TimeoutError()))
    assert forward('http://device', 'self', 'script', {'id':'x'}, None, 'phone')['error_type'] == 'result_unknown'


def test_member_script_values_and_no_hub_path_reads(tmp_path):
    from test_ibl_v2_boundaries import script_module
    module = script_module()
    commands = []
    def exchange(command):
        commands.append(command)
        if command['action'] == 'list':
            return {'script_protocols':['ibl-script/2'], 'items':[{'id':'old'}]}
        return {'success':True, 'stdout':'{"items":[{"n":3}]}'}
    result = module.member_script({'_ibl_edition':2, 'op':'run', 'id':'old', 'args':{'x':'$literal'}},
                                  {'op':'script','action':'run','id':'old'}, exchange, tmp_path)
    assert result['value'] == {'items':[{'n':3}]}
    assert commands[-1]['args'] == {'x':'$literal'}
    assert not module.member_script({'_ibl_edition':2, 'args_file':'/owner/private'}, {}, exchange, tmp_path)['success']


def test_receipt_rechecks_authority_without_an_external_call(tmp_path):
    calls = []
    permitted = [True]
    def authorize():
        if not permitted[0]:
            raise Fault('RECEIPT_ACCESS', 'revoked', kind='permission')
    spec = effect(lambda rt, a: calls.append(a['n']) or a['n'])
    spec = Adapter(spec.contract, spec.run, authorize)
    plan = compile_program('[test:write]{n:1}', {'test:write': spec})
    with Journal(tmp_path, 'p') as j:
        first = Runtime(plan, journal=j).run()
    permitted[0] = False
    with Journal(tmp_path, 'p', first['resume']) as j:
        second = Runtime(plan, journal=j).run()
    assert second['diagnostic']['kind'] == 'permission' and calls == [1]


def test_remote_receiver_is_owner_only_and_does_not_preview_values(monkeypatch):
    import principal as p
    import ibl_engine
    from ibl_remote_call import receive, PROTOCOL
    calls = []
    value = {'error': 'business', 'text': 'a' * 10000}
    monkeypatch.setattr(ibl_engine, 'execute_ibl', lambda *a, **kw: calls.append(a) or {'success':True, 'value':value})
    payload = {'protocol':PROTOCOL, 'node':'self', 'action':'script', 'params':{'id':'demo','args':{}}, 'agent_id':'fixture'}
    result = receive(payload, '.')
    assert result['value'] == value and len(calls) == 1
    with p.narrow(p.member('B',4,'device')):
        assert receive(payload, '.')['error_type'] == 'permission'
    assert len(calls) == 1
    assert receive({**payload, 'params':{'_ibl_context':{}}}, '.')['error_type'] == 'capability'


def test_checkpoint_handle_is_published_before_effect(tmp_path, monkeypatch):
    import thread_context
    import common.spill
    events = []
    monkeypatch.setattr(thread_context, 'get_progress_ticket', lambda: 'ticket')
    monkeypatch.setattr(common.spill, 'ticket_progress', lambda ticket, data: events.append((ticket, data)))
    with Journal(tmp_path, 'p') as journal:
        journal.announce('plan')
        assert events == [('ticket', {'edition':2, 'resume':{'run_id':journal.run_id}, 'plan_hash':'plan'})]


@pytest.mark.parametrize('payload', [
    {'success': False, 'error':'denied', 'error_type':'permission'},
    {'success': True, 'items':[{'n':1}], 'truncated':True},
])
def test_member_legacy_script_preserves_failure_and_partial_evidence(payload, tmp_path):
    from test_ibl_v2_boundaries import script_module
    def exchange(command):
        if command['action'] == 'list':
            return {'script_protocols':['ibl-script/2'], 'items':[{'id':'old'}]}
        return {'success':True, 'stdout':json.dumps(payload)}
    out = script_module().member_script({'_ibl_edition':2, 'op':'run', 'id':'old'},
                                        {'op':'script', 'action':'run', 'id':'old'}, exchange, tmp_path)
    if payload.get('error'):
        assert not out['success'] and out['error_type'] == 'permission'
    else:
        from ibl_v2_adapters import decode_envelope
        with pytest.raises(Fault) as failure:
            decode_envelope(out, {'value_path':'/value'})
        assert failure.value.kind == 'partial'


def test_resume_fingerprint_ignores_library_usage_order():
    # Usage counters update DB ordering without changing the function source.
    defs = {'a':'[def:a](){return 1}', 'b':'[def:b](){return 2}'}
    a = compile_program('[fn:a]{}', definitions=defs)
    b = compile_program('[fn:a]{}', definitions=dict(reversed(list(defs.items()))))
    assert a.fingerprint == b.fingerprint


def test_plan_fingerprint_survives_process_hash_randomization():
    import os
    import subprocess
    import sys
    root = Path(__file__).parents[1]
    program = """import sys
sys.path.insert(0, 'backend')
import boot_paths
from ibl_v2_compile import compile_program
p=compile_program('[def:one](){[fn:a]{}}; [def:two](){[fn:b]{}}; return 3', definitions={'a':'[def:a](){return 1}', 'b':'[def:b](){return 2}'})
print(p.fingerprint)
"""
    hashes = [subprocess.check_output([sys.executable, '-c', program], cwd=root,
              env={**os.environ, 'PYTHONHASHSEED':str(seed)}, text=True, timeout=20).strip()
              for seed in (1, 2, 42)]
    assert len(set(hashes)) == 1


def test_cli_run_uses_shared_resume_entry(tmp_path):
    import subprocess
    import sys
    root = Path(__file__).parents[1]
    source = tmp_path / 'run.ibl'
    output = tmp_path / 'result.json'
    source.write_text('[self:write]{path:"outputs/receipt.txt",content:"first"}')
    cmd = [sys.executable, str(root/'scripts/ibl_v2.py'), 'run', str(source),
           '--project', str(tmp_path), '--output', str(output)]
    subprocess.run(cmd, cwd=root, capture_output=True, check=True, timeout=30)
    first = json.loads(output.read_text())
    target = tmp_path / 'outputs/receipt.txt'
    target.write_text('changed after execution')
    subprocess.run([*cmd, '--resume', first['resume']['run_id']], cwd=root,
                   capture_output=True, check=True, timeout=30)
    assert json.loads(output.read_text())['resumed']
    assert target.read_text() == 'changed after execution'


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, '-q']))
