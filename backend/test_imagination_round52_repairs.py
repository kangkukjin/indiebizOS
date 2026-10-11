"""긴문장 52회차(사무 서류 발행) 수리 회귀 — L52-1 채우기의 머리글·바닥글과 남은 자리,
L52-2 이미 있는 대상에 쓴 사실의 보고, L52-3 워드 문서 그리기, L52-4 PDF 칸의 CJK 대체 글꼴 폭."""
import importlib.util
import json
import shutil
from pathlib import Path

import pytest

TOOLS = Path(__file__).resolve().parent.parent / "data/packages/installed/tools"


def _load(name, rel):
    spec = importlib.util.spec_from_file_location(name, TOOLS / rel)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def office():
    return _load("r52_office_ops", "system_essentials/office_ops.py")


def _template(path):
    import docx
    d = docx.Document()
    p = d.add_paragraph("발급번호: "); p.add_run("{{발급번호}}")
    p = d.add_paragraph("성명: "); p.add_run("{{이름}}").bold = True
    p = d.add_paragraph("강좌: "); p.add_run("{{"); p.add_run("강좌").bold = True; p.add_run("}}")   # run 셋으로 쪼개진 자리
    d.add_paragraph("기간 {{기간}} · 출석 {{출석}}회")                                               # 한 run 에 자리 둘
    t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "강사"; t.cell(0, 1).text = "{{강사}}"
    d.sections[0].header.paragraphs[0].text = "한빛공방 {{기간}}"
    d.sections[0].footer.paragraphs[0].text = "진위 확인 번호 {{발급번호}}"
    d.add_paragraph("대표 {{대표}}")
    d.save(path)


def test_docx_fill_reaches_header_footer_and_reports_what_is_left(office, tmp_path):
    import docx
    tpl, out = tmp_path / "t.docx", tmp_path / "o.docx"
    _template(tpl)
    fields = json.loads(office.fill_op({"path": str(tpl)}, str(tmp_path)))
    where = {f["name"]: f["where"] for f in fields["items"]}
    assert where["발급번호"] == ["body", "footer"] and where["기간"] == ["body", "header"]

    data = {"발급번호": "HB-1", "이름": "김서연", "강좌": "도자기", "기간": "9월", "출석": 7, "강사": "한가람", "없는칸": 1}
    res = json.loads(office.fill_op({"path": str(tpl), "output": str(out), "data": data}, str(tmp_path)))
    assert res["unmatched"] == ["없는칸"]
    assert res["remaining"] == [{"name": "대표", "where": ["body"]}] and "대표" in res["warning"]
    assert "replaced_existing" not in res

    d = docx.Document(out)
    assert d.sections[0].footer.paragraphs[0].text == "진위 확인 번호 HB-1"
    assert d.sections[0].header.paragraphs[0].text == "한빛공방 9월"
    assert d.tables[0].cell(0, 1).text == "한가람"
    body = [p.text for p in d.paragraphs]
    assert body[:4] == ["발급번호: HB-1", "성명: 김서연", "강좌: 도자기", "기간 9월 · 출석 7회"]
    # 글자 모양은 자리표시자가 있던 run 을 잇는다 — 문단 전체를 첫 run 에 몰지 않는다.
    assert [(r.text, r.bold) for r in d.paragraphs[1].runs] == [("성명: ", None), ("김서연", True)]
    assert ("도자기", True) in [(r.text, r.bold) for r in d.paragraphs[2].runs]

    again = json.loads(office.fill_op({"path": str(tpl), "output": str(out), "data": {**data, "대표": "강한빛"}}, str(tmp_path)))
    assert again["remaining"] == [] and "warning" not in again and again["replaced_existing"] is True
    read = json.loads(office.read_docx({"path": str(out)}, str(tmp_path)))
    assert {"where": "footer", "text": "진위 확인 번호 HB-1"} in read["header_footer"]
    assert "진위 확인" not in read["text"]


def _form(path):
    import fitz
    doc = fitz.open(); page = doc.new_page()
    for i, name in enumerate(("name", "course", "memo")):
        w = fitz.Widget(); w.field_name = name; w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
        w.rect = fitz.Rect(72, 72 + 40 * i, 372, 92 + 40 * i); page.add_widget(w)
    doc.save(path)


def test_pdf_fill_declares_latin_widths_on_cjk_fallback_and_reports_empty_fields(office, tmp_path):
    import fitz
    tpl, out = tmp_path / "f.pdf", tmp_path / "o.pdf"
    _form(tpl)
    res = json.loads(office.fill_op({"path": str(tpl), "output": str(out),
                                     "data": {"name": "제갈 O'Neil <공방>", "course": "도자기 물레 입문"}}, str(tmp_path)))
    assert res["filled"] == ["name", "course"] and res["remaining"] == ["memo"]
    doc = fitz.open(out)
    cid_fonts = [x for x in range(1, doc.xref_length())
                 if doc.xref_get_key(x, "Subtype")[1] in ("/CIDFontType0", "/CIDFontType2")]
    assert cid_fonts, "한글 값은 CID 대체 글꼴로 외형이 만들어진다"
    # 폭 선언이 없으면 그 글꼴로 찍히는 빈칸·기호가 전각으로 벌어진다(52회차 L52-4).
    assert all(doc.xref_get_key(x, "W")[0] != "null" or doc.xref_get_key(x, "DW")[0] != "null" for x in cid_fonts)
    words = [w for w in doc[0].get_text("words") if w[4] in ("물레", "입문")]
    if len(words) == 2:   # 낱말 사이 간격이 글자 크기(전각)보다 좁아야 한다
        gap = words[1][0] - words[0][2]
        assert gap < (words[0][3] - words[0][1]) * 0.8, gap


def test_copy_says_when_it_did_not_write_the_requested_name(tmp_path):
    copy_ops = _load("r52_copy_ops", "system_essentials/copy_ops.py")
    src, dest = tmp_path / "a.txt", tmp_path / "b.txt"
    src.write_text("1")
    first = copy_ops.transfer_path(str(src), str(dest))
    assert first["path"] == str(dest) and first["renamed"] == [] and "warning" not in first
    second = copy_ops.transfer_path(str(src), str(dest))
    assert second["path"].endswith("b (2).txt") and second["renamed"]
    assert "b.txt → b (2).txt" in second["warning"] and "renamed" in second["message"]
    assert dest.read_text() == "1"


def test_render_draws_a_word_document_through_the_pdf_lane(office, tmp_path):
    render = _load("r52_render_artifact", "media_producer/render_artifact.py")
    if not render._find_soffice():
        pytest.skip("LibreOffice 없음")
    tpl = tmp_path / "t.docx"
    _template(tpl)
    out = json.loads(render.render_op_pdf({"path": str(tpl), "output_path": "cert"}, str(tmp_path / "shots")))
    assert out["items"][0]["op"] == "pdf" and Path(out["items"][0]["path"]).exists()
    assert out["converted_from"] == str(tpl) and Path(out["pdf_path"]).exists()
    import fitz
    text = fitz.open(out["pdf_path"])[0].get_text().replace(" ", "")   # 추출 글자의 빈칸은 글꼴에 따라 빠진다
    assert "진위확인번호" in text and "성명" in text      # 바닥글까지 인쇄 모양에 있다
    assert out["items"][0]["prescreen"] == ""
    assert not [p for p in (tmp_path / "shots").iterdir() if p.name.startswith("_lo_")]


def test_render_names_the_format_it_cannot_draw(tmp_path):
    render = _load("r52_render_artifact2", "media_producer/render_artifact.py")
    assert ".docx" in render.OFFICE_DOC_EXTS and ".xlsx" not in render.OFFICE_DOC_EXTS
    handler = (TOOLS / "media_producer/handler.py").read_text()
    assert '**{e: "pdf" for e in _ra.OFFICE_DOC_EXTS}' in handler      # op 생략 시 문서는 pdf 길로 추론
    source = (TOOLS / "media_producer/render_artifact.py").read_text()
    assert "Download is starting" in source and "html 로 그릴 수 없습니다" in source


if __name__ == "__main__":
    import sys
    raise SystemExit(pytest.main([__file__] + sys.argv[1:]))
