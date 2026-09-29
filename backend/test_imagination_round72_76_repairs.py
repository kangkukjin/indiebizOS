"""Cross-boundary regressions from imagination reports 72–76. No live services."""
import importlib.util
import json
import sqlite3
import sys
import threading
from datetime import datetime, timedelta
from pathlib import Path

import pytest
import boot_paths  # noqa: F401

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / 'data/packages/installed/tools'


def module(package, name):
    path = TOOLS / package / (name + '.py')
    spec = importlib.util.spec_from_file_location(f'r7276_{package.replace("-", "_")}_{name}', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.mark.parametrize('move', [False, True])
def test_transfer_never_erases_existing_files(tmp_path, move):
    ops = module('system_essentials', 'copy_ops')
    src, dst = tmp_path / 'src', tmp_path / 'dst'
    src.mkdir(); dst.mkdir()
    (src / 'a').write_text('new')
    (dst / 'a').write_text('backup')
    ops.transfer_path(str(src), str(dst), move=move)
    assert (dst / 'a').read_text() == 'backup'
    assert sorted(p.read_text() for p in dst.iterdir()) == ['backup', 'new']
    assert src.exists() != move


@pytest.mark.parametrize('relation', ['same', 'child', 'parent'])
def test_overlap_refused_before_modification(tmp_path, relation):
    ops = module('system_essentials', 'copy_ops')
    src = tmp_path / 'src'; src.mkdir(); (src / 'a').write_text('original')
    dst = {'same': src, 'child': src / 'child', 'parent': tmp_path}[relation]
    with pytest.raises(ValueError):
        ops.transfer_path(str(src), str(dst), move=True)
    assert (src / 'a').read_text() == 'original'


@pytest.fixture
def finance(tmp_path, monkeypatch):
    mod = module('finance-record', 'finance_storage')
    monkeypatch.setattr(mod, 'DATA_DIR', str(tmp_path))
    monkeypatch.setattr(mod, 'DB_PATH', str(tmp_path / 'finance.db'))
    monkeypatch.setattr(mod, 'FILES_DIR', str(tmp_path / 'files'))
    with mod.get_db_connection() as conn:
        owner = conn.execute('SELECT id FROM owners LIMIT 1').fetchone()[0]
        for i in range(1205):
            conn.execute('INSERT INTO transactions (owner_id, tx_type, amount, occurred_at, category, counterparty) VALUES (?,?,?,?,?,?)',
                         (owner, 'expense', 10, datetime.now().isoformat(), '식비', '식당'))
    return mod


def test_finance_full_population_and_no_phantom_owner(finance):
    with finance.get_db_connection() as conn:
        before = conn.execute('SELECT COUNT(*) FROM owners').fetchone()[0]
    assert len(finance.get_transactions()) == 1205
    assert finance.get_summary()['expense'] == 12050
    assert finance.get_transactions(category='교통') == []
    assert finance.get_transactions(owner='없는주체') == []
    with finance.get_db_connection() as conn:
        assert conn.execute('SELECT COUNT(*) FROM owners').fetchone()[0] == before


def test_payment_settlement_is_transfer():
    sync = module('finance-record', 'finance_sync')
    parsed = sync._parse_payment('카드대금 출금', '카드 결제대금 547,601원 출금 완료')
    assert parsed['type'] == 'transfer'
    assert sync._clean_merchant('천안과일도매(오송점) 시 인센티브 (2) 총 보유') == '천안과일도매(오송점)'
    assert sync._merchant_from('테스트상점(지점) / 신용(1234) / 09/29 12:00 / 누적 99,000원') == '테스트상점(지점)'


def test_finance_ingest_accepts_declared_items_without_external_call(monkeypatch):
    from types import SimpleNamespace
    handler = module('finance-record', 'handler')
    received = []
    def capture_source(**kwargs):
        received.append(json.loads(kwargs['text']))
        return {'ok': False, 'error': 'test stops before AI and writes'}
    monkeypatch.setitem(sys.modules, 'ingest_engine', SimpleNamespace(extract_source=capture_source))
    rows = [{'title': '영수증', 'amount': 12000}]
    handler.ingest_finance_info({'items': json.dumps({'items': rows})})
    assert received == [rows]


@pytest.fixture
def calendar(tmp_path, monkeypatch):
    import calendar_manager as cm
    monkeypatch.setattr(cm, 'CALENDAR_CONFIG_PATH', tmp_path / 'calendar.json')
    monkeypatch.setattr(cm, 'DATA_PATH', tmp_path)
    return cm.CalendarManagerBase(log_callback=lambda _: None)


def test_monthly_yearly_rules(calendar):
    base = dict(id='a', action='test', enabled=True, time='09:00', created_at='2025-01-01T00:00:00')
    monthly = dict(base, repeat='monthly', day=15)
    assert calendar._should_run_task(monthly, datetime(2026, 9, 15, 10))
    assert not calendar._should_run_task(monthly, datetime(2026, 9, 16, 10))
    annual = dict(base, repeat='yearly', date='2020-09-15')
    assert calendar._should_run_task(annual, datetime(2026, 9, 15, 10))
    assert not calendar._should_run_task(annual, datetime(2026, 10, 15, 10))


def test_timer_tick_deleted_and_double_claim(calendar):
    called = []
    calendar.register_action('fake', lambda task: called.append(task['id']))
    event = calendar.add_event(title='test', repeat='none', action='fake',
                               execute_at=(datetime.now() - timedelta(seconds=1)).isoformat())
    threads = [threading.Thread(target=calendar._execute_task, args=(event,), kwargs={'due_only': True}) for _ in range(8)]
    for t in threads: t.start()
    for t in threads: t.join()
    assert called == [event['id']]
    calendar.delete_event(event['id'])
    calendar._execute_task(event, due_only=True)
    assert len(called) == 1


@pytest.mark.parametrize('config', [dict(repeat='weekly', time='09:00', weekdays=['mon', 'fri']),
                                    dict(repeat='monthly', time='9:00', date='2026-01-31'),
                                    dict(repeat='yearly', time='09:00', date='2024-02-29')])
def test_shared_calendar_normalization(config):
    from calendar_rules import normalize_schedule_config
    result = normalize_schedule_config(config)
    assert not result.get('error'), result


def test_failed_delivery_does_not_advance_since_and_empty_initializes(tmp_path, monkeypatch):
    from execution_commit import CommitScope, bind_scope
    value_semantics = module('data-ops', 'dataops_value_semantics')
    mod = module('data-ops', 'since_ops')
    def connect():
        conn = sqlite3.connect(tmp_path / 'since.db')
        conn.execute('CREATE TABLE IF NOT EXISTS since_seen (stream TEXT, k TEXT, watched TEXT, first_seen TEXT, last_seen TEXT, PRIMARY KEY(stream,k))')
        return conn
    def since(rows):
        return mod.op_since({'items': rows}, {'key': 'stream', 'by': 'id'},
                            lambda x: (x['items'], x),
                            lambda env, items, **kw: {'items': items},
                            lambda *a: {'error': a}, lambda *a: {'error': a}, value_semantics, connect)
    first = CommitScope()
    with bind_scope(first):
        assert since([])['items'] == []
    first.commit()
    failed = CommitScope()
    with bind_scope(failed):
        assert since([{'id': 1}])['items'][0]['_since'] == 'new'
    retry = CommitScope()
    with bind_scope(retry):
        assert len(since([{'id': 1}])['items']) == 1
    retry.commit()
    assert since([{'id': 1}])['items'] == []


@pytest.mark.parametrize('limit', [1, 10, 100, 1223, 1500])
def test_price_compaction_truth_and_cardinality(limit):
    from common.response_formatter import compact_price_series
    rows = [{'date': str(i), 'close': i} for i in range(1223)]
    compact, truncated = compact_price_series(rows, limit)
    assert len(compact) == min(limit, len(rows))
    assert truncated == (len(compact) < len(rows))
    assert compact[-1] == rows[-1]


def test_native_numeric_scalar_at_wire_boundary():
    np = pytest.importorskip('numpy')
    from common.expression_ir import pack, unpack
    assert unpack(pack({'price': np.float64(10.5), 'count': np.int64(4)})) == {'price': 10.5, 'count': 4}


@pytest.mark.parametrize('tool,fn', [('tool_naver', '_trade_types'), ('tool_zigbang', '_sales_types')])
def test_listing_trade_selection(tool, fn):
    mod = module('real-estate', tool)
    select = getattr(mod, fn)
    assert select('lease', None) == select('rent', '전세')
    with pytest.raises(ValueError): select('garbage', None)


def test_crypto_snapshot_has_currency():
    mod = module('investment', 'handler')
    row = mod._attach_quote_items({'success': True, 'data': {'current_price_krw': 0, 'current_price_usd': 2}})['items'][0]
    assert row['current_price'] == 0 and row['currency'] == 'KRW'


def test_spreadsheet_records_have_cells(tmp_path):
    mod = module('system_essentials', 'office_ops')
    result = json.loads(mod.spreadsheet({'path': str(tmp_path / 'report.xlsx'),
                                         'rows': [{'name': 'a', 'amount': 3}, {'name': 'b', 'amount': 7}]},
                                        str(tmp_path), lambda *a: None))
    assert result.get('success'), result
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / 'report.xlsx')
    assert list(wb.active.values) == [('name', 'amount'), ('a', 3), ('b', 7)]


def test_storage_subtree_full_rollup_and_set_note(tmp_path, monkeypatch):
    mod = module('pc-manager', 'storage_db')
    monkeypatch.setattr(mod, 'SCANS_DIR', str(tmp_path / 'index'))
    monkeypatch.setattr(mod, 'SCANS_JSON', str(tmp_path / 'index/scans.json'))
    root = tmp_path / 'source'; child = root / 'child'; child.mkdir(parents=True)
    for n in range(25): (child / f'a.ext{n}').write_text('abc')
    assert mod.scan_directory(str(root), 'volume')['success']
    summary = mod.get_summary(str(child))
    assert summary['file_count'] == 25 and len(summary['items']) == 25
    assert summary['folders'][0]['total_size'] == 75
    assert mod.add_annotation('volume', str(child), 'one')['success']
    assert mod.add_annotation('volume', str(child), 'two')['success']
    assert [x['note'] for x in mod.get_annotations('volume')['items']] == ['two']
    assert not mod.add_annotation('volume', str(root / 'missing'), 'note')['success']


def test_notifications_suppress_identical_failure_but_report_changes():
    from calendar_actions import CalendarActionsMixin
    check = CalendarActionsMixin._should_notify_result
    task = {}
    assert check(task, {'error': 'bad'}, 100)
    assert not check(task, {'error': 'bad'}, 200)
    assert check(task, {'error': 'different'}, 201)
    assert check(task, {'success': True}, 202)
    assert check(task, {'error': 'different'}, 203)


def test_pipe_contracts_and_nested_check():
    from ibl_v2_compile import compile_program
    from ibl_v2_adapters import load_registry
    registry = load_registry(str(ROOT))
    assert registry['table:chart'].contract['pipe_input'] == 'items'
    assert registry['self:notify_user'].contract['pipe_input'] == 'message'
    bad = compile_program('[self:trigger]{op:"create",name:"test",cron:"bad",do:"return 1"}', registry)
    assert any(x['code'] == 'ARGUMENT_CONTRACT' for x in bad.issues)
    nested = compile_program('[self:schedule]{minutes:1,do:"return $undefined"}', registry)
    assert any('UNBOUND' in x['message'] for x in nested.issues)
    for source in ('[self:schedule]{time:"09:00",do:"return 1"}',
                   '[self:manage_events]{op:"create",title:"test",date:"2026-03-25 14:00"}'):
        assert not compile_program(source, registry).issues
    alternate = compile_program('[self:schedule]{minutes:1,code:"return $undefined"}', registry)
    assert any('UNBOUND' in x['message'] for x in alternate.issues)
    literal = compile_program('return "$row.price"', registry)
    assert any(x['code'] == 'LITERAL_DOLLAR' for x in literal.preflight['warnings'])


def test_document_media_metadata_and_structured_read(tmp_path):
    from PIL import Image
    img = tmp_path / 'image.png'
    Image.new('RGB', (40, 30), 'red').save(img)
    mod = module('data-ops', 'doc_build')
    blocks = [{'type': 'heading', 'text': 'Figure', 'level': 1}, {'type': 'image', 'src': str(img)}]
    html = json.loads(mod.render_document({'format': 'html', 'title': 'Report', 'blocks': blocks}, str(tmp_path)))
    assert html['success'], html
    assert 'data:image/png;base64,' in Path(html['path']).read_text()
    fmt = mod._fmt
    fmt._doc_blocks_to_docx(blocks, 'Report', str(tmp_path / 'a.docx'), 'Subtitle')
    fmt._doc_blocks_to_pptx(blocks, 'Report', str(tmp_path / 'a.pptx'), 'Subtitle')
    from docx import Document
    from pptx import Presentation
    doc = Document(tmp_path / 'a.docx')
    assert 'Subtitle' in [p.text for p in doc.paragraphs]
    assert len(doc.inline_shapes) == 1
    prs = Presentation(tmp_path / 'a.pptx')
    assert len(prs.slides) == 2
    assert prs.slides[0].placeholders[1].text == 'Subtitle'
    assert prs.slides[1].shapes.title.text == 'Figure'
    from ibl_document_value import document_value
    value = document_value({'text': '{"month":9}', 'blocks': [], 'structured_data': {'month': 9}})
    assert value['data']['month'] == 9
    images = mod._image_blocks([{'data': {'path': str(img)}}], src_field='data.path')
    assert images[0]['src'] == str(img)


def test_engine_only_commits_successful_complete_observations():
    from execution_commit import current_scope
    from ibl_v2_adapters import Adapter
    from ibl_v2_compile import compile_program
    from ibl_v2_runtime import Runtime
    committed = []
    def observe(rt, args):
        current_scope().actions['observation'] = lambda: committed.append('done')
        return {'items': []}
    contract = {'version': 1, 'params': {}, 'required': [], 'result': 'Record', 'effects': ['write_external']}
    registry = {'t:observe': Adapter(contract, observe)}
    failed = Runtime(compile_program('[t:observe]{}; return 1 / 0', registry)).run()
    assert not failed['success'] and committed == []
    success = Runtime(compile_program('return [t:observe]{}', registry)).run()
    assert success['success'] and committed == ['done']


def test_since_resume_preserves_receipt_and_cannot_rewind_newer_checkpoint(tmp_path):
    from execution_commit import bind_scope, CommitScope
    from ibl_v2_adapters import Adapter
    from ibl_v2_compile import compile_program
    from ibl_v2_runtime import Runtime
    from ibl_run_journal import Journal
    mod = module('data-ops', 'since_ops')
    semantics = module('data-ops', 'dataops_value_semantics')
    def connect():
        conn = sqlite3.connect(tmp_path / 'observations.db')
        conn.execute('CREATE TABLE IF NOT EXISTS since_seen (stream TEXT, k TEXT, watched TEXT, first_seen TEXT, last_seen TEXT, PRIMARY KEY(stream,k))')
        return conn
    def observe(rt, args):
        return mod.op_since({'items': [{'id': 1, 'price': args['price']}]},
                            {'key': 'stream', 'by': 'id', 'watch': ['price']},
                            lambda x: (x['items'], x), lambda env, items, **kw: {'items': items},
                            lambda *a: {}, lambda *a: {}, semantics, connect)
    contract = {'version': 1, 'params': {'price': 'Number'}, 'required': ['price'],
                'result': 'Record', 'effects': ['write_external'], 'deferred_observation': True}
    reg = {'t:observe': Adapter(contract, observe)}
    plan = compile_program('return [t:observe]{price:10}', reg)
    with Journal(tmp_path / 'journal', 'first') as journal:
        run_id = journal.run_id
        first = Runtime(plan, journal=journal).run()
    assert first['success']
    # Later run observes a changed value. Resuming the first run keeps its
    # receipt but must not restore price=10 in the current checkpoint.
    newer = CommitScope()
    with bind_scope(newer, '2099-01-01T00:00:00'):
        assert observe(None, {'price': 20})['items'][0]['_since'] == 'changed'
    newer.commit()
    with Journal(tmp_path / 'journal', 'first', resume={'run_id': run_id}) as journal:
        resumed = Runtime(plan, journal=journal).run()
    assert resumed['success'] and resumed['value'] == first['value']
    with connect() as conn:
        assert json.loads(conn.execute('SELECT watched FROM since_seen').fetchone()[0])['price'] == 20


def test_storage_scan_reports_inaccessible_file_and_retains_other_results(tmp_path, monkeypatch):
    mod = module('pc-manager', 'storage_db')
    monkeypatch.setattr(mod, 'SCANS_DIR', str(tmp_path / 'index'))
    monkeypatch.setattr(mod, 'SCANS_JSON', str(tmp_path / 'index/scans.json'))
    root = tmp_path / 'source'; root.mkdir()
    (root / 'good.txt').write_text('good')
    (root / 'locked').mkdir(); (root / 'locked' / 'x.txt').write_text('x')
    # 접근 실패의 예 = 잠긴 폴더(74회차 후속: 깨진 링크는 lstat 로 링크 하나로 센다 — 실패가 아니다).
    import os
    os.chmod(root / 'locked', 0o000)
    try:
        result = mod.scan_directory(str(root))
    finally:
        os.chmod(root / 'locked', 0o755)
    assert result['success'] and result['file_count'] == 1
    assert result['error_count'] == 1 and result['source_complete'] is False
    assert mod.get_summary(str(root))['source_complete'] is False
    assert mod.get_summary_all()['source_complete'] is False


def test_chart_selects_explicit_columns_and_preserves_month_labels(tmp_path, monkeypatch):
    from types import SimpleNamespace
    handler = module('visualization', 'handler')
    captured = {}
    monkeypatch.setitem(handler._RENDERERS, 'line', lambda params, data_file: captured.update(params) or {'success': True})
    ctx = SimpleNamespace(tool_name='chart', output_dir=lambda: str(tmp_path))
    out = handler.execute({'chart_type': 'line', 'items': [{'ignored': 99, 'month': '2026-09', 'amount': 4}],
                           'x': 'month', 'y': 'amount'}, ctx)
    assert out['success'] and captured['data'] == [{'month': '2026-09', 'amount': 4}]
    line = module('visualization', 'tool_line')
    figures = []
    common = SimpleNamespace(COLORS=['red'], coerce_x=lambda values: (values, False),
                             apply_bands_plotly=lambda *a: None,
                             save_plotly_figure=lambda fig, *a: figures.append(fig) or {})
    out = line._create_with_plotly([{'x': '2026-08', 'y': 3}, {'x': '2026-09', 'y': 4}],
                                  'title', None, None, None, 'html', str(tmp_path), common, False, 2)
    assert out['success'] and figures[0].layout.xaxis.type == 'category'


@pytest.mark.parametrize('fmt', ['pdf', 'png'])
def test_browser_document_embeds_local_image(tmp_path, fmt):
    from PIL import Image
    path = tmp_path / 'source.png'
    Image.new('RGB', (80, 60), 'red').save(path)
    mod = module('data-ops', 'doc_build')
    result = json.loads(mod.render_document({'format': fmt, 'title': 'Image report',
                      'blocks': [{'type': 'image', 'src': str(path)}]}, str(tmp_path)))
    assert result['success'] and result['format'] == fmt, result
    if fmt == 'pdf':
        import fitz
        with fitz.open(result['path']) as doc:
            assert sum(len(page.get_images()) for page in doc) >= 1
    else:
        pixels = Image.open(result['path']).convert('RGB')
        assert sum(1 for r, g, b in pixels.getdata() if r > 200 and g < 30 and b < 30) >= 100


def test_json_reader_exposes_structure_and_time_declares_offset(tmp_path):
    from types import SimpleNamespace
    import re
    mod = module('system_essentials', 'handler')
    path = tmp_path / 'data.json'; path.write_text('{"month":9,"items":[{"n":3}]}')
    ctx = SimpleNamespace(tool_name='read_file', project_path=str(tmp_path), agent_id=None)
    raw = json.loads(mod.execute({'path': str(path), 'blocks': True}, ctx))
    assert raw['structured_data']['items'] == [{'n': 3}]
    ctx.tool_name = 'get_current_time'
    assert re.search(r'[+-]\d{4}$', mod.execute({}, ctx))


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
