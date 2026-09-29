"""상상훈련 73회차 잔여 수리 — 조건 값 언어 개정(G73-1)·선언-읽기 관문(B73-2)·경로 관문 규칙 2(B73-6)·
차트 날짜 눈금(F73-3)·pptx 빈 자리표(F73-4). 외부 서비스 없음."""
import importlib.util
import subprocess
import sys
from pathlib import Path

import pytest
import boot_paths  # noqa: F401

from common.expression_ir import Fault, unpack
from ibl_v2_adapters import load_registry
from ibl_v2_compile import compile_program
from ibl_v2_runtime import Runtime

ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / 'data/packages/installed/tools'
SCRIPTS = ROOT / 'scripts'


@pytest.fixture(scope='module')
def registry():
    return load_registry(str(ROOT))


def value(code, registry, inputs=None):
    plan = compile_program(code, registry, inputs)
    assert not plan.issues, plan.report()
    result = Runtime(plan, inputs).run()
    assert result['success'], result
    return unpack(result['value_wire']['data'])


def issue_codes(code, registry):
    return [i['code'] for i in compile_program(code, registry).issues]


# ── G73-1 조건 값 `조건 ? 값1 : 값2` ─────────────────────────────────────────

@pytest.mark.parametrize('code,expected', [
    ('return 5 > 3 ? "A" : "B"', 'A'),
    ('return 1 > 3 ? "A" : "B"', 'B'),
    ('$s = 85.6\nreturn $s >= 90 ? "A" : $s >= 80 ? "B" : "C"', 'B'),       # 오른쪽 결합 사슬
    ('return true or false ? 1 : 2', 1),                                    # (a or b) ? …
    ('$x = -2\nreturn $x > 0\n  ? "양"\n  : "음"', '음'),                    # 여러 줄 표기
    ("$x = 3\nreturn f'부호: ${$x > 0 ? \"양\" : \"음\"}'", '부호: 양'),       # 보간 안
    ('$m = {통화:"USD", 금액:10}\nreturn $m.통화 == "KRW" ? $m.금액 : $m.금액 * 1400', 14000),
    ('$v = null\nreturn $v == null ? "없음" : $v', '없음'),
])
def test_conditional_value(registry, code, expected):
    assert value(code, registry) == expected


def test_unselected_branch_is_not_evaluated(registry):
    # 0 으로 나누기 가지는 고르지 않았으므로 실행되지 않는다.
    assert value('$d = 0\nreturn $d != 0 ? 10 / $d : 0', registry) == 0


def test_grades_in_pure_compute_callback(registry):
    code = ('$r = [{s:94.1},{s:85.6},{s:68.15}] >> [table:compute]{set:($r)=>{등급: '
            '$r.s >= 90 ? "A" : $r.s >= 80 ? "B" : "C"}}\nreturn $r')
    rows = value(code, registry)
    rows = rows.get('items', rows) if isinstance(rows, dict) else rows
    assert [r['등급'] for r in rows] == ['A', 'B', 'C']


def test_condition_must_be_bool(registry):
    assert 'TYPE' in issue_codes('return 1 ? "a" : "b"', registry)


def test_branch_must_be_pure(registry):
    assert 'PURE_EXPRESSION' in issue_codes('return true ? [self:read]{path:"x"} : 1', registry)


def test_missing_colon_names_the_form(registry):
    with pytest.raises(Fault) as info:
        compile_program('return true ? "a"', registry)
    assert info.value.code == 'SYNTAX' and '조건 ? 참일 때 값 : 거짓일 때 값' in str(info.value)


def test_fallback_operator_still_lexes(registry):
    # `??` 는 한 토큰(조합 연산자) — 조건 값의 `?` 와 섞이지 않는다.
    assert 'SYNTAX' not in issue_codes('return [1] ?? [2]', registry)


# ── B73-2 선언-읽기 관문 ────────────────────────────────────────────────────

def _declared_reads():
    sys.path.insert(0, str(SCRIPTS))
    import iblbuild_declared_reads
    return iblbuild_declared_reads


def test_declared_reads_gate_is_clean_on_live_tree():
    import yaml
    gate = _declared_reads()
    data = yaml.safe_load((ROOT / 'data/ibl_nodes.yaml').read_text(encoding='utf-8'))
    issues, stats = gate.validate_declared_reads(data, ROOT)
    assert issues == [] and stats['checked'] > 0


def test_declared_reads_flow_tracks_helpers_and_escapes(tmp_path):
    import ast
    gate = _declared_reads()
    handler = '''
_T = {"a": _op_a}
def _arg(d, *names):
    return next((d.get(n) for n in names if d.get(n)), None)
def _title(d):
    return d.get("title")
def _op_a(ti):
    for k in ("x_label", "y_label"):
        ti.get(k)
    return _title(ti)
def _op_alias(ti):
    return _arg(ti, "ticker", "symbol")
def _op_leak(ti):
    return outside(ti)
def execute(tool_input, context):
    fn = _T.get(tool_input.get("op"))
    return fn(tool_input)
'''
    flow = gate._Flow({'handler': ast.parse(handler)})
    reads, complete = flow.flow('handler', '_op_a', 0)
    assert complete and {'x_label', 'y_label', 'title'} <= reads
    reads, complete = flow.flow('handler', '_op_alias', 0)          # 별칭 도우미의 상수 = 읽기(동적 키라 판정 불가)
    assert {'ticker', 'symbol'} <= reads and complete is False
    assert flow.flow('handler', '_op_leak', 0)[1] is False          # 패키지 밖으로 새면 판정 불가
    reads, complete = flow.flow('handler', 'execute', 0)            # 디스패치 표를 따라간다
    assert complete and 'title' in reads and 'op' in reads


def test_music_artist_is_query_alias():
    import yaml
    music = yaml.safe_load((ROOT / 'data/ibl_nodes.yaml').read_text(encoding='utf-8'))['nodes']['self']['actions']['music']
    assert 'artist' in music['aliases']['query'] and 'artist' not in (music.get('params') or {})


# ── B73-6 경로 관문 규칙 2 ─────────────────────────────────────────────────

GATE = SCRIPTS / 'check_body_path_expansion.py'


@pytest.mark.parametrize('src,bad', [
    ('from pathlib import Path\ndef read_pdf(tool_input):\n    f = tool_input.get("file_path") or tool_input.get("path")\n'
     '    return Path(f)\n', True),
    ('from pathlib import Path\ndef r(tool_input):\n    raw = tool_input.get("path")\n    return Path(str(raw))\n', True),
    ('import os\ndef r(tool_input, pp):\n    return os.path.join(pp, tool_input.get("output"))\n', True),
    ('from pathlib import Path\nfrom runtime_utils import expand_body_path\ndef r(tool_input):\n'
     '    f = tool_input.get("path")\n    f = expand_body_path(f)\n    return Path(f)\n', False),
    ('import os\nfrom runtime_utils import expand_body_path\ndef r(p):\n    one = p.get("file")\n'
     '    one = os.path.abspath(expand_body_path(one))\n    return os.path.isfile(one)\n', False),
    ('from pathlib import Path\ndef r(tool_input):\n    return Path(tool_input.get("path") or "").suffix\n', False),
    ('import os\ndef r(tool_input, repo):\n    raw = tool_input.get("path")\n'
     '    return os.path.join(repo, raw)  # path-ok: 저장소 상대\n', False),
])
def test_path_gate_rule2(tmp_path, src, bad):
    f = tmp_path / 'm.py'
    f.write_text(src, encoding='utf-8')
    r = subprocess.run([sys.executable, str(GATE), '--files', str(f)], capture_output=True, text=True)
    assert (r.returncode == 1) == bad, r.stdout


def test_render_source_path_expands_workspace_token(tmp_path, monkeypatch):
    monkeypatch.setenv('INDIEBIZ_BASE_PATH', str(tmp_path))
    spec = importlib.util.spec_from_file_location('r73_render_artifact', TOOLS / 'media_producer/render_artifact.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod._src_path({'path': '~workspace/outputs/a.html'}) == str(tmp_path / 'outputs/a.html')


# ── F73-3 날짜 눈금 · F73-4 pptx 빈 자리표 ─────────────────────────────────

def _viz_common():
    spec = importlib.util.spec_from_file_location('r73_tool_common', TOOLS / 'visualization/tool_common.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_date_axes_get_korean_tick_format():
    import plotly.graph_objects as go
    from plotly.subplots import make_subplots
    common = _viz_common()
    fig = make_subplots(rows=2, cols=1, shared_xaxes=True)
    dates = ['2026-09-01', '2026-09-02', '2026-09-03']
    fig.add_trace(go.Scatter(x=dates, y=[1, 2, 3]), row=1, col=1)
    fig.add_trace(go.Bar(x=dates, y=[1, 2, 3]), row=2, col=1)
    common.apply_korean_date_ticks(fig)
    for axis in (fig.layout.xaxis, fig.layout.xaxis2):
        assert [s.value for s in axis.tickformatstops][1] == '%-m/%-d'
    months = go.Figure(go.Scatter(x=['2026-07', '2026-08'], y=[1, 2]))
    months.update_xaxes(type='category')
    common.apply_korean_date_ticks(months)
    assert not months.layout.xaxis.tickformatstops                  # 월 범주축은 그대로
    numbers = go.Figure(go.Scatter(x=[1, 2], y=[1, 2]))
    common.apply_korean_date_ticks(numbers)
    assert not numbers.layout.xaxis.tickformatstops


def test_pptx_media_slide_leaves_no_empty_body(tmp_path):
    from PIL import Image
    from pptx import Presentation
    img = tmp_path / 'c.png'
    Image.new('RGB', (40, 30), 'white').save(img)
    spec = importlib.util.spec_from_file_location('r73_doc_formats', TOOLS / 'data-ops/doc_formats.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    out = tmp_path / 't.pptx'
    mod._doc_blocks_to_pptx([{'type': 'heading', 'text': '종가 추이'}, {'type': 'image', 'src': str(img)},
                             {'type': 'heading', 'text': '요점'}, {'type': 'paragraph', 'text': '본문'}],
                            '제목', str(out), meta='부제')
    slides = list(Presentation(str(out)).slides)
    media = slides[1]
    assert media.shapes.title.text == '종가 추이'
    assert [s.shape_type for s in media.shapes if s.has_text_frame and not s.text_frame.text.strip()] == []
    assert any(s.shape_type == 13 for s in media.shapes)             # 그림은 제목 슬라이드에
    assert slides[2].shapes.title.text == '요점'


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, '-q']))
