"""실제 문서 라이브러리 + 회원 파일 전송/감사/취소 경계."""
import boot_paths
import base64
import threading
import io
from pathlib import Path
import pytest
import principal as P
import member_runtime as MR
import member_bridge as bridge
import member_profile as profile


@pytest.fixture
def member(tmp_path):
    token = P.set_transport(P.OWNER)
    with P.narrow(P.member('A', 4, 'devA')), MR.turn_scope(tmp_path, 'devA', 'files', threading.Event(), {}):
        yield tmp_path
    P.reset_transport(token)


def exchange(monkeypatch, files, reject=False):
    calls = []
    def run(c, **kw):
        calls.append(c)
        if c['op'] == 'read':
            if c['path'] not in files:
                return {'success': False, 'error': 'member_file_missing'}
            return {'success': True, 'content': base64.b64encode(files[c['path']]).decode()}
        if reject:
            return {'success': False, 'error': 'permission_denied'}
        assert c['op'] == 'write'
        files[c['path']] = base64.b64decode(c['content'])
        return {'success': True, 'path': c['path']}
    monkeypatch.setattr(bridge, 'request', run)
    return calls


def test_read_range_and_missing_file_never_falls_back_to_hub(member, monkeypatch):
    owner = member / 'owner.txt'; owner.write_text('OWNER SECRET')
    files = {'phone.txt': b'one\ntwo\nthree\n'}
    calls = exchange(monkeypatch, files)
    entry = profile.entry('self', 'read')
    assert bridge.execute(entry, {'path': 'phone.txt', 'start_line': 2, 'end_line': 2}) == 'two\n'
    assert bridge.execute(entry, {'path': 'phone.txt', 'limit': 0}) == ''
    result = bridge.execute(entry, {'path': 'phone.txt', 'tail': 1, 'offset': 0})
    assert result['success'] is False
    result = bridge.execute(entry, {'path': str(owner)})
    assert result['error'] == 'member_file_missing'
    assert owner.read_text() == 'OWNER SECRET'


def docx_bytes():
    from docx import Document
    doc = Document(); doc.add_paragraph('Hello {{name}}')
    stream = io.BytesIO(); doc.save(stream); return stream.getvalue()


def test_docx_fill_receipt_and_file_reference(member, monkeypatch):
    files = {'template.docx': docx_bytes()}; calls = exchange(monkeypatch, files)
    result = bridge.execute(profile.entry('self', 'fill'), {'path': 'template.docx', 'data': {'name': 'MEMBER A'}, 'output': 'done.docx'})
    assert result['saved'] is True and result['path'] == 'done.docx'
    assert [c['op'] for c in calls] == ['read', 'write']
    from docx import Document
    assert Document(io.BytesIO(files['done.docx'])).paragraphs[0].text == 'Hello MEMBER A'
    ref = result['files'][0]['ref']
    result = bridge.execute(profile.entry('self', 'read'), {'path': ref})
    assert 'MEMBER A' in str(result) and str(member) not in str(result)
    bad = bridge.execute(profile.entry('self', 'read'), {'path': 'member-file:other-turn'})
    assert bad['error_type'] == 'input'


def test_fill_denied_does_not_claim_delivery(member, monkeypatch):
    files = {'template.docx': docx_bytes()}; exchange(monkeypatch, files, reject=True)
    result = bridge.execute(profile.entry('self', 'fill'), {'path': 'template.docx', 'data': {'name': 'x'}})
    assert result == {'success': False, 'error': 'permission_denied'}
    assert list(files) == ['template.docx']
    assert not MR.current().get('files')


def test_pdf_pages_and_xlsx_real_parsers(member, monkeypatch):
    import fitz
    from openpyxl import Workbook
    doc=fitz.open()
    for text in ('PAGE ONE', 'PAGE TWO'):
        doc.new_page().insert_text((40,40),text)
    pdf=doc.tobytes(); doc.close()
    wb=Workbook(); wb.active.append(['Name','Count']); wb.active.append(['A',3])
    buf=io.BytesIO(); wb.save(buf)
    exchange(monkeypatch, {'a.pdf': pdf, 'a.xlsx': buf.getvalue()})
    entry=profile.entry('self','read')
    result=bridge.execute(entry, {'path':'a.pdf','pages':'2'})
    assert 'PAGE TWO' in result['text'] and 'PAGE ONE' not in result['text']
    result=bridge.execute(entry, {'path':'a.xlsx'})
    assert result['success'] and result['table']['columns']==['Name','Count'] and result['table']['rows']==[['A',3]],result


def test_changed_adapter_is_closed_and_body_binding_still_required(member, monkeypatch):
    monkeypatch.setattr(profile,'_package_open',lambda *a,**kw:True)
    monkeypatch.setattr(bridge,'connected',lambda d:False)
    assert profile.gate('self','read',{})['error_type']=='no_body'
    monkeypatch.setattr(profile,'fingerprint_ok',lambda *a,**kw:False)
    assert not profile.visible('self','read',{})
    assert profile.gate('self','read',{})['error_type']=='permission'


def test_cancelled_conversion_cannot_start(member, monkeypatch):
    MR.current()['cancel'].set()
    monkeypatch.setattr(bridge,'request',lambda *a,**kw:pytest.fail('device called after cancel'))
    assert bridge.execute(profile.entry('self','read'),{'path':'a.docx'})['error_type']=='cancelled'


def test_radio_and_android_declarations_bind_only_member_device(member, monkeypatch):
    sent=[]
    monkeypatch.setattr(bridge,'request',lambda c,**kw:sent.append(c) or {'success':True})
    bridge.execute(profile.entry('limbs','radio'),{'op':'play','stream_url':'https://example.test/live','mode':'host','device_id':'owner'})
    bridge.execute(profile.entry('limbs','android'),{'op':'tap','x':2,'y':3,'device_id':'owner'})
    bridge.execute(profile.entry('sense','here'),{'device_id':'owner'})
    assert sent==[{'op':'media','action':'play','url':'https://example.test/live'},
                  {'op':'accessibility','action':'tap','x':2,'y':3},{'op':'location'}]


def test_actual_ibl_inline_file_and_body_ref(member, monkeypatch):
    from member_runner import MemberRunner
    monkeypatch.setattr(profile, '_package_open', lambda *a, **kw: True)
    monkeypatch.setattr(bridge, 'connected', lambda d: True)
    sent=[]
    monkeypatch.setattr(bridge, 'request', lambda c, **kw: sent.append(c) or {'success': True, 'path': c.get('path')})
    runner=MemberRunner.__new__(MemberRunner); runner.project_path=member
    result=runner._member_tool('execute_ibl', {
        'edition':1, 'code':'[self:write]{path:"local.txt",content:"$file:0"}', 'files':['inline $text stays literal'],
        'files_from':['/owner/forbidden']})
    assert sent==[{'op':'write','path':'local.txt','content':'inline $text stays literal'}],result
    MR.current().setdefault('files',{})['member-file:owned']='local2.txt'
    result=runner._member_tool('execute_ibl', {'code':'[self:write]{path:"member-file:owned",content:"B"}'})
    assert sent[-1]['path']=='local2.txt',result


def test_file_like_body_text_is_not_a_reference(member):
    from member_files import resolve_references
    assert resolve_references({'content':'member-file:literal'})=={'content':'member-file:literal'}


def test_explicit_docx_format_still_checks_archive(member, monkeypatch):
    import zipfile
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as archive:
        archive.writestr('_rels/.rels','<Relationships><Relationship TargetMode="External" Type="image" Target="file:///owner/secret" /></Relationships>')
    exchange(monkeypatch, {'extensionless':buf.getvalue()})
    result=bridge.execute(profile.entry('self','read'),{'path':'extensionless','format':'docx'})
    assert result['error_type']=='conversion'


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, '-q']))
