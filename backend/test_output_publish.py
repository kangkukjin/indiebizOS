"""Collision policies and failed production preserve existing artifacts."""
import boot_paths  # noqa: F401
import importlib.util
import json
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
import pytest
from output_publish import publish_file

ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / 'data/packages/installed/tools/system_essentials'

def module(name):
    spec = importlib.util.spec_from_file_location('audit_' + name, PKG / (name + '.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_failed_writer_keeps_existing_and_cleans_temporary(tmp_path):
    dest = tmp_path / 'result.txt'
    dest.write_text('old')
    def fail(temp):
        temp.write_text('partial')
        raise OSError('disk failure')
    with pytest.raises(OSError):
        publish_file(dest, fail, if_exists='overwrite')
    assert dest.read_text() == 'old'
    assert list(tmp_path.iterdir()) == [dest]


@pytest.mark.parametrize('policy', ['fail', 'overwrite', 'rename'])
def test_completed_publication_policy(tmp_path, policy):
    dest = tmp_path / 'result.txt'
    dest.write_text('old')
    if policy == 'fail':
        with pytest.raises(FileExistsError):
            publish_file(dest, lambda p: pytest.fail('should not produce'), if_exists=policy)
    else:
        _, receipt = publish_file(dest, lambda p: p.write_text('new'), if_exists=policy)
        assert Path(receipt['path']).read_text() == 'new'
        assert dest.read_text() == ('new' if policy == 'overwrite' else 'old')
        assert receipt['replaced_existing'] == (policy == 'overwrite')


def test_rename_race_keeps_every_file(tmp_path):
    def write(i):
        return publish_file(tmp_path / 'x.txt', lambda p: p.write_text(str(i)), if_exists='rename')[1]
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(write, range(12)))
    assert len({r['path'] for r in results}) == 12
    assert {p.read_text() for p in tmp_path.iterdir()} == {str(i) for i in range(12)}


def test_fill_failure_and_explicit_policy(tmp_path, monkeypatch):
    office = module('office_ops')
    source, target = tmp_path / 'template.docx', tmp_path / 'filled.docx'
    from docx import Document
    doc = Document(); doc.add_paragraph('{{name}}'); doc.save(source)
    params = {'path': str(source), 'output': str(target), 'data': {'name': 'first'}}
    assert json.loads(office.fill_op(params, str(tmp_path)))['success']
    first = target.read_bytes()
    assert not json.loads(office.fill_op({**params, 'if_exists': 'fail'}, str(tmp_path)))['success']
    params['data'] = {'name': 'second'}
    out = json.loads(office.fill_op({**params, 'if_exists': 'rename'}, str(tmp_path)))
    assert out['renamed'] and target.read_bytes() == first
    assert Document(out['path']).paragraphs[0].text == 'second'
    def fail(*args):
        Path(args[2]).write_bytes(b'broken')
        raise OSError('failed creation')
    monkeypatch.setattr(office, '_fill_docx', fail)
    assert not json.loads(office.fill_op(params, str(tmp_path)))['success']
    assert target.read_bytes() == first


def test_copy_overwrite_failure_keeps_old_destination(tmp_path, monkeypatch):
    copy = module('copy_ops')
    source, dest = tmp_path / 'source', tmp_path / 'dest'
    source.write_text('new'); dest.write_text('old')
    out = copy.transfer_path(source, dest, if_exists='overwrite')
    assert dest.read_text() == 'new' and out['replaced'] == [str(dest)]
    import shutil
    def fail(a, b):
        Path(b).write_text('incomplete')
        raise OSError('copy failed')
    monkeypatch.setattr(shutil, 'copy2', fail)
    with pytest.raises(OSError):
        copy.transfer_path(source, dest, if_exists='overwrite')
    assert dest.read_text() == 'new' and source.read_text() == 'new'

def test_sheet_output_policy_preserves_source_and_publishes_actual_path(tmp_path):
    from test_document_sheet_transit_2026_09_15 import load, params
    import openpyxl
    source, target = tmp_path / 'source.xlsx', tmp_path / 'out.xlsx'
    wb = openpyxl.Workbook(); wb.active['A1'] = 7; wb.save(source)
    original = source.read_bytes()
    target.write_bytes(original)
    sheet = load('sheet_range_ops')
    snap = sheet.op_range(params(source, range='A1'))
    args = params(source, range='A1', values=[[12]], output=str(target), expected_sha256=snap['sha256'])
    with pytest.raises(ValueError, match='이미'):
        sheet.op_range_write(args)
    out = sheet.op_range_write({**args, 'if_exists': 'rename'})
    assert out['path'] != str(target) and target.read_bytes() == original
    assert openpyxl.load_workbook(out['path']).active['A1'].value == 12
    out = sheet.op_range_write({**args, 'if_exists': 'overwrite'})
    assert out['path'] == str(target) and openpyxl.load_workbook(target).active['A1'].value == 12
    assert source.read_bytes() == original


def test_copy_then_append_rerun_does_not_double_rows(tmp_path):
    import openpyxl
    copy, sheets = module('copy_ops'), module('sheet_ops')
    template, dest = tmp_path / 'template.xlsx', tmp_path / 'report.xlsx'
    wb = openpyxl.Workbook(); wb.active.append(['name', 'qty']); wb.save(template)
    for _ in range(2):
        result = copy.transfer_path(template, dest, if_exists='overwrite')
        appended = sheets.op_append({'path': result['path'], 'items': [{'name': 'A', 'qty': 1}, {'name': 'B', 'qty': 2}],
                                     '_project_path': str(tmp_path), '_path_guard': lambda *_: None})
        assert appended['success']
        assert openpyxl.load_workbook(dest).active.max_row == 3
    assert openpyxl.load_workbook(template).active.max_row == 1


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__]))
