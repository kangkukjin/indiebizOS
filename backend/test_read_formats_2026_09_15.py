"""[self:read] 확장 형식(hwp·hwpx·pptx·epub) 회귀 + 이미지변환·압축 등록 스크립트 계약 (2026-09-15).

hwp 는 OLE 컨테이너를 순수 파이썬으로 만들 수 없어 레코드 파서 단위로 고정한다(실파일 검증은 수동 완료).
"""
import importlib.util
import json
import struct
import subprocess
import sys
import zipfile
import zlib
from pathlib import Path

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / "data/packages/installed/tools/system_essentials"
SCRIPTS = ROOT / "data/scripts"


def load(name):
    spec = importlib.util.spec_from_file_location("feature_" + name, PKG / (name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.fixture(scope="module")
def dre():
    return load("doc_read_extra")


# ---------------------------------------------------------------- hwp 레코드 파서

def _rec(tag, level, payload: bytes) -> bytes:
    return struct.pack("<I", (tag & 0x3FF) | ((level & 0x3FF) << 10) | (len(payload) << 20)) + payload


def _para(text: str, level: int) -> bytes:
    return _rec(dre_mod._TAG_PARA_TEXT, level, (text + "\r").encode("utf-16-le"))


dre_mod = load("doc_read_extra")


def test_hwp_para_text_skips_controls():
    # 인라인 컨트롤(코드 9 탭 = 8 wchar)·확장 컨트롤(코드 11 = 8 wchar)·문단 끝 13
    codes = [ord("가"), 9, 0, 0, 0, 0, 0, 0, 0, ord("나"), 11] + [0] * 7 + [ord("다"), 13]
    payload = struct.pack(f"<{len(codes)}H", *codes)
    assert dre_mod._hwp_para_text(payload) == "가\t나다"


def test_hwp_section_blocks_paragraphs_and_table():
    T = dre_mod
    data = b"".join([
        _para("첫 문단", 0),
        _rec(T._TAG_CTRL_HEADER, 1, b" lbt" + b"\0" * 8),
        _rec(T._TAG_TABLE, 2, struct.pack("<IHH", 0, 2, 2) + b"\0" * 8),
        _rec(T._TAG_LIST_HEADER, 2, b"\0" * 6), _para("항목", 4),
        _rec(T._TAG_LIST_HEADER, 2, b"\0" * 6), _para("값", 4),
        _rec(T._TAG_LIST_HEADER, 2, b"\0" * 6), _para("납기", 4),
        _rec(T._TAG_LIST_HEADER, 2, b"\0" * 6), _para("9월", 4),
        _para("끝 문단", 0),
    ])
    blocks: list = []
    T._hwp_section_blocks(data, blocks)
    assert [b["type"] for b in blocks] == ["paragraph", "table", "paragraph"]
    assert blocks[1]["columns"] == ["항목", "값"] and blocks[1]["rows"] == [["납기", "9월"]]
    assert blocks[2]["text"] == "끝 문단"


def test_hwp_rejects_non_ole(dre, tmp_path):
    p = tmp_path / "x.hwp"
    p.write_bytes(b"not ole")
    r = json.loads(dre.read_hwp({"path": str(p)}, str(tmp_path)))
    assert r["success"] is False and "HWP v5" in r["error"]


# ---------------------------------------------------------------- hwpx

HWPX_SECTION = """<?xml version="1.0" encoding="UTF-8"?>
<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">
  <hp:p><hp:run><hp:t>제출 서류 안내</hp:t></hp:run></hp:p>
  <hp:p><hp:run><hp:tbl>
    <hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>서류</hp:t></hp:run></hp:p></hp:subList></hp:tc>
           <hp:tc><hp:subList><hp:p><hp:run><hp:t>부수</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr>
    <hp:tr><hp:tc><hp:subList><hp:p><hp:run><hp:t>신청서</hp:t></hp:run></hp:p></hp:subList></hp:tc>
           <hp:tc><hp:subList><hp:p><hp:run><hp:t>1</hp:t></hp:run></hp:p></hp:subList></hp:tc></hp:tr>
  </hp:tbl></hp:run><hp:run><hp:t>표 뒤 문장</hp:t></hp:run></hp:p>
  <hp:p><hp:run><hp:t>마감 </hp:t></hp:run><hp:run><hp:t>9월 30일</hp:t></hp:run></hp:p>
</hs:sec>"""


def _make_hwpx(path: Path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", HWPX_SECTION)
        z.writestr("BinData/image1.png", b"\x89PNG\r\n\x1a\n")


def test_hwpx_blocks_order_and_table(dre, tmp_path):
    p = tmp_path / "공문.hwpx"
    _make_hwpx(p)
    r = json.loads(dre.read_hwpx({"path": str(p)}, str(tmp_path)))
    assert r["success"], r
    assert [b["type"] for b in r["blocks"]] == ["paragraph", "table", "paragraph", "paragraph"]
    assert r["blocks"][1] == {"type": "table", "columns": ["서류", "부수"], "rows": [["신청서", "1"]]}
    assert r["blocks"][3]["text"] == "마감 9월 30일"
    assert r["metadata"]["total_images"] == 1 and (tmp_path / "공문_images/image1.png").exists()
    assert "제출 서류 안내" in r["text"] and "신청서\t1" in r["text"]


def test_hwpx_offset_limit_slices_text_but_keeps_blocks(dre, tmp_path):
    p = tmp_path / "a.hwpx"
    _make_hwpx(p)
    r = json.loads(dre.read_hwpx({"path": str(p), "offset": 2, "limit": 1, "extract_images": False}, str(tmp_path)))
    assert r["text"] == "표 뒤 문장"
    assert r["metadata"]["truncated"] is True and r["metadata"]["next_offset"] == 3
    assert len(r["blocks"]) == 4


def test_hwpx_rejects_zip_without_sections(dre, tmp_path):
    p = tmp_path / "b.hwpx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "x")
    r = json.loads(dre.read_hwpx({"path": str(p)}, str(tmp_path)))
    assert r["success"] is False and "section" in r["error"]


# ---------------------------------------------------------------- pptx

def _make_pptx(path: Path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "분기 실적"
    s1.placeholders[1].text_frame.text = "매출 12억"
    s1.placeholders[1].text_frame.add_paragraph().text = "영업이익 3억"
    s1.notes_slide.notes_text_frame.text = "발표 시 강조"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "표"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "항목", "값"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "매출", "12"
    s3 = prs.slides.add_slide(prs.slide_layouts[6])
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "셋째"
    prs.save(str(path))


def test_pptx_slides_headings_tables_notes(dre, tmp_path):
    p = tmp_path / "deck.pptx"
    _make_pptx(p)
    r = json.loads(dre.read_pptx({"path": str(p)}, str(tmp_path)))
    assert r["success"], r
    types = [b["type"] for b in r["blocks"]]
    assert types[0] == "heading" and r["blocks"][0]["text"] == "슬라이드 1 — 분기 실적"
    assert {"type": "paragraph", "text": "매출 12억"} in r["blocks"]
    assert {"type": "paragraph", "text": "[노트] 발표 시 강조"} in r["blocks"]
    table = next(b for b in r["blocks"] if b["type"] == "table")
    assert table["columns"] == ["항목", "값"] and table["rows"] == [["매출", "12"]]
    assert r["metadata"]["total_slides"] == 3
    # 제목 문단은 heading 에만(중복 없음)
    assert sum(1 for b in r["blocks"] if b.get("text") == "분기 실적") == 0


def test_pptx_pages_range(dre, tmp_path):
    p = tmp_path / "deck.pptx"
    _make_pptx(p)
    r = json.loads(dre.read_pptx({"path": str(p), "pages": "3"}, str(tmp_path)))
    assert r["metadata"]["extracted_slides"] == [3]
    assert [b["text"] for b in r["blocks"]] == ["슬라이드 3", "셋째"]
    bad = json.loads(dre.read_pptx({"path": str(p), "pages": "9"}, str(tmp_path)))
    assert bad["success"] is False and "pages" in bad["error"]


# ---------------------------------------------------------------- epub

def _make_epub(path: Path):
    opf = """<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/" version="3.0">
<metadata><dc:title>시험 책</dc:title><dc:creator>지은이</dc:creator></metadata>
<manifest><item id="c1" href="c1.xhtml" media-type="application/xhtml+xml"/><item id="c2" href="c2.xhtml" media-type="application/xhtml+xml"/>
<item id="css" href="s.css" media-type="text/css"/></manifest>
<spine><itemref idref="c2"/><itemref idref="c1"/></spine></package>"""
    c1 = "<html><head><title>첫째</title><style>p{}</style></head><body><h1>1장</h1><p>본문 &amp; 하나</p><p>본문  둘</p></body></html>"
    c2 = "<html><head><title>둘째 장</title></head><body><p>제목 없는 장의 문단</p></body></html>"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml", '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/c1.xhtml", c1)
        z.writestr("OEBPS/c2.xhtml", c2)
        z.writestr("OEBPS/s.css", "p{}")


def test_epub_spine_order_headings_and_entities(dre, tmp_path):
    p = tmp_path / "book.epub"
    _make_epub(p)
    r = json.loads(dre.read_epub({"path": str(p)}, str(tmp_path)))
    assert r["success"], r
    assert r["metadata"]["title"] == "시험 책" and r["metadata"]["creator"] == "지은이"
    assert r["metadata"]["total_chapters"] == 2
    texts = [b["text"] for b in r["blocks"]]
    # spine 순서(c2 → c1), 제목 없는 장은 <title> 로 heading 보충, 있는 장은 자기 h1
    assert texts == ["챕터 1 — 둘째 장", "제목 없는 장의 문단", "1장", "본문 & 하나", "본문 둘"]
    assert r["blocks"][2] == {"type": "heading", "level": 1, "text": "1장"}


def test_epub_pages_selects_chapter(dre, tmp_path):
    p = tmp_path / "book.epub"
    _make_epub(p)
    r = json.loads(dre.read_epub({"path": str(p), "pages": "2"}, str(tmp_path)))
    assert r["metadata"]["extracted_chapters"] == [2]
    assert r["blocks"][0]["text"] == "1장"


def test_directory_is_rejected(dre, tmp_path):
    r = json.loads(dre.read_epub({"path": str(tmp_path)}, str(tmp_path)))
    assert r["success"] is False and "폴더" in r["error"]


# ---------------------------------------------------------------- handler 분기(확장자 → read_<fmt>)

def test_handler_dispatches_new_extensions(tmp_path):
    handler = load("handler")
    p = tmp_path / "공문.hwpx"
    _make_hwpx(p)

    class Ctx:
        tool_name = "read_op"
        project_path = str(tmp_path)
        project_id = "test"
        agent_id = "test"
        task_id = "t"

    out = handler.execute({"path": str(p), "extract_images": False}, Ctx())
    r = json.loads(out) if isinstance(out, str) else out
    assert r["success"] and r["metadata"]["format"] == "hwpx"
    out = handler.execute({"path": str(p), "format": "epub"}, Ctx())
    r = json.loads(out) if isinstance(out, str) else out
    assert r["success"] is False  # 명시 format 이 우선 — hwpx 를 epub 로 읽으려 하면 정직 실패


# ---------------------------------------------------------------- 등록 스크립트: 이미지변환 · 압축

def _run(script: str, args: dict) -> dict:
    proc = subprocess.run([sys.executable, str(SCRIPTS / script)], input=json.dumps(args), capture_output=True, text=True, cwd=str(SCRIPTS))
    assert proc.stdout.strip(), proc.stderr
    return json.loads(proc.stdout)


def test_registry_has_both_scripts():
    import yaml
    reg = yaml.safe_load((SCRIPTS / "registry.yaml").read_text())
    for sid in ("이미지변환", "압축"):
        assert sid in reg and (SCRIPTS / reg[sid]["file"]).exists()
        assert reg[sid]["interpreter"] == "python"


def test_image_script_resize_convert_strip(tmp_path):
    from PIL import Image
    src = tmp_path / "big.jpg"
    im = Image.new("RGB", (3000, 2000), (200, 30, 30))
    ex = im.getexif(); ex[0x0110] = "TestCam"
    im.save(src, quality=95, exif=ex.tobytes())
    r = _run("이미지변환.py", {"paths": [str(src)], "max_side": 1200, "format": "webp", "strip_exif": True, "out_dir": str(tmp_path / "out")})
    assert r["count"] == 1 and not r["errors"], r
    it = r["items"][0]
    assert it["format"] == "WEBP" and (it["width"], it["height"]) == (1200, 800) and it["exif_stripped"] is True
    assert Path(it["out"]).exists() and src.exists()  # 원본 보존
    # 두 번째 호출은 덮어쓰기 거절
    r2 = _run("이미지변환.py", {"paths": [str(src)], "max_side": 1200, "format": "webp", "out_dir": str(tmp_path / "out")})
    assert r2["success"] is False and "overwrite" in r2["error"]
    # info 는 변환 없이
    r3 = _run("이미지변환.py", {"op": "info", "path": str(src)})
    assert r3["items"][0]["has_exif"] is True and r3["items"][0]["width"] == 3000


def test_image_script_max_kb_and_aspect(tmp_path):
    from PIL import Image
    import random
    src = tmp_path / "noise.png"
    rnd = random.Random(1)
    img = Image.new("RGB", (1600, 1600))
    img.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(1600 * 1600)])
    img.save(src)
    r = _run("이미지변환.py", {"path": str(src), "format": "jpg", "max_kb": 120, "aspect": "3:4", "out_dir": str(tmp_path / "o")})
    it = r["items"][0]
    assert it["bytes"] <= 120 * 1024 and abs(it["width"] / it["height"] - 0.75) < 0.02, it


def test_archive_script_pack_list_unpack_and_zip_slip(tmp_path):
    proj = tmp_path / "proj"
    (proj / "sub").mkdir(parents=True)
    (proj / "a.pdf").write_text("a"); (proj / "sub/b.txt").write_text("b"); (proj / ".DS_Store").write_text("x")
    r = _run("압축.py", {"op": "pack", "paths": [str(proj)], "output": str(tmp_path / "p.zip")})
    assert sorted(i["name"] for i in r["items"]) == ["proj/a.pdf", "proj/sub/b.txt"]
    lst = _run("압축.py", {"op": "list", "path": str(tmp_path / "p.zip")})
    assert lst["count"] == 2 and lst["total_size"] == 2
    un = _run("압축.py", {"op": "unpack", "path": str(tmp_path / "p.zip"), "out_dir": str(tmp_path / "un"), "include": ["*.pdf"]})
    assert [i["name"] for i in un["items"]] == ["proj/a.pdf"] and (tmp_path / "un/proj/a.pdf").exists()
    again = _run("압축.py", {"op": "unpack", "path": str(tmp_path / "p.zip"), "out_dir": str(tmp_path / "un")})
    assert again["skipped_existing"] == ["proj/a.pdf"] and [i["name"] for i in again["items"]] == ["proj/sub/b.txt"]
    evil = tmp_path / "evil.zip"
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("../../evil.txt", "x"); z.writestr("ok.txt", "y")
    ev = _run("압축.py", {"op": "unpack", "path": str(evil), "out_dir": str(tmp_path / "ev")})
    assert ev["rejected_unsafe"] == ["../../evil.txt"] and not (tmp_path / "evil.txt").exists()
    assert (tmp_path / "ev/ok.txt").exists()


def test_archive_script_bad_op_and_missing():
    r = _run("압축.py", {"op": "explode"})
    assert r["success"] is False
    r = _run("압축.py", {"op": "list", "path": "/nonexistent/x.zip"})
    assert r["success"] is False


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
