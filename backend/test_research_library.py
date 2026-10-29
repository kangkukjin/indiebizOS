"""불변 원문·재색인·자료 접근 인용·프로젝트 범위 계약."""
import boot_paths  # noqa: F401
import json
import pytest
from research_store import ResearchStore
from research_library import ResearchLibrary
from research_context import scope


@pytest.fixture
def lib(tmp_path, monkeypatch):
    import notebook_store
    monkeypatch.setattr(notebook_store, 'semantic_available', lambda: False)
    return ResearchLibrary(ResearchStore(tmp_path, 'library-one'))


def test_versions_keep_old_bytes_and_duplicate_capture_is_idempotent(lib):
    a = lib.capture('처음 확보한 교육 자료입니다.', title='논문', origin='https://example.org/p', run='r', call='c')
    again = lib.capture('처음 확보한 교육 자료입니다.', title='논문', origin='https://example.org/p', run='r', call='c')
    b = lib.capture('수정된 교육 자료입니다. 반론이 추가되었습니다.', title='논문', origin='https://example.org/p')
    assert a == again and a['source_ref'] == b['source_ref'] and a['version_ref'] != b['version_ref']
    assert '처음' in lib.read(a['version_ref'])['text']
    assert len(lib.versions(a['source_ref'])['items']) == 2


def test_file_snapshot_survives_removal_and_index_can_rebuild(lib, tmp_path):
    path = tmp_path / 'source.txt'
    path.write_text('교육 정책은 작은 학급의 교육 효과와 경제적 비용을 함께 평가합니다. ' * 20)
    receipt = lib.add_file(path)
    path.unlink()
    assert '교육 정책' in lib.read(receipt['version_ref'])['text']
    lib.index_pending()
    hits = lib.search('교육')['items']
    assert hits and hits[0]['version_ref'] == receipt['version_ref']
    lib.exclude(receipt['source_ref'])
    assert not lib.search('교육')['items']
    lib.exclude(receipt['source_ref'], False)
    assert lib.search('교육')['items']


def test_unsupported_preserves_original_and_reports_failure(lib, tmp_path):
    path = tmp_path / 'sample.bin'; path.write_bytes(b'\x00\x01\x02')
    receipt = lib.add_file(path)
    version = lib.version(receipt['version_ref'])
    assert lib.store.blob(version['blob']) == b'\x00\x01\x02'
    assert version['index_state'] == 'unsupported'
    assert version['error']


def test_citation_requires_actual_read_and_retains_extraction_revision(lib):
    receipt = lib.capture('교육 효과를 주장한 원문 자료입니다. 통제집단의 한계가 있습니다.', title='논문')
    vid = receipt['version_ref']
    marker = f'[자료:{vid}:문단 1]'
    assert not lib.citations(marker, 'run1')
    with scope({'project_uid': lib.store.uid, 'run_id': 'run1'}):
        content = lib.read(vid)
    loc = lib.paragraphs(vid)[0][0][0]
    marker = f'[자료:{vid}:{loc}]'
    cited = lib.citations(marker, 'run1')
    assert cited and cited[0]['extraction_revision'] == 1
    lib._extraction(vid, [(loc, '추출기 변경으로 달라진 문장입니다.')])
    assert '통제집단' in lib.citations(marker, 'run1')[0]['quote']
    assert '통제집단' in lib.read(vid, revision=1)['text']
    with pytest.raises(ValueError):
        lib.read(vid, revision=99)
    assert not lib.citations(marker, 'run2')


def test_project_scoped_search_and_references(lib, tmp_path):
    other = tmp_path / 'other'; other.mkdir()
    foreign = ResearchLibrary(ResearchStore(other, 'library-two'))
    a = lib.capture('교육 프로젝트에만 있는 비공개 연구자료입니다.', title='교육')
    b = foreign.capture('경제 프로젝트에만 있는 다른 비공개 연구자료입니다.', title='경제')
    lib.index_pending(); foreign.index_pending()
    assert not foreign.search('교육')['items']
    with pytest.raises(ValueError):
        foreign.read(a['version_ref'])
    assert foreign.search('경제')['items'][0]['version_ref'] == b['version_ref']


def test_search_later_paragraph_registers_exact_evidence_without_reread(lib):
    from research_questions import ResearchQuestions
    tid = lib.store.create_thread('문단별 근거')['id']
    local = ResearchLibrary(lib.store, tid)
    paragraphs = [('¶13', '배경 설명입니다.'), ('¶14', '조사 방법입니다.'),
                  ('¶15', '교역량이 증가했다. 다만 인과관계는 미확인이다.')]
    receipt = local.capture('원본', title='교역 연구', paragraphs=paragraphs)
    local.index_pending()
    run, _ = lib.store.begin_run(tid, '교역량 조사', 'paragraph-search')
    with scope({'project_uid': lib.store.uid, 'thread_id': tid, 'run_id': run}):
        hit = local.search('교역량이', 1)['items'][0]
    assert hit['loc'] == '¶15'
    assert hit['text'] == paragraphs[2][1]
    assert hit['citation'] == f"[자료:{receipt['version_ref']}:¶15]"
    ref = {'version_ref': hit['version_ref'], 'extraction_revision': hit['extraction_revision'],
           'locator': hit['loc'], 'quote': '교역량이 증가했다.'}
    questions = ResearchQuestions(lib.store, tid)
    saved = questions.apply([{'expected_revision': 0, 'text': '교역량 변화?',
                              'supporting': [ref]}], origin='agent', run=run)['items'][0]
    assert saved['supporting'][0]['observed_in'] == [{'run_id': run}]
    # 검색되지 않은 이웃 문단까지 읽었다고 인정하지 않는다.
    with pytest.raises(ValueError, match='읽은 뒤'):
        questions.apply([{'expected_revision': 0, 'text': '방법?', 'supporting': [
            {**ref, 'locator': '¶14', 'quote': paragraphs[1][1]}]}], origin='agent', run=run)


def test_citation_chunks_preserve_whitespace_and_partial_read_receipts(lib):
    from research_evidence import validate_reference
    from research_library import citation_chunks
    from notebook_store import CHUNK_HARD_MAX
    body = ('긴문단 자료입니다.\n\t공백과 줄바꿈을 보존한다.  ' * 90) + ' 마지막미열람표식'
    chunks = citation_chunks([('p.1', body), ('p.1', '같은 쪽의 다른 문단')])
    assert len(chunks) > 2
    assert all(text in body or text == '같은 쪽의 다른 문단' for _, text in chunks)
    assert all(0 < len(text) <= CHUNK_HARD_MAX for _, text in chunks)
    tid = lib.store.create_thread('긴 문단')['id']
    local = ResearchLibrary(lib.store, tid)
    receipt = local.capture(body, title='긴문단', paragraphs=[('p.1', body)])
    local.index_pending()
    run, _ = lib.store.begin_run(tid, '부분 조회', 'long-paragraph')
    with scope({'project_uid': lib.store.uid, 'thread_id': tid, 'run_id': run}):
        hit = local.search('긴문단', 1)['items'][0]
    ref = {'version_ref': receipt['version_ref'], 'extraction_revision': 1,
           'locator': hit['loc'], 'quote': hit['text']}
    with lib.store.connect() as db:
        assert validate_reference(lib.store, db, tid, ref, run=run)['quote'] == hit['text']
        unseen = '마지막미열람표식' if '마지막미열람표식' not in hit['text'] else body[:1000]
        with pytest.raises(ValueError, match='읽은 뒤'):
            validate_reference(lib.store, db, tid, {**ref, 'quote': unseen}, run=run)


def test_legacy_index_rebuild_preserves_extractions_and_old_access(lib, monkeypatch):
    import research_library
    from notebook_store import build_chunks
    paragraphs = [('¶1', '이전 배경.'), ('¶2', '도서관 예산의 실제 근거.')]
    receipt = lib.capture('불변 원본', title='도서관', paragraphs=paragraphs)
    # 배포 전 문단 합성 색인과 읽기 기록을 만든다.
    with monkeypatch.context() as patch:
        patch.setattr(research_library, 'citation_chunks', build_chunks)
        lib.index_pending()
        with scope({'project_uid': lib.store.uid, 'run_id': 'old-run'}):
            old = lib.search('도서관', 1)['items'][0]
    assert old['loc'] == '¶1' and paragraphs[1][1] in old['text']
    with lib.store.connect(True) as db:
        before = [tuple(r) for r in db.execute('SELECT * FROM research_extractions')]
        db.execute("DELETE FROM research_migrations WHERE name='citation-paragraph-chunks-v1'")
    # 준비 중인 옛 색인은 현행 판본의 검색 결과로 내보내지 않는다.
    pending = lib.search('도서관')
    assert pending['items'] == [] and pending['pending'] == 1
    lib.index_pending()
    new = lib.search('도서관', 1)['items'][0]
    assert new['loc'] == '¶2' and new['text'] == paragraphs[1][1]
    assert new['extraction_revision'] == 1
    assert lib.store.blob(lib.version(receipt['version_ref'])['blob']) == '불변 원본'.encode()
    with lib.store.connect() as db:
        assert [tuple(r) for r in db.execute('SELECT * FROM research_extractions')] == before
        assert db.execute('SELECT locator,quote FROM research_access WHERE run_id=?',
                          ('old-run',)).fetchone()[:] == ('¶1', old['text'])
    lib.prepare_index()
    assert lib.search('도서관')['pending'] == 0


def test_capture_precedes_reduction_and_save_failure_closes_gate(lib, monkeypatch):
    import research_capture
    import research_context
    monkeypatch.setattr(research_capture, 'schedule_index', lambda path: None)
    tid = lib.store.create_thread('교육')['id']
    run, _ = lib.store.begin_run(tid, '조사', '1')
    ctx = dict(project_uid=lib.store.uid, epoch=lib.store.epoch, path=str(lib.store.path), thread_id=tid, run_id=run, mode='research')
    payload = {'text': '확보한 전체 본문 문장입니다. ' * 2000, 'token': 'secret-value'}
    with scope(ctx):
        result = research_context.observe({'router':'handler', 'tool':'read'}, {}, payload)
        assert result['research_materials']
    sources = lib.sources()['items']
    v = lib.version(sources[0]['version_ref'])
    raw = lib.store.blob(v['blob']).decode()
    assert 'secret-value' not in raw and payload['text'] in raw
    monkeypatch.setattr(research_context, '_capture', lambda *args: (_ for _ in ()).throw(OSError('disk full')))
    with scope(ctx):
        with pytest.raises(OSError): research_context.observe({'router':'handler'}, {}, payload)
        with pytest.raises(RuntimeError): research_context.before({'router':'handler'}, {})


def test_library_mode_blocks_external_dispatch(lib):
    from research_context import before
    with scope({'mode':'library'}):
        with pytest.raises(PermissionError): before({'router':'handler'}, {'query':'web'})
        before({'research_local':True}, {'op':'search'})
        with pytest.raises(PermissionError): before({'research_local':True}, {'op':'add'})


def test_derived_script_selection_reuses_original_material_without_new_json_source(lib, monkeypatch):
    import research_capture
    monkeypatch.setattr(research_capture, 'schedule_index', lambda path: None)
    tid = lib.store.create_thread('발췌')['id']
    ctx = dict(project_uid=lib.store.uid, epoch=lib.store.epoch, path=str(lib.store.path),
               thread_id=tid, run_id='run', mode='research')
    with scope(ctx):
        ref = lib.capture('원문의 첫 문단\n\n선택할 둘째 문단', title='원문', thread=tid,
                          paragraphs=[('¶1', '원문의 첫 문단'), ('¶2', '선택할 둘째 문단')],
                          acquired_result={'items': [{'text': '원문의 첫 문단', 'paragraph_index': 1},
                                                     {'text': '선택할 둘째 문단', 'paragraph_index': 2}]})
        selected = {'items': [{'text': '선택할 둘째 문단', 'paragraph_index': 2}],
                    'research_materials': [ref]}
        wrapper = {'success': True, 'script_protocol': 'registered-json/1', 'value': selected}
        out = research_capture.capture(ctx, {'tool': 'script_op'}, {}, wrapper, 'script')
        assert out == wrapper
        protected = {**wrapper, 'token': 'private-token'}
        shown = research_capture.capture(ctx, {'tool': 'script_op'}, {}, protected, 'sanitized')
        assert shown['token'] == '[인증정보 제외]'
        with lib.store.connect() as db:
            assert db.execute('SELECT count(*) FROM research_versions').fetchone()[0] == 1
            assert db.execute('SELECT count(*) FROM research_access').fetchone()[0] == 0
        # 영수증을 붙인 새 문장은 원본으로 위장하거나 버리지 않는다.
        selected['items'] = [{'text': '새로 생성한 판단 문장'}]
        out = research_capture.capture(ctx, {'tool': 'script_op'}, {}, wrapper, 'changed')
        assert out['research_materials'][0]['version_ref'] != ref['version_ref']
        version = lib.version(out['research_materials'][0]['version_ref'])
        assert lib.store.blob(version['blob']).decode() == '새로 생성한 판단 문장'


def test_citation_review_binds_current_run_quotes_without_inventing_reads(lib):
    receipt = lib.capture('원문', title='자료', paragraphs=[('¶27', '수요는 788이다.')])
    vid = receipt['version_ref']
    with scope({'project_uid': lib.store.uid, 'run_id': 'current'}):
        lib.access(vid, 1, '¶27', '수요는 788이다.')
    report = lib.citation_review(f'[자료:{vid}:¶27] [자료:{vid}:¶1] [자료:missing:¶1]', 'current')
    assert report['valid'][0]['quote'] == '수요는 788이다.'
    assert len(report['invalid']) == 2
    assert report['invalid'][0]['read_locations'] == ['¶27']
    assert lib.citation_review(f'[자료:{vid}:¶27]', 'other')['invalid']
    with lib.store.connect() as db:
        assert db.execute('SELECT count(*) FROM research_access').fetchone()[0] == 1


def test_citation_review_bounds_large_composite_locations(lib):
    vid = lib.capture('본문', title='자료')['version_ref']
    with scope({'project_uid': lib.store.uid, 'run_id': 'run'}):
        for i in range(10):
            lib.access(vid, 1, f'¶{i}', '긴 문장 ' * 2000)
    report = lib.citation_review(' '.join(f'[자료:{vid}:¶{i}]' for i in range(10)), 'run')
    assert report['evidence_chars'] == 16000
    assert len(report['valid']) == 10 and not report['invalid']
    assert all(row['truncated'] for row in report['valid'])


def test_legacy_script_selection_reindexes_without_rewriting_old_citations(lib):
    raw = json.dumps({'script_protocol': 'registered-json/1', 'log': '/execution/log',
                      'value': {'text': '선택하지 않은 전문', 'research_materials': [{'version_ref': 'old'}],
                                'items': [{'text': '보존한 첫 문단'}, {'text': '보존한 둘째 문단'}]}}, ensure_ascii=False)
    vid = lib.capture(raw, title='script', kind='tool:script_op:run', mime='application/json',
                      paragraphs=[('¶1', raw)])['version_ref']
    with scope({'project_uid': lib.store.uid, 'run_id': 'old-run'}):
        lib.access(vid, 1, '¶1', raw[:40])
    lib.prepare_index()
    assert lib.paragraphs(vid) == ([['¶1', '보존한 첫 문단'], ['¶2', '보존한 둘째 문단']], 2)
    assert lib.paragraphs(vid, 1)[0] == [['¶1', raw]]
    assert lib.store.blob(lib.version(vid)['blob']).decode() == raw
    assert lib.citations(f'[자료:{vid}:¶1]', 'old-run')[0]['extraction_revision'] == 1
    lib.prepare_index()
    assert lib.paragraphs(vid)[1] == 2


def test_deleted_derived_index_rebuilds_from_saved_bytes(lib):
    receipt = lib.capture('교육 자료의 원문은 색인 삭제와 별개로 남아 있어야 합니다. ' * 10, title='보존')
    lib.index_pending()
    assert lib.search('교육')['items']
    for suffix in ('', '-wal', '-shm'):
        lib.index.db_path.with_name(lib.index.db_path.name + suffix).unlink(missing_ok=True)
    lib.index_pending()
    assert lib.search('교육')['items'][0]['version_ref'] == receipt['version_ref']


def test_capture_copies_expiring_source_and_ignores_display_changes(lib, monkeypatch):
    import research_capture
    from common.spill import spill_write
    monkeypatch.setattr(research_capture, 'schedule_index', lambda *a: None)
    tid = lib.store.create_thread('교육')['id']; run, _ = lib.store.begin_run(tid, '조사', '1')
    ctx=dict(project_uid=lib.store.uid, epoch=lib.store.epoch, path=str(lib.store.path), thread_id=tid, run_id=run)
    full='모델에는 일부만 보여도 획득한 교육 원문 전체는 남깁니다. ' * 100
    ref=spill_write(json.dumps({'text':full},ensure_ascii=False))['ref']
    config={'router':'handler','tool':'crawl','research_coverage':'full'}
    result={'text':full[:100],'source_ref':ref,'_display':{'max_chars':10}}
    first=research_capture.capture(ctx,config,{'url':'https://example.org/p'},result,'1')
    result['_display']={'max_chars':20}
    second=research_capture.capture(ctx,config,{'url':'https://example.org/p'},result,'2')
    assert first['research_materials'][0]['version_ref']==second['research_materials'][0]['version_ref']
    assert len(first['research_materials'])==2
    from pathlib import Path
    Path(ref['path']).unlink()
    saved=lib.version(first['research_materials'][1]['version_ref'])
    assert full in lib.store.blob(saved['blob']).decode()




def test_thread_library_filters_before_ranking_versions_and_exclusion(lib):
    a = lib.store.create_thread('교육')['id']
    b = lib.store.create_thread('경제')['id']
    first, second = ResearchLibrary(lib.store, a), ResearchLibrary(lib.store, b)
    old = first.capture('교육 공동 자료의 첫 판본입니다. ' * 20, title='공동 주소', origin='https://example.org/shared')
    new = second.capture('교육 공동 자료의 수정 판본입니다. ' * 20, title='공동 주소', origin='https://example.org/shared')
    for i in range(40):
        second.capture('교육 교육 교육 다른 대화의 자료입니다. ' * 20, title=f'외부 {i}')
    lib.index_pending(limit=100)
    assert first.sources()['items'][0]['version_ref'] == old['version_ref']
    assert [v['version_ref'] for v in first.versions(old['source_ref'])['items']] == [old['version_ref']]
    assert first.search('교육', 1)['items'][0]['version_ref'] == old['version_ref']
    with pytest.raises(ValueError): first.read(new['version_ref'])
    first.exclude(old['source_ref'])
    assert not first.search('교육')['items']
    assert second.search('교육')['items']
    assert not second.version(new['version_ref'])['excluded']
    with scope({'project_uid':lib.store.uid,'thread_id':a,'run_id':'scope'}):
        assert not ResearchLibrary(lib.store).sources()['items']
        with pytest.raises(PermissionError): ResearchLibrary(lib.store, b)
        with pytest.raises(PermissionError): first.capture('혼입', title='거부', thread=b)


def test_docx_uses_existing_converter_without_chat_truncation(lib, tmp_path):
    from docx import Document
    document = Document(); document.add_paragraph('문서 첨부로 전달하는 교육 자료입니다. ' * 7000)
    path = tmp_path / '보고서.docx'; document.save(path)
    receipt = lib.add_file(path)
    paragraphs, _ = lib.paragraphs(receipt['version_ref'])
    assert sum(len(t) for _, t in paragraphs) > 100000
    assert lib.version(receipt['version_ref'])['index_state'] == 'pending'



def test_web_snapshot_indexes_body_and_preserves_raw_and_old_citations(lib, tmp_path):
    from notebook_store import _paragraphs_with_loc
    body = '교육 자료의 깨끗한 기사 본문입니다. ' * 100
    payload = {'text': body, 'url': 'https://example.org/article',
               '_page_structure': {'html': '<nav>메뉴전용잡음</nav>' * 100}}
    raw = json.dumps(payload, ensure_ascii=False)
    path = tmp_path / 'crawl.json'
    path.write_text(raw)
    # 이전 버전이 JSON 전체를 추출한 경우를 재현한다.
    receipt = lib.capture(raw, title='crawl.json', origin=str(path), kind='file', mime='application/json',
                          paragraphs=_paragraphs_with_loc(raw))
    old_version = receipt['version_ref']
    lib.index_pending()
    assert lib.store.blob(lib.version(old_version)['blob']).decode() == raw
    assert '메뉴전용잡음' in lib.read(old_version, revision=1)['text']
    current = lib.read(old_version)
    assert current['extraction_revision'] == 2 and '기사 본문' in current['text']
    assert '메뉴전용잡음' not in current['text']
    lib.prepare_index()
    assert lib.read(old_version)['extraction_revision'] == 2  # 재시작/중복 실행 안전
    assert not lib.search('메뉴전용잡음')['items']
    assert lib.search('기사 본문')['items']
    # 새 수집은 처음부터 본문을 추출한다.
    new_path = tmp_path / 'new.json'; new_path.write_text(raw + ' ')
    new = lib.add_file(new_path)
    assert '메뉴전용잡음' not in lib.read(new['version_ref'])['text']
    ordinary = tmp_path / 'user.json'
    ordinary.write_text(json.dumps({'text': '사용자 본문', 'other': '다른 중요 필드'}, ensure_ascii=False))
    uploaded = lib.add_file(ordinary)
    assert '다른 중요 필드' in lib.read(uploaded['version_ref'])['text']



def test_explicit_add_in_agent_turn_inherits_capture_run(lib, tmp_path):
    tid = lib.store.create_thread('교육')['id']
    run, _ = lib.store.begin_run(tid, '자료 보존', 'request')
    path = tmp_path / 'article.txt'; path.write_text('교육 기사 본문')
    with scope({'project_uid': lib.store.uid, 'thread_id': tid, 'run_id': run}):
        ResearchLibrary(lib.store).add_file(path)
    ResearchLibrary(lib.store, tid).capture('사용자의 수동 추가', title='수동')
    with lib.store.connect() as db:
        rows = db.execute('SELECT run_id FROM research_captures ORDER BY rowid').fetchall()
    assert [r[0] for r in rows] == [run, 'manual']


def test_small_pages_preserve_text_revision_and_only_visible_evidence(lib):
    body = '가나다라마바사0123456789' * 100
    vid = lib.capture(body, title='긴 문단')['version_ref']
    paras, revision = lib.paragraphs(vid)
    loc = paras[0][0]
    expected = '\n\n'.join(f'[자료:{vid}:{place}] {text}' for place, text in paras)
    with scope({'project_uid': lib.store.uid, 'run_id': 'small-page'}):
        first = lib.read(vid, limit=80)
    citation = lib.citations(f'[자료:{vid}:{loc}]', 'small-page')[0]
    assert citation['quote'] in first['text']
    assert 0 < len(citation['quote']) < len(body)
    assert citation['extraction_revision'] == revision
    lib._extraction(vid, [(loc, '바뀐 추출 내용')])
    parts = [first['text']]
    page = first
    while page['next_cursor'] is not None:
        params = page['pages'][0]
        assert params['max_chars'] == 80 and params['extraction_revision'] == revision
        page = lib.read(vid, offset=int(params['cursor']), limit=params['max_chars'],
                        revision=params['extraction_revision'])
        assert len(page['text']) <= 80
        parts.append(page['text'])
    assert ''.join(parts) == expected
    assert page['total_chars'] == len(expected)
    assert page['pages'] == []
    assert lib.read(vid)['extraction_revision'] == revision + 1


@pytest.mark.parametrize('limit', [0, -1, 12001, True, 3.5, '3500'])
def test_read_rejects_invalid_character_budget(lib, limit):
    vid = lib.capture('원문', title='검사')['version_ref']
    with pytest.raises(ValueError, match='max_chars'):
        lib.read(vid, limit=limit)


def test_read_default_page_and_end_cursor(lib):
    vid = lib.capture('본문' * 7000, title='기본 분량')['version_ref']
    page = lib.read(vid)
    assert len(page['text']) == 12000 and page['next_cursor'] == '12000'
    assert page['max_chars'] == 12000
    end = lib.read(vid, offset=page['total_chars'])
    assert end['text'] == '' and end['next_cursor'] is None
    with pytest.raises(ValueError, match='cursor'):
        lib.read(vid, offset=-1)


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
