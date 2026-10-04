"""원자료 크기를 표시 예산과 분리: 실제 파일·IBL·보관 참조·모델 입력 경계."""
import base64
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import boot_paths  # noqa: F401
import pytest

from test_structured_file_read import reader  # noqa: F401
from test_ibl_general_capabilities import boundary  # noqa: F401
from test_repair_capability_parity import repair_handler, setup  # noqa: F401
from test_member_files_2026_09_14 import member  # noqa: F401
from thread_context import repair_workspace_scope

ROOT = Path(__file__).resolve().parents[1]


def module(package, name):
    path = ROOT / 'data/packages/installed/tools' / package / (name + '.py')
    spec = importlib.util.spec_from_file_location('capacity_' + name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def make_file(path):
    if path.suffix == '.xlsx':
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['id', 'amount', 'label'])
        for i in range(250):
            sheet.append([i, 1, '마지막' if i == 249 else '내용' * 100])
        book.save(path)
        book.close()
    elif path.suffix == '.docx':
        from docx import Document
        doc = Document()
        for i in range(301):
            doc.add_paragraph('마지막' if i == 300 else '내용' * 100)
        doc.save(path)
    else:
        import zipfile
        paragraphs = ''.join('<p><t>' + ('마지막' if i == 300 else '내용' * 100) + '</t></p>'
                             for i in range(301))
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('Contents/section0.xml', '<sec>' + paragraphs + '</sec>')


@pytest.mark.parametrize('option,count', [('', 250), (',max_rows:0', 250), (',max_rows:11', 10)])
def test_xlsx_aggregate_reads_full_source_or_explicit_rows(reader, tmp_path, option, count):
    make_file(tmp_path / 'source.xlsx')
    _, result = reader('$d=[self:read]{path:"source.xlsx"' + option + '}; '
                       'return {count:len($d.data.table.rows),last:$d.data.table.rows[-1][0]}')
    assert result['success'] and result['source_complete'], result.get('error')
    assert result['value'] == {'count': count, 'last': count - 1}


@pytest.mark.parametrize('fmt', ['xlsx', 'docx', 'hwpx'])
def test_office_display_is_bounded_but_reference_keeps_tail(reader, tmp_path, fmt):
    path = tmp_path / ('source.' + fmt)
    make_file(path)
    raw, result = reader('return [self:read]{path:' + json.dumps(str(path)) + '}')
    assert result['success'] and result['source_complete'], result.get('error')
    assert len(raw) < 20_000
    path.unlink()
    _, recovered = reader('return $d.text', inputs={'d': {'$ref': result['result_ref']['id']}})
    # Returning the full string previews again, so resolve the stored typed value directly.
    from model_result_view import resolve_input_refs
    values, _ = resolve_input_refs({'d': {'$ref': result['result_ref']['id']}})
    assert '마지막' in values['d']['text']
    assert len(values['d']['text']) > 20_000
    assert recovered['success'] and recovered['source_complete']


@pytest.mark.parametrize('option', [',limit:2', ',max_blocks:2'])
def test_document_explicit_selection_still_selects_text(reader, tmp_path, option):
    make_file(tmp_path / 'source.docx')
    _, result = reader('$d=[self:read]{path:"source.docx"' + option + '}; return $d.text')
    assert result['success'] and result['source_complete']
    assert result['value'] == ('내용' * 100 + '\n\n' + '내용' * 100)


@pytest.mark.parametrize('fmt', ['docx', 'xlsx'])
def test_member_and_repair_office_reads_keep_tail(repair_handler, setup, tmp_path, fmt):
    root, candidate, _, _ = setup
    name = 'source.' + fmt
    make_file(candidate / name)
    with repair_workspace_scope(str(candidate)):
        raw = repair_handler.execute({'path': str(root / name), 'blocks': True},
            SimpleNamespace(tool_name='read_op', project_path=str(root), agent_id='owner'))
    assert '마지막' in json.loads(raw)['text']
    assert not (root / name).exists()
    member = module('system_essentials', 'member_documents')
    body = (candidate / name).read_bytes()
    out = member.read_document({'path': name}, {},
        lambda _: {'success': True, 'content': base64.b64encode(body).decode()}, tmp_path)
    assert '마지막' in out['text']
    assert not out.get('truncated') and not out.get('metadata', {}).get('truncated')


@pytest.mark.parametrize('param', ['max_rows', 'max_blocks'])
@pytest.mark.parametrize('value', [-1, 1.5, True, 'bad'])
def test_invalid_range_does_not_become_complete_default(reader, tmp_path, param, value):
    fmt = 'xlsx' if param == 'max_rows' else 'docx'
    make_file(tmp_path / ('source.' + fmt))
    _, result = reader('return [self:read]{path:"source.' + fmt + '",' + param + ':' + json.dumps(value) + '}')
    assert not result['success']


@pytest.mark.parametrize('limit', [None, 100, 70000])
def test_member_crawl_preserves_source_and_reference(boundary, monkeypatch, limit):
    from model_result_view import project_result
    web = module('web', 'member_web')
    body = '본문' * 40000 + '마지막'
    monkeypatch.setattr(web, 'fetch_public', lambda url: (url, ('<body>' + body + '</body>').encode(), 'text/html'))
    args = {'url': 'https://example.org'}
    if limit is not None:
        args['max_length'] = limit
    result = web.execute('crawl_website', args)
    assert result['success'] and not result.get('truncated')
    assert result['items'][0]['text'] == body
    projected = project_result(result)
    assert len(json.dumps(projected, ensure_ascii=False)) < (limit or 60000) + 5000
    stored = json.loads(boundary.read_evidence(projected['result_ref']['id'], 0, None)['text'])
    assert stored['items'][0]['text'] == body
    assert result['_display']['max_chars'] == (60000 if limit is None else limit)


@pytest.mark.parametrize('limit', [0, -1, True, 1.5, '70000'])
def test_member_invalid_display_budget_fails_before_network(monkeypatch, limit):
    web = module('web', 'member_web')
    monkeypatch.setattr(web, 'fetch_public', lambda _: pytest.fail('invalid request fetched'))
    assert not web.execute('crawl_website', {'url': 'https://example.org', 'max_length': limit})['success']


def test_document_author_gets_full_input_and_reports_provider_failure(monkeypatch, tmp_path):
    builder = module('data-ops', 'doc_build')
    content = '자료' * 20000 + '마지막'
    seen = []
    def author(prompt, *args, **kwargs):
        seen.append(prompt)
        return {'blocks': [{'type': 'paragraph', 'text': '마지막'}]}, None
    monkeypatch.setattr('oneshot_facade.oneshot_json', author)
    result = json.loads(builder.structure_document({'content': content}, str(tmp_path)))
    assert result['success'] and content in seen[0]
    def fail(*args, **kwargs):
        raise ValueError('provider context exceeded')
    monkeypatch.setattr('oneshot_facade.oneshot_json', fail)
    failed = json.loads(builder.structure_document({'content': content}, str(tmp_path)))
    assert not failed['success'] and 'context exceeded' in failed['error']


@pytest.mark.parametrize('name', ['slide_native', 'slide_image'])
def test_slide_author_gets_full_input_and_failure_is_not_success(monkeypatch, tmp_path, name):
    slide = module('media_producer', name)
    content = '자료' * 20000 + '마지막'
    seen = []
    def author(ai, prompt, **kwargs):
        seen.append(prompt)
        raise ValueError('provider context exceeded')
    monkeypatch.setattr('consciousness_agent.call_oneshot_provider', author)
    args = {'content': content, 'instruction': '마지막 사실로 슬라이드 작성'}
    if name == 'slide_native':
        monkeypatch.setattr(slide, '_get_author_ai', lambda: object())
        out = slide.create_native_slide(args, str(tmp_path))
    else:
        monkeypatch.setattr(slide, '_get_ai', lambda: object())
        styles = slide._load_styles()
        style = next(s for s in styles.STYLES if styles.is_image_style(s))
        out = slide.create_image_slide(args, str(tmp_path), style)
    assert content in seen[0]
    result = json.loads(out)
    assert not result['success'] and 'context exceeded' in (result.get('error') or result['message'])


def test_member_crawl_ibl_complete_text_and_last_character(reader, member, monkeypatch):
    import member_profile
    from tool_loader import load_tool_handler
    web = module('web', 'member_web')
    body = '본문' * 40000 + '마지막'
    monkeypatch.setattr(web, 'fetch_public', lambda url: (url, ('<body>' + body + '</body>').encode(), 'text/html'))
    handler = load_tool_handler('crawl_website')
    original = handler.load_module
    monkeypatch.setattr(handler, 'load_module', lambda name: web if name == 'member_web' else original(name))
    monkeypatch.setattr(member_profile, '_package_open', lambda *args, **kwargs: True)
    _, result = reader('$d=[sense:crawl]{url:"https://example.org",max_length:100}; '
                       'return {length:len($d.text),last:$d.text[-1]}')
    assert result['success'] and result['source_complete'], result.get('error')
    assert result['value'] == {'length': len(body), 'last': '막'}


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
