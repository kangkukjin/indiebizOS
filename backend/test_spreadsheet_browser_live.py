"""실제 Chromium + 격자 엔진(Univer) + /spreadsheets 경로 — 스프레드시트 앱 인수 시험(2026-10-07, 엑셀 얼굴).
사무 엔진(ONLYOFFICE)은 필요 없다. 플러그인 경로의 인수는 test_spreadsheet_engine_live 가 따로 본다.

흐름: 빈 통합문서(랜딩) → 이름 상자·수식 입력줄로 값·수식 → 자동 초안 → 저장(파일에 값·수식·캐시)
      → `[self:workspace]` 가 접수한 스냅샷·제안 적용을 격자가 소비 → 셀이 바뀌고 영수증이 닫힌다
      → AI 작업창: 제안 미리보기 → 반영(수식) → 저장.
"""
import importlib.util
import socket
import sys
import threading
import time
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
for layer in ("datastore", "services", "surface"):
    sys.path.insert(0, str(HERE / layer))
import api_spreadsheets
from spreadsheet_workspace import SpreadsheetWorkspace


@pytest.mark.system
def test_grid_app_edit_save_queue_and_ai(tmp_path, monkeypatch):
    import uvicorn
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright, expect
    from runtime_utils import get_base_path
    import document_office

    frontend = ROOT / "frontend"
    dependencies = frontend / "node_modules"
    installed = Path(get_base_path()) / "frontend/node_modules"
    assert installed.is_dir(), "프론트엔드 의존성을 먼저 설치해야 합니다"
    dependencies.mkdir(exist_ok=True)
    if dependencies.resolve() != installed.resolve():
        for package in installed.iterdir():
            target = dependencies / package.name
            if not target.exists():
                target.symlink_to(package, target_is_directory=package.is_dir())
    spec = importlib.util.spec_from_file_location("sheet_web_check", ROOT / "data/scripts/webapp_work.py")
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    plan = runner.plan(frontend, ["build"])
    built = runner.run_check(frontend, "build", 240, plan["fingerprint"])
    assert built["ok"], built.get("output_tail", str(built))

    monkeypatch.setattr(document_office, "available", lambda: False)   # 사무 엔진 없이 — 격자가 기본
    workspace = SpreadsheetWorkspace(tmp_path / "office")
    monkeypatch.setattr(api_spreadsheets, "service", lambda: workspace)
    from test_spreadsheet_app_support import mount_spreadsheet_app, set_cell
    app = FastAPI()
    app.include_router(api_spreadsheets.router)
    # AI 대역: 둘째 열에 첫째 열의 10% 수식 — 그대로인 셀은 그대로(수식이면 수식째) 돌려준다.
    mount_spreadsheet_app(app, workspace, tmp_path / "sheets",
                          ai=lambda dock, table: [[row[0], f"=B{i + 2}*0.1"] + list(row[2:]) for i, row in enumerate(table)])

    @app.get("/launcher/auth/session")
    def session():
        return {"authenticated": True, "external": False}

    @app.get("/health")
    def health():
        return {"status": "ok"}
    html = (frontend / "dist/index.html").read_text()
    html = html.replace("<html", '<html data-indiebiz-surface="remote"', 1)

    @app.get("/")
    def index():
        return HTMLResponse(html)

    app.mount("/assets", StaticFiles(directory=frontend / "dist/assets"))
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    server = uvicorn.Server(uvicorn.Config(app, log_level="error"))
    thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
    thread.start()
    try:
        deadline = time.monotonic() + 10
        while not server.started and time.monotonic() < deadline:
            time.sleep(0.02)
        assert server.started
        book = tmp_path / "sheets/무제.xlsx"
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                page = browser.new_page(viewport={"width": 1440, "height": 1000})
                failures = []
                page.on("pageerror", lambda error: failures.append(str(error)))
                page.goto(f"http://127.0.0.1:{port}/#/spreadsheets")
                status = page.get_by_role("status").first
                try:
                    expect(page.get_by_test_id("excel-frame")).to_be_visible(timeout=20000)
                    expect(status).to_contain_text("원본을 읽었습니다", timeout=20000)
                except Exception as exc:
                    raise AssertionError({"page_errors": failures, "body": page.locator("body").inner_text()[:2000]}) from exc
                assert book.exists()
                doc = workspace.open(book)["document"]
                # ── 입력 → 자동 초안(원본은 그대로) → 저장 ──
                set_cell(page, "A1", "매출"); set_cell(page, "B2", "1200"); set_cell(page, "B3", "800"); set_cell(page, "B4", "=SUM(B2:B3)")
                expect(status).to_contain_text("작업 저장됨", timeout=10000)
                assert load_workbook(book).active["B2"].value is None
                page.get_by_role("button", name="저장", exact=True).first.click()
                expect(status).to_contain_text("저장됨 · 원본 파일 기록 확인", timeout=10000)
                saved = load_workbook(book).active
                assert saved["A1"].value == "매출" and saved["B2"].value == 1200 and saved["B4"].value == "=SUM(B2:B3)"
                assert load_workbook(book, data_only=True).active["B4"].value == 2000   # 격자가 계산한 캐시
                # ── `[self:workspace]` 가 접수한 스냅샷·제안 적용을 격자가 소비한다 ──
                def finished(op_id, timeout=20):
                    end = time.monotonic() + timeout
                    while time.monotonic() < end:
                        state = workspace.operation_status(doc["id"], op_id)
                        if state["completed"] or state["status"] in ("failed", "interrupted"):
                            return state
                        time.sleep(0.3)
                    raise AssertionError("편집창이 접수한 작업을 제때 소비하지 않았습니다: " + op_id)
                receipt = workspace.request_snapshot(doc["id"], "snap-1")
                state = finished(receipt["operation_id"])
                snapshot = state["result"]["snapshot"]
                assert state["status"] == "completed" and snapshot["calc_status"] == "fresh" and snapshot["unsaved"] is False
                read = workspace.read(doc["id"], snapshot["id"], "1", "B2:B4")
                assert [c["effective_value"] for c in read["items"]] == [1200, 800, 2000]
                proposal = workspace.propose(doc["id"], snapshot["id"], "1", "B2", [[7]])
                detail = workspace.detail(doc["id"])
                queued = workspace.apply(doc["id"], proposal["id"], session_id=detail["session"]["id"], client_id=detail["session"]["client_id"],
                                         epoch=detail["session"]["engine_epoch"], expected=detail["session"]["session_revision"], operation_id="apply-1")
                state = finished(queued["operation_id"])
                assert state["status"] == "completed" and state["result"]["applied"] is True, state
                set_cell(page, "B2", "7")  # 같은 값 — 이름 상자로 B2 를 고르고 수식 입력줄이 7 을 보인다
                expect(page.get_by_label("수식 입력줄")).to_have_value("7")
                page.get_by_role("button", name="저장", exact=True).first.click()
                expect(status).to_contain_text("저장됨 · 원본 파일 기록 확인", timeout=10000)
                assert load_workbook(book, data_only=True).active["B4"].value == 807
                # ── AI 작업창: 제안 미리보기 → 반영(수식) ──
                page.get_by_label("이름 상자").fill("B2:C3"); page.get_by_label("이름 상자").press("Enter")
                page.get_by_role("button", name="✦ AI", exact=True).click()
                box = page.get_by_role("textbox", name="AI 요청", exact=True)
                box.fill("둘째 열에 10% 수식"); box.press("Enter")
                expect(page.get_by_text("제안 — 바뀌는 셀 2개", exact=False)).to_be_visible(timeout=15000)
                page.get_by_role("button", name="반영", exact=True).click()
                expect(page.get_by_text("제안 — 바뀌는 셀", exact=False)).to_have_count(0, timeout=5000)
                set_cell(page, "C2", "=B2*0.1")  # 같은 수식 — 이름 상자로 C2 를 고르면 수식 입력줄이 반영된 수식을 보인다
                expect(page.get_by_label("수식 입력줄")).to_have_value("=B2*0.1")
                page.get_by_role("button", name="저장", exact=True).first.click()
                expect(status).to_contain_text("저장됨 · 원본 파일 기록 확인", timeout=10000)
                assert load_workbook(book).active["C3"].value == "=B3*0.1"
                assert load_workbook(book, data_only=True).active["C2"].value == pytest.approx(0.7)
                # 이름 상자 입력 중 비동기 제안·계산이 화면을 갱신해도 목적지를 덮지 않는다.
                snapshot = finished(workspace.request_snapshot(doc["id"], "snap-name-edit")["operation_id"])["result"]["snapshot"]
                proposal = workspace.propose(doc["id"], snapshot["id"], "1", "B2", [[9]])
                name_box = page.get_by_label("이름 상자")
                name_box.fill("D2")
                detail = workspace.detail(doc["id"])
                queued = workspace.apply(doc["id"], proposal["id"], session_id=detail["session"]["id"], client_id=detail["session"]["client_id"],
                                         epoch=detail["session"]["engine_epoch"], expected=detail["session"]["session_revision"], operation_id="apply-name-edit")
                assert finished(queued["operation_id"])["result"]["applied"] is True
                expect(name_box).to_have_value("D2")
                name_box.press("Enter")
                bar = page.get_by_label("수식 입력줄")
                bar.fill("=B2*0.2"); bar.press("Enter")
                page.get_by_role("button", name="저장", exact=True).first.click()
                expect(status).to_contain_text("저장됨 · 원본 파일 기록 확인", timeout=10000)
                assert load_workbook(book).active["B2"].value == 9
                assert load_workbook(book, data_only=True).active["C2"].value == pytest.approx(0.9)
                assert load_workbook(book, data_only=True).active["D2"].value == pytest.approx(1.8)
                assert not failures, failures
            finally:
                browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=10)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q", "-m", "system"]))
