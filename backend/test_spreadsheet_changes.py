"""Conditional cancellation and refresh preserve subsequent human edits."""
import io

import boot_paths  # noqa: F401
import pytest
from openpyxl import Workbook, load_workbook

from spreadsheet_workspace import SpreadsheetWorkspace
from spreadsheet_changes import inverse, refresh, same_cells, finish, structure
from spreadsheet_imports import import_csv
from office_sessions import DocumentConflict


def blob(value=1, other='00123', change=None):
    book = Workbook()
    sheet = book.active
    sheet.append([value, other, '=A1*2'])
    if change:
        change(sheet)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def app(tmp_path, monkeypatch):
    import document_office
    monkeypatch.setattr(document_office, 'available', lambda: True)
    return SpreadsheetWorkspace(tmp_path / 'office')


def snap(app, doc, session, data):
    session.update(blob=app.store.blob(data), session_revision=session['session_revision'] + 1)
    app.store.put('session', session)
    return app.snapshot(doc, session['id'], session['client_id'], session['engine_epoch'],
                        session['session_revision'], 'engine-' + str(session['session_revision']), 'stale')


def changed(app, tmp_path):
    path = tmp_path / 'book.xlsx'
    path.write_bytes(blob())
    doc = app.open(path)['document']['id']
    session = app.acquire(doc, 'window')['session']
    before = snap(app, doc, session, blob())
    p = app.propose(doc, before['id'], '1', 'A1', [[7]])
    after = snap(app, doc, session, blob(7))
    op = {'id': 'change1', 'document_id': doc, 'session_id': session['id'], 'proposal_id': p['id'],
          'status': 'completed', 'result': {'applied': True, 'snapshot_id': after['id']}}
    app.store.put('operation', op)
    return doc, session, op


def test_inverse_preserves_other_cell_and_original(app, tmp_path):
    doc, session, op = changed(app, tmp_path)
    now = snap(app, doc, session, blob(7, 'person'))
    p = inverse(app, doc, op['id'], now['id'])
    assert p['range'] == 'A1'
    assert p['kind'] == 'restore_cells'
    assert p['values'] == [[{'value': 1, 'formula': None}]]
    assert load_workbook(tmp_path / 'book.xlsx').active['A1'].value == 1
    assert load_workbook(io.BytesIO(app.store.bytes(now['blob']))).active['B1'].value == 'person'


@pytest.mark.parametrize('change', [lambda s: s.insert_rows(1), lambda s: setattr(s, 'title', 'Moved'),
    lambda s: setattr(s.auto_filter, 'ref', 'A1:C2')])
def test_inverse_rejects_structure_change(app, tmp_path, change):
    doc, session, op = changed(app, tmp_path)
    now = snap(app, doc, session, blob(7, change=change))
    with pytest.raises(DocumentConflict):
        inverse(app, doc, op['id'], now['id'])


def test_inverse_rejects_later_edit_to_same_cell(app, tmp_path):
    doc, session, op = changed(app, tmp_path)
    now = snap(app, doc, session, blob(8))
    with pytest.raises(DocumentConflict, match='후속 편집'):
        inverse(app, doc, op['id'], now['id'])


@pytest.mark.parametrize('a,b', [('00123', '123'), (' A', 'A'), (False, 0), ('', None), ('=A1', '=a1')])
def test_concurrency_identity_is_exact(a, b):
    assert not same_cells([{'value': a}], [{'value': b}])


def test_refresh_keeps_text_ids_and_rejects_human_changes(app, tmp_path):
    source = tmp_path / 'source.csv'
    source.write_text('id,value\n00123,7\n00456,8\n')
    registered = app.open(source)['document']
    imported = import_csv(app, registered['id'], registered['revision_id'], types=['text', 'number'])
    doc, recipe = imported['document']['id'], imported['import']
    session = app.acquire(doc, 'window')['session']
    current = snap(app, doc, session, app.store.bytes(session['blob']))
    source2 = tmp_path / 'new.csv'
    source2.write_text('id,value\n00123,9\n')
    registered2 = app.open(source2)['document']
    p = refresh(app, doc, recipe['id'], registered2['id'], registered2['revision_id'], current['id'])
    assert p['values'] == [['id', 'value'], ['00123', 9], [None, None]]
    book = load_workbook(io.BytesIO(app.store.bytes(current['blob'])))
    book.active['A2'] = 'human'
    stream = io.BytesIO(); book.save(stream)
    edited = snap(app, doc, session, stream.getvalue())
    with pytest.raises(DocumentConflict, match='사람의 편집'):
        refresh(app, doc, recipe['id'], registered2['id'], registered2['revision_id'], edited['id'])


def test_refresh_rejects_type_errors_without_advancing_recipe(app, tmp_path):
    source = tmp_path / 'source.csv'; source.write_text('id,n\n001,2\n')
    d = app.open(source)['document']
    imported = import_csv(app, d['id'], d['revision_id'], types=['text', 'number'])
    doc, recipe = imported['document']['id'], imported['import']
    s = app.acquire(doc, 'window')['session']
    current = snap(app, doc, s, app.store.bytes(s['blob']))
    bad = tmp_path / 'bad.csv'; bad.write_text('id,n\n001,not-number\n')
    d = app.open(bad)['document']
    with pytest.raises(ValueError, match='오류'):
        refresh(app, doc, recipe['id'], d['id'], d['revision_id'], current['id'])
    assert app.store.get('sheet_import', recipe['id']) == recipe


def test_receipt_does_not_adopt_a_racing_human_edit(app, tmp_path):
    doc, session, op = changed(app, tmp_path)
    raced = snap(app, doc, session, blob(99))
    result = {'applied': True, 'snapshot_id': raced['id']}
    finish(app, op, result)
    assert result['applied'] and 'snapshot_id' not in result
    assert 'capture_warning' in result


def test_preflight_rejects_filter_only_change(app, tmp_path):
    from spreadsheet_changes import preflight
    path = tmp_path / 'book.xlsx'; path.write_bytes(blob())
    doc = app.open(path)['document']['id']; s = app.acquire(doc, 'window')['session']
    original = snap(app, doc, s, blob())
    p = app.propose(doc, original['id'], '1', 'A1', [[3]])
    args = dict(session_id=s['id'], client_id=s['client_id'], epoch=s['engine_epoch'], expected=s['session_revision'])
    op = app.apply(doc, p['id'], **args, operation_id='queued-filter')
    assert preflight(app, doc, op['operation_id'], **args)['ready']
    snap(app, doc, s, blob(change=lambda sheet: setattr(sheet.auto_filter, 'ref', 'A1:C1')))
    args['expected'] = s['session_revision']
    with pytest.raises(DocumentConflict, match='필터'):
        preflight(app, doc, op['operation_id'], **args)


def test_inspection_cache_cannot_leak_mutations_or_stale_metadata():
    from spreadsheet_files import inspect
    original = blob()
    first = inspect(original)
    first['sheets'][0]['name'] = 'tampered'
    assert inspect(original)['sheets'][0]['name'] == 'Sheet'
    renamed = blob(change=lambda sheet: setattr(sheet, 'title', 'Current'))
    assert inspect(renamed)['sheets'][0]['name'] == 'Current'


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
