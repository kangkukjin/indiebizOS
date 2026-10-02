"""Real spreadsheet engine selection gate; no user files or Automation API.

Run with INDIEBIZ_OFFICE_LIVE_TEST=1. The fixture serves one generated workbook
and one fixed plugin to the already configured local ONLYOFFICE server.
"""
import io
import json
import os
import secrets
import socket
import subprocess
import threading
import time
import zipfile
from pathlib import Path

import boot_paths  # noqa: F401
import pytest

PLUGIN = r"""
(function () {
'use strict';
function reply(result) { window.top.postMessage({type:'sheet-gate',result:result}, '*'); }
Asc.plugin.init = function () {
 Asc.plugin.callCommand(function () {
  try {
   var sheet=Api.GetActiveSheet();
   sheet.SetName('Transactions');
   sheet.GetRange('A1:C4').SetValue([['Category','Amount','Quantity'],['A',10,2],['B',20,3],['A',30,4]]);
   sheet.GetRange('E1').SetValue('=SUM(B2:B4)');
   sheet.GetRange('E2').SetValue('=SUMIFS(B2:B4,A2:A4,"A")');
   sheet.GetRange('E3').SetValue('=XLOOKUP("B",A2:A4,B2:B4)');
   var chart=sheet.AddChart("'Transactions'!$A$1:$B$4",false,'bar',2,5000000,3000000,6,0,1,0);
   chart.SetTitle('Engine acceptance',12);
   var pivot=Api.InsertPivotNewWorksheet(sheet.GetRange('A1:C4'));
   pivot.AddFields({rows:'Category'});
   pivot.AddDataField('Amount','Total','sum');
   pivot.RefreshTable();
   sheet.GetRange('B2').SetValue(15);
   pivot.RefreshTable();
   sheet.SetActive();
   return {edited:true,sum:sheet.GetRange('E1').GetValue(),sumif:sheet.GetRange('E2').GetValue(),lookup:sheet.GetRange('E3').GetValue()};
  } catch(e) { return {error:String(e)}; }
 },false,true,function(r){reply(r);});
};
Asc.plugin.button=function(){};
})();
"""


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit local engine acceptance')
def test_spreadsheet_plugin_roundtrip(tmp_path, monkeypatch):
    import httpx
    import jwt
    import uvicorn
    import document_office
    from fastapi import FastAPI, Request
    from fastapi.responses import HTMLResponse, Response
    from openpyxl import Workbook, load_workbook
    from playwright.sync_api import sync_playwright

    cfg = document_office.settings()
    assert cfg, 'Local ONLYOFFICE configuration required'
    seed = io.BytesIO()
    Workbook().save(seed)
    key, ticket = secrets.token_hex(16), secrets.token_urlsafe(24)
    callbacks = []
    app = FastAPI()
    from fastapi.middleware.cors import CORSMiddleware
    app.add_middleware(CORSMiddleware, allow_origins=[cfg['url']], allow_methods=['GET'])
    sock = socket.socket()
    sock.bind(('0.0.0.0', 0))
    port = sock.getsockname()[1]
    callback_host = __import__('urllib.parse', fromlist=['urlsplit']).urlsplit(cfg['callback_origin']).hostname
    base = f'http://{callback_host}:{port}'
    browser_base = f'http://127.0.0.1:{port}'
    guid = 'asc.{928741E2-8E71-4410-9AEC-B802813D50BA}'

    @app.get('/seed/{capability}')
    def content(capability: str):
        assert secrets.compare_digest(capability, ticket)
        return Response(seed.getvalue(), media_type='application/octet-stream')

    @app.post('/callback/{capability}')
    async def callback(capability: str, request: Request):
        assert secrets.compare_digest(capability, ticket)
        body = await request.json()
        claims = jwt.decode(body.get('token') or request.headers.get('authorization','').removeprefix('Bearer '), cfg['secret'], algorithms=['HS256'], leeway=5)
        payload = claims.get('payload', claims)
        assert payload['key'] == key
        if payload.get('status') in (2, 6):
            data = document_office.download(payload['url'], cfg)
            callbacks.append(data)
        return {'error': 0}

    @app.get('/plugin/config.json')
    def plugin_config():
        return {'name':'Spreadsheet gate','guid':guid,'baseUrl':browser_base+'/plugin/',
                'variations':[{'description':'gate','url':'index.html','isViewer':False,'EditorsSupport':['cell'],
                               'isVisual':False,'isModal':False,'isInsideMode':False,'initDataType':'none','initData':'','buttons':[]}]}

    @app.get('/plugin/index.html')
    def plugin_html():
        return HTMLResponse('<script src="'+cfg['url']+'/sdkjs-plugins/v1/plugins.js"></script><script src="plugin.js"></script>')

    @app.get('/plugin/plugin.js')
    def plugin_js():
        return Response(PLUGIN, media_type='application/javascript')

    @app.get('/')
    def index():
        options = {'documentType':'cell','width':'100%','height':'100%','document':{'fileType':'xlsx','key':key,
                   'title':'Engine gate.xlsx','url':base+'/seed/'+ticket},
                   'editorConfig':{'lang':'en','mode':'edit','callbackUrl':base+'/callback/'+ticket,
                   'user':{'id':'gate','name':'Acceptance'},'plugins':{'autostart':[guid],'pluginsData':[browser_base+'/plugin/config.json']}}}
        options['token'] = jwt.encode(options, cfg['secret'], algorithm='HS256')
        return HTMLResponse('<html style="height:100%"><body style="height:100%;margin:0"><div id="editor"></div>'
             +'<script src="'+cfg['url']+'/web-apps/apps/api/documents/api.js"></script><script>'
             +'window.gate=null;window.addEventListener("message",e=>{if(e.data&&e.data.type==="sheet-gate")window.gate=e.data.result});'
             +'window.editor=new DocsAPI.DocEditor("editor",'+json.dumps(options)+');</script></body></html>')

    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets':[sock]}, daemon=True)
    thread.start()
    try:
        deadline = time.monotonic()+10
        while not server.started and time.monotonic()<deadline:
            time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                page = browser.new_page(viewport={'width':1500,'height':1000})
                page.goto(browser_base)
                try:
                    page.wait_for_function('window.gate !== null', timeout=60000)
                except Exception:
                    print('ENGINE UI:', '\n'.join(f.locator('body').inner_text()[:1500] for f in page.frames))
                    raise
                result = page.evaluate('window.gate')
                assert result.get('edited'), result
                # Formula values returned inside a mutation can precede recalculation.
                # The exported workbook, not this immediate return, decides the gate.
                page.keyboard.press('Control+s')
                command = {'c':'forcesave','key':key,'userdata':'gate-save'}
                command['token'] = jwt.encode(command, cfg['secret'], algorithm='HS256')
                deadline = time.monotonic()+60
                while not callbacks and time.monotonic()<deadline:
                    response = httpx.post(cfg['url']+'/coauthoring/CommandService.ashx',json=command,timeout=20,trust_env=False)
                    assert response.json().get('error') in (0,4), response.text
                    page.wait_for_timeout(1000)
                assert callbacks, 'No engine save callback'
                saved = tmp_path/'roundtrip.xlsx'
                saved.write_bytes(callbacks[-1])
                formulas = load_workbook(saved, data_only=False)
                values = load_workbook(saved, data_only=True)
                assert formulas['Transactions']['E1'].value == '=SUM(B2:B4)'
                assert values['Transactions']['E1'].value == 65
                assert values['Transactions']['E2'].value == 45
                assert values['Transactions']['E3'].value == 20
                assert len(formulas['Transactions']._charts) == 1
                assert any(sheet._pivots for sheet in formulas)
                with zipfile.ZipFile(saved) as archive:
                    assert any(name.startswith('xl/pivotTables/') for name in archive.namelist())
                out = tmp_path/'independent'
                out.mkdir()
                run = subprocess.run(['soffice','-env:UserInstallation='+ (tmp_path/'lo-profile').as_uri(),
                    '--headless','--convert-to','xlsx','--outdir',str(out),str(saved)],capture_output=True,text=True,timeout=90)
                assert run.returncode == 0, run.stderr
                independent = load_workbook(out/saved.name, data_only=True)
                assert independent['Transactions']['E1'].value == 65
                assert independent['Transactions']['E2'].value == 45
                assert independent['Transactions']['E3'].value == 20
                print(json.dumps({'engine':'ONLYOFFICE 9.3.1','plugin':result,'saved_sum':65,'pivot_parts':True,
                    'editable_chart':True,'independent':'LibreOffice','source':'synthetic fixture'}))
            finally:
                browser.close()
    finally:
        server.should_exit=True
        thread.join(timeout=10)
        sock.close()


if __name__ == '__main__':
    import sys
    import pytest
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
