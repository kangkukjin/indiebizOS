"""Conversion copies preserve source ownership and reject active file content."""
import io
import zipfile

import boot_paths  # noqa: F401
import pytest
from openpyxl import Workbook, load_workbook

import spreadsheet_formats as formats
from office_sessions import DocumentConflict, DocumentUnsupported
from spreadsheet_workspace import SpreadsheetWorkspace


def xlsx(value=12, template=False):
    book = Workbook(); book.template = template
    book.active['A1'] = '00123'; book.active['B1'] = value
    book.active['C1'] = '=B1*2'
    output = io.BytesIO(); book.save(output)
    return output.getvalue()


def odf(extra='', files=None, format='ods'):
    output = io.BytesIO()
    content = ('<office:document-content xmlns:office="'+formats.ODF+'office:1.0" '
               'xmlns:table="'+formats.ODF+'table:1.0" xmlns:xlink="http://www.w3.org/1999/xlink">'
               '<office:body><office:spreadsheet><table:table table:name="Sheet1">'+extra+
               '</table:table></office:spreadsheet></office:body></office:document-content>')
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('mimetype', formats.MIMES[format])
        archive.writestr('content.xml', content)
        archive.writestr('META-INF/manifest.xml', '<manifest/>')
        for name, value in (files or {}).items():
            archive.writestr(name, value)
    return output.getvalue()


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setattr(formats.document_office, 'available', lambda: True)
    return SpreadsheetWorkspace(tmp_path/'office')


def opened(app, tmp_path):
    source = tmp_path/'original.xlsx'; source.write_bytes(xlsx())
    return source, app.open(source)['document']


def test_copy_records_pinned_draft_not_original(app, tmp_path, monkeypatch):
    source, doc = opened(app, tmp_path)
    session = app.acquire(doc['id'], 'one')['session']
    session.update(blob=app.store.blob(xlsx(37)), state='draft', session_revision=1)
    app.store.put('session', session)
    seen = []
    def converter(app, document_id, data, source_format, output_format):
        seen.append(load_workbook(io.BytesIO(data)).active['B1'].value)
        return odf()
    monkeypatch.setattr(formats, '_convert_bytes', converter)
    result = formats.convert(app, doc['id'], 'ods', doc['revision_id'], session['id'],
                             session['client_id'], session['engine_epoch'], 1)
    assert seen == [37]
    assert load_workbook(source).active['B1'].value == 12
    assert result['document']['id'] != doc['id']
    assert result['document']['provenance']['unsaved'] is True
    assert result['conversion_report']['status'] == 'partial'
    assert result['conversion_report']['unverified']
    assert app.detail(doc['id'])['session']['blob'] == session['blob']


def test_revision_and_writer_conflicts_before_conversion(app, tmp_path, monkeypatch):
    _, doc = opened(app, tmp_path)
    monkeypatch.setattr(formats, '_convert_bytes', lambda *a: pytest.fail('must not convert'))
    with pytest.raises(DocumentConflict):
        formats.convert(app, doc['id'], 'ods', 'stale')
    app.acquire(doc['id'], 'one')
    with pytest.raises((DocumentConflict, ValueError, TypeError)):
        formats.convert(app, doc['id'], 'ods', doc['revision_id'])


def test_changed_session_during_conversion_is_not_registered(app, tmp_path, monkeypatch):
    _, doc = opened(app, tmp_path)
    session = app.acquire(doc['id'], 'one')['session']
    def converter(*args):
        session['session_revision'] += 1
        app.store.put('session', session)
        return odf()
    monkeypatch.setattr(formats, '_convert_bytes', converter)
    with pytest.raises(DocumentConflict):
        formats.convert(app, doc['id'], 'ods', doc['revision_id'], session['id'],
                        session['client_id'], session['engine_epoch'], 0)
    assert len(app.list()) == 1


@pytest.mark.parametrize('extra,files', [
    ('<table:table-source xlink:href="https://example.org/private"/>', {}),
    ('<office:script/>', {}),
    ('', {'Basic/macro.xml': '<script/>'}),
    ('<table:table-cell table:formula="of:=WEBSERVICE(&quot;https://x&quot;)"/>', {}),
    ('', {'META-INF/evil.xml': '<encryption-data/>'}),
])
def test_odf_active_inputs_rejected(extra, files):
    with pytest.raises(DocumentUnsupported):
        formats.inspect(odf(extra, files), 'ods')


def test_template_identity_and_rejection_of_renamed_format():
    assert formats.inspect(xlsx(template=True), 'xltx')['sheets'] == 1
    with pytest.raises(DocumentUnsupported):
        formats.inspect(xlsx(template=True), 'xlsx')
    with pytest.raises(DocumentUnsupported):
        formats.inspect(odf(format='ots'), 'ods')


def test_converter_content_ticket_and_expiry(app, tmp_path):
    _, doc = opened(app, tmp_path)
    app.store.put('sheet_conversion', {'id':'conversion', 'document_id':doc['id'],
        'blob':doc['source_sha256'], 'token':'secret', 'expires':9999999999})
    assert formats.content(app, 'conversion', 'secret') == app.store.bytes(doc['source_sha256'])
    with pytest.raises(PermissionError):
        formats.content(app, 'conversion', 'wrong')
    row = app.store.get('sheet_conversion', 'conversion'); row['expires'] = 0
    app.store.put('sheet_conversion', row)
    with pytest.raises(PermissionError):
        formats.content(app, 'conversion', 'secret')


def test_odf_internal_chart_does_not_count_as_workbook_sheet():
    payload=odf('<table:table-cell xlink:href="./Object 1"/>',
                {'Object 1/content.xml':'<table xmlns="'+formats.ODF+'table:1.0" xmlns:xlink="http://www.w3.org/1999/xlink" xlink:href=".."/>'})
    assert formats.inspect(payload,'ods')['sheets']==1


def test_missing_libreoffice_is_explicit(monkeypatch):
    monkeypatch.setattr(formats.shutil,'which',lambda name:None)
    with pytest.raises(DocumentUnsupported,match='LibreOffice'):
        formats._libreoffice(xlsx(),'xlsx','ods')


def flat_odf(extra=''):
    return ('<office:document xmlns:office="'+formats.ODF+'office:1.0" '
            'xmlns:table="'+formats.ODF+'table:1.0" xmlns:of="urn:oasis:names:tc:opendocument:xmlns:of:1.2" '
            'xmlns:xlink="http://www.w3.org/1999/xlink" office:mimetype="'+formats.MIMES['ods']+'">'
            '<office:body><office:spreadsheet><table:table table:name="Sheet">'+extra+
            '</table:table></office:spreadsheet></office:body></office:document>').encode()


def boolean_row(value='false', formula='', repeat=1):
    return ('<table:table-row><table:table-cell office:value-type="boolean" '
            f'office:boolean-value="{value}" table:number-columns-repeated="{repeat}" '+
            (f'table:formula="{formula}"' if formula else '')+'/></table:table-row>')


def test_fods_identity_and_external_reference_guards():
    assert formats.inspect(flat_odf(), 'fods')['sheets'] == 1
    for extra in ('<office:script/>', '<table:table-source xlink:href="file:///etc/passwd"/>',
                  '<table:table-source xlink:href="../source.ods"/>',
                  '<table:table-cell table:formula="of:=WEBSERVICE(&quot;https://x&quot;)"/>'):
        with pytest.raises(DocumentUnsupported):
            formats.inspect(flat_odf(extra), 'fods')
    with pytest.raises(DocumentUnsupported):
        formats.inspect(b'<document/>', 'fods')
    with pytest.raises(Exception):
        formats.inspect(b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x>&x;</x>', 'fods')


def test_fods_empty_library_is_inert_but_code_and_library_links_are_rejected():
    empty = ('<office:script xmlns:script="'+formats.ODF+'script:1.0" '
             'xmlns:ooo="http://openoffice.org/2004/office" script:language="ooo:Basic">'
             '<ooo:libraries/></office:script>')
    assert formats.inspect(flat_odf(empty), 'fods')['sheets'] == 1
    for active in (empty.replace('<ooo:libraries/>', '<ooo:libraries><ooo:library/></ooo:libraries>'),
                   empty.replace('<ooo:libraries/>', '<ooo:libraries xlink:href="file:///macro"/>'),
                   empty.replace('<ooo:libraries/>', '<ooo:libraries>code</ooo:libraries>'),
                   empty.replace('<ooo:libraries/>', 'code<ooo:libraries/>')):
        with pytest.raises(DocumentUnsupported):
            formats.inspect(flat_odf(active), 'fods')

def test_fods_chart_parent_is_internal_only_in_embedded_document():
    chart = ('<draw:object xmlns:draw="'+formats.ODF+'drawing:1.0" '
             'xmlns:chart="'+formats.ODF+'chart:1.0">'
             '<office:document><office:body><office:chart><chart:chart xlink:href=".."/>'
             '</office:chart></office:body></office:document></draw:object>')
    assert formats.inspect(flat_odf(chart), 'fods')['sheets'] == 1
    for invalid in (chart.replace('href=".."', 'href="../other.ods"'),
                    '<chart:chart xmlns:chart="'+formats.ODF+'chart:1.0" xlink:href=".."/>'):
        with pytest.raises(DocumentUnsupported):
            formats.inspect(flat_odf(invalid), 'fods')

def test_boolean_conversion_preserves_original_formulas_and_other_zip_parts():
    from spreadsheet_conversion_types import preserve_booleans
    from lxml import etree
    book = Workbook(); book.active['A1'] = False; book.active['B1'] = '=FALSE()'
    source = io.BytesIO(); book.save(source)
    raw = io.BytesIO()
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(raw, 'w') as after:
        for item in before.infolist():
            content = before.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(content)
                for cell in root.iter(formats.spreadsheet_files.NS+'c'):
                    cell.set('t', 'b')
                    formula = cell.find(formats.spreadsheet_files.NS+'f')
                    if formula is None:
                        formula = etree.SubElement(cell, formats.spreadsheet_files.NS+'f')
                    formula.text = 'FALSE()'
                    cell.find(formats.spreadsheet_files.NS+'v').text = '0'
                content = etree.tostring(root)
            after.writestr(item, content)
    fixed, count = preserve_booleans(source.getvalue(), 'xlsx', raw.getvalue(), 'xlsx')
    assert count == 1
    result = load_workbook(io.BytesIO(fixed))
    assert result.active['A1'].value is False and result.active['A1'].data_type == 'b'
    assert result.active['B1'].value == '=FALSE()'
    with zipfile.ZipFile(raw) as before, zipfile.ZipFile(io.BytesIO(fixed)) as after:
        for name in before.namelist():
            if name != 'xl/worksheets/sheet1.xml':
                assert before.read(name) == after.read(name)


def test_repeated_odf_booleans_are_preserved_without_rewriting_other_formulas():
    from spreadsheet_conversion_types import preserve_booleans, TABLE
    from lxml import etree
    source = flat_odf(boolean_row(repeat=3))
    output = flat_odf(boolean_row(formula='of:=FALSE()', repeat=3))
    fixed, count = preserve_booleans(source, 'fods', output, 'fods')
    assert count == 3
    root = etree.fromstring(fixed)
    cell = next(root.iter(TABLE+'table-cell'))
    assert TABLE+'formula' not in cell.attrib
    assert root.nsmap['of'] == 'urn:oasis:names:tc:opendocument:xmlns:of:1.2'
    # The original formula must remain a formula, even when its result is boolean.
    assert preserve_booleans(output, 'fods', output, 'fods') == (output, 0)


@pytest.mark.parametrize('output', [
    flat_odf(boolean_row(value='true', formula='of:=TRUE()')),
    flat_odf(boolean_row(formula='of:=1=2')),
    flat_odf(),
    flat_odf(boolean_row(formula='of:=FALSE()', repeat=2)),
])
def test_boolean_mismatch_or_ambiguous_repeat_fails_closed(output):
    from spreadsheet_conversion_types import preserve_booleans
    with pytest.raises(DocumentUnsupported):
        preserve_booleans(flat_odf(boolean_row()), 'fods', output, 'fods')


def test_embedded_table_rows_do_not_shift_workbook_boolean_coordinates():
    from spreadsheet_conversion_types import preserve_booleans, TABLE
    from lxml import etree
    nested = ('<table:table-row><table:table-cell><table:table table:name="chart-data">'+
              boolean_row(value='true')+'</table:table></table:table-cell></table:table-row>')
    source = flat_odf(nested+'<table:table-row-group>'+boolean_row()+'</table:table-row-group>')
    output = flat_odf(nested+'<table:table-row-group>'+boolean_row(formula='of:=FALSE()')+'</table:table-row-group>')
    result, repaired = preserve_booleans(source, 'fods', output, 'fods')
    assert repaired == 1
    root = etree.fromstring(result)
    rows = list(root.iter(TABLE+'table-row'))
    assert rows[1][0].get('{'+formats.ODF+'office:1.0}boolean-value') == 'true'
    assert rows[2][0].get(TABLE+'formula') is None

def test_excessive_boolean_repetition_is_bounded():
    from spreadsheet_conversion_types import preserve_booleans
    payload = flat_odf(boolean_row(repeat=16384).replace('<table:table-row>',
        '<table:table-row table:number-rows-repeated="1048576">'))
    with pytest.raises(DocumentUnsupported, match='100,000'):
        preserve_booleans(payload, 'fods', payload, 'fods')


def test_missing_output_boolean_never_registers_copy(app, tmp_path, monkeypatch):
    payload = xlsx(False)
    source = tmp_path/'bool.xlsx'; source.write_bytes(payload)
    doc = app.open(source)['document']
    monkeypatch.setattr(formats, '_convert_bytes', lambda *a: odf())
    with pytest.raises(DocumentUnsupported):
        formats.convert(app, doc['id'], 'ods', doc['revision_id'])
    assert len(app.list()) == 1 and source.read_bytes() == payload


@pytest.mark.system
@pytest.mark.skipif(__import__('os').environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1',
                    reason='explicit local conversion test')
def test_live_fods_roundtrip_boolean_formula_and_namespace():
    from spreadsheet_conversion_types import preserve_booleans
    book = Workbook(); book.active['A1'] = False; book.active['B1'] = True
    book.active['C1'] = '=FALSE()'; book.active['D1'] = '=1+2'
    book.active['A2'] = '00123'; book.create_sheet('한글')['A1'] = '보존'
    source = io.BytesIO(); book.save(source)
    flat = formats._libreoffice(source.getvalue(), 'xlsx', 'fods')
    assert formats.inspect(flat, 'fods')['sheets'] == 2
    flat, repaired = preserve_booleans(source.getvalue(), 'xlsx', flat, 'fods')
    assert repaired == 2
    result = formats._libreoffice(flat, 'fods', 'xlsx')
    formats.inspect(result, 'xlsx')
    result, _ = preserve_booleans(flat, 'fods', result, 'xlsx')
    reopened = load_workbook(io.BytesIO(result))
    assert reopened.active['A1'].value is False and reopened.active['B1'].value is True
    assert reopened.active['C1'].value == '=FALSE()' and reopened.active['D1'].value == '=1+2'
    assert reopened.active['A2'].value == '00123' and reopened['한글']['A1'].value == '보존'
    assert load_workbook(io.BytesIO(result), data_only=True).active['D1'].value == 3


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__]))
