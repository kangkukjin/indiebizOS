"""Release function matrix evaluated by ONLYOFFICE and reopened by LibreOffice."""
import json
import os

import boot_paths  # noqa: F401
import pytest
from openpyxl import load_workbook
import test_spreadsheet_engine_live as gate

# Each result gets four free rows for array spill; fixed inputs avoid locale guesses.
CASES = [
 ('SUM', '=SUM(1,2,3)', 6), ('AVERAGE', '=AVERAGE(1,2,3)', 2),
 ('MIN', '=MIN(3,1,2)', 1), ('MAX', '=MAX(3,1,2)', 3),
 ('COUNT', '=COUNT(1,2,"x")', 2), ('COUNTA', '=COUNTA(1,"x")', 2),
 ('ROUND', '=ROUND(1.234,2)', 1.23), ('SUMIF', '=SUMIF(A2:A4,"A",B2:B4)', 45),
 ('SUMIFS', '=SUMIFS(B2:B4,A2:A4,"A")', 45), ('COUNTIF', '=COUNTIF(A2:A4,"A")', 2),
 ('COUNTIFS', '=COUNTIFS(A2:A4,"A",B2:B4,">20")', 1),
 ('SUBTOTAL', '=SUBTOTAL(9,B2:B4)', 65), ('AGGREGATE', '=AGGREGATE(9,0,B2:B4)', 65),
 ('IF', '=IF(1<2,7,8)', 7), ('IFS', '=IFS(1>2,1,1<2,2)', 2),
 ('AND', '=AND(TRUE,TRUE)', True), ('OR', '=OR(FALSE,TRUE)', True), ('NOT', '=NOT(FALSE)', True),
 ('IFERROR', '=IFERROR(1/0,9)', 9), ('ISBLANK', '=ISBLANK(Z999)', True), ('ISNUMBER', '=ISNUMBER(7)', True),
 ('VLOOKUP', '=VLOOKUP("B",A2:B4,2,FALSE)', 20),
 ('HLOOKUP', '=HLOOKUP("b",{"a","b";1,2},2,FALSE)', 2),
 ('XLOOKUP', '=XLOOKUP("B",A2:A4,B2:B4)', 20),
 ('INDEX', '=INDEX(B2:B4,2)', 20), ('MATCH', '=MATCH("B",A2:A4,0)', 2),
 ('OFFSET', '=OFFSET(B2,1,0)', 20), ('INDIRECT', '=INDIRECT("B3")', 20),
 ('LEFT', '=LEFT("abcdef",2)', 'ab'), ('RIGHT', '=RIGHT("abcdef",2)', 'ef'),
 ('MID', '=MID("abcdef",2,3)', 'bcd'), ('LEN', '=LEN("한글")', 2),
 ('TRIM', '=TRIM(" a  b ")', 'a b'), ('SUBSTITUTE', '=SUBSTITUTE("abc","b","x")', 'axc'),
 ('TEXT', '=TEXT(12.5,"0.00")', '12.50'), ('VALUE', '=VALUE("12.5")', 12.5),
 ('CONCAT', '=CONCAT("a","b")', 'ab'), ('TEXTJOIN', '=TEXTJOIN("-",TRUE,"a","b")', 'a-b'),
 ('DATE', '=DATE(2024,2,29)', 45351), ('YEAR', '=YEAR(DATE(2024,2,29))', 2024),
 ('MONTH', '=MONTH(DATE(2024,2,29))', 2), ('DAY', '=DAY(DATE(2024,2,29))', 29),
 ('TODAY', '=TODAY()>DATE(2020,1,1)', True), ('NOW', '=NOW()>=TODAY()', True),
 ('EDATE', '=DAY(EDATE(DATE(2024,1,31),1))', 29),
 ('EOMONTH', '=DAY(EOMONTH(DATE(2024,2,1),0))', 29),
 ('NETWORKDAYS', '=NETWORKDAYS(DATE(2024,1,1),DATE(2024,1,5))', 5),
 ('WORKDAY', '=DAY(WORKDAY(DATE(2024,1,5),1))', 8),
 ('FILTER', '=FILTER(B2:B4,A2:A4="A")', 15), ('SORT', '=SORT({3;1;2})', 1),
 ('UNIQUE', '=UNIQUE({1;1;2})', 1), ('SEQUENCE', '=SEQUENCE(3)', 1),
 ('TRANSPOSE', '=TRANSPOSE({1,2,3})', 1), ('SUMPRODUCT', '=SUMPRODUCT({1,2},{3,4})', 11),
 ('LET', '=LET(x,2,x*3)', 6), ('MEDIAN', '=MEDIAN(1,3,8)', 3),
 ('STDEV.S', '=STDEV.S(1,2,3)', 1), ('VAR.S', '=VAR.S(1,2,3)', 1),
 ('CORREL', '=CORREL({1,2,3},{2,4,6})', 1), ('RANK.EQ', '=RANK.EQ(2,{1,2,3},0)', 2),
 ('PMT', '=PMT(0,10,100)', -10), ('NPV', '=NPV(0.1,110)', 100), ('IRR', '=IRR({-100;110})', .1),
]


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit engine test')
def test_required_functions(tmp_path, monkeypatch):
    code = '\n'.join('sheet.GetRange('+json.dumps('H'+str(i*5+1))+').SetValue('+json.dumps(formula)+');'
                     for i, (_, formula, _) in enumerate(CASES))
    monkeypatch.setattr(gate, 'PLUGIN', gate.PLUGIN.replace('sheet.SetActive();', code+'\nsheet.SetActive();'))
    gate.test_spreadsheet_plugin_roundtrip(tmp_path, monkeypatch)
    saved = load_workbook(tmp_path/'roundtrip.xlsx', data_only=True)['Transactions']
    independent = load_workbook(tmp_path/'independent/roundtrip.xlsx', data_only=True)['Transactions']
    observations = []
    for i, (name, formula, expected) in enumerate(CASES):
        value = saved['H'+str(i*5+1)].value
        equal = value == pytest.approx(expected) if type(expected) in (int, float) else type(value) is type(expected) and value == expected
        observations.append({'function':name,'formula':formula,'engine':value,'expected':expected,
                             'passed':bool(equal),'libreoffice':independent['H'+str(i*5+1)].value})
    print('FUNCTION_MATRIX '+json.dumps(observations, ensure_ascii=False, default=str))
    assert all(r['passed'] for r in observations), [r for r in observations if not r['passed']]


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
