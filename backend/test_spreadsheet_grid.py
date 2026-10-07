"""격자 엔진 I/O — grid.v1 ↔ XLSX 왕복, capabilities 의 엔진 판정, grid-capture 초안과 스냅샷 계산 상태.
docs/SPREADSHEET_APP_ON_IBL_PLAN_2026_10_07.md §2·§4."""
import datetime
import io

import boot_paths  # noqa: F401
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

import spreadsheet_files as files
import spreadsheet_grid as grid
from office_sessions import DocumentConflict
from spreadsheet_workspace import SpreadsheetWorkspace


def book(with_chart=False):
    wb = Workbook(); ws = wb.active; ws.title = '매출'
    ws.append(['날짜', '거래처', '수입', '지출', '잔액'])
    ws.append([datetime.date(2026, 10, 1), '한빛', 1200000, 0, '=C2-D2'])
    ws.append([datetime.date(2026, 10, 2), '00123', 0, 48500, '=E2+C3-D3'])
    ws['A1'].font = Font(b=True, color='FFFFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='FF217346')
    ws['C2'].number_format = '#,##0'; ws['A2'].number_format = 'yyyy-mm-dd'
    ws['B3'].data_type = 's'
    ws.merge_cells('G1:H2'); ws.column_dimensions['B'].width = 20; ws.freeze_panes = 'A2'
    ws['F1'] = '=LET(x,2,x*3)'
    if with_chart:
        chart = BarChart(); chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=3)); ws.add_chart(chart, 'J2')
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def test_grid_roundtrip_keeps_values_formulas_styles_and_layout():
    base = book()
    g = grid.to_grid(base)
    s0 = g['sheets'][0]
    assert s0['cells']['A2'] == {'v': 46296.0, 't': 'n', 's': {'n': {'pattern': 'yyyy-mm-dd'}}}
    assert s0['cells']['B3'] == {'v': '00123', 't': 's'}
    assert s0['cells']['E2']['f'] == '=C2-D2' and s0['cells']['F1']['f'] == '=LET(x,2,x*3)'
    assert s0['cells']['A1']['s']['bl'] == 1 and s0['cells']['A1']['s']['bg'] == {'rgb': '#217346'}
    assert s0['freeze'] == {'rows': 1, 'cols': 0} and s0['merges'] == ['G1:H2'] and s0['cols']['1']['w'] == 145
    # 격자가 계산한 값을 들고 돌아온다 → 파일에 캐시로 심긴다
    s0['cells']['E2']['v'] = 1200000; s0['cells']['E3']['v'] = 1151500; s0['cells']['F1']['v'] = 6
    s0['cells']['I1'] = {'v': '새 값', 't': 's', 's': {'it': 1, 'cl': {'rgb': '#c00000'}, 'n': {'pattern': '@'}}}
    out = grid.from_grid(base, g)
    assert files.calculation_state(out)['status'] == 'fresh'
    wb = load_workbook(io.BytesIO(out)); ws = wb['매출']
    assert ws['F1'].value == '=_xlfn.LET(x,2,x*3)'   # 엑셀이 파일에서 요구하는 접두사
    assert ws['B3'].value == '00123' and ws['B3'].data_type == 's'
    assert ws['I1'].font.i and ws['I1'].font.color.rgb == 'FFC00000'
    assert ws.freeze_panes == 'A2' and [str(r) for r in ws.merged_cells.ranges] == ['G1:H2']
    cached = load_workbook(io.BytesIO(out), data_only=True)['매출']
    assert cached['E3'].value == 1151500 and cached['F1'].value == 6
    assert grid.to_grid(out)['sheets'][0]['cells']['F1']['f'] == '=LET(x,2,x*3)'  # 접두사는 다시 벗긴다
    assert grid.from_grid(base, g) == out  # 같은 격자 = 같은 바이트(초안 해시가 헛되이 바뀌지 않는다)


def test_from_grid_rejects_external_formulas_and_bad_cells():
    base = book(); g = grid.to_grid(base)
    bad = {**g, 'sheets': [{**g['sheets'][0], 'cells': {'A1': {'f': '=WEBSERVICE("http://x")'}}}]}
    with pytest.raises(ValueError):
        grid.from_grid(base, bad)
    with pytest.raises(ValueError):
        grid.from_grid(base, {**g, 'sheets': [{**g['sheets'][0], 'cells': {'1A': {'v': 1}}}]})


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    import document_office
    monkeypatch.setattr(document_office, 'available', lambda: False)
    return SpreadsheetWorkspace(tmp_path / 'office')


def test_capabilities_pick_grid_or_office_by_parts(workspace, tmp_path, monkeypatch):
    import document_office
    plain = tmp_path / 'plain.xlsx'; plain.write_bytes(book())
    charted = tmp_path / 'chart.xlsx'; charted.write_bytes(book(with_chart=True))
    a = workspace.open(plain)['capabilities']
    assert a['engine'] == 'grid' and a['edit_native'] and a['grid'] and not a['office'] and a['empty'] is False
    b = workspace.open(charted)['capabilities']
    assert b['engine'] == 'none' and not b['edit_native'] and 'charts' in b['grid_blockers']
    monkeypatch.setattr(document_office, 'available', lambda: True)
    c = workspace.capabilities(workspace.open(charted)['document']['id'])
    assert c['engine'] == 'office' and c['edit_native'] and c['office_available']
    blank = Workbook(); buf = io.BytesIO(); blank.save(buf)
    empty = tmp_path / 'empty.xlsx'; empty.write_bytes(buf.getvalue())
    assert workspace.open(empty)['capabilities']['empty'] is True


def args(s):
    return dict(session_id=s['id'], client_id=s['client_id'], epoch=s['engine_epoch'], expected=s['session_revision'])


def test_grid_capture_draft_then_snapshot_is_fresh_and_save_writes_file(workspace, tmp_path):
    path = tmp_path / 'book.xlsx'; path.write_bytes(book())
    doc = workspace.open(path)['document']
    s = workspace.acquire(doc['id'], 'window1')['session']
    g = workspace.grid(doc['id'])['grid']
    g['sheets'][0]['cells']['E2']['v'] = 1200000; g['sheets'][0]['cells']['E3']['v'] = 1151500; g['sheets'][0]['cells']['F1']['v'] = 6
    g['sheets'][0]['cells']['C2'] = {'v': 1300000, 't': 'n'}; g['sheets'][0]['cells']['E2']['v'] = 1300000; g['sheets'][0]['cells']['E3']['v'] = 1251500
    r = workspace.grid_capture(doc['id'], **args(s), operation_id='cap1', grid=g)
    s = r['session']
    assert s['state'] == 'draft' and s['session_revision'] == 1 and s['engine_id'] == 'grid'
    assert workspace.grid_capture(doc['id'], session_id=s['id'], client_id=s['client_id'], epoch=s['engine_epoch'], expected=0, operation_id='cap1', grid=g)['session']['session_revision'] == 1  # 멱등
    with pytest.raises(DocumentConflict):
        workspace.grid_capture(doc['id'], session_id=s['id'], client_id=s['client_id'], epoch=s['engine_epoch'], expected=0, operation_id='cap2', grid=g)  # 낡은 개정 번호
    snap = workspace.snapshot(doc['id'], **args(s), engine_state='{"grid":1}', calculation='fresh')
    assert snap['calc_status'] == 'fresh' and snap['unsaved'] is True and snap['engine_id'] == 'grid'
    read = workspace.read(doc['id'], snap['id'], '1', 'C2:E2')
    assert [c['effective_value'] for c in read['items']] == [1300000, 0, 1300000]
    workspace.save(doc['id'], **args(s), operation_id='save1', expected_revision=doc['revision_id'])
    assert load_workbook(path, data_only=True)['매출']['E3'].value == 1251500
    assert workspace.grid(doc['id'])['grid']['sheets'][0]['cells']['C2']['v'] == 1300000


if __name__ == '__main__':
    import sys
    sys.exit(pytest.main([__file__, '-q']))
