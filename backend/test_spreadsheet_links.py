"""Explicit source refresh, pinned calculation evidence and human edit fences."""
import io
from pathlib import Path

import boot_paths  # noqa: F401
import pytest
from openpyxl import Workbook, load_workbook
from document_workspace import DocumentWorkspace
from office_sessions import DocumentConflict
from office_store import digest
from resource_links import ResourceLinks
from spreadsheet_workspace import SpreadsheetWorkspace


def book(value, name='매출'):
    wb = Workbook()
    wb.active.title = name
    wb.active.append(['품목', '금액'])
    wb.active.append(['00123', value])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def args(session):
    return dict(session_id=session['id'], client_id=session['client_id'],
                epoch=session['engine_epoch'], expected=session['session_revision'])


@pytest.fixture
def setup(tmp_path, monkeypatch):
    import document_office
    monkeypatch.setattr(document_office, 'available', lambda: True)
    sheets = SpreadsheetWorkspace(tmp_path / 'office')
    path = tmp_path / 'book.xlsx'
    path.write_bytes(book(10))
    source = sheets.open(path)['document']
    s = sheets.acquire(source['id'], 'sheet-window')['session']
    snap = sheets.snapshot(source['id'], **args(s), engine_state='initial', calculation='fresh')
    report = sheets.create_report(source['id'], snap['id'], '1', 'A1:B2', 'report', linked=True)
    documents = DocumentWorkspace(sheets.store.root)
    target = report['document']
    session = documents.acquire(target['id'], 'document-window')['session']
    return sheets, documents, source, target, s, session, report['reference'], path


def selection(documents, target, session):
    snap = documents.snapshot(target['id'], **args(session))
    text = documents.detail(target['id'])['text']
    start = text.index('| A |')
    end = text.index('\n계산 상태:')
    return dict(snapshot_id=snap['id'], start=start, end=end,
                selected_sha256=digest(text[start:end].encode()))


def changed_snapshot(sheets, source, session, value=25, name='매출', calculation='fresh'):
    session.update(blob=sheets.store.blob(book(value, name)), session_revision=session['session_revision'] + 1, state='draft')
    sheets.store.put('session', session)
    return sheets.snapshot(source['id'], **args(session), engine_state='changed', calculation=calculation)


def test_live_snapshot_refresh_preserves_source_and_report_human_edits(setup):
    sheets, docs, source, target, s, t, ref, path = setup
    links = ResourceLinks(docs)
    old_file = path.read_bytes()
    original = docs.detail(target['id'])['text']
    t = docs.draft(target['id'], **args(t), operation_id='human', text=original+'\n사람의 해설 유지\n')['session']
    selected = selection(docs, target, t)
    snap = changed_snapshot(sheets, source, s, name='이름 변경')
    p = links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_snapshot_id=snap['id'])
    assert '| 00123 | 25 |' in p['replacement']
    assert links.links(target['id'])[0]['snapshot_id'] == ref['snapshot_id'], 'proposal is not application'
    assert docs.detail(target['id'])['text'].endswith('사람의 해설 유지\n')
    result = docs.apply(target['id'], p['id'], **args(t), operation_id='apply')
    assert result['text'].endswith('사람의 해설 유지\n')
    assert '| 00123 | 25 |' in result['text']
    refs = links.links(target['id'])
    assert len(refs) == 1 and refs[0]['id'] == ref['id']
    assert refs[0]['selector']['sheet'] == '이름 변경'
    assert refs[0]['snapshot_id'] == snap['id'] and refs[0]['provenance']['unsaved']
    assert refs[0]['previous_source']['snapshot_id'] == ref['snapshot_id']
    docs.save(target['id'], **args(result['session']), operation_id='save', expected_revision=target['revision_id'])
    assert '사람의 해설 유지' in Path(target['source_uri']).read_text()
    assert path.read_bytes() == old_file and load_workbook(path).active['B2'].value == 10


def test_refresh_does_not_overwrite_edit_after_proposal(setup):
    sheets, docs, source, target, s, t, ref, _ = setup
    snap = changed_snapshot(sheets, source, s)
    p = ResourceLinks(docs).refresh_sheet_proposal(target['id'], ref['id'], **selection(docs, target, t), source_snapshot_id=snap['id'])
    changed = docs.draft(target['id'], **args(t), operation_id='human', text='사람이 새로 작성')['session']
    with pytest.raises(DocumentConflict):
        docs.apply(target['id'], p['id'], **args(changed), operation_id='apply')
    assert docs.detail(target['id'])['text'] == '사람이 새로 작성'
    assert ResourceLinks(docs).links(target['id'])[0]['snapshot_id'] == ref['snapshot_id']


def test_stale_cross_resource_and_copy_refresh_are_rejected(setup):
    sheets, docs, source, target, s, t, ref, _ = setup
    links = ResourceLinks(docs)
    selected = selection(docs, target, t)
    snap = changed_snapshot(sheets, source, s, calculation='stale')
    with pytest.raises(DocumentConflict, match='계산'):
        links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_snapshot_id=snap['id'])
    p = links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_snapshot_id=snap['id'], allow_stale=True)
    assert p['resource_reference']['provenance']['calculation_state'] == 'stale'
    with pytest.raises(DocumentConflict, match='열린 시트'):
        links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_revision_id=source['revision_id'], allow_stale=True)
    with pytest.raises(PermissionError):
        links.refresh_sheet_proposal(source['id'], ref['id'], **selected, source_snapshot_id=snap['id'])
    bad = {**snap, 'id': 'wrong', 'document_id': target['id']}
    sheets.store.put('sheet_snapshot', bad)
    with pytest.raises(PermissionError):
        links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_snapshot_id='wrong')
    ref['linked'] = False
    sheets.store.put('resource_link', ref)
    with pytest.raises(DocumentConflict, match='고정 사본'):
        links.refresh_sheet_proposal(target['id'], ref['id'], **selected, source_snapshot_id=snap['id'])


def test_closed_source_refresh_and_deleted_source_report(setup):
    sheets, docs, source, target, s, t, ref, path = setup
    sheets.close(source['id'], **args(s))
    path.write_bytes(book(30))
    links = ResourceLinks(docs)
    assert links.status(ref['id'])['source_changed']
    revision = links.refresh(source['id'], source['revision_id'])
    p = links.refresh_sheet_proposal(target['id'], ref['id'], **selection(docs, target, t), source_revision_id=revision['revision_id'], allow_stale=True)
    assert '| 00123 | 30 |' in p['replacement']
    assert p['resource_reference']['provenance']['calculation_state'] == 'stored_cache_unverified'
    path.unlink()
    assert not links.status(ref['id'])['source_available']
    assert docs.detail(target['id'])['text'] and links.links(target['id'])


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, *(__import__('sys').argv[1:])]))
