"""Two actual app windows, live ONLYOFFICE and a saved Markdown consumer."""
import os
import re
from pathlib import Path
import socket
import threading
import time

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit local engine acceptance')
def test_live_linked_report_refresh(tmp_path, monkeypatch):
    import api_documents
    import api_spreadsheets
    import document_office
    import uvicorn
    from document_workspace import DocumentWorkspace
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from openpyxl import Workbook, load_workbook
    from playwright.sync_api import sync_playwright, expect
    from resource_links import ResourceLinks
    from spreadsheet_workspace import SpreadsheetWorkspace
    from urllib.parse import urlsplit

    sheets = SpreadsheetWorkspace(tmp_path / 'office')
    documents = DocumentWorkspace(sheets.store.root)
    source = tmp_path / '매출.xlsx'
    book = Workbook(); book.active.title = '매출'
    book.active.append(['품목', '금액']); book.active.append(['00123', 10])
    book.save(source); original = source.read_bytes()
    monkeypatch.setattr(api_spreadsheets, 'service', lambda: sheets)
    monkeypatch.setattr(api_documents, 'service', lambda: documents)
    app = FastAPI(); app.include_router(api_spreadsheets.router); app.include_router(api_documents.router)
    @app.get('/launcher/auth/session')
    def auth():
        return {'authenticated': True, 'external': False}
    @app.get('/health')
    def health():
        return {'status': 'ok'}
    html = (ROOT / 'frontend/dist/index.html').read_text().replace('<html', '<html data-indiebiz-surface="remote"', 1)
    @app.get('/')
    def index():
        return HTMLResponse(html)
    app.mount('/assets', StaticFiles(directory=ROOT / 'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('0.0.0.0', 0)); port = sock.getsockname()[1]
    cfg = document_office.settings(); assert cfg
    cfg = {**cfg, 'callback_origin': f"http://{urlsplit(cfg['callback_origin']).hostname}:{port}"}
    monkeypatch.setattr(document_office, 'settings', lambda: cfg)
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
    try:
        deadline = time.monotonic() + 10
        while not server.started and time.monotonic() < deadline:
            time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            context = browser.new_context(viewport={'width': 1600, 'height': 1100})
            page = context.new_page()
            # Replace only Electron's native window opener with a real browser
            # popup. Editing, snapshots, proposal review and saving are unmocked.
            try:
                page.goto(f'http://127.0.0.1:{port}/#/spreadsheets')
                page.get_by_label('로컬 파일 경로').fill(str(source))
                page.get_by_role('button', name='파일 열기', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=60000)
                expect(page.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                page.get_by_label('범위', exact=True).fill('A1:B2')
                page.get_by_role('button', name='최신 범위 읽기', exact=True).click()
                expect(page.locator('.sheet-review pre').first).to_contain_text('effective_value', timeout=30000)
                page.evaluate("window.electron = {openToolWindow: name => window.open('#/'+name, '_blank')};")
                page.get_by_label('보고서에 원본 연결 유지').check()
                with page.expect_popup() as popup, page.expect_response(lambda r: r.url.endswith('/report') and r.request.method == 'POST') as response:
                    page.get_by_role('button', name='문서에 표 보고서 만들기', exact=True).click()
                report = response.value.json(); assert response.value.ok
                ref = report['reference']; assert ref['linked']
                docpage = popup.value
                docpage.get_by_label('로컬 파일 경로').fill(report['document']['source_uri'])
                docpage.get_by_role('button', name='파일 열기', exact=True).click()
                editor = docpage.get_by_label('문서 원문', exact=True)
                expect(editor).to_be_editable(timeout=20000)
                before = editor.input_value()
                editor.fill(before + '\n사람의 해설 유지\n')
                # An unsaved change in the real spreadsheet editor.
                page.get_by_label('범위', exact=True).fill('B2')
                page.get_by_role('button', name='최신 범위 읽기', exact=True).click()
                expect(page.locator('.sheet-review pre').first).to_contain_text('effective_value', timeout=30000)
                page.get_by_text('직접 변경안 작성', exact=True).click()
                page.get_by_label('행·열 값 (JSON)').fill('[[25]]')
                page.get_by_role('button', name='변경안 만들기', exact=True).click()
                page.get_by_role('button', name='변경안 적용', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_contain_text('변경 묶음 적용됨', timeout=30000)
                assert source.read_bytes() == original
                editor.evaluate("el => {el.focus(); el.setSelectionRange(el.value.indexOf('| A |'), el.value.indexOf('\\n계산 상태:'));}")
                docpage.get_by_role('button', name='선택 고정', exact=True).click()
                refresh = docpage.get_by_role('button', name='선택 구간 갱신 제안', exact=True)
                expect(refresh).to_be_enabled(timeout=15000)
                refresh.click()
                expect(docpage.get_by_role('textbox', name='교체할 문구', exact=True)).to_have_value('| A | B |\n| --- | --- |\n| 품목 | 금액 |\n| 00123 | 25 |\n', timeout=40000)
                assert editor.input_value() == before + '\n사람의 해설 유지\n', 'proposal must not change draft'
                docpage.get_by_role('button', name='제안 적용', exact=True).click()
                expect(editor).to_have_value(re.compile(r'.*\| 00123 \| 25 \|.*사람의 해설 유지.*', re.S))
                docpage.get_by_role('button', name='원본 저장', exact=True).click()
                expect(docpage.locator('footer[role=status]')).to_contain_text('저장', timeout=15000)
                stored = Path(report['document']['source_uri']).read_text()
                assert '| 00123 | 25 |' in stored and '사람의 해설 유지' in stored
                references = ResourceLinks(documents).links(report['document']['id'])
                assert len(references) == 1 and references[0]['id'] == ref['id']
                assert references[0]['provenance']['unsaved']
                assert references[0]['provenance']['calculation_state'] == 'fresh'
                assert source.read_bytes() == original and load_workbook(source).active['B2'].value == 10
                print('LINK_ACCEPTANCE', {'report': report['document']['source_uri'], 'reference': references[0]['id'],
                                           'snapshot': references[0]['snapshot_id'], 'value': 25, 'source_saved_value': 10})
            finally:
                browser.close()
    finally:
        server.should_exit = True; thread.join(15); sock.close()


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, *(__import__('sys').argv[1:])]))
