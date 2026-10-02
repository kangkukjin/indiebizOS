"""Printable templates and live ONLYOFFICE PDF/save acceptance on synthetic files."""
import io
import os
import re
import socket
import threading
import time
from pathlib import Path

import boot_paths  # noqa: F401
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('template,title', [('print_quotation', '견적서'), ('invoice', '청구서')])
def test_print_template(tmp_path, template, title):
    from spreadsheet_workspace import SpreadsheetWorkspace
    workspace = SpreadsheetWorkspace(tmp_path / 'office')
    result = workspace.create('양식', template)
    book = load_workbook(result['document']['source_uri'])
    sheet = book[title]
    assert sheet['A1'].value == title
    assert sheet['F27'].value == '=SUM(F25:F26)'
    assert sheet['D26'].value == 0, 'Tax rate must be chosen by the user'
    assert sheet['D10'].protection.locked is False
    assert sheet['F10'].protection.locked is True
    assert sheet.protection.sheet and not sheet.protection.password
    assert sheet.print_title_rows == '$1:$9'
    assert sheet.oddFooter.center.text == '&P 쪽'
    assert '$A$1:$F$31' in sheet.print_area
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert len(sheet.data_validations.dataValidation) == 1
    assert 'A1:F2' in sheet.merged_cells
    assert book['사용 안내']['A6'].value


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit local app acceptance')
@pytest.mark.parametrize('multipage', [False, True, 'structure', 'manual'], ids=['singlepage', 'multipage', 'structure', 'manual'])
def test_printable_form_app_save_pdf(tmp_path, monkeypatch, multipage):
    import api_spreadsheets
    import document_office
    import pymupdf
    import uvicorn
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from openpyxl.drawing.image import Image
    from PIL import Image as Raster, ImageDraw
    from playwright.sync_api import sync_playwright, expect
    from spreadsheet_workspace import SpreadsheetWorkspace
    from urllib.parse import urlsplit

    workspace = SpreadsheetWorkspace(tmp_path / 'office')
    source = Path(workspace.create('인쇄 인수', 'print_quotation')['document']['source_uri'])
    # A generated fixture logo avoids accessing or modifying user files.
    logo = Raster.new('RGB', (96, 32), '#245B48')
    ImageDraw.Draw(logo).text((8, 8), 'TEST', fill='white')
    png = io.BytesIO(); logo.save(png, format='PNG'); png.seek(0)
    book = load_workbook(source)
    book.active.add_image(Image(png), 'E3')
    structural = multipage in ('structure', 'manual')
    total_cell = 'F28' if structural else 'F27'
    expected_total = 66000 if structural else 60500
    if multipage is True:
        # Tall item rows force natural page breaks without altering formulas.
        for row in range(10, 25):
            book.active.row_dimensions[row].height = 85
    # Register the completed fixture for the first time; editing an already
    # registered source correctly triggers the external-write conflict fence.
    source = tmp_path / '견적서.xlsx'
    book.save(source)
    original = source.read_bytes()
    monkeypatch.setattr(api_spreadsheets, 'service', lambda: workspace)
    app = FastAPI(); app.include_router(api_spreadsheets.router)
    @app.get('/launcher/auth/session')
    def auth():
        return {'authenticated': True, 'external': False}
    @app.get('/health')
    def health():
        return {'status': 'ok'}
    html = (ROOT / 'frontend/dist/index.html').read_text().replace('<html', '<html data-indiebiz-surface="remote"', 1)
    @app.get('/')
    def index():
        return HTMLResponse(html)
    app.mount('/assets', StaticFiles(directory=ROOT / 'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('0.0.0.0', 0)); port = sock.getsockname()[1]
    cfg = document_office.settings(); assert cfg
    cfg = {**cfg, 'callback_origin': f"http://{urlsplit(cfg['callback_origin']).hostname}:{port}"}
    monkeypatch.setattr(document_office, 'settings', lambda: cfg)
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
    thread.start()
    try:
        deadline = time.monotonic() + 10
        while not server.started and time.monotonic() < deadline:
            time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            page = browser.new_page(viewport={'width': 1600, 'height': 1100})
            # Observe the public download event without changing production code.
            page.add_init_script('''
                let sdk;
                Object.defineProperty(window, 'DocsAPI', {configurable:true,
                  get(){return sdk}, set(value){
                    sdk=value;
                    let Original=value.DocEditor;
                    Object.defineProperty(value,'DocEditor',{configurable:true,
                      set(ctor){Original=ctor;}, get(){return function(id,config){
                        config.events.onDownloadAs=e=>{window.acceptanceDownload=e.data;};
                        const editor=new Original(id,config);
                        window.acceptanceEditor=editor;
                        return editor;
                      };}});
                  }});
            ''')
            try:
                page.goto(f'http://127.0.0.1:{port}/#/spreadsheets')
                expect(page.get_by_label('템플릿', exact=True).locator('option[value=invoice]')).to_have_count(1)
                page.get_by_label('로컬 파일 경로').fill(str(source))
                page.get_by_role('button', name='파일 열기', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=60000)
                expect(page.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                plugin = next(f for f in page.frames if '/plugin/' in f.url and 'index.html' in f.url)
                edited = plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                    var sheet=Api.GetActiveSheet();
                    sheet.GetRange('B4').SetValue('한글 거래처');
                    sheet.GetRange('B10:E11').SetValue([['설계 작업','건',2,12500],['자료 정리','시간',3,10000]]);
                    sheet.GetRange('D26').SetValue(0.1);
                    return {total:sheet.GetRange('F27').GetValue(), values:sheet.GetRange('B10:F11').GetValue(), buyer:sheet.GetRange('B4').GetValue()};
                },false,true,resolve))''')
                assert edited['buyer'] == '한글 거래처', edited
                if structural:
                    editor_frame = next(f for f in page.frames if '/spreadsheeteditor/' in f.url)
                    editor_frame.get_by_role('tab', name='보호', exact=True).click()
                    editor_frame.get_by_text(re.compile(r'^시트\s*보호$')).click()
                    plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                        Api.GetActiveSheet().GetRange('12:12').Select();
                    },false,false,resolve))''')
                    editor_frame.locator('#area_id').focus()
                    page.keyboard.press('Control+Shift+Equal')
                    editor_frame.get_by_role('radio', name='전체 행', exact=True).click()
                    editor_frame.get_by_text('OK', exact=True).click()
                    inserted = plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                        var sheet=Api.GetActiveSheet();
                        sheet.GetRange('A11:F11').Copy(sheet.GetRange('A12:F12'));
                        sheet.GetRange('B12:E12').SetValue([['추가 품목','건',1,5000]]);
                        return sheet.GetRange('F26').GetFormula();
                    },false,true,resolve))''')
                    assert inserted.replace(' ', '') == '=SUM(F10:F25)', inserted
                # callCommand batches recalculation until its command closes.
                calculated = plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                    return Api.GetActiveSheet().GetRange('TOTAL_CELL').GetValue();
                },false,false,resolve))'''.replace('TOTAL_CELL', total_cell))
                assert calculated == str(expected_total), calculated
                if multipage == 'manual':
                    plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                        Api.GetActiveSheet().GetRange('A20').Select();
                    },false,false,resolve))''')
                    editor_frame.get_by_role('tab', name='레이아웃', exact=True).click()
                    editor_frame.get_by_text('나누기', exact=True).click()
                    editor_frame.get_by_text('페이지 나누기 삽입', exact=True).click()
                assert source.read_bytes() == original
                page.get_by_role('button', name='원본 저장', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_have_text('저장됨 · 원본 파일 기록 확인', timeout=60000)
                saved = load_workbook(source)
                sheet = saved['견적서']
                assert sheet['B4'].value == '한글 거래처'
                assert sheet['F10'].data_type == 'f' and sheet[total_cell].data_type == 'f'
                assert sheet.protection.sheet == (not structural)
                assert not sheet['D10'].protection.locked
                if structural:
                    assert '$A$1:$F$32' in sheet.print_area, sheet.print_area
                    assert sheet['F26'].value == '=SUM(F10:F25)'
                    if multipage == 'manual':
                        assert any(b.id == 19 and b.man for b in sheet.row_breaks.brk), sheet.row_breaks
                    assert sheet['B12'].value == '추가 품목'
                    assert sheet['A12'].value == '=ROW()-9'
                    assert sheet['F12'].data_type == 'f'
                    assert sheet['F12'].protection.locked and not sheet['D12'].protection.locked
                    cached = load_workbook(source, data_only=True)['견적서']
                    assert [cached.cell(row, 1).value for row in range(10, 26)] == list(range(1, 17))
                assert 'A1:F2' in sheet.merged_cells and sheet._images
                # ONLYOFFICE stores validations in OOXML x14 extensions;
                # openpyxl warns and drops that extension from its read model.
                import zipfile
                from xml.etree import ElementTree as ET
                with zipfile.ZipFile(source) as archive:
                    xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
                validations = [e for e in xml.iter() if e.tag.rsplit('}', 1)[-1] == 'dataValidation']
                from openpyxl.worksheet.cell_range import MultiCellRange
                unit_ranges = []
                for validation in validations:
                    if validation.get('type') == 'list' and '개,건,시간,일,월,식' in ''.join(validation.itertext()):
                        refs = validation.get('sqref') or next(
                            e.text for e in validation.iter() if e.tag.rsplit('}', 1)[-1] == 'sqref')
                        unit_ranges.append(MultiCellRange(refs))
                last_item_row = 25 if structural else 24
                assert all(any(f'C{row}' in refs for refs in unit_ranges)
                           for row in range(10, last_item_row + 1)), unit_ranges
                assert not any(f'C{last_item_row + 1}' in refs for refs in unit_ranges)
                assert load_workbook(source, data_only=True)['견적서'][total_cell].value == expected_total
                page.evaluate("window.acceptanceEditor.downloadAs('pdf')")
                editor_frame = next(f for f in page.frames if '/spreadsheeteditor/' in f.url)
                if structural:
                    for combo in editor_frame.get_by_role('combobox').all():
                        if combo.input_value() == '활성 시트':
                            combo.click()
                            break
                    editor_frame.get_by_text('모든 시트', exact=True).click()
                editor_frame.get_by_text('저장 및 다운로드', exact=True).click()
                page.wait_for_function('window.acceptanceDownload && window.acceptanceDownload.url', timeout=60000)
                pdf = document_office.download(page.evaluate('window.acceptanceDownload.url'), cfg)
                assert pdf.startswith(b'%PDF-'), 'Engine must return an actual PDF'
                (tmp_path / 'form.pdf').write_bytes(pdf)
                with pymupdf.open(stream=pdf, filetype='pdf') as doc:
                    texts = [p.get_text() for p in doc]
                    print('PRINT_PDF', {'pages': len(doc), 'text': texts})
                    if multipage is True:
                        assert len(doc) > 1, 'Tall rows must produce multiple pages'
                        assert all('견적서' in t and '품목 / 내용' in t for t in texts), texts
                    elif multipage == 'manual':
                        assert len(doc) == 3, 'Manual break must split the form before the help sheet'
                        assert all('견적서' in t and '품목 / 내용' in t for t in texts[:2])
                        assert '66,000' in texts[1] and '사용 안내' in texts[2]
                    elif structural:
                        assert len(doc) == 2, 'Workbook must include the form and its help sheet'
                        assert '사용 안내' in texts[1]
                    else:
                        assert len(doc) == 1, 'Selected form sheet must fit one A4 page'
                    assert '견적서' in texts[0] and '한글 거래처' in texts[0]
                    assert f'{expected_total:,}' in ''.join(texts)
                    if not structural:
                        assert all('사용 안내' not in t for t in texts)
                    assert any(i['width'] == 96 and i['height'] == 32 for i in doc[0].get_image_info()), 'Fixture logo, including inline PDF images, must survive'
                    assert all(t.strip() for t in texts), 'No blank pages'
                    for index, text in enumerate(texts, 1):
                        if not structural or index == 1 or (multipage == 'manual' and index == 2):
                            assert f'{index} 쪽' in text, 'PDF page number must match its actual position'
                        doc[index - 1].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(
                            str(tmp_path / f'form-{index}.png'))
                    doc[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(str(tmp_path / 'form.png'))
                import subprocess
                out = tmp_path / 'independent'; out.mkdir()
                saved_bytes = source.read_bytes()
                reopened = subprocess.run(['soffice', '-env:UserInstallation=' + (tmp_path / 'lo-profile').as_uri(),
                                           '--headless', '--convert-to', 'xlsx', '--outdir', str(out), str(source)],
                                          capture_output=True, text=True, timeout=90)
                assert reopened.returncode == 0, reopened.stderr
                independent = load_workbook(out / source.name, data_only=True)['견적서']
                assert independent[total_cell].value == expected_total
                independent_structure = load_workbook(out / source.name)['견적서']
                assert independent_structure._images
                assert independent_structure.protection.sheet == (not structural)
                assert independent_structure.data_validations.dataValidation
                if structural:
                    assert '$A$1:$F$32' in independent_structure.print_area
                    assert [independent.cell(row, 1).value for row in range(10, 26)] == list(range(1, 17))
                    assert all(any(f'C{row}' in rule.sqref for rule in independent_structure.data_validations.dataValidation)
                               for row in range(10, 26))
                    if multipage == 'manual':
                        assert any(b.id == 19 and b.man for b in independent_structure.row_breaks.brk)
                assert source.read_bytes() == saved_bytes, 'Independent consumer must not overwrite app output'
                print('PRINT_ACCEPTANCE', {'source': str(source), 'pdf': str(tmp_path / 'form.pdf'),
                                          'image': str(tmp_path / 'form.png'), 'total': expected_total})
            finally:
                print('APP_ALERTS', page.get_by_role('alert').all_text_contents())
                for frame in page.frames:
                    if '/spreadsheeteditor/' in frame.url:
                        print('EDITOR_UI', frame.locator('body').inner_text()[-6500:])
                browser.close()
    finally:
        server.should_exit = True
        thread.join(timeout=15)
        sock.close()


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
