"""Recovery preserves original/draft bytes across fenced claims and process death."""
import errno
import io
import os
from pathlib import Path
import subprocess
import sys

import boot_paths  # noqa: F401
import pytest
from openpyxl import load_workbook
from office_sessions import DocumentConflict
from spreadsheet_workspace import SpreadsheetWorkspace
from test_spreadsheet_workspace import args, data, opened, workspace  # noqa: F401


def draft(app, session, value=20):
    session.update(blob=app.store.blob(data(value)), session_revision=1, state='draft')
    app.store.put('session', session)


def test_reclaim_fences_stale_window_and_stale_review(workspace, tmp_path):
    path, doc, old = opened(workspace, tmp_path)
    original = path.read_bytes()
    draft(workspace, old)
    with pytest.raises(DocumentConflict, match='초안이 바뀌었'):
        workspace.reclaim(doc['id'], 'new', old['engine_epoch'], expected_revision=0)
    new = workspace.reclaim(doc['id'], 'new', old['engine_epoch'], expected_revision=1)['session']
    assert new['blob'] == old['blob'] and new['engine_epoch'] != old['engine_epoch']
    with pytest.raises(DocumentConflict):
        workspace.save(doc['id'], **args(old), operation_id='stale', expected_revision=doc['revision_id'])
    with pytest.raises(DocumentConflict):
        workspace.reclaim(doc['id'], 'third', old['engine_epoch'], expected_revision=1)
    assert path.read_bytes() == original


def test_restore_keeps_pre_restore_draft_as_selectable_version(workspace, tmp_path):
    path, doc, session = opened(workspace, tmp_path)
    original = path.read_bytes()
    draft(workspace, session)
    r = workspace.restore(doc['id'], doc['revision_id'], **args(session), operation_id='restore')
    candidates = [v for v in workspace.versions(doc['id']) if v.get('label') == '복구 직전 초안']
    assert len(candidates) == 1 and candidates[0]['blob'] == session['blob']
    assert path.read_bytes() == original
    recovered = workspace.restore(doc['id'], candidates[0]['id'], **args(r['session']), operation_id='restore-draft')
    assert recovered['session']['blob'] == session['blob']
    assert path.read_bytes() == original


@pytest.mark.parametrize('failure', ['allocation', 'publication'])
def test_disk_full_retains_original_and_allows_explicit_retry(workspace, tmp_path, monkeypatch, failure):
    import office_sessions
    path, doc, session = opened(workspace, tmp_path)
    original = path.read_bytes()
    draft(workspace, session)
    def no_space(*a, **kw):
        raise OSError(errno.ENOSPC, 'No space left on device')
    with monkeypatch.context() as patch:
        if failure == 'allocation':
            patch.setattr(office_sessions.tempfile, 'mkstemp', no_space)
        else:
            patch.setattr(office_sessions.os, 'replace', no_space)
        with pytest.raises(OSError) as caught:
            workspace.save(doc['id'], **args(session), operation_id='failed', expected_revision=doc['revision_id'])
        assert caught.value.errno == errno.ENOSPC
    assert path.read_bytes() == original
    assert workspace.store.bytes(session['blob'])
    result = SpreadsheetWorkspace(workspace.store.root).recover(doc['id'])
    assert result['items'] == [{'state': 'not_written', 'draft_preserved': True}]
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('.indiebiz-document-*'))
    workspace.save(doc['id'], **args(session), operation_id='retry', expected_revision=doc['revision_id'])
    assert load_workbook(path).active['B1'].value == 20


def _writer(root, doc_id, stage):
    import office_sessions
    app = SpreadsheetWorkspace(root)
    detail = app.detail(doc_id)
    real_replace = office_sessions.os.replace
    def blocked(src, dst):
        if stage == 'after':
            real_replace(src, dst)
        print('publication-boundary', flush=True)
        import time
        time.sleep(120)
        if stage == 'before':
            real_replace(src, dst)
    office_sessions.os.replace = blocked
    app.save(doc_id, **args(detail['session']), operation_id='crash', expected_revision=detail['document']['revision_id'])


@pytest.mark.parametrize('stage', ['before', 'after'])
def test_process_killed_during_save_recovers_without_rewriting(workspace, tmp_path, stage):
    path, doc, session = opened(workspace, tmp_path)
    original = path.read_bytes()
    draft(workspace, session)
    cmd = 'import boot_paths; from test_spreadsheet_recovery import _writer; import sys; _writer(*sys.argv[1:])'
    child = subprocess.Popen([sys.executable, '-c', cmd, str(workspace.store.root), doc['id'], stage],
        env={**os.environ, 'PYTHONPATH': str(Path(__file__).parent)}, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        # A bounded readiness wait; terminate the real child at the durable save boundary.
        import selectors
        with selectors.DefaultSelector() as selector:
            selector.register(child.stdout, selectors.EVENT_READ)
            assert selector.select(20), 'save child did not reach publication'
            assert child.stdout.readline().strip() == 'publication-boundary'
        child.kill(); child.wait(timeout=10)
        assert child.returncode < 0
        before_recovery = path.read_bytes()
        assert (before_recovery == original) == (stage == 'before')
        restarted = SpreadsheetWorkspace(workspace.store.root)
        result = restarted.recover(doc['id'])
        assert result['items'][0]['state'] == ('not_written' if stage == 'before' else 'saved')
        assert path.read_bytes() == before_recovery
        assert restarted.store.bytes(restarted.detail(doc['id'])['session']['blob']) == workspace.store.bytes(session['blob'])
        assert restarted.recover(doc['id'])['items'] == []
        assert load_workbook(path).active['B1'].value == (10 if stage == 'before' else 20)
    finally:
        if child.poll() is None:
            child.kill(); child.wait(timeout=10)
        child.stdout.close(); child.stderr.close()


def test_recovery_does_not_overwrite_external_change(workspace, tmp_path, monkeypatch):
    import office_sessions
    path, doc, session = opened(workspace, tmp_path)
    draft(workspace, session)
    with monkeypatch.context() as patch:
        patch.setattr(office_sessions.tempfile, 'mkstemp', lambda **kw: (_ for _ in ()).throw(OSError(errno.ENOSPC, 'full')))
        with pytest.raises(OSError):
            workspace.save(doc['id'], **args(session), operation_id='failed', expected_revision=doc['revision_id'])
    external = data(99); path.write_bytes(external)
    result = workspace.recover(doc['id'])
    assert result['items'] == [{'state': 'conflict', 'draft_preserved': True}]
    assert path.read_bytes() == external
    assert load_workbook(io.BytesIO(workspace.store.bytes(session['blob']))).active['B1'].value == 20


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
