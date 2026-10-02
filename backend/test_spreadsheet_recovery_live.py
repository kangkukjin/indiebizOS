"""Real browser crash, new owner recovery and offline reconnection."""
import os
import re
from pathlib import Path
import socket
import threading
import time

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit local engine acceptance')
def test_live_crash_and_offline_recovery(tmp_path, monkeypatch):
    import api_documents
    import api_spreadsheets
    import document_office
    import uvicorn
    from document_workspace import DocumentWorkspace
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from openpyxl import Workbook, load_workbook
    from playwright.sync_api import sync_playwright, expect
    from resource_links import ResourceLinks
    from spreadsheet_workspace import SpreadsheetWorkspace
    from urllib.parse import urlsplit

    sheets = SpreadsheetWorkspace(tmp_path / 'office')
    documents = DocumentWorkspace(sheets.store.root)
    source = tmp_path / '매출.xlsx'
    book = Workbook(); book.active.title = '매출'
    book.active.append(['품목', '금액', '불리언', '계산']); book.active.append(['00123', 10, False, '=B2*2'])
    book.save(source); original = source.read_bytes()
    monkeypatch.setattr(api_spreadsheets, 'service', lambda: sheets)
    monkeypatch.setattr(api_documents, 'service', lambda: documents)
    app = FastAPI(); app.include_router(api_spreadsheets.router); app.include_router(api_documents.router)
    @app.get('/launcher/auth/session')
    def auth():
        return {'authenticated': True, 'external': False}
    @app.get('/health')
    def health():
        return {'status': 'ok'}
    html = (ROOT / 'frontend/dist/index.html').read_text().replace('<html', '<html data-indiebiz-surface="remote"', 1)
    @app.get('/')
    def index():
        return HTMLResponse(html)
    app.mount('/assets', StaticFiles(directory=ROOT / 'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('0.0.0.0', 0)); port = sock.getsockname()[1]
    cfg = document_office.settings(); assert cfg
    cfg = {**cfg, 'callback_origin': f"http://{urlsplit(cfg['callback_origin']).hostname}:{port}"}
    monkeypatch.setattr(document_office, 'settings', lambda: cfg)
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
    try:
        deadline = time.monotonic() + 10
        while not server.started and time.monotonic() < deadline:
            time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            context = browser.new_context(viewport={'width': 1600, 'height': 1100})
            page = context.new_page()
            # Editing, snapshots, recovery and saving use the real app and engine.
            try:
                page.goto(f'http://127.0.0.1:{port}/#/spreadsheets')
                page.get_by_label('로컬 파일 경로').fill(str(source))
                page.get_by_role('button', name='파일 열기', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=60000)
                expect(page.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                page.get_by_label('범위', exact=True).fill('B2')
                page.get_by_role('button', name='최신 범위 읽기', exact=True).click()
                expect(page.locator('.sheet-review pre').first).to_contain_text('effective_value', timeout=30000)
                page.get_by_text('직접 변경안 작성', exact=True).click()
                page.get_by_label('행·열 값 (JSON)').fill('[[25]]')
                page.get_by_role('button', name='변경안 만들기', exact=True).click()
                expect(page.get_by_role('button', name='변경안 적용', exact=True)).to_be_visible()
                page.get_by_role('button', name='변경안 적용', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_contain_text('변경 묶음 적용됨', timeout=40000)
                page.get_by_role('button', name='작업 저장', exact=True).click()
                expect(page.locator('.sheet-editor [role=status]')).to_have_text('복구 초안 저장됨', timeout=30000)
                doc = sheets.list()[0]
                old = sheets.detail(doc['id'])['session']
                assert source.read_bytes() == original
                # Kill only this test's browser process tree. CDP Page.crash
                # itself can hang and is not a reliable completion boundary.
                import psutil
                roots = [p for p in psutil.Process().children(recursive=True)
                         if '--remote-debugging-pipe' in p.cmdline()]
                ids = {p.pid for p in roots}
                roots = [p for p in roots if not any(a.pid in ids for a in p.parents())]
                assert len(roots) == 1
                victims = roots[0].children(recursive=True) + roots
                for process in victims:
                    try:
                        process.kill()
                    except psutil.NoSuchProcess:
                        pass
                _, alive = psutil.wait_procs(victims, timeout=10)
                assert not alive
                browser = pw.chromium.launch(headless=True)
                fresh_context = browser.new_context(viewport={'width': 1600, 'height': 1100})
                newpage = fresh_context.new_page()
                newpage.goto(f'http://127.0.0.1:{port}/#/spreadsheets')
                newpage.get_by_label('로컬 파일 경로').fill(str(source))
                newpage.get_by_role('button', name='파일 열기', exact=True).click()
                recovery = newpage.get_by_role('region', name='스프레드시트 복구')
                expect(recovery).to_be_visible(timeout=20000)
                expect(recovery.get_by_role('button', name='복구 상태 다시 확인')).to_be_enabled()
                if sheets.detail(doc['id'])['session']['state'] == 'recovering':
                    choices = [v for v in sheets.versions(doc['id']) if v.get('label') == '마지막 확인 초안']
                    assert choices
                    recovery.get_by_label('복구할 버전').select_option(choices[-1]['id'])
                recovery.get_by_role('checkbox').check()
                recovery.get_by_role('button', name='선택한 초안으로 다시 연결').click()
                expect(newpage.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=60000)
                expect(newpage.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                from office_sessions import DocumentConflict
                from test_spreadsheet_workspace import args
                with pytest.raises(DocumentConflict):
                    sheets.save(doc['id'], **args(old), operation_id='old-window', expected_revision=doc['revision_id'])
                assert source.read_bytes() == original
                newpage.get_by_label('범위', exact=True).fill('A2:D2')
                with newpage.expect_response(lambda r: r.url.endswith('/snapshot') and r.request.method == 'POST') as response:
                    newpage.get_by_role('button', name='최신 범위 읽기', exact=True).click()
                snapshot = response.value.json()
                assert response.value.ok and snapshot['calc_status'] == 'fresh'
                values = sheets.read(doc['id'], snapshot['id'], '1', 'A2:D2')['items']
                assert values[0]['entered_value'] == '00123'
                assert values[1]['effective_value'] == 25
                assert values[2]['entered_value'] is False
                assert values[3]['effective_value'] == 50
                expect(newpage.locator('.sheet-review pre').first).to_contain_text('effective_value', timeout=30000)
                # The new recovery surface also works when API requests have failed.
                fresh_context.set_offline(True)
                newpage.get_by_role('button', name='연결·저장 복구', exact=True).click()
                expect(recovery.get_by_role('alert')).to_be_visible(timeout=15000)
                assert source.read_bytes() == original
                fresh_context.set_offline(False)
                recovery.get_by_role('button', name='복구 상태 다시 확인').click()
                expect(recovery.get_by_role('button', name='복구 상태 다시 확인')).to_be_enabled(timeout=15000)
                recovery.get_by_role('checkbox').check()
                recovery.get_by_role('button', name='선택한 초안으로 다시 연결').click()
                expect(newpage.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=60000)
                expect(newpage.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                newpage.get_by_role('button', name='원본 저장', exact=True).click()
                expect(newpage.locator('.sheet-editor [role=status]')).to_contain_text('저장됨 · 원본 파일 기록 확인', timeout=30000)
                saved = load_workbook(source)
                assert saved.active['A2'].value == '00123'
                assert saved.active['B2'].value == 25 and saved.active['C2'].value is False
                assert saved.active['D2'].value == '=B2*2'
                assert load_workbook(source, data_only=True).active['D2'].value == 50
                import subprocess
                out = tmp_path / 'independent'; out.mkdir()
                saved_bytes = source.read_bytes()
                checked = subprocess.run(['soffice', '-env:UserInstallation=' + (tmp_path / 'lo-profile').as_uri(),
                    '--headless', '--convert-to', 'xlsx', '--outdir', str(out), str(source)],
                    capture_output=True, text=True, timeout=90)
                assert checked.returncode == 0, checked.stderr
                independent = load_workbook(out / source.name, data_only=True)
                assert independent.active['A2'].value == '00123'
                assert independent.active['B2'].value == 25 and independent.active['C2'].value is False
                assert independent.active['D2'].value == 50
                assert source.read_bytes() == saved_bytes
                print('RECOVERY_ACCEPTANCE', {'browser_process_killed': True, 'libreoffice_reopened': True, 'old_writer_fenced': True,
                    'offline_reconnected': True, 'confirmed_draft': 25, 'formula_cache': 50,
                    'original_preserved_until_save': True})
                fresh_context.close()
            finally:
                browser.close()
    finally:
        server.should_exit = True; thread.join(15); sock.close()


if __name__ == '__main__':
    raise SystemExit(pytest.main([__file__, *(__import__('sys').argv[1:])]))
