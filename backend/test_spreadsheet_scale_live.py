"""Measured scale acceptance against a real app and engine; synthetic files only."""
import hashlib
import json
import os
import platform
import socket
import threading
import time
from pathlib import Path

import boot_paths  # noqa: F401
import pytest

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.system
@pytest.mark.skipif(os.environ.get('INDIEBIZ_OFFICE_LIVE_TEST') != '1', reason='explicit scale measurement')
def test_scale_measurements(tmp_path, monkeypatch):
    import psutil
    import api_spreadsheets
    import document_office
    import uvicorn
    from fastapi import FastAPI
    from fastapi.responses import HTMLResponse
    from fastapi.staticfiles import StaticFiles
    from playwright.sync_api import sync_playwright, expect
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from spreadsheet_workspace import SpreadsheetWorkspace
    from urllib.parse import urlsplit

    sources = []
    for rows, formulas, charts in [(10000, 2000, 0), (100000, 50000, 5)]:
        book = Workbook(write_only=True); sheet = book.create_sheet('Benchmark')
        for r in range(1, rows + 1):
            sheet.append([r % 100 for _ in range(19)] + [f'=A{r}+B{r}' if r <= formulas else r % 100])
        for i in range(charts):
            chart = BarChart(); chart.add_data(Reference(sheet, min_col=1, max_col=2, min_row=1, max_row=20))
            sheet.add_chart(chart, f'V{i*16+1}')
        path = tmp_path / f'scale-{rows}.xlsx'; book.save(path)
        sources.append((path, rows, formulas, charts, hashlib.sha256(path.read_bytes()).hexdigest()))
    workspace = SpreadsheetWorkspace(tmp_path / 'office')
    monkeypatch.setattr(api_spreadsheets, 'service', lambda: workspace)
    app = FastAPI(); app.include_router(api_spreadsheets.router)
    @app.get('/launcher/auth/session')
    def auth(): return {'authenticated': True, 'external': False}
    @app.get('/health')
    def health(): return {'status': 'ok'}
    html = (ROOT/'frontend/dist/index.html').read_text().replace('<html', '<html data-indiebiz-surface="remote"', 1)
    @app.get('/')
    def index(): return HTMLResponse(html)
    app.mount('/assets', StaticFiles(directory=ROOT/'frontend/dist/assets'))
    sock = socket.socket(); sock.bind(('0.0.0.0', 0)); port = sock.getsockname()[1]
    cfg = document_office.settings(); assert cfg
    monkeypatch.setattr(document_office, 'settings', lambda: {**cfg, 'callback_origin': f"http://{urlsplit(cfg['callback_origin']).hostname}:{port}"})
    server = uvicorn.Server(uvicorn.Config(app, log_level='error'))
    thread = threading.Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True); thread.start()
    measures = []
    environment = {'platform': platform.platform(), 'machine': platform.machine(),
        'host_memory_bytes': psutil.virtual_memory().total, 'engine': 'ONLYOFFICE 9.3.1',
        'memory_scope': 'test process plus Chromium descendants; engine VM excluded',
        'input_scope': 'official plugin command roundtrip, not keyboard/IME paint latency'}
    try:
        deadline = time.monotonic() + 10
        while not server.started and time.monotonic() < deadline: time.sleep(.05)
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=True)
            try:
                for path, rows, formulas, charts, original in sources:
                    page = browser.new_page(viewport={'width': 1600, 'height': 1100})
                    result = {'rows': rows, 'columns': 20, 'formulas': formulas, 'charts': charts,
                              'source_bytes': path.stat().st_size, 'first_display_target_s': 3 if rows == 10000 else 10}
                    peak = [0]; stop = threading.Event()
                    def sample():
                        while not stop.wait(.05):
                            processes = [psutil.Process()] + psutil.Process().children(recursive=True)
                            rss = 0
                            for process in processes:
                                try: rss += process.memory_info().rss
                                except psutil.Error: pass
                            peak[0] = max(peak[0], rss)
                    sampler = threading.Thread(target=sample, daemon=True); sampler.start()
                    try:
                        page.goto(f'http://127.0.0.1:{port}/#/spreadsheets')
                        page.get_by_label('로컬 파일 경로').fill(str(path))
                        started = time.perf_counter()
                        page.get_by_role('button', name='파일 열기', exact=True).click()
                        expect(page.locator('.sheet-editor [role=status]')).to_have_text('편집 준비 완료', timeout=120000)
                        result['first_display_s'] = round(time.perf_counter() - started, 3)
                        result['display_target_met'] = result['first_display_s'] <= result['first_display_target_s']
                        expect(page.get_by_role('button', name='계산·스냅샷', exact=True)).to_be_enabled(timeout=20000)
                        plugin = next(f for f in page.frames if '/plugin/' in f.url and 'index.html' in f.url)
                        timings = []
                        for i in range(20):
                            started = time.perf_counter()
                            plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                                Api.GetActiveSheet().GetRange('A1').SetValue(42);
                                return {value:Api.GetActiveSheet().GetRange('A1').GetValue2()};
                            },false,true,resolve))''')
                            timings.append((time.perf_counter() - started) * 1000)
                        result['command_p95_ms'] = round(sorted(timings)[18], 2)
                        started = time.perf_counter()
                        plugin.evaluate('''() => new Promise(resolve => Asc.plugin.callCommand(function(){
                            Api.RecalculateAllFormulas(); return true;
                        },false,true,resolve))''')
                        result['recalculate_command_s'] = round(time.perf_counter() - started, 3)
                        result['calculation_completion_verified'] = False
                        result['status'] = 'measured'
                    except Exception as exc:
                        result.update(status='failed', error=str(exc))
                    finally:
                        stop.set(); sampler.join(timeout=2)
                        result['peak_test_and_browser_rss_bytes'] = peak[0]
                        result['original_preserved'] = hashlib.sha256(path.read_bytes()).hexdigest() == original
                        measures.append(result); page.close()
            finally: browser.close()
    finally:
        server.should_exit = True; thread.join(timeout=10); sock.close()
    print('SCALE_EVIDENCE ' + json.dumps({'environment': environment, 'measurements': measures}, ensure_ascii=False))
    assert all(r['original_preserved'] for r in measures)
    assert all(r['status'] == 'measured' for r in measures), measures
    # Performance target misses remain explicit evidence, not failures hidden by retries.


if __name__ == '__main__':
    import sys
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
