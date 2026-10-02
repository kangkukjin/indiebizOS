"""Shared office ownership, spreadsheet types and conditional publication."""
import io
import zipfile

import boot_paths  # noqa: F401
import pytest
from openpyxl import Workbook, load_workbook

from spreadsheet_workspace import SpreadsheetWorkspace, SpreadsheetEngine
from office_sessions import DocumentConflict, DocumentUnsupported
from spreadsheet_files import projection, validate


def data(value=10):
    book=Workbook(); sheet=book.active
    sheet['A1']='00123'; sheet['B1']=value; sheet['C1']=False; sheet['D1']='=B1*2'
    buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()


@pytest.fixture
def workspace(tmp_path,monkeypatch):
    import document_office
    monkeypatch.setattr(document_office,'available',lambda:True)
    return SpreadsheetWorkspace(tmp_path/'office')


def opened(workspace,tmp_path):
    path=tmp_path/'book.xlsx';path.write_bytes(data())
    doc=workspace.open(path)['document'];detail=workspace.acquire(doc['id'],'window1')
    return path,doc,detail['session']


def args(s):
    return dict(session_id=s['id'],client_id=s['client_id'],epoch=s['engine_epoch'],expected=s['session_revision'])


def test_projection_preserves_identifiers_boolean_zero_formula():
    rows=projection(data(0),'1','A1:E1')['items']
    assert rows[0]['entered_value']=='00123'
    assert rows[1]['entered_value']==0 and rows[1]['value_type']=='number'
    assert rows[2]['entered_value'] is False
    assert rows[3]['formula']=='=B1*2' and rows[3]['effective_value'] is None
    assert rows[4]['entered_value'] is None and rows[4]['value_type']=='blank'


def test_common_registration_and_writer_fence(workspace,tmp_path):
    from document_workspace import DocumentWorkspace
    path,doc,s=opened(workspace,tmp_path)
    other=DocumentWorkspace(workspace.store.root)
    assert other.open(path)['document']['id']==doc['id']
    with pytest.raises(DocumentConflict):
        workspace.acquire(doc['id'],'window2')
    assert path.read_bytes()==data() or load_workbook(path).active['B1'].value==10


def test_proposal_fences_revision_and_replay(workspace,tmp_path):
    path,doc,s=opened(workspace,tmp_path)
    snap=workspace.snapshot(doc['id'],**args(s),engine_state='synthetic engine state',calculation='stale')
    p=workspace.propose(doc['id'],snap['id'],'1','B1',[[22]])
    queued=workspace.apply(doc['id'],p['id'],**args(s),operation_id='one')
    assert queued['status']=='queued'
    s['session_revision']+=1;workspace.store.put('session',s)
    with pytest.raises(DocumentConflict):
        workspace.apply(doc['id'],p['id'],**args(s),operation_id='two')
    assert load_workbook(path).active['B1'].value==10


def test_cross_resource_snapshot_rejected(workspace,tmp_path):
    _,doc,s=opened(workspace,tmp_path)
    snap=workspace.snapshot(doc['id'],**args(s),engine_state='state',calculation='stale')
    other=tmp_path/'other.xlsx';other.write_bytes(data())
    other_id=workspace.open(other)['document']['id']
    with pytest.raises(PermissionError):
        workspace.read(other_id,snap['id'],'1','A1')


def test_external_change_preserves_draft_and_original(workspace,tmp_path):
    path,doc,s=opened(workspace,tmp_path)
    s.update(blob=workspace.store.blob(data(20)),session_revision=1,state='draft');workspace.store.put('session',s)
    external=data(99);path.write_bytes(external)
    with pytest.raises(DocumentConflict):
        workspace.save(doc['id'],**args(s),operation_id='save',expected_revision=doc['revision_id'])
    assert path.read_bytes()==external
    assert workspace.store.bytes(s['blob'])==data(20) or load_workbook(io.BytesIO(workspace.store.bytes(s['blob']))).active['B1'].value==20


def test_save_idempotence_and_restore(workspace,tmp_path):
    path,doc,s=opened(workspace,tmp_path)
    s.update(blob=workspace.store.blob(data(20)),session_revision=1,state='draft');workspace.store.put('session',s)
    a=workspace.save(doc['id'],**args(s),operation_id='save',expected_revision=doc['revision_id'])
    b=workspace.save(doc['id'],**args(s),operation_id='save',expected_revision=doc['revision_id'])
    assert a==b and load_workbook(path).active['B1'].value==20
    restored=workspace.restore(doc['id'],doc['revision_id'],**args(s),operation_id='restore')
    assert restored['session']['engine_epoch']!=s['engine_epoch']
    assert load_workbook(path).active['B1'].value==20


def test_external_resource_is_blocked():
    source=io.BytesIO(data());result=io.BytesIO()
    with zipfile.ZipFile(source) as original,zipfile.ZipFile(result,'w') as changed:
        for name in original.namelist():changed.writestr(name,original.read(name))
        changed.writestr('xl/externalLinks/externalLink1.xml','<externalLink/>')
    with pytest.raises(DocumentUnsupported):validate(result.getvalue())


def test_unordered_close_is_a_recovery_candidate_not_a_newer_draft(workspace,tmp_path,monkeypatch):
    import document_office
    import jwt
    cfg={'url':'http://127.0.0.1:8093','callback_origin':'http://localhost:8765','secret':'x'*40}
    monkeypatch.setattr(document_office,'settings',lambda:cfg)
    path,doc,s=opened(workspace,tmp_path)
    adapter=SpreadsheetEngine(workspace)
    setup=adapter.config(doc['id'],**args(s),browser_origin='http://localhost',browser_api_origin='http://localhost:8765')
    s=setup['session'];confirmed=workspace.store.blob(data(20))
    s.update(blob=confirmed,captured_sequence=1,capture_requests={'capture':1})
    workspace.store.put('session',s)
    monkeypatch.setattr(document_office,'download',lambda url,cfg:data(10))
    body={'key':s['engine_key'],'status':2,'filetype':'xlsx','url':cfg['url']+'/closing'}
    body['token']=jwt.encode(body,cfg['secret'],algorithm='HS256')
    assert adapter.callback(s['id'],s['ticket'],body)=={'error':0}
    current=workspace.store.get('session',s['id'])
    assert current['blob']==confirmed and current['state']=='recovering'
    recovery=[v for v in workspace.versions(doc['id']) if v.get('is_recovery')]
    assert len(recovery)==2 and {r['blob'] for r in recovery} >= {confirmed}
    with pytest.raises(DocumentConflict):
        workspace.save(doc['id'],**args(current),operation_id='unsafe',expected_revision=doc['revision_id'])
    chosen=next(v for v in recovery if v['blob']==confirmed)
    restored=workspace.restore(doc['id'],chosen['id'],**args(current),operation_id='restore-confirmed')['session']
    assert restored['state']=='draft'
    assert load_workbook(path).active['B1'].value==10


def test_report_pins_unsaved_snapshot_and_does_not_rewrite_source(workspace,tmp_path):
    from resource_links import ResourceLinks
    path,doc,s=opened(workspace,tmp_path)
    s.update(blob=workspace.store.blob(data(20)),session_revision=1,state='draft');workspace.store.put('session',s)
    snap=workspace.snapshot(doc['id'],**args(s),engine_state='fixture state',calculation='stale')
    with pytest.raises(DocumentConflict):
        workspace.create_report(doc['id'],snap['id'],'1','A1:C1','reject-stale')
    report=workspace.create_report(doc['id'],snap['id'],'1','A1:C1','report',allow_stale=True)
    assert report['reference']['snapshot_id']==snap['id']
    assert report['reference']['provenance']['unsaved']
    assert '00123' in __import__('pathlib').Path(report['document']['source_uri']).read_text()
    assert load_workbook(path).active['B1'].value==10
    assert workspace.create_report(doc['id'],snap['id'],'1','A1:C1','report',allow_stale=True)==report
    assert not ResourceLinks(workspace).status(report['reference']['id'])['source_changed']


def test_engine_type_and_namespace(workspace,tmp_path,monkeypatch):
    import document_office
    monkeypatch.setattr(document_office,'settings',lambda:{'url':'http://127.0.0.1:8093','callback_origin':'http://localhost:8765','secret':'x'*40})
    _,doc,s=opened(workspace,tmp_path)
    setup=SpreadsheetEngine(workspace).config(doc['id'],**args(s),browser_origin='http://localhost',browser_api_origin='http://localhost:8765')
    assert setup['config']['documentType']=='cell'
    assert setup['config']['editorConfig']['customization']['macrosMode']=='disable'
    assert '/spreadsheets/engine-io/' in setup['config']['document']['url']
    assert setup['config']['editorConfig']['plugins']['autostart']==[SpreadsheetEngine.plugin_guid]


if __name__ == '__main__':
    import sys
    import pytest
    raise SystemExit(pytest.main([__file__, *sys.argv[1:]]))
