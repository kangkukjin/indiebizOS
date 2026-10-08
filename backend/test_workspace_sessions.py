"""범용 작업 공간(workspace_sessions)·[self:workspace] 회귀 (2026-10-05).

docs/APP_COMPOSITION_ON_IBL_PLAN_2026_10_05.md §3-a·§3-b. 문서(원문·사무 투영)·시트(저장본 투영)·
코딩(git worktree)이 한 계약으로 열리고, 세션 배관은 언어에 나오지 않으며, 작성 창의 세션은 빼앗지 않는다.
실 엔진(ONLYOFFICE·RHWP)·모델 호출 없음.
"""
from pathlib import Path

import boot_paths  # noqa: F401
import pytest
import yaml

from coding_git import git
from coding_store import CodingStore
from office_sessions import DocumentConflict, DocumentUnsupported
from workspace_sessions import Workspace, operation_key

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def ws(tmp_path):
    return Workspace(tmp_path / "state", coding_store=CodingStore(tmp_path / "cstate"))


def test_text_document_flow_hides_session_plumbing(ws, tmp_path):
    path = tmp_path / "memo.md"
    path.write_text("첫 문단입니다.\n둘째 문단.\n", encoding="utf-8")
    opened = ws.open(path)
    r = opened["resource"]
    assert opened["kind"] == "document" and opened["session"] is None and opened["capabilities"]["propose"] is True
    assert ws.open(path)["resource"] == r                      # 같은 경로 = 같은 자료
    whole = ws.read(r)
    assert whole["text"].startswith("첫 문단") and whole["snapshot"] and whole["unsaved"] is False
    part = ws.read(r, {"start": 0, "end": 8})
    assert part["text"] == "첫 문단입니다." and part["selected_sha256"]
    p = ws.propose(r, {"start": 0, "end": 8, "selected_sha256": part["selected_sha256"]}, replacement="고친 문단입니다.")
    applied = ws.apply(r, p["proposal"])                        # 창이 없으니 행위자 신원으로 세션을 얻는다
    assert applied["applied"] is True and applied["text"].startswith("고친 문단") and applied["session_revision"] == 1
    assert ws.apply(r, p["proposal"])["applied"] is True        # 같은 호출 = 같은 작업 ID = 멱등
    with pytest.raises(DocumentConflict):
        ws.close(r)                                             # 초안이 남아 있으면 닫지 않는다
    saved = ws.save(r)
    assert saved["state"] == "saved" and path.read_text(encoding="utf-8").startswith("고친 문단")
    assert saved["revision"] != opened["revision"]
    versions = ws.versions(r)["items"]
    assert len(versions) == 2 and all(v["id"] for v in versions)
    first = [v for v in versions if v["parent"] is None][0]
    restored = ws.restore(r, first["id"])
    assert restored["restored"] == first["id"] and ws.read(r)["text"].startswith("첫 문단")
    assert ws.read(r)["unsaved"] is True
    ws.save(r)
    assert ws.close(r)["closed"] is True


def test_ui_session_is_never_stolen_but_reads_and_proposals_work(ws, tmp_path):
    path = tmp_path / "draft.txt"
    path.write_text("hello world", encoding="utf-8")
    r = ws.open(path)["resource"]
    held = ws.documents().acquire(r, "window-1")["session"]
    ws.documents().draft(r, held["id"], "window-1", held["engine_epoch"], 0, "ui-draft", "hello there")
    assert ws.open(path)["session"]["held_by"] == "window-1"
    snap = ws.read(r)
    assert snap["text"] == "hello there" and snap["unsaved"] is True   # 창의 초안을 읽는다(빼앗지 않고)
    p = ws.propose(r, {"start": 6, "end": 11}, replacement="friend")
    with pytest.raises(DocumentConflict, match="다른 창"):
        ws.apply(r, p["proposal"])
    with pytest.raises(DocumentConflict, match="다른 창"):
        ws.save(r)
    applied = ws.apply(r, p["proposal"], client="window-1")
    assert applied["text"] == "hello friend"
    stale = ws.propose(r, {"start": 0, "end": 5, "selected_sha256": "0" * 64}, replacement="x") if False else None
    with pytest.raises(DocumentConflict):
        ws.propose(r, {"start": 0, "end": 5, "selected_sha256": "0" * 64}, replacement="x")
    assert stale is None


def test_office_document_is_a_read_projection_not_an_edit_address(ws, tmp_path):
    from docx import Document
    doc = Document(); doc.add_paragraph("사무 문서 본문"); path = tmp_path / "office.docx"; doc.save(path)
    opened = ws.open(path)
    r = opened["resource"]
    assert opened["kind"] == "document" and opened["capabilities"]["propose"] is False
    read = ws.read(r)
    assert read["projection"] is True and "사무 문서 본문" in (read.get("text") or str(read))
    with pytest.raises(DocumentUnsupported):
        ws.propose(r, {"start": 0, "end": 3}, replacement="x")


def test_closed_sheet_reads_saved_projection_and_honest_apply(ws, tmp_path):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active; sheet.title = "자료"
    sheet.append(["항목", "값"]); sheet.append(["a", 1]); path = tmp_path / "book.xlsx"; book.save(path)
    opened = ws.open(path)
    r = opened["resource"]
    assert opened["kind"] == "sheet"
    snap = ws.snapshot(r)
    assert snap["snapshot"] and snap["unsaved"] is False and snap["source"] == "saved_file"
    read = ws.read(r, {"sheet": 0, "range": "A1:B2"})
    assert read["table"]["items"][0]["entered_value"] == "항목" and read["unsaved"] is False
    p = ws.propose(r, {"sheet": 0, "range": "B2:B2"}, values=[[2]])
    assert p["proposal"] and p["affected_cells"] == 1
    caps = ws.capabilities(r)
    if not caps["edit_native"]:
        with pytest.raises(DocumentUnsupported):
            ws.apply(r, p["proposal"])                          # 엔진이 없으면 성공으로 포장하지 않는다
    assert ws.versions(r)["items"][0]["parent"] is None


@pytest.fixture
def repo(tmp_path):
    repo = tmp_path / "repo"; repo.mkdir()
    git(repo, "init", "-b", "main")
    git(repo, "config", "user.name", "Workspace Test")
    git(repo, "config", "user.email", "workspace-test@example.invalid")
    (repo / "a.txt").write_text("line1\nline2\nline3\n")
    git(repo, "add", "."); git(repo, "commit", "-m", "seed")
    return repo


def test_code_workspace_shares_the_contract(ws, repo):
    """코딩 자료 = 프로젝트 폴더(2026-10-07). 저장은 기록(커밋), 복구는 기록으로. 자세한 회귀는 test_coding_projects."""
    opened = ws.open(repo)
    r = opened["resource"]
    assert opened["kind"] == "code" and r.startswith("coding_") and opened["capabilities"]["save"] is True
    assert ws.open(repo)["resource"] == r
    assert (repo / "목표.md").is_file()          # 목표 문서 없는 폴더를 가져오면 추론한 목표 문서가 생긴다
    files = ws.read(r)
    assert any(f["path"] == "a.txt" for f in files["items"])
    piece = ws.read(r, {"path": "a.txt", "start_line": 2, "end_line": 2})
    assert piece["text"] == "line2" and piece["fingerprint"]
    p = ws.propose(r, {"path": "a.txt", "start_line": 2, "end_line": 2}, replacement="LINE2")
    applied = ws.apply(r, p["proposal"])
    assert applied["applied"] is True and applied["path"] == "a.txt"
    assert ws.read(r, {"path": "a.txt"})["text"] == "line1\nLINE2\nline3"
    diff = ws.read(r, {"diff": True})
    assert diff["paths"] == ["a.txt", "목표.md"] and "+LINE2" in diff["patch"]
    with pytest.raises(ValueError, match="message"):
        ws.save(r)
    saved = ws.save(r, message="줄 수정")
    assert saved["state"] == "committed" and saved["commit"] and set(saved["paths"]) == {"a.txt", "목표.md"}
    assert ws.versions(r)["items"][0]["label"] == "줄 수정"
    with pytest.raises(DocumentUnsupported):
        ws.export(r, "copy.txt")
    assert ws.close(r)["closed"] is True

def test_operation_key_is_stable_and_blind_to_plumbing():
    a = operation_key("apply", "res", {"proposal": "p1", "expected": 3})
    assert a == operation_key("apply", "res", {"expected": 3, "proposal": "p1"})
    assert a != operation_key("apply", "res", {"proposal": "p1", "expected": 4})


def test_vocabulary_absorbed_session_ops_into_one_word():
    src = yaml.safe_load((ROOT / "data/packages/installed/tools/system_essentials/ibl_actions.yaml").read_text(encoding="utf-8"))
    actions = src["self"]["actions"] if "self" in src else src["nodes"]["self"]["actions"]
    workspace_ops = set(actions["workspace"]["ops"]["values"])
    assert workspace_ops == {"open", "snapshot", "read", "propose", "apply", "save", "export", "versions",
                             "restore", "close", "capabilities", "recover"}
    assert set(actions["document"]["ops"]["values"]) == {"inspect", "edit"}
    assert set(actions["sheet"]["ops"]["values"]) == {"find", "append", "update", "range", "range_write", "calculate"}
    from resource_links import package_module
    handler = package_module("system_essentials", "handler")
    assert set(handler._OP_DISPATCHERS["workspace_op"]) == workspace_ops
    assert set(handler._OP_DISPATCHERS["document_op"]) == {"inspect", "edit"}
    assert "open" not in handler._OP_DISPATCHERS["sheet_op"]
    for retired in ("essentials_document_workspace.py", "essentials_spreadsheet_workspace.py"):
        assert not (ROOT / "data/packages/installed/tools/system_essentials" / retired).exists()


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))
