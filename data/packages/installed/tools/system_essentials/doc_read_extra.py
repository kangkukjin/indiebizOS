"""doc_read_extra.py — [self:read] 확장 형식 (2026-09-15): hwp(v5)·hwpx·pptx·epub.

read_docx 와 같은 봉투를 낸다 — {success, metadata, text, blocks[, images]}.
  text   = 블록(문단·표)을 offset/limit/max_blocks 로 자른 본문
  blocks = 문서 IR 전체 (paragraph/heading/table) — [table:document] 로 흐른다
pages  = pptx 슬라이드·epub 챕터의 1-기반 범위('1-5,8'), pdf 와 같은 자.

의존: hwp → olefile(순수 파이썬), pptx → python-pptx, hwpx·epub → 표준 라이브러리(zip+xml).
암호·배포용 hwp 는 정직 거절(복호화 미지원). hwp v3(옛 형식)도 거절.
"""
import html
import io
import json
import os
import re
import struct
import zipfile
import zlib
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

from runtime_utils import expand_body_path


# ---------------------------------------------------------------- 공통 봉투

def _resolve(tool_input: dict, project_path: str):
    raw = tool_input.get("file_path") or tool_input.get("path") or tool_input.get("target") or ""
    if not raw:
        return None, json.dumps({"success": False, "error": "file_path가 제공되지 않았습니다."}, ensure_ascii=False)
    path = Path(expand_body_path(raw))
    if not path.is_absolute():
        path = Path(project_path) / path
    if not path.exists():
        return None, json.dumps({"success": False, "error": f"파일을 찾을 수 없습니다: {path}"}, ensure_ascii=False)
    if not path.is_file():
        return None, json.dumps({"success": False, "error": f"파일이 아닙니다(폴더): {path} — 폴더 목록은 [self:list]"}, ensure_ascii=False)
    return path, None


def _int(v, default):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def _block_text(b: dict) -> str:
    if b.get("type") == "table":
        rows = [b.get("columns") or []] + list(b.get("rows") or [])
        return "\n".join("\t".join(str(c) for c in r) for r in rows)
    return b.get("text", "")


def _finish(path: Path, blocks: list, tool_input: dict, extra_meta: dict, images: list | None = None) -> str:
    """docx 와 같은 부분 읽기 슬라이싱(offset/limit/max_blocks) + 봉투."""
    offset = max(0, _int(tool_input.get("offset", 0), 0))
    limit_raw = tool_input.get("limit")
    limit = _int(limit_raw, None) if limit_raw is not None else None
    max_blocks = _int(tool_input.get("max_blocks", 300), 300)

    total = len(blocks)
    if limit is not None and limit >= 0:
        eff = limit
    elif max_blocks > 0:
        eff = max_blocks
    else:
        eff = None
    start = min(offset, total)
    end = total if eff is None else min(start + eff, total)
    sliced = blocks[start:end]
    text = "\n\n".join(_block_text(b) for b in sliced)

    metadata = {
        "filename": path.name,
        "total_blocks": total,
        "offset": start,
        "returned_blocks": len(sliced),
        "truncated": end < total,
        **extra_meta,
    }
    if metadata["truncated"]:
        metadata["next_offset"] = end
        metadata["hint"] = f"전체 {total}블록 중 {start}~{end - 1}만 반환됨. 다음 호출에 offset={end}로 이어 읽기."
    res = {"success": True, "metadata": metadata, "text": text}
    if blocks:
        res["blocks"] = blocks
    if images:
        res["images"] = images
    return json.dumps(res, ensure_ascii=False)


def _range_indices(spec, total: int) -> list:
    """pdf 의 pages 와 같은 자: '2' · '1-5,8' → 0-기반 인덱스 목록. 없으면 전체."""
    if spec in (None, "", "all"):
        return list(range(total))
    from importlib.util import spec_from_file_location, module_from_spec
    s = spec_from_file_location("fs_read_range", Path(__file__).with_name("fs_read_range.py"))
    mod = module_from_spec(s)
    s.loader.exec_module(mod)
    return list(mod.pdf_page_indices(spec, total))


def _table_block(grid: list) -> dict | None:
    grid = [[str(c).strip() for c in r] for r in grid if r]
    if not grid:
        return None
    ncol = max(len(r) for r in grid)
    grid = [r + [""] * (ncol - len(r)) for r in grid]
    return {"type": "table", "columns": grid[0], "rows": grid[1:]}


def _save_images(path: Path, entries: list[tuple[str, bytes]]) -> tuple[list, Path | None]:
    if not entries:
        return [], None
    images_dir = path.parent / f"{path.stem}_images"
    images_dir.mkdir(exist_ok=True)
    info = []
    for name, data in entries:
        out = images_dir / os.path.basename(name)
        with open(out, "wb") as f:
            f.write(data)
        info.append({"name": out.name, "saved_path": str(out), "size": len(data)})
    return info, images_dir


# ---------------------------------------------------------------- HWP v5 (olefile)

_HWPTAG_BEGIN = 0x10
_TAG_PARA_TEXT = _HWPTAG_BEGIN + 51
_TAG_CTRL_HEADER = _HWPTAG_BEGIN + 55
_TAG_LIST_HEADER = _HWPTAG_BEGIN + 56
_TAG_TABLE = _HWPTAG_BEGIN + 61
# 문자 코드 < 32 의 폭(wchar 수): 확장·인라인 컨트롤은 8, 나머지 1 (hwp5 명세)
_CTRL_EXTENDED = {1, 2, 3, 11, 12, 14, 15, 16, 17, 18, 21, 22, 23}
_CTRL_INLINE = {4, 5, 6, 7, 8, 9, 19, 20}


def _hwp_records(data: bytes):
    pos, n = 0, len(data)
    while pos + 4 <= n:
        hdr = struct.unpack_from("<I", data, pos)[0]
        pos += 4
        tag, level, size = hdr & 0x3FF, (hdr >> 10) & 0x3FF, hdr >> 20
        if size == 0xFFF:
            if pos + 4 > n:
                return
            size = struct.unpack_from("<I", data, pos)[0]
            pos += 4
        yield tag, level, data[pos:pos + size]
        pos += size


def _hwp_para_text(payload: bytes) -> str:
    out = []
    codes = struct.unpack(f"<{len(payload) // 2}H", payload[: len(payload) // 2 * 2])
    i, n = 0, len(codes)
    while i < n:
        c = codes[i]
        if c in _CTRL_EXTENDED or c in _CTRL_INLINE:
            if c == 9:
                out.append("\t")
            elif c == 4:
                out.append(" ")
            i += 8
            continue
        if c == 10:
            out.append("\n")
        elif c == 13 or c < 32:
            pass
        else:
            out.append(chr(c))
        i += 1
    return "".join(out)


def _hwp_section_blocks(data: bytes, blocks: list) -> None:
    """섹션 레코드 → 문단/표 블록. 표: CTRL_HEADER('tbl ') → TABLE(행·열) → LIST_HEADER(셀)마다 문단."""
    table = None  # {"level": ctrl_level, "rows": r, "cols": c, "cells": [[...]]}

    def close_table():
        nonlocal table
        if not table:
            return
        cells = ["\n".join(c).strip() for c in table["cells"]]
        cols = max(1, table["cols"])
        grid = [cells[i:i + cols] for i in range(0, len(cells), cols)]
        tb = _table_block(grid)
        if tb:
            blocks.append(tb)
        table = None

    for tag, level, payload in _hwp_records(data):
        if table and level <= table["level"] and tag != _TAG_TABLE:
            close_table()
        if tag == _TAG_CTRL_HEADER and payload[:4] == b" lbt":
            close_table()
            table = {"level": level, "rows": 0, "cols": 0, "cells": [], "open": False}
        elif tag == _TAG_TABLE and table and not table["open"] and len(payload) >= 8:
            table["rows"], table["cols"] = struct.unpack_from("<HH", payload, 4)
            table["open"] = True
        elif tag == _TAG_LIST_HEADER and table and table["open"] and level == table["level"] + 1:
            table["cells"].append([])
        elif tag == _TAG_PARA_TEXT:
            txt = _hwp_para_text(payload).strip()
            if table and table["open"] and table["cells"]:
                if txt:
                    table["cells"][-1].append(txt)
            elif txt:
                blocks.append({"type": "paragraph", "text": txt})
    close_table()


def read_hwp(tool_input: dict, project_path: str) -> str:
    path, err = _resolve(tool_input, project_path)
    if err:
        return err
    try:
        import olefile
    except ImportError:
        return json.dumps({"success": False, "error": "hwp 읽기에 olefile 이 필요합니다 — [self:install_lib]{name:\"olefile\"}"}, ensure_ascii=False)
    if not olefile.isOleFile(str(path)):
        return json.dumps({"success": False, "error": "HWP v5(OLE) 형식이 아닙니다 — 옛 hwp(v3)·hwpx 가 아닌지 확인 (hwpx 는 format:\"hwpx\")."}, ensure_ascii=False)
    try:
        ole = olefile.OleFileIO(str(path))
    except Exception as e:  # noqa: BLE001
        return json.dumps({"success": False, "error": f"HWP 를 여는 중 문제: {e}"}, ensure_ascii=False)
    try:
        if not ole.exists("FileHeader"):
            return json.dumps({"success": False, "error": "FileHeader 스트림이 없어 HWP v5 로 읽을 수 없습니다."}, ensure_ascii=False)
        hdr = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", hdr, 36)[0] if len(hdr) >= 40 else 0
        compressed, encrypted, distribution = bool(flags & 1), bool(flags & 2), bool(flags & 4)
        ver = hdr[32:36]
        version = f"{ver[3]}.{ver[2]}.{ver[1]}.{ver[0]}" if len(ver) == 4 else "?"
        if encrypted:
            return json.dumps({"success": False, "error": "암호가 걸린 HWP 는 읽지 못합니다(복호화 미지원). 한글에서 암호를 풀어 다시 저장하세요."}, ensure_ascii=False)
        if distribution:
            return json.dumps({"success": False, "error": "배포용(읽기 전용 배포) HWP 는 본문이 암호화돼 읽지 못합니다. 원본 문서를 요청하세요."}, ensure_ascii=False)

        sections = sorted(
            (e for e in ole.listdir() if len(e) == 2 and e[0] == "BodyText" and e[1].startswith("Section")),
            key=lambda e: _int(e[1][7:], 0),
        )
        blocks: list = []
        parse_error = None
        for e in sections:
            data = ole.openstream("/".join(e)).read()
            try:
                if compressed:
                    data = zlib.decompress(data, -15)
                _hwp_section_blocks(data, blocks)
            except Exception as ex:  # noqa: BLE001 — 한 섹션 실패는 부분 결과로
                parse_error = f"{e[1]}: {ex}"
        partial = False
        if not blocks and ole.exists("PrvText"):
            # 본문 파싱이 비면 미리보기 텍스트(앞부분만)로 대체 — 정직 표지
            prv = ole.openstream("PrvText").read().decode("utf-16-le", errors="ignore").strip("\x00 \r\n")
            if prv:
                blocks = [{"type": "paragraph", "text": p.strip()} for p in re.split(r"\r?\n", prv) if p.strip()]
                partial = True

        images_info: list = []
        if tool_input.get("extract_images", True):
            entries = []
            for e in ole.listdir():
                if len(e) == 2 and e[0] == "BinData":
                    raw = ole.openstream("/".join(e)).read()
                    if compressed:
                        try:
                            raw = zlib.decompress(raw, -15)
                        except zlib.error:
                            pass
                    entries.append((e[1], raw))
            images_info, _ = _save_images(path, entries)

        meta = {
            "format": "hwp",
            "hwp_version": version,
            "total_sections": len(sections),
            "total_images": len(images_info),
            "total_paragraphs": sum(1 for b in blocks if b["type"] == "paragraph"),
            "total_tables": sum(1 for b in blocks if b["type"] == "table"),
        }
        if partial:
            meta["partial"] = True
            meta["hint"] = "본문 레코드 파싱이 비어 미리보기 텍스트(PrvText, 앞부분)로 대체했다."
        if parse_error:
            meta["parse_error"] = parse_error
        return _finish(path, blocks, tool_input, meta, images_info)
    finally:
        ole.close()


# ---------------------------------------------------------------- HWPX (zip + OWPML)

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _hwpx_cell_text(tc) -> str:
    parts = []
    for p in tc.iter():
        if _local(p.tag) == "t":
            parts.append("".join(p.itertext()))
        elif _local(p.tag) == "p" and parts and not parts[-1].endswith("\n"):
            parts.append("\n")
    return "".join(parts).strip()


def _hwpx_table(tbl, blocks: list) -> None:
    grid = []
    for tr in (x for x in tbl.iter() if _local(x.tag) == "tr"):
        grid.append([_hwpx_cell_text(tc) for tc in tr if _local(tc.tag) == "tc"])
    tb = _table_block(grid)
    if tb:
        blocks.append(tb)


def _hwpx_para(p, blocks: list) -> None:
    """문단 하나 → 표 밖 텍스트는 문단 블록, 안에 박힌 표(hp:tbl)는 표 블록(문서 순서 유지)."""
    buf: list = []

    def flush():
        txt = "".join(buf).strip()
        buf.clear()
        if txt:
            blocks.append({"type": "paragraph", "text": txt})

    def rec(el):
        for c in el:
            n = _local(c.tag)
            if n == "tbl":
                flush()
                _hwpx_table(c, blocks)
            elif n == "t":
                buf.append("".join(c.itertext()))
            else:
                rec(c)

    rec(p)
    flush()


def _hwpx_walk(el, blocks: list) -> None:
    for child in list(el):
        name = _local(child.tag)
        if name == "p":
            _hwpx_para(child, blocks)
        elif name == "tbl":
            _hwpx_table(child, blocks)
        else:
            _hwpx_walk(child, blocks)


def read_hwpx(tool_input: dict, project_path: str) -> str:
    path, err = _resolve(tool_input, project_path)
    if err:
        return err
    if not zipfile.is_zipfile(str(path)):
        return json.dumps({"success": False, "error": "HWPX(zip) 형식이 아닙니다 — 옛 hwp(v5) 는 format:\"hwp\"."}, ensure_ascii=False)
    blocks: list = []
    try:
        with zipfile.ZipFile(str(path)) as zf:
            names = zf.namelist()
            secs = sorted(
                (n for n in names if re.fullmatch(r"Contents/section\d+\.xml", n)),
                key=lambda n: _int(re.search(r"(\d+)", n).group(1), 0),
            )
            if not secs:
                return json.dumps({"success": False, "error": "Contents/section*.xml 이 없어 HWPX 본문을 찾지 못했습니다."}, ensure_ascii=False)
            for n in secs:
                root = ET.fromstring(zf.read(n))
                _hwpx_walk(root, blocks)
            images_info: list = []
            if tool_input.get("extract_images", True):
                entries = [(n, zf.read(n)) for n in names if n.startswith("BinData/") and not n.endswith("/")]
                images_info, _ = _save_images(path, entries)
    except ET.ParseError as e:
        return json.dumps({"success": False, "error": f"HWPX XML 파싱 실패: {e}"}, ensure_ascii=False)
    except Exception as e:  # noqa: BLE001
        return json.dumps({"success": False, "error": f"HWPX 를 읽는 중 문제: {e}"}, ensure_ascii=False)
    meta = {
        "format": "hwpx",
        "total_sections": len(secs),
        "total_images": len(images_info),
        "total_paragraphs": sum(1 for b in blocks if b["type"] == "paragraph"),
        "total_tables": sum(1 for b in blocks if b["type"] == "table"),
    }
    return _finish(path, blocks, tool_input, meta, images_info)


# ---------------------------------------------------------------- PPTX (python-pptx)

def _pptx_shapes(shapes, out: list, images: list) -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            _pptx_shapes(sh.shapes, out, images)
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                images.append((f"slide_img_{len(images) + 1}.{sh.image.ext}", sh.image.blob))
            except Exception:  # noqa: BLE001
                pass
        if getattr(sh, "has_table", False) and sh.has_table:
            grid = [[c.text for c in row.cells] for row in sh.table.rows]
            tb = _table_block(grid)
            if tb:
                out.append(tb)
            continue
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in para.runs).strip() or para.text.strip()
                if txt:
                    lvl = getattr(para, "level", 0) or 0
                    out.append({"type": "paragraph", "text": ("  " * lvl + "• " + txt) if lvl else txt})


def read_pptx(tool_input: dict, project_path: str) -> str:
    path, err = _resolve(tool_input, project_path)
    if err:
        return err
    try:
        from pptx import Presentation
    except ImportError:
        return json.dumps({"success": False, "error": "pptx 읽기에 python-pptx 가 필요합니다 — [self:install_lib]{name:\"python-pptx\"}"}, ensure_ascii=False)
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        return json.dumps({"success": False, "error": f"PPTX 를 여는 중 문제: {e}"}, ensure_ascii=False)
    slides = list(prs.slides)
    total = len(slides)
    try:
        idx = _range_indices(tool_input.get("pages"), total)
    except Exception as e:  # noqa: BLE001
        return json.dumps({"success": False, "error": f"pages 범위 오류: {e}"}, ensure_ascii=False)
    blocks: list = []
    img_entries: list = []
    for i in idx:
        s = slides[i]
        title = ""
        try:
            if s.shapes.title is not None:
                title = s.shapes.title.text.strip()
        except Exception:  # noqa: BLE001
            title = ""
        blocks.append({"type": "heading", "level": 2, "text": f"슬라이드 {i + 1}" + (f" — {title}" if title else "")})
        body: list = []
        _pptx_shapes(s.shapes, body, img_entries)
        # 제목 문단은 heading 에 이미 실렸으므로 중복 제거
        body = [b for b in body if not (title and b.get("type") == "paragraph" and b.get("text") == title)]
        blocks.extend(body)
        if getattr(s, "has_notes_slide", False) and s.has_notes_slide:
            notes = (s.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                blocks.append({"type": "paragraph", "text": "[노트] " + notes})
    images_info: list = []
    if tool_input.get("extract_images", True) and img_entries:
        images_info, _ = _save_images(path, img_entries)
    meta = {
        "format": "pptx",
        "total_slides": total,
        "extracted_slides": [i + 1 for i in idx],
        "total_images": len(img_entries),
        "extracted_images": len(images_info),
        "total_tables": sum(1 for b in blocks if b["type"] == "table"),
    }
    return _finish(path, blocks, tool_input, meta, images_info)


# ---------------------------------------------------------------- EPUB (zip + OPF + XHTML)

class _XHTMLBlocks(HTMLParser):
    _BLOCK = {"p", "div", "li", "blockquote", "pre", "dd", "dt", "td", "th", "figcaption", "br"}
    _SKIP = {"script", "style", "head", "title", "nav"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list = []
        self._buf: list = []
        self._heading: int | None = None
        self._skip = 0
        self.title = ""
        self._in_title = False

    def _flush(self):
        txt = re.sub(r"[ \t\r\f\v]+", " ", "".join(self._buf)).strip()
        self._buf = []
        if not txt:
            return
        if self._heading:
            self.blocks.append({"type": "heading", "level": min(self._heading, 3), "text": txt})
        else:
            self.blocks.append({"type": "paragraph", "text": txt})

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip += 1
            self._in_title = tag == "title"
            return
        if re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = int(tag[1])
        elif tag in self._BLOCK:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self._skip = max(0, self._skip - 1)
            self._in_title = False
            return
        if re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self._heading = None
        elif tag in self._BLOCK:
            self._flush()

    def handle_data(self, data):
        if self._in_title:
            self.title += data
            return
        if self._skip:
            return
        self._buf.append(data)

    def close(self):
        super().close()
        self._flush()


def read_epub(tool_input: dict, project_path: str) -> str:
    path, err = _resolve(tool_input, project_path)
    if err:
        return err
    if not zipfile.is_zipfile(str(path)):
        return json.dumps({"success": False, "error": "EPUB(zip) 형식이 아닙니다."}, ensure_ascii=False)
    try:
        with zipfile.ZipFile(str(path)) as zf:
            container = ET.fromstring(zf.read("META-INF/container.xml"))
            rootfile = next((e.get("full-path") for e in container.iter() if _local(e.tag) == "rootfile"), None)
            if not rootfile:
                return json.dumps({"success": False, "error": "container.xml 에 rootfile 이 없습니다."}, ensure_ascii=False)
            opf = ET.fromstring(zf.read(rootfile))
            base = rootfile.rsplit("/", 1)[0] + "/" if "/" in rootfile else ""
            manifest = {}
            for it in opf.iter():
                if _local(it.tag) == "item":
                    manifest[it.get("id")] = (it.get("href"), it.get("media-type") or "")
            spine = [ir.get("idref") for ir in opf.iter() if _local(ir.tag) == "itemref"]
            book_title = next(("".join(e.itertext()).strip() for e in opf.iter() if _local(e.tag) == "title"), "")
            creator = next(("".join(e.itertext()).strip() for e in opf.iter() if _local(e.tag) == "creator"), "")
            chapters = []
            for idref in spine:
                href, mt = manifest.get(idref, (None, ""))
                if not href or ("html" not in mt and not re.search(r"\.x?html?$", href, re.I)):
                    continue
                name = base + href.split("#", 1)[0]
                try:
                    raw = zf.read(name)
                except KeyError:
                    continue
                chapters.append((name, raw))
            total = len(chapters)
            try:
                idx = _range_indices(tool_input.get("pages"), total)
            except Exception as e:  # noqa: BLE001
                return json.dumps({"success": False, "error": f"pages 범위 오류: {e}"}, ensure_ascii=False)
            blocks: list = []
            used = 0
            for i in idx:
                name, raw = chapters[i]
                parser = _XHTMLBlocks()
                parser.feed(raw.decode("utf-8", errors="ignore"))
                parser.close()
                if not parser.blocks:
                    continue
                used += 1
                if parser.blocks[0]["type"] != "heading":
                    label = html.unescape(parser.title.strip()) or os.path.basename(name)
                    blocks.append({"type": "heading", "level": 1, "text": f"챕터 {i + 1} — {label}"})
                blocks.extend(parser.blocks)
    except KeyError as e:
        return json.dumps({"success": False, "error": f"EPUB 구조 파일이 없습니다: {e}"}, ensure_ascii=False)
    except ET.ParseError as e:
        return json.dumps({"success": False, "error": f"EPUB XML 파싱 실패: {e}"}, ensure_ascii=False)
    except Exception as e:  # noqa: BLE001
        return json.dumps({"success": False, "error": f"EPUB 를 읽는 중 문제: {e}"}, ensure_ascii=False)
    meta = {
        "format": "epub",
        "title": book_title,
        "creator": creator,
        "total_chapters": total,
        "extracted_chapters": [i + 1 for i in idx],
        "chapters_with_text": used,
    }
    return _finish(path, blocks, tool_input, meta)
