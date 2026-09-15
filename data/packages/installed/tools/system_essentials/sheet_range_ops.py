"""XLSX 범위·서식 편집과 계산 캐시 갱신. 비대상 ZIP 파트는 재직렬화하지 않는다."""
from copy import deepcopy
import importlib.util
import math
import os
from pathlib import Path
import posixpath
import re
import shutil
import subprocess
import tempfile

from lxml import etree
from openpyxl.utils.cell import range_boundaries, get_column_letter

_spec = importlib.util.spec_from_file_location('essentials_ooxml', Path(__file__).with_name('essentials_ooxml.py'))
office = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(office)
S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'s': S}


def tag(n):
    return '{' + S + '}' + n


def _sheets(parts):
    book = office.xml(parts['xl/workbook.xml'])
    rels = office.xml(parts['xl/_rels/workbook.xml.rels'])
    targets = {e.get('Id'): posixpath.normpath('xl/' + e.get('Target', '')) if not e.get('Target', '').startswith('/')
               else e.get('Target').lstrip('/') for e in rels if e.get('TargetMode') != 'External'}
    return {e.get('name'): targets[e.get('{' + R + '}id')] for e in book.findall('s:sheets/s:sheet', NS)}


def _range(params):
    value = params.get('range')
    if not isinstance(value, str) or not re.fullmatch(r'[A-Za-z]{1,3}[1-9]\d*(?::[A-Za-z]{1,3}[1-9]\d*)?', value):
        raise ValueError('range는 A1 또는 A1:C10 형식이어야 합니다. 시트는 sheet로 지정하세요.')
    bounds = range_boundaries(value.upper())
    c1, r1, c2, r2 = bounds
    if c1 > c2 or r1 > r2 or c2 > 16384 or r2 > 1048576 or (c2-c1+1)*(r2-r1+1) > 10000:
        raise ValueError('유효한 Excel 범위(최대 10000셀)가 필요합니다.')
    return bounds


def _context(params):
    path, digest, parts, infos, comment = office.load(params, {'.xlsx', '.xlsm'})
    sheets = _sheets(parts)
    view = office.xml(parts['xl/workbook.xml']).find('s:bookViews/s:workbookView', NS)
    active = int(view.get('activeTab', '0')) if view is not None else 0
    names = list(sheets)
    sn = params.get('sheet') or names[min(max(active, 0), len(names)-1)]
    if sn not in sheets:
        raise ValueError(f'sheet를 찾을 수 없습니다. 사용 가능: {list(sheets)}')
    part = sheets[sn]
    root = office.xml(parts[part])
    return path, digest, parts, infos, comment, sn, part, root


def _strings(parts):
    if 'xl/sharedStrings.xml' not in parts:
        return []
    return [''.join(t.text or '' for t in e.xpath('s:t | s:r/s:t', namespaces=NS))
            for e in office.xml(parts['xl/sharedStrings.xml']).findall(tag('si'))]


def _value(cell, strings):
    if cell is None:
        return None
    t = cell.get('t')
    if t == 'inlineStr':
        return ''.join(n.text or '' for n in cell.findall('.//' + tag('t')))
    v = cell.find(tag('v'))
    if v is None or v.text is None:
        return None
    text = v.text
    if t == 's':
        return strings[int(text)]
    if t == 'b':
        return text == '1'
    if t in ('str', 'e', 'd'):
        return text
    value = float(text)
    return int(value) if value.is_integer() else value


def op_range(params):
    path, digest, parts, _, _, sn, _, root = _context(params)
    c1, r1, c2, r2 = _range(params)
    cells = {c.get('r'): c for c in root.findall('.//s:sheetData/s:row/s:c', NS)}
    strings = _strings(parts)
    rows = []
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            addr = f'{get_column_letter(c)}{r}'
            cell = cells.get(addr)
            f = cell.find(tag('f')) if cell is not None else None
            rows.append({'cell': addr, 'value': _value(cell, strings),
                         'formula': '=' + (f.text or '') if f is not None else None,
                         'formula_type': f.get('t', 'normal') if f is not None else None,
                         'style_id': cell.get('s', '0') if cell is not None else '0'})
    return {'success': True, 'path': str(path), 'sha256': digest, 'sheet': sn, 'items': rows,
            'note': 'value는 저장된 계산 캐시입니다. 최신 계산은 calculate로 확인하세요.'}


def _cell(root, col, row):
    data = root.find(tag('sheetData'))
    if data is None:
        raise ValueError('워크시트에 sheetData가 없습니다.')
    node = next((r for r in data if int(r.get('r')) == row), None)
    if node is None:
        node = etree.Element(tag('row'), r=str(row))
        index = next((i for i, r in enumerate(data) if int(r.get('r')) > row), len(data))
        data.insert(index, node)
    addr = f'{get_column_letter(col)}{row}'
    cell = next((c for c in node if c.get('r') == addr), None)
    if cell is None:
        cell = etree.Element(tag('c'), r=addr)
        index = next((i for i, c in enumerate(node) if range_boundaries(c.get('r'))[0] > col), len(node))
        node.insert(index, cell)
    return cell


def _style(styles, cell, fmt, cache):
    if not fmt:
        return
    old = int(cell.get('s', 0))
    if old in cache:
        cell.set('s', cache[old])
        return
    xfs = styles.find(tag('cellXfs'))
    xf = deepcopy(xfs[old])
    if 'number_format' in fmt:
        formats = styles.find(tag('numFmts'))
        if formats is None:
            formats = etree.Element(tag('numFmts'), count='0')
            styles.insert(0, formats)
        found = next((f for f in formats if f.get('formatCode') == fmt['number_format']), None)
        if found is None:
            next_fmt = max([163] + [int(f.get('numFmtId')) for f in formats]) + 1
            found = etree.SubElement(formats, tag('numFmt'), numFmtId=str(next_fmt), formatCode=fmt['number_format'])
            formats.set('count', str(len(formats)))
        xf.set('numFmtId', found.get('numFmtId'))
        xf.set('applyNumberFormat', '1')
    if 'bold' in fmt:
        fonts = styles.find(tag('fonts'))
        font = deepcopy(fonts[int(xf.get('fontId', 0))])
        for b in font.findall(tag('b')):
            font.remove(b)
        if fmt['bold']:
            font.insert(0, etree.Element(tag('b')))
        xf.set('fontId', str(len(fonts)))
        xf.set('applyFont', '1')
        fonts.append(font)
        fonts.set('count', str(len(fonts)))
    cache[old] = str(len(xfs))
    cell.set('s', cache[old])
    xfs.append(xf)
    xfs.set('count', str(len(xfs)))


def _invalidate(parts):
    for part in _sheets(parts).values():
        root = office.xml(parts[part])
        dirty = False
        for cell in root.findall('.//s:sheetData/s:row/s:c', NS):
            if cell.find(tag('f')) is not None:
                for value in cell.findall(tag('v')):
                    cell.remove(value)
                    dirty = True
        if dirty:
            parts[part] = office.encoded(root)
    book = office.xml(parts['xl/workbook.xml'])
    calc = book.find(tag('calcPr'))
    if calc is None:
        calc = etree.SubElement(book, tag('calcPr'))
    calc.set('fullCalcOnLoad', '1')
    calc.set('forceFullCalc', '1')
    parts['xl/workbook.xml'] = office.encoded(book)
    if 'xl/calcChain.xml' in parts:
        del parts['xl/calcChain.xml']
        for name in ['xl/_rels/workbook.xml.rels', '[Content_Types].xml']:
            root = office.xml(parts[name])
            for e in list(root):
                if (e.get('Type') or '').endswith('/calcChain') or e.get('PartName') == '/xl/calcChain.xml':
                    root.remove(e)
            parts[name] = office.encoded(root)


def op_range_write(params):
    path, digest, parts, infos, comment, sn, part, root = _context(params)
    if params.get('expected_sha256') != digest:
        raise ValueError('range의 sha256을 expected_sha256으로 지정하세요.')
    c1, r1, c2, r2 = _range(params)
    values = params.get('values')
    fmt = params.get('format', {})
    if not isinstance(fmt, dict) or set(fmt) - {'number_format', 'bold'}:
        raise ValueError('format은 number_format/ bold 객체입니다.')
    if 'number_format' in fmt and (not isinstance(fmt['number_format'], str) or len(fmt['number_format']) > 255):
        raise ValueError('number_format은 255자 이하 문자열입니다.')
    if 'bold' in fmt and type(fmt['bold']) is not bool:
        raise ValueError('bold는 boolean입니다.')
    if values is None and not fmt:
        raise ValueError('values 또는 format이 필요합니다.')
    if values is not None:
        if not isinstance(values, list) or len(values) != r2-r1+1 or any(not isinstance(r, list) or len(r) != c2-c1+1 for r in values):
            raise ValueError('values는 range와 크기가 같은 2차원 배열입니다.')
        for row in values:
            for v in row:
                if v is not None and type(v) not in (str, int, float, bool):
                    raise ValueError('셀 값은 문자열/숫자/boolean/null만 지원합니다.')
                if isinstance(v, float) and not math.isfinite(v):
                    raise ValueError('비유한 숫자는 지원하지 않습니다.')
                if isinstance(v, str) and (len(v) > 32767 or v == '='):
                    raise ValueError('셀 문자열은 최대 32767자이고 빈 수식은 허용되지 않습니다.')
    # 병합/배열 수식의 일부를 바꿔 구조를 깨지 않는다.
    for merged in root.findall('s:mergeCells/s:mergeCell', NS):
        a,b,c,d = range_boundaries(merged.get('ref'))
        if not (c < c1 or a > c2 or d < r1 or b > r2):
            raise ValueError('병합 셀과 겹치는 범위는 지원하지 않습니다.')
    if root.find(tag('sheetProtection')) is not None:
        raise ValueError('보호된 시트는 편집하지 않습니다.')
    for cell in root.findall('.//s:sheetData/s:row/s:c', NS):
        f = cell.find(tag('f'))
        if f is not None and f.get('t', 'normal') != 'normal':
            raise ValueError('공유/배열/데이터표 수식이 있는 시트의 범위 편집은 지원하지 않습니다.')
    dest = office.output_path(params, path, '_range')
    styles = office.xml(parts['xl/styles.xml']) if fmt else None
    cache = {}
    for ri, r in enumerate(range(r1, r2+1)):
        for ci, c in enumerate(range(c1, c2+1)):
            cell = _cell(root, c, r)
            if values is not None:
                for child in list(cell):
                    if child.tag in {tag('f'), tag('v'), tag('is')}:
                        cell.remove(child)
                cell.attrib.pop('t', None)
                v = values[ri][ci]
                if isinstance(v, str) and v.startswith('='):
                    etree.SubElement(cell, tag('f')).text = v[1:]
                elif isinstance(v, str):
                    cell.set('t', 'inlineStr')
                    t = etree.SubElement(etree.SubElement(cell, tag('is')), tag('t'))
                    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    t.text = v
                elif v is not None:
                    if type(v) is bool:
                        cell.set('t', 'b')
                    etree.SubElement(cell, tag('v')).text = str(int(v) if type(v) is bool else v)
            _style(styles, cell, fmt, cache)
    dimension = root.find(tag('dimension'))
    if dimension is not None:
        a,b,c,d = range_boundaries(dimension.get('ref', 'A1'))
        dimension.set('ref', f'{get_column_letter(min(a,c1))}{min(b,r1)}:{get_column_letter(max(c,c2))}{max(d,r2)}')
    parts[part] = office.encoded(root)
    if fmt:
        parts['xl/styles.xml'] = office.encoded(styles)
    if values is not None:
        _invalidate(parts)
    office.save(path, digest, dest, parts, infos, comment)
    return {'success': True, 'path': str(dest), 'source_path': str(path), 'sheet': sn,
            'range': params['range'], 'cells_affected': (c2-c1+1)*(r2-r1+1),
            'recalculated': False, 'note': '원본 보존. 값 변경 시 수식 캐시를 비웠습니다. calculate로 재계산하세요.'}


def _soffice():
    # render_artifact의 사용자 설정 이름과 같은 SOFFICE_PATH를 사용한다.
    candidates = [os.environ.get('SOFFICE_PATH'), shutil.which('soffice'), shutil.which('libreoffice'),
                  '/Applications/LibreOffice.app/Contents/MacOS/soffice',
                  r'C:\Program Files\LibreOffice\program\soffice.exe']
    return next((p for p in candidates if p and Path(p).is_file()), None)


def op_calculate(params):
    path, digest, parts, infos, comment = office.load(params, {'.xlsx'})
    if any('externalLink' in n or 'vbaProject' in n for n in parts):
        raise ValueError('외부 통합문서 링크·매크로가 있는 파일은 재계산하지 않습니다.')
    formulas = {}
    for sn, part in _sheets(parts).items():
        root = office.xml(parts[part])
        for cell in root.findall('.//s:sheetData/s:row/s:c', NS):
            f = cell.find(tag('f'))
            if f is not None:
                text = f.text or ''
                if f.get('t', 'normal') != 'normal' or re.search(r'(?i)(WEBSERVICE|DDE|RTD)\s*\(|\[|\|', text):
                    raise ValueError('공유/배열/데이터표·외부 자원 수식은 재계산 지원 밖입니다.')
                formulas[(sn, cell.get('r'))] = text
    binary = _soffice()
    if not binary:
        return {'success': False, 'error_type': 'dependency_missing', 'error': 'LibreOffice가 필요합니다. SOFFICE_PATH 또는 PATH로 soffice를 지정하세요.', 'recalculated': False}
    timeout = params.get('timeout', 60)
    if type(timeout) not in (int, float) or not 1 <= timeout <= 120:
        raise ValueError('timeout은 1~120초입니다.')
    dest = office.output_path(params, path, '_calculated')
    with tempfile.TemporaryDirectory(prefix='indiebiz-calc-') as td:
        temp = Path(td)
        src = temp/'input.xlsx'
        # 계산 엔진이 기존 캐시를 재사용하지 않도록 입력 사본 캐시를 비운다.
        calc_parts = dict(parts)
        _invalidate(calc_parts)
        office.save(path, digest, src, calc_parts, infos, comment)
        profile = temp/'profile'
        profile.mkdir()
        (profile/'user').mkdir()
        (profile/'user'/'registrymodifications.xcu').write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>'
            '</oor:items>', encoding='utf-8')
        out = temp/'out'
        out.mkdir()
        try:
            proc = subprocess.run([binary, '--headless', '--norestore', '--nodefault',
                                   '-env:UserInstallation=' + profile.as_uri(), '--convert-to',
                                   'xlsx:Calc MS Excel 2007 XML', '--outdir', str(out), str(src)],
                                  capture_output=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return {'success': False, 'error_type': 'timeout', 'error': '재계산 제한 시간을 초과했습니다.', 'recalculated': False}
        calculated = out/'input.xlsx'
        if proc.returncode or not calculated.exists():
            return {'success': False, 'error': 'LibreOffice가 계산 결과 파일을 만들지 못했습니다.', 'recalculated': False}
        _, _, result, _, _ = office.load({'path': str(calculated)}, {'.xlsx'})
    cached, errors = {}, []
    strings = _strings(result)
    for sn, part in _sheets(result).items():
        for cell in office.xml(result[part]).findall('.//s:sheetData/s:row/s:c', NS):
            if (sn, cell.get('r')) in formulas:
                value = _value(cell, strings)
                if cell.find(tag('v')) is None:
                    raise ValueError(f'{sn}!{cell.get("r")} 계산 캐시가 없습니다.')
                cached[(sn, cell.get('r'))] = (cell, value)
                if cell.get('t') == 'e':
                    errors.append({'sheet': sn, 'cell': cell.get('r'), 'error': value})
    if set(cached) != set(formulas):
        raise ValueError('계산 후 수식 셀 누락이 있어 결과를 게시하지 않습니다.')
    for sn, part in _sheets(parts).items():
        root = office.xml(parts[part])
        dirty = False
        for cell in root.findall('.//s:sheetData/s:row/s:c', NS):
            key = (sn, cell.get('r'))
            if key not in cached:
                continue
            calculated_cell, value = cached[key]
            cell.attrib.pop('t', None)
            kind = calculated_cell.get('t')
            if kind:
                cell.set('t', 'str' if kind in ('s', 'inlineStr') else kind)
            for v in cell.findall(tag('v')):
                cell.remove(v)
            node = etree.Element(tag('v'))
            node.text = str(value) if kind in ('s', 'inlineStr') else calculated_cell.find(tag('v')).text
            f = cell.find(tag('f'))
            cell.insert(cell.index(f)+1, node)
            dirty = True
        if dirty:
            parts[part] = office.encoded(root)
    office.save(path, digest, dest, parts, infos, comment)
    return {'success': not errors, 'path': str(dest), 'source_path': str(path), 'recalculated': True,
            'engine': 'LibreOffice', 'formula_count': len(formulas), 'errors': errors,
            'items': [{'sheet': sn, 'cell': addr, 'value': value} for (sn, addr), (_, value) in cached.items()],
            **({'error': '수식 계산 오류가 있습니다. errors와 결과 파일을 확인하세요.'} if errors else {}),
            'note': '셀 계산 캐시만 갱신했습니다. 도형·차트 파트와 수식 원문은 보존하며 차트 캐시는 갱신하지 않습니다.'}
