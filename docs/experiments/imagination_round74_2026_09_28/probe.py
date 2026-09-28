"""74회차: 파일·폴더 업무 — 목록·복사·폴더 생성·삭제·폴더 메모·용량·읽기·쓰기 (훈련 턴 · 무수정).

축 = 행동 기준 미조합 메뉴의 self:list·copy·mkdir·delete·folder_note·storage·read·write 와
table:filter/sort/groupby/each 의 조합을 사용자 도메인(다운로드 정리·강의 자료 주차별 정리·
사진/문서 백업·프로젝트 용량·폴더 메모·중복 파일·빈 폴더·숨김·NFD 한글·공백/특수문자 경로) 위에서.
라이브 탐침 = 모델 경로(agent_id·task_id IT74_*). 모든 요청 edition 2·project_id 컨텐츠·origin training.
★쓰기·삭제·복사·폴더 생성은 스크래치 트리(outputs/IT74_*)에서만. self:delete 는 선언상 **영구 삭제**
  (휴지통 아님) — 스크래치 트리 안 경로에만 쓴다. 사용자 실제 폴더(~/Downloads 등)는 list 읽기만.
★self:storage scan·folder_note set 은 저장소 색인(data/packages/storage_scans)에 스크래치 볼륨을
  하나 추가한다 — 탐침 끝에 앱 자신의 REST(DELETE /pcmanager/analyze/scan/{id})로 그 볼륨만 지운다.
★결과 파일 상태는 셸(os.stat·읽기)로 직접 대조한다(거짓 성공 탐지).
사용: .venv/bin/python probe.py before|after [--keep]
"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import unicodedata

HERE = Path(__file__).resolve().parent
BASE = Path('/Users/kangkukjin/Desktop/AI/indiebizOS')
OUT = BASE / 'outputs'
TREE = OUT / 'IT74_tree'
BAK = OUT / 'IT74_backup'
W = '~workspace/outputs/IT74_tree'
WB = '~workspace/outputs/IT74_backup'
SCANS_JSON = BASE / 'data/packages/storage_scans/scans.json'
AGENT = {'agent_id': 'IT74_probe', 'task_id': 'IT74_task'}
LOG = []
PY = str(BASE / '.venv/bin/python')
NFD = lambda s: unicodedata.normalize('NFD', s)  # noqa: E731
OLD = time.mktime((2025, 1, 15, 9, 0, 0, 0, 0, -1))     # 오래된 파일 시각
RECENT = time.mktime((2026, 9, 25, 9, 0, 0, 0, 0, -1))  # 최근 파일 시각


def post(route, payload, method='POST'):
    proc = subprocess.run(['curl', '-sS', '--max-time', '240', '-X', method, f'http://127.0.0.1:8765/{route}',
                           '-H', 'Content-Type: application/json', '--data-binary', '@-'],
                          input=json.dumps(payload), text=True, capture_output=True, check=True)
    return json.loads(proc.stdout)


def ex(code, **extra):
    payload = dict(code='#!ibl edition=2\n' + code, edition=2, project_id='컨텐츠',
                   origin='training', **AGENT, **extra)
    response = post('ibl/execute', payload)
    slim = {k: v for k, v in response.items() if k not in ('evidence', 'source_map', 'recordings', 'value_wire')}
    LOG.append({'request': payload, 'response': slim})
    return response


def codes(r):
    return [i.get('code') for i in r.get('issues') or []]


def issues_text(r):
    return json.dumps([{k: i.get(k) for k in ('code', 'message', 'hint')} for i in r.get('issues') or []],
                      ensure_ascii=False)


def err(r):
    e = r.get('error')
    if isinstance(e, dict):
        return json.dumps({k: e.get(k) for k in ('code', 'message', 'hint')}, ensure_ascii=False)[:400]
    return str(e)[:400] if e else issues_text(r)[:400]


def put(path, data=b'x', mtime=RECENT):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data if isinstance(data, bytes) else data.encode())
    os.utime(path, (mtime, mtime))


def make_pptx(path, title):
    path.parent.mkdir(parents=True, exist_ok=True)
    src = f'''
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import io
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[5])
s.shapes.title.text = {title!r}
buf = io.BytesIO(); Image.new("RGB", (40, 30), (200, 30, 30)).save(buf, "PNG"); buf.seek(0)
s.shapes.add_picture(buf, Inches(1), Inches(2))
prs.save({str(path)!r})
'''
    subprocess.run([PY, '-c', src], check=True)


def make_docx(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    src = f'''
from docx import Document
from docx.shared import Inches
from PIL import Image
import io
d = Document(); d.add_paragraph({text!r})
buf = io.BytesIO(); Image.new("RGB", (20, 20), (0, 90, 200)).save(buf, "PNG"); buf.seek(0)
d.add_picture(buf, width=Inches(1)); d.save({str(path)!r})
'''
    subprocess.run([PY, '-c', src], check=True)


def setup():
    """스크래치 트리 — 다운로드·강의자료·사진·프로젝트 흉내. 사용자 폴더는 건드리지 않는다."""
    cleanup()
    d = TREE / 'Downloads'
    put(d / 'old_report.pdf', b'%PDF-1.4 old' + b'0' * 3000, OLD)
    put(d / '설치파일.dmg', b'\0' * 2_000_000, OLD)
    put(d / 'big_video.mp4', b'\0' * 5_000_000, RECENT)
    put(d / 'report.pdf', b'%PDF-1.4 same' + b'1' * 4000, RECENT)
    put(d / 'report (1).pdf', b'%PDF-1.4 same' + b'1' * 4000, RECENT)    # 내용까지 같은 중복
    put(d / 'report_v2.pdf', b'%PDF-1.4 diff' + b'2' * 4000, RECENT)     # 크기만 같은 다른 파일
    put(d / '.hidden_config', b'secret=1', RECENT)
    put(d / 'file with spaces & (special).txt', '공백과 특수문자', RECENT)
    put(d / NFD('강의계획서_NFD.txt'), 'NFD 이름의 강의계획서', RECENT)       # macOS 다운로드가 흔히 NFD
    (d / 'empty_dir').mkdir(parents=True, exist_ok=True)
    (d / 'nested' / 'deep_empty').mkdir(parents=True, exist_ok=True)
    put(d / 'nested' / 'note.txt', 'nested note', RECENT)
    L = TREE / '강의자료'
    make_pptx(L / '1주차_개요.pptx', '1주차 개요')
    put(L / '2주차_실습.pdf', b'%PDF-1.4 w2', RECENT)
    make_docx(L / '3주차_과제.docx', '3주차 과제 안내')
    put(L / '3주차_참고.txt', '참고 자료', RECENT)
    put(L / '기타_공지.txt', '공지', RECENT)
    P = TREE / '사진'
    put(P / 'IMG_0001.jpg', b'\xff\xd8' + b'a' * 1000, RECENT)
    put(P / 'IMG_0002.jpg', b'\xff\xd8' + b'b' * 1200, RECENT)
    # 이전 백업 — 원본엔 없고 백업에만 남은 사진(예: 폰에서 지운 뒤 백업으로만 남은 것)
    put(BAK / '사진' / 'IMG_0000_only_in_backup.jpg', b'\xff\xd8' + b'z' * 900, OLD)
    put(BAK / 'report.pdf', b'OLD BACKUP CONTENT', OLD)
    J = TREE / '프로젝트A'
    put(J / 'src' / 'main.py', 'print(1)\n' * 200, RECENT)
    put(J / 'assets' / 'logo.png', b'\x89PNG' + b'p' * 300_000, RECENT)
    put(J / 'docs' / 'README.md', '# A\n' * 50, RECENT)
    for i in range(1, 23):                                  # 확장자 22종 — 요약의 상한 관찰용
        put(J / 'misc' / f'f{i:02d}.e{i:02d}', b'm' * (100 + i), RECENT)
    R = TREE / '공유드라이브_읽기전용'                      # 남의 폴더·USB·공유 드라이브 흉내
    make_pptx(R / '외부강의.pptx', '외부 강의')
    os.chmod(R, 0o555)


def cleanup():
    for d in (TREE, BAK):
        if d.exists():
            for root, dirs, _files in os.walk(d):   # t18 의 읽기 전용 폴더 권한 복원
                for x in dirs:
                    os.chmod(os.path.join(root, x), 0o755)
            os.chmod(d, 0o755)
            shutil.rmtree(d)
    for p in OUT.glob('IT74_*'):
        if p.is_file():
            p.unlink()
        elif p.is_dir():
            shutil.rmtree(p)


def scratch_scan_ids():
    try:
        scans = json.loads(SCANS_JSON.read_text())
    except Exception:
        return []
    return [s['id'] for s in scans if str(s.get('root_path', '')).startswith(str(OUT / 'IT74_'))]


def drop_scratch_scans():
    """스크래치 볼륨만 앱 REST 로 지운다(사용자 볼륨 1·2는 손대지 않음)."""
    return [post(f'pcmanager/analyze/scan/{sid}', {}, method='DELETE') for sid in scratch_scan_ids()]


def names(v):
    return sorted(r.get('name') for r in (v or []) if isinstance(r, dict))


# ─────────────────────────────── 과제 ───────────────────────────────

def t01():
    """다운로드 폴더에 뭐가 있는지 — 폴더 먼저, 이름순으로 보여줘 (list → sort)."""
    r = ex(f'$l = [self:list]{{path:"{W}/Downloads"}}\n'
           'return $l >> [table:sort]{by:"name"} >> [table:select]{columns:["name","is_dir","size","mtime"]}')
    v = r.get('value') or []
    got = [x.get('name') for x in v]
    raw = ex(f'return [self:list]{{path:"{W}/Downloads"}} >> [table:select]{{columns:["name"]}}').get('value') or []
    raw_order = [x.get('name') for x in raw]
    hidden = [n for n in got if n.startswith('.')]
    fd = ex(f'$l = [self:list]{{path:"{W}/Downloads"}}\n'
            'return $l >> [table:sort]{by:"name"} >> [table:sort]{by:"is_dir", descending:true} >> [table:select]{columns:["name"]}')
    tup = ex(f'$l = [self:list]{{path:"{W}/Downloads"}}\nreturn sorted($l, ($r)=>[!$r.is_dir, $r.name])')
    first2 = [x.get('name') for x in (fd.get('value') or [])][:2]
    return (r.get('success') is True and len(got) == 11 and got == sorted(got) and first2 == ['empty_dir', 'nested'],
            f"n={len(got)} sorted={got == sorted(got)} folders_first(2-pass)={first2} tuple_key={err(tup)[:120]} hidden={hidden} raw_order_sorted={raw_order == sorted(raw_order)} raw={raw_order}")


def t02():
    """다운로드에서 한 달 넘게 안 건드린 파일만 (list → filter mtime)."""
    r = ex(f'$l = [self:list]{{path:"{W}/Downloads"}}\n'
           'return $l >> [table:filter]{where:($r)=> !$r.is_dir && $r.mtime < "2026-08-28"} >> [table:select]{columns:["name","mtime"]}')
    got = names(r.get('value'))
    return got == ['old_report.pdf', '설치파일.dmg'], f"got={got} err={err(r) if not r.get('success') else ''}"


def t03():
    """다운로드에서 제일 큰 파일 3개 (list → filter 파일 → sort size desc → take)."""
    r = ex(f'return [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir}}'
           ' >> [table:sort]{by:"size", descending:true} >> [table:take]{n:3} >> [table:select]{columns:["name","size"]}')
    got = [(x.get('name'), x.get('size')) for x in r.get('value') or []]
    return [g[0] for g in got] == ['big_video.mp4', '설치파일.dmg', 'report_v2.pdf'] or \
        [g[0] for g in got][:2] == ['big_video.mp4', '설치파일.dmg'], f"got={got}"


def t04():
    """(사용자 실제 폴더, 읽기만) 내 다운로드 폴더 용량 큰 순 5개 — 이름은 보고서에 싣지 않는다."""
    r = ex('return [self:list]{path:"~/Downloads"} >> [table:filter]{where:($r)=> !$r.is_dir}'
           ' >> [table:sort]{by:"size", descending:true} >> [table:take]{n:5} >> [table:select]{columns:["size","mtime","is_dir"]}')
    v = r.get('value') or []
    sizes = [x.get('size') for x in v]
    return r.get('success') is True and sizes == sorted(sizes, reverse=True), f"n={len(v)} sizes={sizes} err={err(r) if not r.get('success') else ''}"


def t05():
    """다운로드에서 '강의' 들어간 파일 — NFD 로 저장된 한글 이름도 찾아야 (list pattern · filter contains)."""
    a = ex(f'return [self:list]{{path:"{W}/Downloads", pattern:"*강의*"}} >> [table:select]{{columns:["name","path"]}}')
    b = ex(f'return [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> contains($r.name, "강의")}} >> [table:select]{{columns:["name"]}}')
    va, vb = a.get('value') or [], b.get('value') or []
    pa = (va[0].get('path') if va else '') or ''
    exists = os.path.exists(pa)
    nfd_kept = pa and unicodedata.is_normalized('NFD', os.path.basename(pa)) and not unicodedata.is_normalized('NFC', os.path.basename(pa))
    return len(va) == 1 and len(vb) == 1 and exists, f"pattern={len(va)} contains={len(vb)} path_exists={exists} nfd_preserved={nfd_kept}"


def t06():
    """숨김 파일(.으로 시작) 빼고 파일이 몇 개인지 (list → filter slice → len)."""
    r = ex(f'$f = [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir && $r.name[0:1] != "."}}\n'
           f'$h = [self:list]{{path:"{W}/Downloads", pattern:".*"}}\n'
           'return {count:len($f), hidden:len($h)}')
    v = r.get('value') or {}
    return v.get('count') == 8 and v.get('hidden') == 1, f"value={v} err={err(r) if not r.get('success') else ''}"


def t07():
    """확장자별 파일 수·용량 (list → compute 확장자 → groupby)."""
    r = ex(f'$f = [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir}}\n'
           '$e = $f >> [table:compute]{set:($r)=>{ext: lower(split($r.name, ".")[-1])}}\n'
           'return $e >> [table:groupby]{by:"ext", agg:{개수:["count","name"], 용량:["sum","size"]}}')
    v = r.get('value') or {}
    items = v.get('items') if isinstance(v, dict) else v
    pdf = [x for x in items or [] if x.get('ext') == 'pdf']
    return bool(pdf) and pdf[0].get('개수') == 4, f"groups={[(x.get('ext'), x.get('개수'), x.get('용량')) for x in items or []]} err={err(r) if not r.get('success') else ''}"


def t08():
    """강의 자료를 주차별 폴더로 정리 — 'N주차_' 앞머리로 폴더를 만들고 그 안에 복사 (each mkdir+copy)."""
    r = ex(f'$f = [self:list]{{path:"{W}/강의자료"}} >> [table:filter]{{where:($r)=> !$r.is_dir && contains($r.name, "주차_")}}\n'
           '$done = $f >> [table:each]{\n'
           '  $wk = split($it.name, "_")[0]\n'
           f'  $d = [self:mkdir]{{path:f"{W}/강의자료_정리/${{$wk}}"}}\n'
           '  $c = [self:copy]{src:$it.path, dest:f"${$d.path}/${$it.name}"}\n'
           '  return {week:$wk, file:$it.name}\n'
           '}\n'
           'return $done')
    tgt = TREE / '강의자료_정리'
    got = sorted(str(p.relative_to(tgt)) for p in tgt.rglob('*') if p.is_file()) if tgt.exists() else []
    want = ['1주차/1주차_개요.pptx', '2주차/2주차_실습.pdf', '3주차/3주차_과제.docx', '3주차/3주차_참고.txt']
    return got == want, f"disk={got} value={r.get('value')} err={err(r) if not r.get('success') else ''}"


def t09():
    """방금 만든 주차 폴더에 무엇이 들어갔는지 폴더별 개수 (list 폴더 → each list → len)."""
    r = ex(f'$w = [self:list]{{path:"{W}/강의자료_정리"}} >> [table:filter]{{where:($r)=> $r.is_dir}} >> [table:sort]{{by:"name"}}\n'
           'return $w >> [table:each]{ $in = [self:list]{path:$it.path}\n return {week:$it.name, n:len($in)} }')
    v = r.get('value') or []
    got = [(x.get('week'), x.get('n')) for x in v]
    return got == [('1주차', 1), ('2주차', 1), ('3주차', 2)], f"got={got} err={err(r) if not r.get('success') else ''}"


def t10():
    """사진 폴더를 백업 폴더로 복사 — 백업에만 남아 있던 옛 사진은 그대로 있어야 (copy 폴더 → 기존 백업)."""
    before = sorted(p.name for p in (BAK / '사진').iterdir())
    r = ex(f'return [self:copy]{{src:"{W}/사진", dest:"{WB}/사진"}}')
    after = sorted(p.name for p in (BAK / '사진').iterdir()) if (BAK / '사진').exists() else []
    kept = 'IMG_0000_only_in_backup.jpg' in after
    return kept and 'IMG_0001.jpg' in after, f"before={before} after={after} msg={str(r.get('value'))[:160]}"


def t11():
    """같은 이름 파일을 백업에 복사 — 덮어쓰지 말고 '(2)' 로 피해야 (선언: 이름 충돌은 (2), 덮어쓰지 않는다)."""
    r = ex(f'return [self:copy]{{src:"{W}/Downloads/report.pdf", dest:"{WB}/report.pdf"}}')
    content = (BAK / 'report.pdf').read_bytes()[:20]
    siblings = sorted(p.name for p in BAK.iterdir() if p.name.startswith('report'))
    return content == b'OLD BACKUP CONTENT' and len(siblings) == 2, f"backup_head={content!r} siblings={siblings} msg={str(r.get('value'))[:140]}"


def t12():
    """공백·특수문자·한글 폴더 이름으로 복사하고 되읽기 (copy → read)."""
    r = ex(f'$c = [self:copy]{{src:"{W}/Downloads/file with spaces & (special).txt", dest:"{WB}/백업 폴더 (9월)/메모 & 목록.txt"}}\n'
           f'return [self:read]{{path:"{WB}/백업 폴더 (9월)/메모 & 목록.txt"}}.text')
    return r.get('value') == '공백과 특수문자', f"value={r.get('value')!r} err={err(r) if not r.get('success') else ''}"


def t13():
    """중복 파일 찾기 — 크기가 같은 파일 묶음을 찾고 내용까지 같은지 확인 (list → groupby size → filter)."""
    r = ex(f'$f = [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir}}\n'
           '$g = $f >> [table:groupby]{by:"size", agg:{n:["count","name"]}}\n'
           '$sizes = $g.items >> [table:filter]{where:($r)=> $r.n > 1} >> [table:each]{ $it.size }\n'
           'return $f >> [table:filter]{where:($r)=> len(intersection([$r.size], $sizes)) > 0} >> [table:sort]{by:"name"}')
    v = r.get('value') or []
    groups = [[x.get('name') for x in v]] if v else []
    # 기대: 크기 같은 3개(report·report (1)·report_v2) 중 내용 같은 건 2개 — 내용 비교 수단이 있는가?
    return r.get('success') is True and groups == [['report (1).pdf', 'report.pdf', 'report_v2.pdf']], \
        f"groups={groups} err={err(r) if not r.get('success') else ''}"


def t14():
    """빈 폴더 찾기 — 다운로드 아래 모든 폴더 중 안이 비어 있는 것 (file_find 폴더 → each list → filter len 0)."""
    r = ex(f'$d = [self:file_find]{{path:"{W}/Downloads", pattern:"*"}}\n'
           '$dirs = $d.items >> [table:filter]{where:($r)=> $r.is_dir}\n'
           '$c = $dirs >> [table:each]{ $in = [self:list]{path:$it.path}\n return {name:$it.name, n:len($in)} }\n'
           'return $c >> [table:filter]{where:($r)=> $r.n == 0} >> [table:select]{columns:["name"]}')
    got = names(r.get('value'))
    return got == ['deep_empty', 'empty_dir'], f"got={got} err={err(r) if not r.get('success') else ''}"


def t15():
    """1GB 넘는 파일이 있으면 나에게 알려줘 (조건 → notify_user, check 만)."""
    r = ex(f'$big = [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir && $r.size > 1000000000}}\n'
           '[if: len($big) > 0] { [self:notify_user]{message:f"큰 파일 ${len($big)}개"} } [else] { "없음" }', check=True)
    return r.get('status') in ('valid', 'incomplete') and not codes(r), f"status={r.get('status')} issues={issues_text(r)}"


def t16():
    """(스크래치) 한 달 넘은 설치파일(.dmg)만 지우기 — 지운 목록 보고 (list → filter → each delete)."""
    r = ex(f'$old = [self:list]{{path:"{W}/Downloads"}} >> [table:filter]{{where:($r)=> !$r.is_dir && contains($r.name, ".dmg") && $r.mtime < "2026-08-28"}}\n'
           '$gone = $old >> [table:each]{ $x = [self:delete]{path:$it.path}\n return $it.name }\n'
           f'$left = [self:list]{{path:"{W}/Downloads", pattern:"*.dmg"}}\n'
           'return {deleted:$gone, left:len($left)}')
    exists = (TREE / 'Downloads' / '설치파일.dmg').exists()
    return (not exists) and (r.get('value') or {}).get('left') == 0, f"value={r.get('value')} disk_exists={exists} err={err(r) if not r.get('success') else ''}"


def t17():
    """지우려는 파일이 이미 없으면 넘어가기 (delete missing_ok) + 없는 경로 삭제의 오류 문구."""
    a = ex(f'return [self:delete]{{path:"{W}/Downloads/없는파일.zip", missing_ok:true}}')
    b = ex(f'return [self:delete]{{path:"{W}/Downloads/없는파일.zip"}}')
    return a.get('success') is True and b.get('success') is False, f"missing_ok={str(a.get('value'))[:80]} plain={err(b)[:200]}"


def t18():
    """강의 슬라이드(pptx)·과제(docx) 내용만 읽기 — 읽기가 강의 폴더에 새 폴더를 만들면 안 되고, 읽기 전용 폴더도 읽혀야 (read)."""
    L = TREE / '강의자료'
    before = sorted(p.name for p in L.iterdir())
    a = ex(f'$x = [self:read]{{path:"{W}/강의자료/1주차_개요.pptx"}}\nreturn len($x.text)')
    b = ex(f'$x = [self:read]{{path:"{L}/3주차_과제.docx"}}\nreturn len($x.text)')              # 절대경로(B73-6 회피)
    b_ws = ex(f'$x = [self:read]{{path:"{W}/강의자료/3주차_과제.docx"}}\nreturn len($x.text)')  # ~workspace 토큰(B73-6)
    after = sorted(p.name for p in L.iterdir())
    new = sorted(set(after) - set(before))
    for n in new:
        shutil.rmtree(L / n)
    c = ex(f'$x = [self:read]{{path:"{W}/강의자료/1주차_개요.pptx", extract_images:false}}\nreturn len($x.text)')
    after_c = sorted(set(p.name for p in L.iterdir()) - set(before))
    ro = ex(f'$x = [self:read]{{path:"{W}/공유드라이브_읽기전용/외부강의.pptx"}}\nreturn len($x.text)')
    ro_off = ex(f'$x = [self:read]{{path:"{W}/공유드라이브_읽기전용/외부강의.pptx", extract_images:false}}\nreturn len($x.text)')
    ok = a.get('success') and b.get('success') and not new and ro.get('success')
    return ok, (f"new_dirs={new} pptx_len={a.get('value')} docx_abs_len={b.get('value')} docx_ws={err(b_ws)[:120] if not b_ws.get('success') else 'ok'}"
                f" | extract_images:false new={after_c} | readonly={err(ro)[:200] if not ro.get('success') else ro.get('value')} readonly_off={ro_off.get('value') if ro_off.get('success') else err(ro_off)[:120]}")


def t19():
    """폴더 목록을 색인 파일(마크다운)로 만들어 그 폴더에 두기 (list → each 줄 → join → write → read)."""
    r = ex(f'$f = [self:list]{{path:"{W}/프로젝트A"}} >> [table:sort]{{by:"name"}}\n'
           '$lines = $f >> [table:each]{ f"- ${$it.name}" }\n'
           f'$w = [self:write]{{path:"{W}/프로젝트A/INDEX.md", content:"# 프로젝트A\\n" + join("\\n", $lines)}}\n'
           f'return [self:read]{{path:"{W}/프로젝트A/INDEX.md"}}.text')
    txt = (TREE / '프로젝트A' / 'INDEX.md').read_text() if (TREE / '프로젝트A' / 'INDEX.md').exists() else ''
    return '- assets' in txt and r.get('value') == txt, f"disk={txt!r} err={err(r) if not r.get('success') else ''}"


def t20():
    """폴더 만들기 경계 — 이미 있는 폴더·파일과 같은 이름·NFD 한글 폴더 (mkdir)."""
    a = ex(f'return [self:mkdir]{{path:"{W}/Downloads/empty_dir"}}')
    b = ex(f'return [self:mkdir]{{path:"{W}/Downloads/report.pdf"}}')
    c = ex(f'return [self:mkdir]{{path:"{W}/새 폴더/하위/더 깊이"}}')
    va = a.get('value') or {}
    ok = va.get('existed') is True and b.get('success') is False and (TREE / '새 폴더/하위/더 깊이').is_dir()
    return ok, f"existing={va} file_clash={err(b)[:160]} nested={(TREE / '새 폴더/하위/더 깊이').is_dir()}"


def t21():
    """없는 폴더·파일 경로를 목록으로 부르면 (list 오류 문구)."""
    a = ex(f'return [self:list]{{path:"{W}/없는폴더"}}')
    b = ex(f'return [self:list]{{path:"{W}/Downloads/report.pdf"}}')
    return a.get('success') is False and b.get('success') is False, f"missing={err(a)[:220]} | file={err(b)[:220]}"


def t22():
    """프로젝트 폴더 용량 보기 — 스캔 후 하위 폴더별 용량 (storage scan → summary)."""
    s = ex(f'return [self:storage]{{op:"scan", path:"{W}/프로젝트A"}}')
    r = ex(f'return [self:storage]{{op:"summary", root_path:"{W}/프로젝트A"}}')
    v = r.get('value') or {}
    items = v.get('items') if isinstance(v, dict) else v
    keys = sorted({k for x in items or [] for k in x})
    sub = ex(f'return [self:storage]{{op:"summary", root_path:"{W}/프로젝트A/assets"}}')
    has_folder = any('folder' in k or 'path' in k for k in keys)
    return has_folder, f"scan={str(s.get('value'))[:120]} keys={keys} rows={[(x.get('extension'), x.get('total_size_mb')) for x in items or []]} sub_err={err(sub)[:220]}"


def t23():
    """폴더 메모 남기기 — '프로젝트A/assets = 로고 원본' 메모를 달고 다시 확인 (folder_note set → detail)."""
    a = ex(f'return [self:folder_note]{{op:"set", root_path:"{W}/프로젝트A", folder_path:"{W}/프로젝트A/assets", note:"IT74 로고 원본 모음"}}')
    a2 = ex(f'return [self:folder_note]{{op:"set", root_path:"{W}/프로젝트A", folder_path:"{W}/프로젝트A/assets", note:"IT74 로고 원본 모음(개정)"}}')
    ghost = ex(f'return [self:folder_note]{{op:"set", root_path:"{W}/프로젝트A", folder_path:"{W}/없는폴더", note:"IT74 유령"}}')
    sub = ex(f'return [self:folder_note]{{op:"set", root_path:"{W}/프로젝트A/src", folder_path:"{W}/프로젝트A/src", note:"IT74 하위 볼륨"}}')
    d = ex(f'$n = [self:folder_note]{{op:"detail", root_path:"{W}/프로젝트A"}}\n'
           'return $n.annotations >> [table:select]{columns:["folder_path","note"]}')
    rows = d.get('value') or []
    return (a.get('success') and len(rows) == 1 and ghost.get('success') is False,
            f"rows={rows} ghost={str(ghost.get('value') or err(ghost))[:120]} sub={err(sub)[:200] if not sub.get('success') else str(sub.get('value'))[:100]}")


def t24():
    """(파이프) 다운로드의 PDF 들만 골라 백업 폴더로 복사 — 교재 형태 `… >> [self:copy]{dest}` (B73-1 재확인)."""
    a = ex(f'return [self:list]{{path:"{W}/Downloads", pattern:"*.pdf"}} >> [self:copy]{{dest:"{WB}/pdf"}}', check=True)
    b = ex(f'$f = [self:list]{{path:"{W}/Downloads", pattern:"*.pdf"}}\n'
           f'return $f >> [table:each]{{ [self:copy]{{src:$it.path, dest:f"{WB}/pdf/${{$it.name}}"}} }}')
    n = len(list((BAK / 'pdf').glob('*.pdf'))) if (BAK / 'pdf').exists() else 0
    return a.get('status') in ('valid', 'incomplete') and not codes(a), f"pipe={a.get('status')} {issues_text(a)[:200]} | each_copied={n}"


TASKS = [globals()[f't{i:02d}'] for i in range(1, 25)]


def main():
    phase = sys.argv[1] if len(sys.argv) > 1 else 'before'
    setup()
    rows = []
    for task in TASKS:
        try:
            ok, detail = task()
        except Exception as exc:  # 탐침의 실패도 기록한다
            ok, detail = False, f'probe error: {type(exc).__name__}: {exc}'
        rows.append({'task': task.__name__, 'intent': task.__doc__, 'ok': bool(ok), 'detail': detail})
        print(('PASS' if ok else 'FAIL'), task.__name__, detail[:500])
    dropped = drop_scratch_scans()
    if '--keep' not in sys.argv:
        cleanup()
    passed = sum(r['ok'] for r in rows)
    print(f'{passed}/{len(rows)}  scratch_scans_dropped={dropped}  residue={[p.name for p in OUT.glob("IT74_*")]}')
    (HERE / f'{phase}.json').write_text(json.dumps({'passed': passed, 'total': len(rows), 'rows': rows, 'log': LOG},
                                                   ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
