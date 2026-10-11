"""52회차 합성 자료 — 공방 강좌 수강 DB(SQLite)·정산 장부 틀(xlsx)·수료증 틀(docx)·납부 안내서 틀(PDF 폼)·oracle.
사용: python3 prepare.py   (seed 고정, 실행자별 입력 사본 3개 + 변형 A 사본)"""
import json, random, shutil, sqlite3
from pathlib import Path
import docx, fitz, openpyxl
from docx.shared import Pt
ROOT = Path(__file__).resolve().parents[5]
OUT = ROOT / 'outputs/long_sentence_imagination/2026-10-11_52회차'
SUR = '김이박최정강조윤장임한오서신권황안송류홍'
GIV = ['서연', '민준', '지우', '하은', '도윤', '수아', '시우', '예린', '주원', '채원', '건우', '다인', '현우', '소율', '유준', '나윤', '선우', '지안', '은호', '가은']
COURSES = [('C01', '도자기 물레 입문', 180000, 8, '한가람'), ('C02', '목공 스툴 만들기', 240000, 8, '오세준'), ('C03', '가죽 카드지갑', 90000, 4, '문하늘'),
           ('C04', '실크스크린 포스터', 120000, 6, '백이든'), ('C05', '금속 반지 세공', 210000, 6, '한가람')]


def build_db(path, month, n_students, seed):
    r = random.Random(seed); path.unlink(missing_ok=True); db = sqlite3.connect(path)
    db.executescript('''create table students(id text primary key, name text, phone text, discount_pct integer);
    create table courses(id text primary key, title text, fee integer, sessions integer, teacher text);
    create table enrollments(id integer primary key, student_id text, course_id text, month text, status text, paid integer, attended integer, note text);''')
    db.executemany('insert into courses values(?,?,?,?,?)', COURSES)
    names = [r.choice(SUR) + r.choice(GIV) for _ in range(n_students)]
    if n_students >= 10:
        names[3] = names[7] = '김서연'                      # 동명이인(다른 id)
        names[5] = '알렉산드라 페트로브나 이바노바-스미르노프'   # 긴 이름(양식 넘침)
        names[8] = "제갈 O'Neil & 아들 <공방>"              # 따옴표·&·꺾쇠
    for i, n in enumerate(names):
        db.execute('insert into students values(?,?,?,?)', (f'S{i + 1:03d}', n, f'010-55{i:02d}-{r.randint(1000, 9999)}', r.choice([0, 0, 0, 10, 20])))
    eid = 0
    for i in range(n_students):
        for c in r.sample(COURSES, r.choice([1, 1, 1, 2, 2, 3])):
            for m in (month, '2026-08'):
                if m != month and r.random() < 0.6: continue
                eid += 1; fee = c[2]; disc = db.execute('select discount_pct from students where id=?', (f'S{i + 1:03d}',)).fetchone()[0]
                due = fee * (100 - disc) // 100; k = r.random()
                status, paid, att, note = 'active', due, r.randint(int(c[3] * 0.5), c[3]), None
                if k < 0.08: status, paid, att, note = 'cancelled', 0, r.randint(0, 1), '환불 완료'
                elif k < 0.26: paid = r.choice([0, due // 2, due - 10000])
                elif k < 0.31: paid, note = None, '입금 확인 중'          # 납부액 미입력 — 0 으로 치면 안 됨
                elif k < 0.35: att = 0                                    # 한 번도 안 나옴
                elif k < 0.38: paid = due + 20000; note = '과입금'        # 더 냄
                db.execute('insert into enrollments values(?,?,?,?,?,?,?,?)', (eid, f'S{i + 1:03d}', c[0], m, status, paid, att, note))
    db.commit(); return db


def oracle(db, month, threshold=80):
    rows = db.execute('''select e.id,s.id,s.name,s.discount_pct,c.id,c.title,c.fee,c.sessions,c.teacher,e.status,e.paid,e.attended,e.note
      from enrollments e join students s on s.id=e.student_id join courses c on c.id=e.course_id where e.month=? order by e.id''', (month,)).fetchall()
    led, cert, notice, unknown, cancelled = [], [], [], [], []
    for (eid, sid, name, disc, cid, title, fee, sess, teacher, status, paid, att, note) in rows:
        if status == 'cancelled': cancelled.append(eid); continue
        due = fee * (100 - disc) // 100; rate = round(att * 100 / sess, 1)
        row = {'enroll_id': eid, 'student_id': sid, 'name': name, 'course': title, 'fee': fee, 'discount_pct': disc, 'due': due, 'paid': paid,
               'unpaid': None if paid is None else due - paid, 'attended': att, 'sessions': sess, 'rate': rate, 'teacher': teacher}
        led.append(row)
        if paid is None: unknown.append(eid)
        elif paid < due: notice.append(eid)
        if paid is not None and paid >= due and rate >= threshold: cert.append(eid)
    known = [x for x in led if x['paid'] is not None]
    return {'month': month, 'threshold': threshold, 'enrollments_in_month': len(rows), 'cancelled': cancelled, 'ledger_rows': len(led),
            'sum_due': sum(x['due'] for x in led), 'sum_paid_known': sum(x['paid'] for x in known), 'sum_unpaid_known': sum(x['unpaid'] for x in known),
            'paid_unknown': unknown, 'certificates': cert, 'notices': notice, 'overpaid': [x['enroll_id'] for x in known if x['unpaid'] < 0], 'rows': led}


def xlsx_template(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '정산'
    ws.append(['접수번호', '수강생ID', '이름', '강좌', '수강료', '할인율', '청구액', '납부액', '미납액', '출석', '회차', '출석률'])
    for c in 'ABCDEFGHIJKL': ws.column_dimensions[c].width = 14
    s2 = wb.create_sheet('요약'); s2.append(['항목', '값'])
    for k, f in (('행 수', '=COUNTA(정산!A2:A2000)'), ('청구액 합계', '=SUM(정산!G2:G2000)'), ('납부액 합계', '=SUM(정산!H2:H2000)'), ('미납액 합계', '=SUM(정산!I2:I2000)'), ('납부 미확인 건수', '=COUNTBLANK(정산!H2:H2000)-COUNTBLANK(정산!A2:A2000)')): s2.append([k, f])
    wb.save(path)


def docx_template(path):
    d = docx.Document(); d.add_heading('수 료 증', 0)
    p = d.add_paragraph('발급번호: '); p.add_run('{{발급번호}}')
    p = d.add_paragraph('성명: '); p.add_run('{{이름}}').bold = True
    p = d.add_paragraph('강좌: '); p.add_run('{{'); r = p.add_run('강좌'); r.bold = True; p.add_run('}}')     # 워드에서 손본 틀처럼 자리표시자가 run 셋으로 쪼개짐
    d.add_paragraph('기간: {{기간}}  /  출석: {{출석}}회 중 {{회차}}회 과정 (출석률 {{출석률}}%)')
    t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = '지도 강사'; t.cell(0, 1).text = '{{강사}}'; t.cell(1, 0).text = '발급일'; t.cell(1, 1).text = '{{발급일}}'
    d.sections[0].footer.paragraphs[0].text = '한빛공방 · 이 증서의 진위는 발급번호 {{발급번호}} 로 확인합니다.'
    d.add_paragraph('위 사람은 한빛공방의 위 강좌를 성실히 마쳤기에 이 증서를 드립니다.'); d.add_paragraph('한빛공방 대표  {{대표}}')
    d.save(path)


def pdf_template(path):
    doc = fitz.open(); pg = doc.new_page(width=595, height=842)
    pg.insert_text((72, 80), 'Payment Notice / Hanbit Workshop', fontsize=16)
    y = 130
    for name, label, w in (('notice_no', 'Notice No.', 200), ('name', 'Name', 300), ('course', 'Course', 300), ('due', 'Amount due (KRW)', 150), ('paid', 'Paid (KRW)', 150), ('unpaid', 'Balance (KRW)', 150), ('deadline', 'Pay by', 150)):
        pg.insert_text((72, y + 14), label, fontsize=11)
        wd = fitz.Widget(); wd.field_name = name; wd.field_type = fitz.PDF_WIDGET_TYPE_TEXT; wd.rect = fitz.Rect(220, y, 220 + w, y + 20); wd.text_fontsize = 11; pg.add_widget(wd); y += 34
    wd = fitz.Widget(); wd.field_name = 'half_paid'; wd.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX; wd.rect = fitz.Rect(220, y, 236, y + 16); pg.add_widget(wd)
    pg.insert_text((72, y + 13), 'Partially paid', fontsize=11); doc.save(path)


def main():
    for sub in ('trainer', 'agent', '_oracle'): shutil.rmtree(OUT / sub, ignore_errors=True)
    (OUT / 'runs').mkdir(parents=True, exist_ok=True); (OUT / '_oracle').mkdir(); exp = {}
    for who, mode, month, n, seed in (('trainer', 'base', '2026-09', 70, 52), ('trainer', 'a', '2026-09', 6, 7), ('trainer', 'b', '2026-09', 70, 52), ('agent', 'base', '2026-09', 70, 52)):
        inp = OUT / who / mode / 'input'; inp.mkdir(parents=True); (OUT / who / mode / 'out').mkdir()
        db = build_db(inp / 'academy.db', month, n, seed); xlsx_template(inp / 'ledger_template.xlsx'); docx_template(inp / 'certificate_template.docx'); pdf_template(inp / 'notice_form.pdf')
        (inp / 'job.json').write_text(json.dumps({'month': month, 'period': '2026.09.01 ~ 2026.09.30', 'issue_date': '2026-10-11', 'pay_by': '2026-10-25', 'attendance_pct_min': 70 if mode == 'b' else 80, 'director': '강한빛'}, ensure_ascii=False, indent=1))
        o = oracle(db, month, 70 if mode == 'b' else 80); exp[f'{who}_{mode}'] = o; db.close()
        print(who, mode, {k: (len(v) if isinstance(v, list) else v) for k, v in o.items() if k != 'rows'})
    (OUT / '_oracle/expected.json').write_text(json.dumps(exp, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main()
