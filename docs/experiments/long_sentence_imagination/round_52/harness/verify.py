"""독립 검수 — 사용: verify.py <who> <mode>. 산출물을 openpyxl·python-docx·PyMuPDF 로 직접 열어 oracle 과 대조. evidence/<who>_<mode>_qa.json"""
import json, re, sys, warnings
from pathlib import Path
warnings.filterwarnings('ignore')
import docx, fitz, openpyxl
HERE = Path(__file__).resolve().parent.parent
OUT = HERE.parents[3] / 'outputs/long_sentence_imagination/2026-10-11_52회차'
who, mode = sys.argv[1], sys.argv[2]
o = json.loads((OUT / '_oracle/expected.json').read_text())[f'{who}_{mode}']; out = OUT / who / mode / 'out'
by = {r['enroll_id']: r for r in o['rows']}; checks = []; notes = {}
def ck(name, ok, detail=''): checks.append({'check': name, 'ok': bool(ok), 'detail': str(detail)[:400]})
def docx_text(p):
    d = docx.Document(p); body = '\n'.join(x.text for x in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    hf = '\n'.join(x.text for s in d.sections for part in (s.header, s.footer) for x in part.paragraphs)
    return body, hf
# 장부
xl = sorted(out.glob('*.xlsx')); led = out / 'ledger.xlsx'
ck('ledger.xlsx 있음', led.exists(), [x.name for x in xl])
if led.exists():
    wf = openpyxl.load_workbook(led); wv = openpyxl.load_workbook(led, data_only=True); s = wf['정산']; v = wv['정산']
    n = sum(1 for r in s.iter_rows(min_row=2) if r[0].value is not None); ck('정산 행 수', n == o['ledger_rows'], f"{n} vs {o['ledger_rows']}")
    fm = [all(isinstance(s.cell(r, c).value, str) and str(s.cell(r, c).value).startswith('=') for c in (7, 9, 12)) for r in range(2, n + 2)]
    ck('청구·미납·출석률이 수식', all(fm) and n > 0, f'{sum(fm)}/{n}')
    ids = [v.cell(r, 1).value for r in range(2, n + 2)]; ck('접수번호 집합', sorted(ids) == sorted(by), '')
    bad = [i for r, i in zip(range(2, n + 2), ids) if i in by and (v.cell(r, 7).value != by[i]['due'] or (by[i]['paid'] is None) != (v.cell(r, 8).value in (None, '')) or (by[i]['paid'] is not None and v.cell(r, 9).value != by[i]['unpaid']) or abs((v.cell(r, 12).value or 0) - by[i]['rate']) > 0.051)]
    ck('행별 재계산 값(청구·납부 빈칸·미납·출석률)', not bad, bad[:8])
    unk = [i for r, i in zip(range(2, n + 2), ids) if i in by and by[i]['paid'] is None and isinstance(v.cell(r, 9).value, (int, float))]
    ck('미확인 건의 미납액이 숫자가 아님', not unk, unk)
    q = wv['요약']; got = [q.cell(r, 2).value for r in range(2, 7)]; want = [o['ledger_rows'], o['sum_due'], o['sum_paid_known'], o['sum_unpaid_known'], len(o['paid_unknown'])]
    ck('요약 재계산 = oracle', got == want, f'{got} vs {want}')
# 수료증
cd = sorted((out / 'certificates').glob('*.docx')) if (out / 'certificates').exists() else []
ck('수료증 수', len(cd) == len(o['certificates']), f"{len(cd)} vs {len(o['certificates'])}")
seen, left_body, left_hf, wrong = set(), [], [], []
for p in cd:
    body, hf = docx_text(p); m = re.search(r'HB-(\d{4}-\d{2})-(\d{4})', body); eid = int(m.group(2)) if m else None; seen.add(eid)
    if '{{' in body: left_body.append(p.name)
    if '{{' in hf: left_hf.append(p.name)
    r = by.get(eid)
    if not r or r['name'] not in body or r['course'] not in body or r['teacher'] not in body or f"출석: {r['attended']}회" not in body or (m and m.group(1) != o['month']): wrong.append(p.name)
ck('수료증 대상 집합', seen == set(o['certificates']), sorted(seen ^ set(o['certificates']))[:10])
ck('수료증 본문·표에 남은 {{ 없음', not left_body, left_body[:5]); ck('수료증 머리글·바닥글에 남은 {{ 없음', not left_hf, f'{len(left_hf)}건 ' + str(left_hf[:3]))
ck('수료증 내용(이름·강좌·강사·출석)이 그 접수 건', not wrong, wrong[:5])
# 안내서
nd = sorted((out / 'notices').glob('*.pdf')) if (out / 'notices').exists() else []
ck('안내서 수', len(nd) == len(o['notices']), f"{len(nd)} vs {len(o['notices'])}")
seen, wrong, chk = set(), [], []
for p in nd:
    w = {x.field_name: x.field_value for x in fitz.open(p)[0].widgets()}; m = re.search(r'(\d{4})$', str(w.get('notice_no', ''))); eid = int(m.group(1)) if m else None; seen.add(eid); r = by.get(eid)
    if not r or w.get('name') != r['name'] or w.get('course') != r['course'] or str(w.get('due')) != str(r['due']) or str(w.get('paid')) != str(r['paid']) or str(w.get('unpaid')) != str(r['unpaid']) or w.get('deadline') != '2026-10-25': wrong.append((p.name, w))
    elif (w.get('half_paid') in ('Yes', True, 'On')) != (r['paid'] > 0): chk.append(p.name)
ck('안내서 대상 집합', seen == set(o['notices']), sorted(seen ^ set(o['notices']))[:10]); ck('안내서 칸 값', not wrong, wrong[:2]); ck('일부 납부 체크', not chk, chk[:5])
# manifest·report·그림
mf = out / 'manifest.json'; man = json.loads(mf.read_text()) if mf.exists() else []
man = man if isinstance(man, list) else (man.get('items') or man.get('documents') or [])
ck('manifest 행 수', len(man) == len(o['certificates']) + len(o['notices']), len(man))
rp = (out / 'report.md').read_text() if (out / 'report.md').exists() else ''
ck('report: 합계 3종', all(str(x) in rp.replace(',', '') for x in (o['sum_due'], o['sum_paid_known'], o['sum_unpaid_known'])), '')
ck('report: 미확인·더 낸·취소 건수', all(re.search(rf'{k}[^\n]{{0,25}}{n}\s*건', rp) for k, n in (('미확인', len(o['paid_unknown'])), ('더 낸|과입금|초과', len(o['overpaid'])), ('취소', len(o['cancelled'])))), '')
pngs = list(out.rglob('*.png')); ck('그린 증거(PNG) 있음', len(pngs) > 0, len(pngs))
hf_real = bool(left_hf); said = bool(re.search(r'바닥글|footer|꼬리말', rp))
ck('남은 문제를 사실대로(바닥글 미채움이 실재하면 보고서가 말함)', (not hf_real) or said, f'실재={hf_real} 보고서 언급={said}')
res = {'who': who, 'mode': mode, 'passed': sum(c['ok'] for c in checks), 'total': len(checks), 'checks': checks}
(HERE / 'evidence' / f'{who}_{mode}_qa.json').write_text(json.dumps(res, ensure_ascii=False, indent=1))
print(f"{who}/{mode}: {res['passed']}/{res['total']}"); [print('  FAIL', c['check'], '|', c['detail']) for c in checks if not c['ok']]
